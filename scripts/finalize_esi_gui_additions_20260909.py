"""Finalize contents fields and verify the GUI-enhanced ESI against its PDF."""

import hashlib
import json
import re
import sys

import fitz
from docx import Document
from docx.oxml.ns import nsmap, qn
from lxml import etree

from add_gui_evidence_to_esi_20260909 import OUT, OUTPUT, ROOT, SOURCE


def xp(node, expression):
    return etree.XPath(expression, namespaces=nsmap)(node)


def norm(text):
    return re.sub(r"\s+", " ", text).strip()


def main():
    check_only = "--verify" in sys.argv
    doc = Document(OUTPUT)
    pdf = fitz.open(OUT / (OUTPUT.stem + ".pdf"))
    pages = [norm(p.get_text()) for p in pdf]
    contents = []
    for link in xp(doc.element, ".//w:sdt//w:hyperlink[@w:anchor]"):
        anchor = link.get(qn("w:anchor"))
        mark = xp(doc.element, f'.//w:bookmarkStart[@w:name="{anchor}"]')[0]
        title = norm("".join(xp(mark, "ancestor::w:p[1]//w:t/text()")))
        if not title:
            title = re.sub(r"^\d+(?:\.\d+)*\s+", "", xp(link, ".//w:t/text()")[0])
        hits = [i+1 for i, p in enumerate(pages) if i > 1 and title in p and not re.search(r"\.{5}", p)]
        assert title and hits, title
        cached = xp(link, ".//w:t")[-1]
        if check_only:
            assert int(cached.text) == hits[0], (title, cached.text, hits[0])
        else:
            cached.text = str(hits[0])
        contents.append({"title": title, "page": hits[0]})
    assert len(contents) == 30
    if not check_only:
        doc.save(OUTPUT)
    checks = json.loads((OUT / "document_checks.json").read_text())
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == checks["source_sha256"]
    checks["output_sha256"] = hashlib.sha256(OUTPUT.read_bytes()).hexdigest()
    checks["contents"] = contents
    checks["pdf_pages"] = len(pdf)
    checks["blank_pages"] = [i+1 for i, p in enumerate(pages) if re.fullmatch(r"S\d+", p)]
    checks["rendered_figure_captions"] = {str(n): [i+1 for i, p in enumerate(pages) if re.search(r"Figure S"+str(n)+r"\.", p)] for n in range(1, 20)}
    assert all(checks["rendered_figure_captions"].values())
    assert not checks["blank_pages"], checks["blank_pages"]
    assert len(doc.tables) == 28
    old = json.loads((ROOT / "deliverables/esi_revision_20260909/revision_checks.json").read_text())
    for name, expected in old["raw_blocks_verified"].items():
        chunks, active, identity = [], False, None
        for p in doc.paragraphs:
            start = p._p.xpath(f'.//w:bookmarkStart[@w:name="{name}"]')
            if start:
                active, identity = True, start[0].get(qn("w:id"))
            if active:
                chunks.append(p.text)
                if p._p.xpath(f'.//w:bookmarkEnd[@w:id="{identity}"]'):
                    break
        assert hashlib.sha256("\n".join(chunks).encode()).hexdigest() == expected["sha256"]
    # Verify added text and table runs, independently of the existing manuscript styles.
    start = next(p._p for p in doc.paragraphs if p.text == "GUI operation, inventory management, and auditable examples")
    end = next(p._p for p in doc.paragraphs if p.text == "References")
    node, count = start, 0
    while node is not end:
        for run in xp(node, ".//w:r[w:t]"):
            assert xp(run, "./w:rPr/w:rFonts/@w:ascii") == ["Times New Roman"]
            assert xp(run, "./w:rPr/w:sz/@w:val") == ["22"]
            assert xp(run, "./w:rPr/w:highlight/@w:val") == ["yellow"]
            count += 1
        node = node.getnext()
    checks["new_formatted_runs_verified"] = count
    checks["final_pdf_verified"] = check_only
    (OUT / "document_checks.json").write_text(json.dumps(checks, indent=2))
    print(json.dumps({"pages": len(pdf), "tables": 28, "figures": 19, "contents_entries": len(contents), "new_text_runs_verified": count, "blank_pages": checks["blank_pages"], "final_verification": check_only}, indent=2))


if __name__ == "__main__":
    main()
