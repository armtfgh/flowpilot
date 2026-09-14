"""Revise the current ESI without overwriting it or altering archived experiments."""

from __future__ import annotations

import hashlib
import json
import shutil
from copy import deepcopy
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "manuscript/esi.docx"
OUTPUT = ROOT / "manuscript/esi_revised_20260909.docx"
OUT = ROOT / "deliverables/esi_revision_20260909"
BENCH = ROOT / "deliverables/manuscript_benchmark_visualizations_20260825"
FIGURES = ROOT / "deliverables/flowpilot_esi_revision_20260902/figures"
CAMPAIGN = ROOT / "ablation_results/manuscript_benchmark/alternative_frontier_three_repeat_20260824"
FP = CAMPAIGN / "generation/photochemical_oxidation/claude_opus_flowpilot/repeat_01"
OS = CAMPAIGN / "generation/photochemical_oxidation/claude_opus_one_shot/repeat_01"
AUDIT = BENCH / "prompt_audit_claude_opus_46"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def font(run, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    run.bold = bold
    run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")


def text(paragraph, value, bold_prefix=None):
    # Retain heading bookmarks and paragraph-level formatting, not stale run text.
    for child in list(paragraph._p):
        if child.tag not in (qn("w:pPr"), qn("w:bookmarkStart"), qn("w:bookmarkEnd")):
            paragraph._p.remove(child)
    if bold_prefix and value.startswith(bold_prefix):
        font(paragraph.add_run(bold_prefix), True)
        font(paragraph.add_run(value[len(bold_prefix):]))
    else:
        font(paragraph.add_run(value))


def find(doc, prefix):
    return next(p for p in doc.paragraphs if p.text.strip().startswith(prefix))


def before(anchor, value, bold_prefix=None):
    p = anchor.insert_paragraph_before()
    p.style = "Normal"
    text(p, value, bold_prefix)
    return p


def after(anchor, value, bold_prefix=None):
    node = OxmlElement("w:p")
    anchor._p.addnext(node)
    p = Paragraph(node, anchor._parent)
    p.style = "Normal"
    text(p, value, bold_prefix)
    return p


def replace_block(start, stop, value, bookmark, expected):
    node = start._p.getnext()
    while node is not None and node is not stop._p:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt
    paragraphs = []
    for line in value.split("\n"):
        p = before(stop, line)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1
        p.paragraph_format.keep_with_next = False
        paragraphs.append(p)
    marker = OxmlElement("w:bookmarkStart")
    marker.set(qn("w:name"), bookmark)
    marker.set(qn("w:id"), str(20000 + len(expected)))
    paragraphs[0]._p.insert(1 if paragraphs[0]._p.pPr is not None else 0, marker)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), marker.get(qn("w:id")))
    paragraphs[-1]._p.append(end)
    expected[bookmark] = value


def populate_table(table, headers, rows, widths):
    template = deepcopy(table.rows[1]._tr)
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    for row in rows:
        table._tbl.append(deepcopy(template))
    for row, values in zip(table.rows, [headers] + rows):
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Inches(width)
            for extra in list(cell.paragraphs)[1:]:
                extra._p.getparent().remove(extra._p)
            p = cell.paragraphs[0]
            text(p, str(value))
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = False
            if row is table.rows[0]:
                for run in p.runs:
                    run.bold = True
        props = row._tr.get_or_add_trPr()
        if props.find(qn("w:cantSplit")) is None:
            props.append(OxmlElement("w:cantSplit"))
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for run in [r for c in table.rows[0].cells for p in c.paragraphs for r in p.runs]:
        run.bold = True
    header_props = table.rows[0]._tr.get_or_add_trPr()
    if header_props.find(qn("w:tblHeader")) is None:
        header_props.append(OxmlElement("w:tblHeader"))


def figure(doc, number, path, caption, width, own_page=False):
    cap = find(doc, f"Figure S{number}.")
    node = cap._p.getprevious()
    while node is not None and not node.xpath(".//w:drawing"):
        node = node.getprevious()
    assert node is not None
    p = Paragraph(node, cap._parent)
    text(p, "")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = own_page
    p.add_run().add_picture(str(path), width=Inches(width))
    text(cap, f"Figure S{number}. {caption}", f"Figure S{number}. ")
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(0)
    cap.paragraph_format.line_spacing = 1
    cap.paragraph_format.keep_with_next = False
    cap.paragraph_format.keep_together = True
    if own_page:
        previous = p._p.getprevious()
        while previous is not None and previous.tag == qn("w:p"):
            if previous.xpath(".//w:t | .//w:drawing | .//w:sectPr | .//w:bookmarkStart"):
                break
            older = previous.getprevious()
            previous.getparent().remove(previous)
            previous = older
        node = cap._p.getnext()
        while node is not None:
            if node.tag == qn("w:p") and (node.xpath(".//w:t") or node.xpath(".//w:drawing")):
                Paragraph(node, cap._parent).paragraph_format.page_break_before = True
                break
            node = node.getnext()


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    files = OUT / "source_artifacts"
    files.mkdir(exist_ok=True)
    source_hash = digest(SOURCE)
    backup = OUT / "esi_before_revision.docx"
    if not backup.exists():
        shutil.copy2(SOURCE, backup)
    assert digest(backup) == source_hash
    doc = Document(SOURCE)
    expected = {}
    public = json.loads((FP / "input_public.json").read_text())
    result = json.loads((FP / "result.json").read_text())
    canonical = result["final_design"]["parameters"]
    events = [json.loads(line) for line in (FP / "llm_events.jsonl").read_text().splitlines() if line.strip()]
    raw_report = (ROOT / "deliverables/case_study_comparison_presentation_20260828/flowpilot_exact_final_response.txt").read_text()
    raw_os = json.loads((OS / "raw_response.json").read_text())["text"]
    raw_sys = (AUDIT / "verbatim_prompts/oneshot_system_prompt.txt").read_text()
    raw_user = (AUDIT / "verbatim_prompts/oneshot_user_prompt.txt").read_text()
    protocol = public["protocol"] + "\n\nOBJECTIVE:\n" + public["objective"]
    assert len(protocol) == 894 and len(raw_sys) + len(raw_user) == 6799
    assert len(raw_report) == 8274 and len(raw_os) == 17149

    # Keep the exact outputs, including their original defects, as auditable evidence.
    for path, name in [(FP / "result.json", "flowpilot_result.json"), (OS / "result.json", "oneshot_result.json"),
                       (FP / "input_inventory.json", "inventory.json"), (FP / "metrics.json", "flowpilot_legacy_metrics.json"),
                       (OS / "metrics.json", "oneshot_legacy_metrics.json")]:
        shutil.copy2(path, files / name)
    for name, value in [("flowpilot_user_input.txt", protocol), ("oneshot_system.txt", raw_sys),
                        ("oneshot_user.txt", raw_user), ("flowpilot_raw_narrative.txt", raw_report), ("oneshot_raw_response.txt", raw_os)]:
        (files / name).write_text(value)
    (files / "flowpilot_final_contract.json").write_text(json.dumps(result["final_design"], ensure_ascii=False, indent=2) + "\n")

    figure(doc, 7, FIGURES / "Figure_S07.png",
           "Case-specific benchmark scores for (a) CuAAC, (b) photochemical oxidation and (c) hydrogenolysis. Points and horizontal error bars show means and sample SDs across three generation repeats. Red denotes one-shot (OS); teal denotes FlowPilot (FP). Both Qwen models have 27B parameters. All 30 SDs are nonzero; small whiskers may be obscured by markers. Values beside the plots are rounded means.", 5.30, True)
    figure(doc, 9, FIGURES / "Figure_S09.png",
           "Critical flags for (a) CuAAC, (b) photochemical oxidation and (c) hydrogenolysis. Each model has adjacent one-shot (OS) and FlowPilot (FP) blocks with repeats 1-3. Rows identify criteria. Cell numbers count judges flagging a critical issue, not distinct physical errors. Pale cells have no flag; gray cells with a dash are not applicable. All 90 campaigns and 990 campaign-criterion cells are retained. No critical flag does not establish experimental success.", 6.05, True)
    figure(doc, 10, FIGURES / "Figure_S10.png",
           "FlowPilot resource use and quality per generation cost. (a) Tokens, (b) estimated generation cost from the stored token-price schedule and (c) observed runtime: mean +/- sample SD across nine campaigns per model. (d-f) Mean benchmark-score/ USD ratios for CuAAC, photochemical oxidation and hydrogenolysis, respectively, with sample SDs across three repeats, on logarithmic axes. Ratios are calculated per campaign before averaging. Judge-model costs are excluded. Runtime is environment-dependent; token pricing is not a measurement of total local deployment cost.", 6.25, True)
    text(find(doc, "Figure S6."),
         "Figure S6. Retrieval alignment and source-exclusion controls. (a) Top-five family match rates, with query counts. (b) Rank change across 1,600 query-result pairs; 401 (25.1%) change rank. The count axis is logarithmic; positive changes indicate promotion. (c) Mean component scores and nonzero percentages among pairs with the query field available. (d) Recorded families for an iridium query: Ir, iridium; Ru, ruthenium; ?, unassigned family, not an established mismatch. (e) Exclusion of normalized held-out source identifiers before final ranking. These metrics assess metadata alignment, not reaction success or absence from model pretraining.", "Figure S6. ")
    topology = BENCH / "figures_revised/topology_atlas/individual/photochemical_oxidation/claude_opus_46/repeat_01.png"
    figure(doc, 13, topology,
           "Archived inventory-assigned topology for the Claude Opus 4.6 FlowPilot photochemical-oxidation run, repeat 01. The diagram represents the final structured design, not the unreconciled LLM narrative reproduced below. Its 1.0 mL reactor, 0.12427 mL/min liquid flow and 1.856819 mL/min air feed at inlet/STP correspond to a nominal V_R/Q_L time of 8.046994 min and 1.4 oxygen equivalents. This nominal time is not a measured liquid contact time. The recorded disposition is SCREEN, not experimentally validated execution. The exact one-shot response is preserved in the following subsection and the companion files.", 6.45)

    # Expand existing tables rather than inventing new benchmark outcomes.
    rubric_path = CAMPAIGN / "frozen/outcome_rubric.json"
    rubric = json.loads(rubric_path.read_text())
    shutil.copy2(rubric_path, files / "outcome_rubric.json")
    rows = [[r["criterion_id"] + "\n" + r["name"], r["question"], "; ".join(r["critical_conditions"]),
             {"ALWAYS": "Always", "GAS_REAGENT_OR_GAS_LIQUID_PROCESS": "Gas use", "MORE_THAN_ONE_REACTIVE_STAGE": "Multi-stage"}[r["applicability"]]] for r in rubric["criteria"]]
    populate_table(doc.tables[8], ["Criterion", "Exact evaluation question", "Critical-error conditions", "Applies"], rows, [1.25, 2.2, 2.25, 0.8])
    text(find(doc, "Table S9."), "Table S9. Complete frozen outcome questions, critical-error triggers and applicability rules. Gas use denotes a gas reagent or gas-liquid process; multi-stage denotes more than one reactive stage.", "Table S9. ")
    modules = pd.read_csv(BENCH / "core_tables/module_module_summary.csv")
    assert len(modules) == 15
    rows = [[r.condition, f"{r.mean_score_0_1:.3f}", f"{r.score_sd:.3f}", f"{r.executable_rate:.0%}",
             f"{r.mean_llm_calls:.1f}", str(int(r.critical_flags))] for r in modules.itertuples()]
    populate_table(doc.tables[9], ["Condition", "Mean score", "Between-case SD", "Exec. rate", "LLM calls", "Critical flags"], rows, [2.2, 0.75, 1.0, 0.85, 0.8, 0.9])
    text(find(doc, "Table S10."), "Table S10. All 15 internal module-ablation conditions. Means and sample SDs summarize three chemistry cases with one generation per condition-case cell; they are not repeatability estimates. Exec. rate denotes the archived executable-outcome fraction, not experimentally validated success.", "Table S10. ")

    replacements = {
        "Major pipeline boundaries": "Major pipeline boundaries are represented by validated data objects rather than prose alone (Figure S1b). The principal chain is DesignInputPackage > BatchRecord > ChemistryPlan > FlowProposal > DesignCalculations > inventory-assigned ProcessTopology > FinalDesignContract. Natural-language rationales may accompany these objects, but downstream code consumes typed fields and performs boundary validation.",
        "For each chemistry, the corresponding frozen": "For each chemistry, the frozen design input reproduced in the three case-study subsections below was inserted after FROZEN DESIGN INPUT:, followed by the common response contract.",
        "Two asymmetries summarize": "The example distinguishes user-facing input, internal computation, and returned artifacts. The recorded text counts are 894 characters for protocol plus objective and 6,799 for the complete one-shot prompt; the 4,342-character inventory is additional shared input, not absent from the FlowPilot task. FlowPilot used six calls and 50,175 tokens, versus one call and 7,430 tokens for one-shot. Its structured final design passed the archived deterministic checks and retained SCREEN disposition. However, its generated narrative was not synchronized to that final design and contained different conditions and unsupported assertions. The structured result and narrative must therefore be evaluated separately; neither is evidence of wet-lab validation.",
        "This section reproduces, verbatim": "This section preserves the exact archived inputs and provider responses for a matched Claude Opus 4.6 photochemical-oxidation run (repeat 01), together with the separately stored final structured design. It illustrates this benchmark's input-assembly and artifact-handling arrangements, not an inherent inability of other LLM interfaces to accept uploaded inventories or structured prompts. The comparison is descriptive, concerns a single run, and is not a usability trial or a token-matched experiment.",
        "Read as a whole, the figure": "Figure S11 separates user-composed text from system-assembled prompts. The numerical QA values in that historical figure (0.96 and 0.63) are legacy heuristic scores, not the equal-weight multi-LLM benchmark scores in Section 3. Their weights were 0.10 formal validity, 0.25 engineering integrity, 0.15 process completeness, 0.15 safety adequacy, 0.10 evidence provenance, 0.20 decision assurance and 0.05 actionability calibration. Trace-dependent components make that legacy metric unsuitable as architecture-neutral evidence of superiority. For this exact pair, the independent fixed-criteria scores were 0.935898 (FlowPilot) and 0.916667 (one-shot). The legacy scores are retained for provenance, not substituted for those benchmark results.",
        "The complete text composed by the user:": "The archived protocol plus objective contains 894 characters. The inventory is a separate 4,342-character input. The following transcription and the companion flowpilot_user_input.txt preserve the source text; line wrapping is a display operation.",
        "The complete single prompt supplied": "The one-shot call used claude-opus-4-6, temperature 0.2 and recorded seed 1007535521. Its system and user messages contain 6,799 characters in total, with 2,095 recorded input tokens. The two messages below are restored from the archived prompt files, without inserting spaces inside JSON keys. Their system/user labels are editorial and not included in the character count. The protocol, objective and inventory are shared with FlowPilot.",
        "FlowPilot returned a validated process flow diagram": "The final structured FlowPilot artifact contains the inventory-realized SCREEN design shown in Figure S13 and summarized in Table S22. The 8,274-character generated report reproduced afterward is a distinct archived LLM narrative and does not agree with those final parameters. It is preserved as historical evidence, not endorsed as an executable procedure. The 17,149-character one-shot response is reproduced in the following subsection.",
        "The complete raw provider response in the one-shot": "The 17,149-character one-shot provider response below is preserved exactly, including its code fence and errors. The archived FlowProposal parser reported six type-validation errors. The response contains conflicting channel-gas values (0.136 and 0.167 mL/min) and reports 0.235 oxygen equivalents. These are findings about the raw response, not a claim that every parsed benchmark field was scored identically. The final-outcome judge score and the legacy parser/QA diagnostics are distinct records.",
    }
    for prefix, value in replacements.items():
        text(find(doc, prefix), value)
    text(find(doc, "Figure S11."),
         "Figure S11. Historical prompt-footprint comparison for one matched Claude Opus 4.6 run. (a) User-facing protocol plus objective versus the complete one-shot prompt; the structured inventory is additional shared input. (b) Six FlowPilot internal prompts total 112,754 model-facing characters. (c) Archived heuristic QA diagnostics (0.96 versus 0.63), distinct from the independent multi-LLM outcome score. Validation claims apply to the final structured artifact, not the unsynchronized generated narrative. (d) Recorded computation: 50,175 tokens and 273 s versus 7,430 tokens and 100 s. Exact inputs and outputs are reproduced in the following named subsections and companion files. This is a single-run artifact comparison, not an experimental or human-usability validation.", "Figure S11. ")
    table17 = doc.tables[16]
    text(table17.rows[6].cells[1].paragraphs[0], "Inventory-assigned final SCREEN contract and Figure S13 topology. The separately generated report is unsynchronized and is not executable authority.")
    text(table17.rows[7].cells[0].paragraphs[0], "Legacy heuristic QA (not the multi-LLM score)")
    text(table17.rows[1].cells[2].paragraphs[0], "6,799 characters in the archived system/user prompt: protocol, objective, constraints, inventory and output contract. This count describes the tested interface, not a universal typing requirement.")

    judge_models = {}
    for family in ("qwen", "openai", "claude"):
        telemetry = json.loads((CAMPAIGN / "judgments" / family / "N2-8B193490/telemetry.json").read_text())[0]
        judge_models[family] = {k: telemetry.get(k) for k in ("provider", "model", "temperature", "max_tokens")}
    assert judge_models["openai"]["model"] == "gpt-5.4-2026-03-05"
    after(find(doc, "Three independent judge families"),
          "The archived judge configuration for the matched example was Qwen /models/Qwen3.6-27B, OpenAI gpt-5.4-2026-03-05 and Anthropic claude-sonnet-4-6, at temperature 0.0. GPT-5.4 is identified here only as an evaluator; it is not one of the five retained generator models. Exact prompt schemas, seeds, responses and telemetry are retained in the companion source records. A recorded seed does not guarantee that a provider honors deterministic sampling.")
    after(find(doc, "The user interface exposes separate model routes"),
          "This ESI distinguishes the historical benchmark configuration from the current software. The archived worked example reports liquid-only geometric-volume time, inlet/STP apparent time and pressure-corrected apparent time as different fields. Current GUI defaults must not be used to retroactively relabel these measurements or regenerate archived benchmark scores. The held-out photochemical reference was machine-extracted and not page-verified during the campaign; internally inconsistent published-flow fields were explicitly withheld from numerical reference scoring.")
    before(find(doc, "RAW BATCH PROTOCOL:"), "Source of the hydrogenolysis benchmark case: Tu et al., reference 1. Source exclusion is a retrieval control, not proof of absence from pretraining.")
    after(find(doc, "Benchmark title: Photocatalytic aerobic"), "Source of the photochemical benchmark case: Thomson et al., reference 2. The reaction name is not the article title. The paper's supporting information includes Fmoc-L-methionine oxidation; the campaign's extracted numerical facts have not been revalidated page by page in this document revision.")
    after(find(doc, "Benchmark title: Cu/C-catalyzed"), "Source of the CuAAC benchmark case: Fuchs et al., reference 3. Frozen case inputs and inventory constraints are retained unchanged.")

    # Fix stale cross-references outside raw archival blocks before restoring those blocks.
    for p in doc.paragraphs:
        if "Figure S1A" in p.text or "Figure S1B" in p.text:
            text(p, p.text.replace("Figure S1A", "Figure S1a").replace("Figure S1B", "Figure S1b"))
    raw_input_intro = find(doc, "The archived protocol plus objective")
    raw_prompt_heading = find(doc, "One-shot engineered prompt")
    raw_prompt_intro = find(doc, "The one-shot call used")
    fp_heading = find(doc, "FlowPilot returned design")
    os_heading = find(doc, "One-shot raw response")
    os_intro = find(doc, "The 17,149-character one-shot")
    refs = find(doc, "References")
    replace_block(raw_input_intro, raw_prompt_heading, protocol, "archived_fp_input", expected)
    replace_block(raw_prompt_intro, fp_heading, "[SYSTEM MESSAGE]\n" + raw_sys + "\n[USER MESSAGE]\n" + raw_user, "archived_os_prompt", expected)
    replace_block(find(doc, "Figure S13."), os_heading, raw_report, "archived_fp_narrative", expected)
    replace_block(os_intro, refs, raw_os, "archived_os_response", expected)

    # Add the final-versus-narrative distinction adjacent to the raw report, not as a hidden correction.
    raw_start = next(p for p in doc.paragraphs if p._p.xpath('.//w:bookmarkStart[@w:name="archived_fp_narrative"]'))
    before(raw_start,
           "Artifact reconciliation. The final contract is nominally 8.046994 min at 0.12427 mL/min in 1.0 mL; its air feed is 1.856819 mL/min at STP (1.4 O2 equivalents). The narrative instead specifies 5 min, 0.2 mL/min, 1.6094 mL and 0.4483 mL/min air. The latter corresponds to approximately 0.210 O2 equivalents, not its claimed 1.0. The final geometric-volume time is not measured contact time. Statements such as 'no thermal runaway risk' in the raw report are unverified model assertions. This discrepancy documents a historical narrative-synchronization defect and limits claims of fully consistent user-facing output.")
    cap = before(raw_start, "Table S22. Final structured parameters and contemporaneous generated narrative for the same archived FlowPilot run.", "Table S22. ")
    table = doc.add_table(rows=2, cols=3)
    table.style = doc.tables[16].style
    cap._p.addnext(table._tbl)
    rows = [["Reactor geometric volume (mL)", "1.0", "1.6094"], ["Liquid flow (mL/min)", "0.12427", "0.2"],
            ["Reported residence time (min)", "8.046994; nominal V_R/Q_L", "5.0; claimed liquid-holdup time"],
            ["Air at inlet/STP (mL/min)", "1.856819", "0.4483"], ["O2 equivalents", "1.4", "Claims 1.0; arithmetic gives 0.210"],
            ["Temperature / BPR", "21 C / 3 bar gauge", "21 C / 3 bar gauge"],
            ["Validation status", "Archived final checks: ready; disposition SCREEN", "NOT engine-validated; LOW confidence"],
            ["Authority", "Final structured contract", "Historical raw narrative; not operating instructions"]]
    populate_table(table, ["Field", "Final structured record", "Raw LLM narrative"], rows, [1.8, 2.35, 2.35])
    before(raw_start, "Archived raw narrative follows unchanged:", "Archived raw narrative follows unchanged:")

    references = [
        "1. Tu, J.; Sang, L.; Cheng, H.; Ai, N.; Zhang, J. Continuous Hydrogenolysis of N-Diphenylmethyl Groups in a Micropacked-Bed Reactor. Organic Process Research & Development 2020, 24, 59-66. https://doi.org/10.1021/acs.oprd.9b00416",
        "2. Thomson, C. G.; Banks, C.; Allen, M.; Barker, G.; Coxon, C. R.; Lee, A.-L.; Vilela, F. Expanding the Tool Kit of Automated Flow Synthesis: Development of In-line Flash Chromatography Purification. Journal of Organic Chemistry 2021, 86, 14079-14094. https://doi.org/10.1021/acs.joc.1c01151",
        "3. Fuchs, M.; Goessler, W.; Pilger, C.; Kappe, C. O. Mechanistic Insights into Copper(I)-Catalyzed Azide-Alkyne Cycloadditions using Continuous Flow Conditions. Advanced Synthesis & Catalysis 2010, 352, 323-328. https://doi.org/10.1002/adsc.200900726",
    ]
    last = refs
    for value in references:
        last = after(last, value)

    # Remove the stray test footnote, without disturbing other parts or drawings.
    for node in list(doc.element.xpath(".//w:footnoteReference[@w:id='1']")):
        node.getparent().remove(node)
    for part in doc.part.package.parts:
        if str(part.partname) == "/word/footnotes.xml":
            from lxml import etree
            root = etree.fromstring(part.blob)
            for node in list(root):
                if node.get(qn("w:id")) == "1":
                    root.remove(node)
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    doc.save(OUTPUT)
    assert digest(SOURCE) == source_hash

    # Round-trip verification of every restored raw block, including whitespace.
    check = Document(OUTPUT)
    verified = {}
    for name, original in expected.items():
        active = False
        strings = []
        identity = None
        for p in check.paragraphs:
            markers = p._p.xpath(f'.//w:bookmarkStart[@w:name="{name}"]')
            if markers:
                active = True
                identity = markers[0].get(qn("w:id"))
            if active:
                strings.append(p.text)
                if p._p.xpath(f'.//w:bookmarkEnd[@w:id="{identity}"]'):
                    break
        actual = "\n".join(strings)
        assert actual == original, name
        verified[name] = {"characters": len(actual), "sha256": hashlib.sha256(actual.encode()).hexdigest()}
    assert len(check.tables) == 22
    assert len(check.tables[9].rows) == 16
    assert len(check.element.xpath(".//m:oMath")) == len(Document(SOURCE).element.xpath(".//m:oMath"))
    report = {"source_sha256": source_hash, "output_sha256": digest(OUTPUT), "raw_blocks_verified": verified,
              "tables": len(check.tables), "module_conditions": 15, "judge_configuration_example": judge_models,
              "source_unchanged": digest(SOURCE) == source_hash, "final_parameters": canonical,
              "remaining": ["Full campaign-wide judge/snapshot audit beyond the matched example", "Source protocol page-by-page verification and handbook bibliography", "Main-manuscript cross-reference audit not performed in this revision", "Council-section relocation deferred to avoid unapproved figure renumbering", "Prospective wet-lab additions require supplied measurements"]}
    (OUT / "revision_checks.json").write_text(json.dumps(report, indent=2) + "\n")
    (OUT / "source_artifacts/SHA256SUMS.txt").write_text("\n".join(f"{digest(p)}  {p.name}" for p in sorted(files.iterdir()) if p.is_file() and p.name != "SHA256SUMS.txt") + "\n")
    print(OUTPUT)
    print(json.dumps({"source_unchanged": True, "raw_blocks": len(verified), "tables": 22}, indent=2))


if __name__ == "__main__":
    main()
