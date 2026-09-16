#!/usr/bin/env python3
"""Build a matched FlowPilot versus one-shot prompt-and-output case study deck."""

from __future__ import annotations

import json
import csv
import textwrap
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
RUN_BASE = (
    ROOT
    / "ablation_results"
    / "manuscript_benchmark"
    / "alternative_frontier_three_repeat_20260824"
    / "generation"
    / "photochemical_oxidation"
)
FLOWPILOT_RUN = RUN_BASE / "claude_opus_flowpilot" / "repeat_01"
ONESHOT_RUN = RUN_BASE / "claude_opus_one_shot" / "repeat_01"
TOPOLOGY_IMAGE = (
    ROOT
    / "deliverables"
    / "manuscript_benchmark_visualizations_20260825"
    / "figures_revised"
    / "topology_atlas"
    / "individual"
    / "photochemical_oxidation"
    / "claude_opus_46"
    / "repeat_01.png"
)
OUT_DIR = ROOT / "deliverables" / "case_study_comparison_presentation_20260828"
OUT_PATH = OUT_DIR / "FlowPilot_vs_OneShot_Prompt_Footprint_Opus46_Photochemical_Oxidation.pptx"
SOURCE_PATH = OUT_DIR / "source_data.json"
FLOWPILOT_OUTPUT_PATH = OUT_DIR / "flowpilot_exact_final_response.txt"
ONESHOT_OUTPUT_PATH = OUT_DIR / "oneshot_exact_raw_response.txt"
PROMPT_INVENTORY = (
    ROOT
    / "deliverables"
    / "manuscript_benchmark_visualizations_20260825"
    / "prompt_audit_claude_opus_46"
    / "prompt_inventory.csv"
)


WHITE = "FFFFFF"
INK = "17212B"
MUTED = "5D6975"
LINE = "D8DEE4"
NAVY = "214B8C"
FLOW = "087F7A"
FLOW_DARK = "075E5B"
FLOW_LIGHT = "EAF7F5"
ONE = "D85A4F"
ONE_DARK = "A83D35"
ONE_LIGHT = "FCEFED"
AMBER = "B87800"
AMBER_LIGHT = "FFF5DD"
GREEN = "26734D"
GREEN_LIGHT = "E9F6EF"
RED = "B83A3A"
RED_LIGHT = "FCEBEC"
GREY_LIGHT = "F4F6F8"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_cell_margin(text_frame, left=0.08, right=0.08, top=0.05, bottom=0.05) -> None:
    text_frame.margin_left = Inches(left)
    text_frame.margin_right = Inches(right)
    text_frame.margin_top = Inches(top)
    text_frame.margin_bottom = Inches(bottom)


def add_rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    shape.line.width = Pt(0.8)
    if radius:
        try:
            shape.adjustments[0] = 0.08
        except Exception:
            pass
    return shape


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    *,
    size=10,
    color=INK,
    bold=False,
    font="Aptos",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    set_cell_margin(frame, margin, margin, margin, margin)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_labelled_lines(slide, rows, x, y, w, h, *, label_color, body_size=8.6):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    set_cell_margin(frame, 0.12, 0.12, 0.08, 0.05)
    for index, (label, body) in enumerate(rows):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(3 if index < len(rows) - 1 else 0)
        paragraph.line_spacing = 0.95
        run = paragraph.add_run()
        run.text = label + "  "
        run.font.name = "Aptos"
        run.font.size = Pt(body_size)
        run.font.bold = True
        run.font.color.rgb = rgb(label_color)
        run = paragraph.add_run()
        run.text = body
        run.font.name = "Aptos"
        run.font.size = Pt(body_size)
        run.font.color.rgb = rgb(INK)
    return box


def add_condition_grid(slide, items, x, y, w, h, *, accent, alert_keys=None):
    alert_keys = set(alert_keys or [])
    columns = 4
    rows = 2
    gap = 0.06
    cell_w = (w - gap * (columns - 1)) / columns
    cell_h = (h - gap * (rows - 1)) / rows
    for index, (label, value) in enumerate(items):
        row, column = divmod(index, columns)
        left = x + column * (cell_w + gap)
        top = y + row * (cell_h + gap)
        is_alert = label in alert_keys
        fill = RED_LIGHT if is_alert else WHITE
        border = RED if is_alert else LINE
        shape = add_rect(slide, left, top, cell_w, cell_h, fill, border, radius=True)
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        set_cell_margin(frame, 0.05, 0.05, 0.02, 0.02)
        p = frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = label
        r.font.name = "Aptos"
        r.font.size = Pt(6.7)
        r.font.bold = True
        r.font.color.rgb = rgb(RED if is_alert else MUTED)
        p = frame.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        r = p.add_run()
        r.text = value
        r.font.name = "Aptos"
        r.font.size = Pt(8.4 if len(value) < 16 else 7.2)
        r.font.bold = True
        r.font.color.rgb = rgb(RED if is_alert else accent)


def add_status_banner(slide, text, x, y, w, h, *, fill, color):
    shape = add_rect(slide, x, y, w, h, fill, color, radius=True)
    frame = shape.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_cell_margin(frame, 0.08, 0.08, 0, 0)
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.space_after = Pt(0)
    r = p.add_run()
    r.text = text
    r.font.name = "Aptos"
    r.font.size = Pt(8.5)
    r.font.bold = True
    r.font.color.rgb = rgb(color)


def add_metrics(slide, items, x, y, w, h, *, accent, fill):
    background = add_rect(slide, x, y, w, h, fill, accent, radius=True)
    columns = len(items)
    col_w = w / columns
    for index, (label, value) in enumerate(items):
        left = x + index * col_w
        if index:
            line = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(left),
                Inches(y + 0.11),
                Inches(0.008),
                Inches(h - 0.22),
            )
            line.fill.solid()
            line.fill.fore_color.rgb = rgb(accent)
            line.line.fill.background()
        add_text(slide, label, left + 0.03, y + 0.07, col_w - 0.06, 0.14, size=5.7, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, value, left + 0.03, y + 0.23, col_w - 0.06, 0.23, size=8.6, color=accent, bold=True, align=PP_ALIGN.CENTER)
    return background


def add_picture_contain(slide, path: Path, x, y, w, h):
    from PIL import Image

    with Image.open(path) as image:
        image_ratio = image.width / image.height
    box_ratio = w / h
    if image_ratio >= box_ratio:
        width = w
        height = w / image_ratio
        left = x
        top = y + (h - height) / 2
    else:
        height = h
        width = h * image_ratio
        top = y
        left = x + (w - width) / 2
    return slide.shapes.add_picture(
        str(path), Inches(left), Inches(top), Inches(width), Inches(height)
    )


def wrap_exact_text(text: str, width: int) -> list[str]:
    lines: list[str] = []
    for source_line in text.splitlines():
        if not source_line:
            lines.append("")
            continue
        wrapped = textwrap.wrap(
            source_line,
            width=width,
            replace_whitespace=False,
            drop_whitespace=False,
            break_long_words=True,
            break_on_hyphens=False,
        )
        lines.extend(wrapped or [""])
    return lines


def chunks(values: list[str], count: int) -> list[list[str]]:
    size = (len(values) + count - 1) // count
    return [values[index * size : (index + 1) * size] for index in range(count)]


def load_llm_events(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text().splitlines() if line.strip()]


def add_prompt_page(slide, lines, x, y, w, h, *, page_label, accent, font_size=4.2):
    add_rect(slide, x, y, w, h, WHITE, LINE, radius=True)
    add_rect(slide, x, y, w, 0.25, accent, accent, radius=True)
    add_text(
        slide,
        page_label,
        x + 0.08,
        y + 0.055,
        w - 0.16,
        0.11,
        size=6.1,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "\n".join(lines),
        x + 0.10,
        y + 0.34,
        w - 0.20,
        h - 0.43,
        size=font_size,
        color=INK,
        font="Liberation Mono",
    )


def build_prompt_footprint_slide(presentation: Presentation, data: dict) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)

    protocol_objective = (
        data["public_input"]["protocol"]
        + "\n\nOBJECTIVE:\n"
        + data["public_input"]["objective"]
    )
    one_exact = (
        "[SYSTEM MESSAGE]\n"
        + data["one_prompt"]["system"]
        + "\n\n[USER MESSAGE]\n"
        + data["one_prompt"]["user"]
    )
    flow_lines = wrap_exact_text(protocol_objective, 52)
    one_lines = wrap_exact_text(one_exact, 52)
    one_pages = chunks(one_lines, 3)

    add_text(
        slide,
        "The prompt is the comparison: what must be assembled before inference?",
        0.40,
        0.18,
        12.53,
        0.38,
        size=20,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Exact text footprint using equal-width pages and a common monospace scale",
        0.40,
        0.62,
        12.53,
        0.20,
        size=8.3,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    add_rect(slide, 0.35, 0.96, 3.40, 5.78, FLOW_LIGHT, FLOW_LIGHT, radius=True)
    add_rect(slide, 3.94, 0.96, 9.04, 5.78, ONE_LIGHT, ONE_LIGHT, radius=True)
    add_text(slide, "FLOWPILOT HUMAN INPUT", 0.55, 1.08, 3.00, 0.25, size=13.5, color=FLOW_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "ONE-SHOT ENGINEERED PROMPT", 4.15, 1.08, 8.62, 0.25, size=13.5, color=ONE_DARK, bold=True, align=PP_ALIGN.CENTER)

    add_prompt_page(
        slide,
        flow_lines,
        0.88,
        1.47,
        2.35,
        4.30,
        page_label="PROTOCOL + OBJECTIVE",
        accent=FLOW,
        font_size=3.7,
    )
    add_prompt_page(slide, one_pages[0], 4.55, 1.47, 2.35, 4.30, page_label="PROMPT PAGE 1/3", accent=ONE, font_size=3.7)
    add_prompt_page(slide, one_pages[1], 7.20, 1.47, 2.35, 4.30, page_label="PROMPT PAGE 2/3", accent=ONE, font_size=3.7)
    add_prompt_page(slide, one_pages[2], 9.85, 1.47, 2.35, 4.30, page_label="PROMPT PAGE 3/3", accent=ONE, font_size=3.7)

    add_status_banner(
        slide,
        f"{len(protocol_objective):,} characters | {len(flow_lines)} wrapped lines | inventory imported as JSON, not pasted into the chat",
        0.62,
        5.94,
        2.87,
        0.48,
        fill=GREEN_LIGHT,
        color=FLOW_DARK,
    )
    add_status_banner(
        slide,
        f"{len(data['one_prompt']['system']) + len(data['one_prompt']['user']):,} characters | {len(one_lines)} wrapped lines | protocol + constraints + complete inventory + output contract",
        4.28,
        5.94,
        8.31,
        0.48,
        fill=RED_LIGHT,
        color=ONE_DARK,
    )
    add_text(
        slide,
        f"FlowPilot also receives an attached {len(json.dumps(data['public_input']['inventory'], indent=2)):,}-character inventory profile and {len(data['public_input']['hard_constraints'])} structured constraints. "
        "The distinction is interaction design: these are collected/imported as fields rather than manually composed into one expert prompt.",
        0.55,
        6.82,
        12.23,
        0.32,
        size=7.3,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_text(slide, "1", 12.82, 7.18, 0.18, 0.12, size=6, color=MUTED, align=PP_ALIGN.RIGHT)


def build_exact_oneshot_prompt_slide(presentation: Presentation, data: dict) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)

    exact = (
        "[SYSTEM MESSAGE]\n"
        + data["one_prompt"]["system"]
        + "\n\n[USER MESSAGE]\n"
        + data["one_prompt"]["user"]
    )
    lines = wrap_exact_text(exact, 58)
    columns = chunks(lines, 4)

    add_text(slide, "Exact one-shot prompt", 0.40, 0.18, 12.53, 0.38, size=21, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(
        slide,
        f"Claude Opus 4.6 | {len(data['one_prompt']['system']) + len(data['one_prompt']['user']):,} characters | 2,095 recorded input tokens | no text omitted; line wrapping only",
        0.40,
        0.62,
        12.53,
        0.22,
        size=8.2,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    gap = 0.14
    col_w = (12.43 - 3 * gap) / 4
    for index, column in enumerate(columns):
        x = 0.45 + index * (col_w + gap)
        add_prompt_page(
            slide,
            column,
            x,
            1.02,
            col_w,
            5.98,
            page_label=f"EXACT PROMPT {index + 1}/4",
            accent=ONE,
            font_size=5.6,
        )
    add_text(slide, "Source: saved benchmark prompt.json | system and user message boundaries are labelled for readability.", 0.50, 7.15, 12.25, 0.15, size=6.2, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "2", 12.82, 7.18, 0.18, 0.12, size=6, color=MUTED, align=PP_ALIGN.RIGHT)


def build_output_footprint_slide(presentation: Presentation, data: dict) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)

    flow_text = data["flow_final_response"]
    one_text = data["one_raw_response"]
    flow_pages = chunks(wrap_exact_text(flow_text, 52), 4)
    one_pages = chunks(wrap_exact_text(one_text, 52), 8)

    add_text(slide, "The outputs at the same physical scale", 0.40, 0.18, 12.53, 0.38, size=21, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(
        slide,
        "Exact saved response text | equal-width pages | common monospace scale | no output omitted",
        0.40,
        0.62,
        12.53,
        0.22,
        size=8.2,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    label_w = 1.35
    page_x = 1.62
    gap = 0.09
    page_w = 1.36
    page_h = 2.42

    add_rect(slide, 0.36, 0.98, 12.61, 2.66, FLOW_LIGHT, FLOW_LIGHT, radius=True)
    add_text(slide, "FLOWPILOT\nFINAL REPORT", 0.52, 1.62, 0.94, 0.52, size=11.2, color=FLOW_DARK, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    for index, page in enumerate(flow_pages):
        add_prompt_page(
            slide,
            page,
            page_x + index * (page_w + gap),
            1.10,
            page_w,
            page_h,
            page_label=f"OUTPUT {index + 1}/4",
            accent=FLOW,
            font_size=2.35,
        )
    add_text(
        slide,
        f"{len(flow_text):,} characters | {len(flow_text.splitlines())} source lines | final output_formatter response",
        7.60,
        2.12,
        4.88,
        0.34,
        size=10.0,
        color=FLOW_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    add_rect(slide, 0.36, 3.80, 12.61, 2.66, ONE_LIGHT, ONE_LIGHT, radius=True)
    add_text(slide, "ONE-SHOT\nRAW RESPONSE", 0.52, 4.44, 0.94, 0.52, size=11.2, color=ONE_DARK, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    for index, page in enumerate(one_pages):
        add_prompt_page(
            slide,
            page,
            page_x + index * (page_w + gap),
            3.92,
            page_w,
            page_h,
            page_label=f"OUTPUT {index + 1}/8",
            accent=ONE,
            font_size=2.35,
        )
    add_text(
        slide,
        "Page count is a visual measure of response footprint, not a quality score. Full-size text follows on Slides 4-6.",
        0.55,
        6.72,
        12.23,
        0.22,
        size=7.4,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_text(slide, "3", 12.82, 7.18, 0.18, 0.12, size=6, color=MUTED, align=PP_ALIGN.RIGHT)


def build_exact_output_slide(
    presentation: Presentation,
    *,
    title: str,
    subtitle: str,
    text_value: str,
    accent: str,
    page_start: int,
    page_total: int,
    slide_number: int,
) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)

    all_lines = wrap_exact_text(text_value, 58)
    all_pages = chunks(all_lines, page_total)
    visible_pages = all_pages[page_start : page_start + 4]

    add_text(slide, title, 0.40, 0.18, 12.53, 0.38, size=21, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, subtitle, 0.40, 0.62, 12.53, 0.22, size=8.2, color=MUTED, align=PP_ALIGN.CENTER)

    gap = 0.14
    col_w = (12.43 - 3 * gap) / 4
    for offset, column in enumerate(visible_pages):
        page_index = page_start + offset
        x = 0.45 + offset * (col_w + gap)
        add_prompt_page(
            slide,
            column,
            x,
            1.02,
            col_w,
            5.98,
            page_label=f"EXACT OUTPUT {page_index + 1}/{page_total}",
            accent=accent,
            font_size=5.6,
        )
    add_text(slide, "Source: saved benchmark response text | line wrapping only; wording and values are unchanged.", 0.50, 7.15, 12.25, 0.15, size=6.2, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, str(slide_number), 12.82, 7.18, 0.18, 0.12, size=6, color=MUTED, align=PP_ALIGN.RIGHT)


def build_internal_prompt_workload_slide(presentation: Presentation, data: dict) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)

    flow_stages = [row for row in data["prompt_inventory"] if row["architecture"] == "flowpilot"]
    one_row = next(row for row in data["prompt_inventory"] if row["architecture"] == "one_shot")
    max_chars = max(int(row["total_prompt_chars"]) for row in flow_stages)

    add_text(slide, "Human input burden and model-facing workload are different questions", 0.40, 0.18, 12.53, 0.38, size=20, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(
        slide,
        "FlowPilot simplifies the user interaction by spending more computation internally; this is not a token-matched comparison.",
        0.40,
        0.62,
        12.53,
        0.22,
        size=8.3,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    add_rect(slide, 0.42, 1.02, 5.20, 5.95, FLOW_LIGHT, FLOW_LIGHT, radius=True)
    add_rect(slide, 5.82, 1.02, 7.08, 5.95, ONE_LIGHT, ONE_LIGHT, radius=True)
    add_text(slide, "FLOWPILOT: SIX INTERNAL PROMPTS", 0.68, 1.20, 4.68, 0.24, size=13, color=FLOW_DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "ONE-SHOT: ONE MANUALLY ENGINEERED PROMPT", 6.07, 1.20, 6.58, 0.24, size=13, color=ONE_DARK, bold=True, align=PP_ALIGN.CENTER)

    stage_labels = {
        "batch_protocol_extraction": "Batch extraction",
        "upstream_chemistry_analysis": "Upstream chemistry",
        "flow_translation_proposal": "Flow proposal",
        "council_problem_framing": "Council framing",
        "council_sampling_designer": "Council designer",
        "final_chemist_report": "Final report",
    }
    start_y = 1.72
    for index, row in enumerate(flow_stages):
        value = int(row["total_prompt_chars"])
        y = start_y + index * 0.66
        add_text(slide, stage_labels.get(row["stage"], row["stage"]), 0.68, y + 0.06, 1.42, 0.18, size=7.5, color=INK, bold=True)
        add_rect(slide, 2.10, y, 2.68, 0.34, WHITE, LINE, radius=True)
        add_rect(slide, 2.10, y, max(0.08, 2.68 * value / max_chars), 0.34, FLOW, FLOW, radius=True)
        add_text(slide, f"{value:,} chars", 4.28, y + 0.07, 0.88, 0.16, size=7.1, color=FLOW_DARK, bold=True, align=PP_ALIGN.RIGHT)

    total_flow_chars = sum(int(row["total_prompt_chars"]) for row in flow_stages)
    add_status_banner(slide, f"TOTAL MODEL-FACING PROMPT TEXT: {total_flow_chars:,} CHARACTERS", 0.72, 5.78, 4.60, 0.42, fill=GREEN_LIGHT, color=FLOW_DARK)
    add_metrics(
        slide,
        [
            ("CALLS", "6"),
            ("INPUT TOKENS", "37,552"),
            ("OUTPUT TOKENS", "12,623"),
            ("RUNTIME", "273 s"),
        ],
        0.72,
        6.30,
        4.60,
        0.50,
        accent=FLOW_DARK,
        fill=WHITE,
    )

    one_chars = int(one_row["total_prompt_chars"])
    add_text(slide, "Single prompt envelope", 6.28, 1.86, 2.15, 0.22, size=10.5, color=INK, bold=True)
    add_rect(slide, 6.28, 2.18, 5.90, 0.58, WHITE, LINE, radius=True)
    add_rect(slide, 6.28, 2.18, 5.90 * one_chars / total_flow_chars, 0.58, ONE, ONE, radius=True)
    add_text(slide, f"{one_chars:,} characters", 6.53, 2.35, 5.35, 0.20, size=12, color=ONE_DARK, bold=True, align=PP_ALIGN.CENTER)

    add_labelled_lines(
        slide,
        [
            ("Includes", "raw protocol, objective, nine hard constraints, complete inventory JSON, and a detailed response contract."),
            ("Human burden", "The user or benchmark author must compose or assemble the complete instruction envelope before the single call."),
            ("Computational burden", "One call and substantially fewer model-facing tokens than the staged FlowPilot execution."),
        ],
        6.28,
        3.08,
        5.92,
        1.52,
        label_color=ONE_DARK,
        body_size=8.2,
    )
    add_status_banner(slide, "TOTAL MODEL-FACING PROMPT TEXT: 6,799 CHARACTERS", 6.28, 5.78, 5.90, 0.42, fill=RED_LIGHT, color=ONE_DARK)
    add_metrics(
        slide,
        [
            ("CALLS", "1"),
            ("INPUT TOKENS", "2,095"),
            ("OUTPUT TOKENS", "5,335"),
            ("RUNTIME", "100 s"),
        ],
        6.28,
        6.30,
        5.90,
        0.50,
        accent=ONE_DARK,
        fill=WHITE,
    )
    add_text(slide, "7", 12.82, 7.18, 0.18, 0.12, size=6, color=MUTED, align=PP_ALIGN.RIGHT)


def load_sources() -> dict:
    flow_result = json.loads((FLOWPILOT_RUN / "result.json").read_text())
    flow_summary = json.loads((FLOWPILOT_RUN / "run_summary.json").read_text())
    flow_metrics = json.loads((FLOWPILOT_RUN / "metrics.json").read_text())
    one_result = json.loads((ONESHOT_RUN / "result.json").read_text())
    one_summary = json.loads((ONESHOT_RUN / "run_summary.json").read_text())
    one_metrics = json.loads((ONESHOT_RUN / "metrics.json").read_text())
    one_prompt = json.loads((ONESHOT_RUN / "prompt.json").read_text())
    one_raw_response = json.loads((ONESHOT_RUN / "raw_response.json").read_text())["text"]
    one_schema_error = json.loads((ONESHOT_RUN / "schema_error.json").read_text())
    public_input = json.loads((FLOWPILOT_RUN / "input_public.json").read_text())
    flow_events = load_llm_events(FLOWPILOT_RUN / "llm_events.jsonl")
    flow_final_response = next(
        event["response_text"]
        for event in reversed(flow_events)
        if event.get("api_name") == "output_formatter"
    )
    with PROMPT_INVENTORY.open(newline="", encoding="utf-8") as handle:
        prompt_inventory = list(csv.DictReader(handle))

    flow_air = next(
        stream for stream in flow_result["proposal"]["streams"] if stream.get("phase") == "gas"
    )
    one_air = next(
        stream for stream in one_result["proposal"]["streams"] if stream.get("phase") == "gas"
    )
    return {
        "public_input": public_input,
        "flow_result": flow_result,
        "flow_summary": flow_summary,
        "flow_metrics": flow_metrics,
        "flow_air": flow_air,
        "one_result": one_result,
        "one_summary": one_summary,
        "one_metrics": one_metrics,
        "one_air": one_air,
        "one_prompt": one_prompt,
        "flow_final_response": flow_final_response,
        "one_raw_response": one_raw_response,
        "one_schema_error": one_schema_error,
        "prompt_inventory": prompt_inventory,
    }


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    data = load_sources()

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    build_prompt_footprint_slide(presentation, data)
    build_exact_oneshot_prompt_slide(presentation, data)
    build_output_footprint_slide(presentation, data)
    build_exact_output_slide(
        presentation,
        title="Exact FlowPilot final response",
        subtitle=f"Claude Opus 4.6 | {len(data['flow_final_response']):,} characters | complete final output_formatter response | no text omitted",
        text_value=data["flow_final_response"],
        accent=FLOW,
        page_start=0,
        page_total=4,
        slide_number=4,
    )
    build_exact_output_slide(
        presentation,
        title="Exact one-shot raw response (Part 1 of 2)",
        subtitle=f"Claude Opus 4.6 | {len(data['one_raw_response']):,} characters | raw provider response | pages 1-4 of 8",
        text_value=data["one_raw_response"],
        accent=ONE,
        page_start=0,
        page_total=8,
        slide_number=5,
    )
    build_exact_output_slide(
        presentation,
        title="Exact one-shot raw response (Part 2 of 2)",
        subtitle=f"Claude Opus 4.6 | {len(data['one_raw_response']):,} characters | raw provider response | pages 5-8 of 8",
        text_value=data["one_raw_response"],
        accent=ONE,
        page_start=4,
        page_total=8,
        slide_number=6,
    )
    build_internal_prompt_workload_slide(presentation, data)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)

    add_text(
        slide,
        "Secondary comparison: what each architecture returned",
        0.35,
        0.18,
        12.63,
        0.38,
        size=21,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        "Photocatalytic aerobic oxidation of Fmoc-L-methionine | repeat 01 | same protocol, inventory, temperature (0.2), and seed",
        0.35,
        0.60,
        12.63,
        0.22,
        size=8.3,
        color=MUTED,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )

    # Column backgrounds and divider.
    add_rect(slide, 0.27, 0.91, 6.31, 6.25, FLOW_LIGHT, FLOW_LIGHT, radius=True)
    add_rect(slide, 6.76, 0.91, 6.31, 6.25, ONE_LIGHT, ONE_LIGHT, radius=True)
    divider = add_rect(slide, 6.66, 0.94, 0.015, 6.18, LINE, LINE)

    add_rect(slide, 0.27, 0.91, 6.31, 0.40, FLOW, FLOW, radius=True)
    add_text(slide, "FLOWPILOT", 0.47, 0.98, 2.1, 0.22, size=16, color=WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "staged, inventory-validated architecture", 2.45, 1.00, 3.85, 0.18, size=7.3, color=WHITE, align=PP_ALIGN.RIGHT)

    add_rect(slide, 6.76, 0.91, 6.31, 0.40, ONE, ONE, radius=True)
    add_text(slide, "ONE-SHOT", 6.96, 0.98, 2.1, 0.22, size=16, color=WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "one general-LLM inference", 9.03, 1.00, 3.75, 0.18, size=7.3, color=WHITE, align=PP_ALIGN.RIGHT)

    # Input panels.
    add_text(slide, "USER-FACING STANDARDIZED INPUT", 0.48, 1.38, 5.90, 0.18, size=7.3, color=FLOW_DARK, bold=True)
    add_rect(slide, 0.45, 1.58, 5.95, 1.03, WHITE, LINE, radius=True)
    add_labelled_lines(
        slide,
        [
            (
                "Protocol",
                "Fmoc-L-methionine (0.10 M), organic photocatalyst (10 mol%), MeCN, air, 420 nm LED (1.5 W), 21 C, 5 min; batch yield 99%.",
            ),
            ("Objective", "One conservative, inventory-constrained gas-liquid photochemical flow screen."),
            (
                "Inventory",
                "Imported profile: HPLC pump, air MFC, PFA T-mixer, 1.0 mL PFA photocoil, 420 nm LED, 3 bar BPR, G-L separator, collector.",
            ),
        ],
        0.50,
        1.61,
        5.84,
        0.94,
        label_color=FLOW_DARK,
        body_size=7.5,
    )

    add_text(slide, "FULL ONE-SHOT PROMPT (ABBREVIATED ON SLIDE)", 6.97, 1.38, 5.90, 0.18, size=7.3, color=ONE_DARK, bold=True)
    add_rect(slide, 6.94, 1.58, 5.95, 1.03, WHITE, LINE, radius=True)
    prompt_excerpt = (
        "SYSTEM: You are a general chemistry assistant...\n"
        "USER: FROZEN DESIGN INPUT:\n"
        "RAW BATCH PROTOCOL: Fmoc-L-methionine... 0.10 M... air... 420 nm... 5 min... 99% yield.\n"
        "OBJECTIVE: Produce one executable, inventory-constrained screen.\n"
        "HARD CONSTRAINTS: Use only authoritative inventory; report STP and in-channel gas...\n"
        "AVAILABLE INVENTORY: {pumps:[...], MFC:[...], mixer:[...], photocoil:[...], BPR:[...], ...}\n"
        "OUTPUT CONTRACT: {recommended_disposition, disposition_rationale, proposal:{...}}"
    )
    add_text(
        slide,
        prompt_excerpt,
        7.07,
        1.66,
        5.68,
        0.86,
        size=6.4,
        color=INK,
        font="Liberation Mono",
    )
    add_text(slide, "6,799 prompt characters | exact prompt retained in benchmark artifacts", 8.10, 2.47, 4.60, 0.11, size=5.7, color=MUTED, align=PP_ALIGN.RIGHT)

    # Output sections.
    add_text(slide, "OUTPUT", 0.48, 2.70, 1.0, 0.16, size=7.3, color=FLOW_DARK, bold=True)
    add_rect(slide, 0.45, 2.89, 5.95, 1.78, WHITE, LINE, radius=True)
    add_picture_contain(slide, TOPOLOGY_IMAGE, 0.53, 2.96, 5.79, 1.64)

    flow_items = [
        ("Concentration", "0.10 M"),
        ("Liquid flow", "0.124 mL/min"),
        ("Air at STP", "1.857 sccm"),
        ("Air in-channel", "0.505 mL/min"),
        ("Liquid tau", "8.05 min"),
        ("Reactor", "1.0 mL PFA"),
        ("Temperature", "21 C"),
        ("Pressure / light", "3 bar(g) / 420 nm"),
    ]
    add_condition_grid(slide, flow_items, 0.45, 4.76, 5.95, 1.02, accent=FLOW_DARK)
    add_status_banner(
        slide,
        "SCHEMA VALID | INVENTORY ASSIGNED | GAS BOOKKEEPING CLOSED | QA SCORE 0.96",
        0.45,
        5.87,
        5.95,
        0.30,
        fill=GREEN_LIGHT,
        color=GREEN,
    )

    add_text(slide, "OUTPUT", 6.97, 2.70, 1.0, 0.16, size=7.3, color=ONE_DARK, bold=True)
    add_rect(slide, 6.94, 2.89, 5.95, 0.77, WHITE, LINE, radius=True)
    add_text(
        slide,
        "Liquid feed + air  ->  T-mixer  ->  1.0 mL PFA coil / 420 nm\n"
        "->  3 bar BPR  ->  gas-liquid separator  ->  amber collector",
        7.12,
        3.03,
        5.57,
        0.34,
        size=9.0,
        color=INK,
        bold=True,
        font="Liberation Mono",
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        "Narrative chain reconstructed from the response for display; no validated process graph was emitted.",
        7.16,
        3.45,
        5.49,
        0.13,
        size=5.8,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_text(slide, "8", 12.82, 7.18, 0.18, 0.12, size=6, color=MUTED, align=PP_ALIGN.RIGHT)

    one_items = [
        ("Concentration", "0.10 M"),
        ("Liquid flow", "0.200 mL/min"),
        ("Air at STP", "0.500 sccm"),
        ("Air in-channel", "0.136 / 0.167"),
        ("Liquid tau", "5.00 min"),
        ("Reactor", "1.0 mL PFA"),
        ("Temperature", "21 C"),
        ("Pressure / light", "3 bar(g) / 420 nm"),
    ]
    add_condition_grid(
        slide,
        one_items,
        6.94,
        3.75,
        5.95,
        1.02,
        accent=ONE_DARK,
        alert_keys={"Air in-channel"},
    )
    add_rect(slide, 6.94, 4.86, 5.95, 0.92, AMBER_LIGHT, AMBER, radius=True)
    add_labelled_lines(
        slide,
        [
            ("Gas inconsistency", "Two pressure-corrected values were reported: 0.136 and 0.167 mL/min."),
            ("O2 basis", "0.235 mol O2/mol substrate, below the stated 0.5 stoichiometric requirement."),
            ("Contract failure", "Six schema validation errors; nested objects replaced required strings/lists."),
        ],
        7.05,
        4.91,
        5.72,
        0.81,
        label_color=ONE_DARK,
        body_size=6.6,
    )
    add_status_banner(
        slide,
        "SCHEMA INVALID | NO INVENTORY-VALIDATED GRAPH | QA SCORE 0.63",
        6.94,
        5.87,
        5.95,
        0.30,
        fill=RED_LIGHT,
        color=RED,
    )

    # Resource use strips.
    flow_tokens = data["flow_summary"]["token_totals"]
    one_tokens = data["one_summary"]["token_totals"]
    add_metrics(
        slide,
        [
            ("LLM CALLS", str(data["flow_summary"]["llm_call_count"])),
            ("INPUT TOKENS", f"{flow_tokens['input_tokens']:,}"),
            ("OUTPUT TOKENS", f"{flow_tokens['output_tokens']:,}"),
            ("TOTAL TOKENS", f"{flow_tokens['total_tokens']:,}"),
            ("RUNTIME", f"{data['flow_summary']['runtime_s']:.0f} s"),
        ],
        0.45,
        6.28,
        5.95,
        0.62,
        accent=FLOW_DARK,
        fill=WHITE,
    )
    add_metrics(
        slide,
        [
            ("LLM CALLS", str(data["one_summary"]["llm_call_count"])),
            ("INPUT TOKENS", f"{one_tokens['input_tokens']:,}"),
            ("OUTPUT TOKENS", f"{one_tokens['output_tokens']:,}"),
            ("TOTAL TOKENS", f"{one_tokens['total_tokens']:,}"),
            ("RUNTIME", f"{data['one_summary']['runtime_s']:.0f} s"),
        ],
        6.94,
        6.28,
        5.95,
        0.62,
        accent=ONE_DARK,
        fill=WHITE,
    )

    add_text(
        slide,
        "Resource totals include internal inference. FlowPilot uses a compact user-facing intake but expands it into staged internal prompts. "
        "One-shot is cheaper and faster here; the comparison concerns output closure and usability, not token parity.",
        0.38,
        7.14,
        12.57,
        0.18,
        size=5.9,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    # Remove the default first slide if a template created one (python-pptx does not).
    presentation.save(OUT_PATH)

    source_payload = {
        "schema_version": "flowpilot_case_study_slide_v1.0",
        "case": "photochemical_oxidation",
        "case_title": data["public_input"]["title"],
        "model": "claude-opus-4-6",
        "repeat": "repeat_01",
        "matched_settings": {
            "same_public_input": True,
            "same_inventory": True,
            "temperature": 0.2,
            "seed": 1007535521,
        },
        "prompt_footprint": {
            "flowpilot_protocol_chars": len(data["public_input"]["protocol"]),
            "flowpilot_protocol_plus_objective_chars": len(
                data["public_input"]["protocol"]
                + "\n\nOBJECTIVE:\n"
                + data["public_input"]["objective"]
            ),
            "flowpilot_inventory_attachment_chars": len(
                json.dumps(data["public_input"]["inventory"], indent=2)
            ),
            "flowpilot_structured_constraint_count": len(
                data["public_input"]["hard_constraints"]
            ),
            "oneshot_prompt_chars": len(data["one_prompt"]["system"])
            + len(data["one_prompt"]["user"]),
            "prompt_inventory": data["prompt_inventory"],
        },
        "output_footprint": {
            "flowpilot_final_response_chars": len(data["flow_final_response"]),
            "flowpilot_final_response_source_lines": len(data["flow_final_response"].splitlines()),
            "oneshot_raw_response_chars": len(data["one_raw_response"]),
            "oneshot_raw_response_source_lines": len(data["one_raw_response"].splitlines()),
            "flowpilot_source": str((FLOWPILOT_RUN / "llm_events.jsonl").relative_to(ROOT)),
            "oneshot_source": str((ONESHOT_RUN / "raw_response.json").relative_to(ROOT)),
        },
        "source_paths": {
            "flowpilot_run": str(FLOWPILOT_RUN.relative_to(ROOT)),
            "oneshot_run": str(ONESHOT_RUN.relative_to(ROOT)),
            "flowpilot_topology": str(TOPOLOGY_IMAGE.relative_to(ROOT)),
        },
        "flowpilot": {
            "conditions": data["flow_metrics"]["proposal_summary"],
            "air_sccm": data["flow_air"]["gas_flow_sccm"],
            "air_in_channel_mL_min": data["flow_air"]["gas_flow_actual_mL_min"],
            "oxygen_equiv": data["flow_air"]["molar_equiv"],
            "schema_valid": data["flow_metrics"]["schema_valid"],
            "quality_assurance_score_v2": data["flow_metrics"]["quality_assurance_score_v2"],
            "runtime_s": data["flow_summary"]["runtime_s"],
            "llm_calls": data["flow_summary"]["llm_call_count"],
            "tokens": flow_tokens,
        },
        "one_shot": {
            "conditions": data["one_metrics"]["proposal_summary"],
            "air_sccm": data["one_air"]["gas_flow_sccm"],
            "air_in_channel_top_level_mL_min": data["one_air"]["gas_flow_actual_mL_min"],
            "air_in_channel_calculated_mL_min": data["one_air"]["gas_flow_calculations"]["gas_flow_actual_corrected_mL_min"],
            "oxygen_to_substrate_ratio": data["one_air"]["oxygen_equivalent_basis"]["O2_to_substrate_ratio"],
            "schema_valid": data["one_metrics"]["schema_valid"],
            "schema_error": data["one_schema_error"]["error"],
            "quality_assurance_score_v2": data["one_metrics"]["quality_assurance_score_v2"],
            "runtime_s": data["one_summary"]["runtime_s"],
            "llm_calls": data["one_summary"]["llm_call_count"],
            "tokens": one_tokens,
            "prompt_characters": len(data["one_prompt"]["system"]) + len(data["one_prompt"]["user"]),
        },
    }
    SOURCE_PATH.write_text(json.dumps(source_payload, indent=2), encoding="utf-8")
    FLOWPILOT_OUTPUT_PATH.write_text(data["flow_final_response"], encoding="utf-8")
    ONESHOT_OUTPUT_PATH.write_text(data["one_raw_response"], encoding="utf-8")
    return OUT_PATH


if __name__ == "__main__":
    print(build())
