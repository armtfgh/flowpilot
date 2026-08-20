from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
PROFILE_PATH = ROOT / "khu_inventory_flowpilot.json"

SW = 13.333
SH = 7.5

INK = "17212B"
MUTED = "5F6B76"
LINE = "D7DEE3"
PANEL = "F5F7F8"
TEAL = "0F766E"
TEAL_LIGHT = "E8F4F2"
GREEN = "2A9D8F"
AMBER = "D49328"
AMBER_LIGHT = "FFF4D8"
RED = "C94B4B"
RED_LIGHT = "FCEBEC"
BLUE = "2F6B9A"
BLUE_LIGHT = "EAF2F8"
WHITE = "FFFFFF"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_fill(shape, color: str, transparency: int = 0) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.fill.transparency = transparency


def set_line(shape, color: str, width: float = 1.0) -> None:
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 16,
    color: str = INK,
    bold: bool = False,
    font: str = "Aptos",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0,
    fit: bool = False,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    if fit:
        frame.fit_text(font_family=font, max_size=int(size))
    return box


def add_rich_text(slide, runs, x, y, w, h, *, size=14, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.space_after = Pt(0)
    for item in runs:
        run = paragraph.add_run()
        run.text = item[0]
        run.font.name = "Aptos"
        run.font.size = Pt(item[3] if len(item) > 3 else size)
        run.font.bold = item[1]
        run.font.color.rgb = rgb(item[2])
    return box


def rect(slide, x, y, w, h, *, fill=WHITE, line=LINE, radius=False, width=1.0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_fill(shape, fill)
    set_line(shape, line, width)
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def circle(slide, x, y, d, *, fill=WHITE, line=LINE, width=1.0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)
    )
    set_fill(shape, fill)
    set_line(shape, line, width)
    return shape


def add_header(slide, number: str, title: str, subtitle: str) -> None:
    add_text(slide, "FlowPilot", 0.55, 0.27, 1.4, 0.28, size=12, bold=True, color=TEAL)
    add_text(slide, number, 12.25, 0.27, 0.5, 0.28, size=11, color=MUTED, align=PP_ALIGN.RIGHT)
    add_text(slide, title, 0.55, 0.72, 12.1, 0.55, size=28, bold=True)
    add_text(slide, subtitle, 0.55, 1.28, 12.1, 0.38, size=13.5, color=MUTED)
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(0.55), Inches(1.72), Inches(12.78), Inches(1.72),
    )
    set_line(line, LINE, 1.0)


def add_footer(slide, text: str) -> None:
    add_text(slide, text, 0.55, 7.15, 12.2, 0.18, size=8.5, color=MUTED)


def add_number_badge(slide, number: str, x: float, y: float, color: str) -> None:
    circle(slide, x, y, 0.42, fill=color, line=color)
    add_text(
        slide, number, x, y + 0.01, 0.42, 0.38,
        size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_arrow(slide, x1, y1, x2, y2, color=TEAL, width=2.0):
    arrow = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2),
    )
    set_line(arrow, color, width)
    arrow.line.end_arrowhead = True
    return arrow


def add_file_token(slide, label: str, x: float, y: float, color: str):
    shape = rect(slide, x, y, 0.75, 0.48, fill=WHITE, line=color, radius=True, width=1.3)
    add_text(
        slide, label, x, y + 0.02, 0.75, 0.4,
        size=9.5, color=color, bold=True, align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    return shape


def load_counts() -> dict[str, int]:
    profile = json.loads(PROFILE_PATH.read_text(encoding="utf-8"))
    inventory = profile["lab_inventory"]
    return {
        "Reactors": len(inventory.get("reactors") or []),
        "Pumps": len(inventory.get("pumps") or []),
        "Tubing": len(inventory.get("tubing") or []),
        "Gas hardware": len(inventory.get("gas_hardware") or []),
        "BPR settings": len(inventory.get("BPR_available") or []),
    }


def build_slide_1(prs: Presentation, counts: dict[str, int]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    add_header(
        slide,
        "01",
        "Inventory-aware design: the latest FlowPilot addition",
        "A laboratory profile is built once, versioned, and enforced in every batch-to-flow translation.",
    )

    steps = [
        (
            "1", "Collect", TEAL, TEAL_LIGHT,
            "Accept laboratory PDFs, PPTX/DOCX files, pasted JSON, or structured forms.",
        ),
        (
            "2", "Normalize + freeze", BLUE, BLUE_LIGHT,
            "Convert equipment and operating constraints into one reproducible profile (v2.0).",
        ),
        (
            "3", "Enforce", RED, RED_LIGHT,
            "Snap designs to listed hardware and block unresolved reactor trains, mixers, BPRs, or limits.",
        ),
    ]
    for idx, (num, heading, color, light, body) in enumerate(steps):
        y = 2.05 + idx * 1.36
        rect(slide, 0.55, y, 5.65, 1.08, fill=light, line=light, radius=True)
        add_number_badge(slide, num, 0.82, y + 0.32, color)
        add_text(slide, heading, 1.42, y + 0.18, 1.9, 0.3, size=16, bold=True, color=color)
        add_text(slide, body, 1.42, y + 0.51, 4.42, 0.43, size=11.5, color=INK)

    add_text(slide, "Example: KHU laboratory profile", 6.62, 2.03, 5.85, 0.3, size=16, bold=True)
    rect(slide, 6.62, 2.43, 6.15, 3.74, fill=PANEL, line=LINE, radius=True)
    add_rich_text(
        slide,
        [
            ("khu_laboratory_inventory", True, INK, 13),
            ("   v1  |  schema v2.0", False, MUTED, 10.5),
        ],
        6.92, 2.68, 5.5, 0.35,
    )
    metrics = list(counts.items())
    for idx, (label, value) in enumerate(metrics):
        row = idx // 3
        col = idx % 3
        x = 6.92 + col * 1.78
        y = 3.24 + row * 1.02
        rect(slide, x, y, 1.55, 0.78, fill=WHITE, line=LINE, radius=True)
        add_text(slide, str(value), x + 0.12, y + 0.10, 0.52, 0.36, size=21, bold=True, color=TEAL)
        add_text(slide, label, x + 0.12, y + 0.49, 1.3, 0.18, size=8.5, color=MUTED)

    rect(slide, 6.92, 5.30, 5.48, 0.56, fill=RED_LIGHT, line=RED, radius=True, width=1.2)
    add_text(slide, "HARD CONSTRAINT", 7.10, 5.46, 1.35, 0.18, size=8.5, bold=True, color=RED)
    add_text(slide, "No inline degasser  |  BPR not confirmed", 8.48, 5.42, 3.65, 0.22, size=10.5, bold=True)

    rect(slide, 0.55, 6.44, 12.22, 0.48, fill=TEAL_LIGHT, line=TEAL_LIGHT, radius=True)
    add_rich_text(
        slide,
        [
            ("Design consequence  ", True, TEAL, 11.5),
            ("The agent may propose chemistry, but only inventory-compatible parameters can be shown as an executable design.", False, INK, 11.5),
        ],
        0.82, 6.57, 11.7, 0.23,
    )
    add_footer(slide, "Latest release | Inventory Manager + deterministic design disposition")


def add_stage_box(slide, x, y, w, h, title, body, *, color, fill):
    rect(slide, x, y, w, h, fill=fill, line=color, radius=True, width=1.3)
    add_text(slide, title, x + 0.18, y + 0.15, w - 0.36, 0.27, size=13, bold=True, color=color)
    add_text(slide, body, x + 0.18, y + 0.50, w - 0.36, h - 0.62, size=9.5, color=INK)


def build_slide_2(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)
    add_header(
        slide,
        "02",
        "Inventory management: end-to-end workflow",
        "The profile is generated independently, then injected as a frozen hard-constraint layer before design execution.",
    )

    add_text(slide, "BUILD ONCE", 0.58, 1.95, 1.1, 0.2, size=9, bold=True, color=TEAL)
    add_text(slide, "REUSE FOR EVERY DESIGN", 6.43, 1.95, 2.1, 0.2, size=9, bold=True, color=BLUE)

    # Build lane
    add_stage_box(
        slide, 0.55, 2.26, 1.72, 1.58,
        "Lab sources", "Equipment lists\nand constraints", color=TEAL, fill=TEAL_LIGHT,
    )
    add_file_token(slide, "PDF", 0.72, 3.13, TEAL)
    add_file_token(slide, "PPTX", 1.38, 3.13, TEAL)
    add_file_token(slide, "FORM", 1.05, 3.60, TEAL)

    add_arrow(slide, 2.31, 3.05, 2.64, 3.05, TEAL)
    add_stage_box(
        slide, 2.70, 2.26, 1.92, 1.58,
        "Input collector", "LLM extraction +\ndeterministic fallback", color=TEAL, fill=WHITE,
    )
    circle(slide, 3.30, 3.08, 0.55, fill=TEAL, line=TEAL)
    add_text(slide, "AI", 3.30, 3.14, 0.55, 0.24, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    add_arrow(slide, 4.66, 3.05, 4.99, 3.05, TEAL)
    add_stage_box(
        slide, 5.05, 2.26, 1.92, 1.58,
        "Normalize", "Schema validation", color=TEAL, fill=WHITE,
    )
    for i, label in enumerate(("Pumps", "Reactors", "Limits")):
        rect(slide, 5.28, 3.13 + i * 0.22, 1.45, 0.17, fill=PANEL, line=LINE, radius=True)
        add_text(slide, label, 5.38, 3.14 + i * 0.22, 1.1, 0.13, size=7.2, color=MUTED)

    add_arrow(slide, 7.01, 3.05, 7.34, 3.05, TEAL)
    add_stage_box(
        slide, 7.40, 2.26, 2.12, 1.58,
        "Versioned JSON", "Equipment + constraints", color=BLUE, fill=BLUE_LIGHT,
    )
    add_text(slide, "{ profile_id\n  lab_inventory\n  operating_limits }", 7.65, 3.05, 1.65, 0.60, size=8.0, color=BLUE, font="Liberation Mono")

    # Reuse lane and design core
    down = slide.shapes.add_connector(
        MSO_CONNECTOR.ELBOW,
        Inches(8.46), Inches(3.87), Inches(8.46), Inches(4.44),
    )
    set_line(down, BLUE, 2.0)
    down.line.end_arrowhead = True

    add_stage_box(
        slide, 5.05, 4.52, 4.47, 1.30,
        "FlowPilot design core",
        "Standardized intake  →  upstream chemistry  →  engineering + council",
        color=BLUE, fill=WHITE,
    )
    # Small inventory rail entering the core
    rect(slide, 5.33, 5.29, 3.91, 0.30, fill=BLUE_LIGHT, line=BLUE_LIGHT, radius=True)
    add_text(slide, "Frozen inventory context is injected into every agent prompt", 5.48, 5.37, 3.62, 0.15, size=8.2, color=BLUE, bold=True)

    add_arrow(slide, 9.57, 5.17, 9.93, 5.17, BLUE)
    add_stage_box(
        slide, 9.99, 4.52, 2.78, 1.30,
        "Deterministic gate",
        "Exact hardware match\n+ geometry closure\n+ safety and capability checks",
        color=RED, fill=RED_LIGHT,
    )

    # Decision outputs
    trunk = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(11.38), Inches(5.84), Inches(11.38), Inches(6.04),
    )
    set_line(trunk, MUTED, 1.6)
    branch = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(9.47), Inches(6.04), Inches(11.72), Inches(6.04),
    )
    set_line(branch, MUTED, 1.6)
    for x in (9.47, 11.72):
        stem = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(x), Inches(6.04), Inches(x), Inches(6.19),
        )
        set_line(stem, MUTED, 1.6)
        stem.line.end_arrowhead = True
    rect(slide, 8.46, 6.21, 2.02, 0.60, fill=TEAL_LIGHT, line=TEAL, radius=True, width=1.3)
    add_text(slide, "SCREEN", 8.64, 6.32, 0.72, 0.20, size=11, bold=True, color=TEAL)
    add_text(slide, "exact listed hardware", 9.33, 6.34, 0.95, 0.18, size=7.7, color=INK)
    rect(slide, 10.67, 6.21, 2.10, 0.60, fill=RED_LIGHT, line=RED, radius=True, width=1.3)
    add_text(slide, "BLOCK", 10.85, 6.32, 0.62, 0.20, size=11, bold=True, color=RED)
    add_text(slide, "no diagram or recipe", 11.46, 6.34, 1.08, 0.18, size=7.7, color=INK)

    add_text(slide, "Example", 0.58, 4.56, 0.72, 0.20, size=9, bold=True, color=MUTED)
    rect(slide, 0.55, 4.87, 3.98, 1.95, fill=PANEL, line=LINE, radius=True)
    add_text(slide, "KHU profile + two-stage O₂ process", 0.83, 5.10, 3.45, 0.25, size=12.5, bold=True)
    add_text(slide, "Available", 0.83, 5.54, 0.82, 0.18, size=8.5, bold=True, color=TEAL)
    add_text(slide, "PFA/FEP coils, pumps, O₂ MFC", 1.61, 5.52, 2.55, 0.22, size=9.5)
    add_text(slide, "Missing", 0.83, 5.91, 0.82, 0.18, size=8.5, bold=True, color=RED)
    add_text(slide, "confirmed BPR, mixer, serial train", 1.61, 5.89, 2.55, 0.22, size=9.5)
    add_text(slide, "Disposition", 0.83, 6.28, 0.82, 0.18, size=8.5, bold=True, color=RED)
    add_text(slide, "BLOCK — unresolved hardware", 1.61, 6.26, 2.55, 0.22, size=9.5, bold=True)

    add_footer(slide, "Authority order: measured evidence > hard constraints > protocol facts > hypotheses > model inference")


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    prs.core_properties.title = "FlowPilot Inventory Management"
    prs.core_properties.subject = "Latest release: inventory-aware batch-to-flow design"
    prs.core_properties.author = "FlowPilot"
    counts = load_counts()
    build_slide_1(prs, counts)
    build_slide_2(prs)
    output = OUT / "FlowPilot_Inventory_Management.pptx"
    prs.save(output)
    return output


if __name__ == "__main__":
    print(build())
