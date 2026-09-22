"""Preserve existing Word packages while extending the two connected-process cases."""
from __future__ import annotations

import csv
import importlib.util
import io
import json
import re
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Inches, Pt
from lxml import etree as E
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "manuscript"
OUT = BASE / "revision_20260922"
FIG = OUT / "figures"
F5 = ROOT / "outputs/figure5_pure_oxygen_physics_20260918/slides"
F6 = ROOT / "outputs/khu_revised_six_20260915/collaborator_slides_20260915_163845"
spec = importlib.util.spec_from_file_location("prior_revision", BASE / "revise_text_only.py")
prior = importlib.util.module_from_spec(spec)
spec.loader.exec_module(prior)
W, NS = prior.W, dict(prior.NS)
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P = "http://schemas.openxmlformats.org/package/2006/relationships"
NS.update(r=R, wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing")


def text(node):
    return "".join(node.xpath(".//w:t/text()", namespaces=NS))


def rows(path):
    with path.open(newline="") as f:
        return list(csv.DictReader(f))


def num(value):
    if value in (None, "", "None"):
        return "-"
    return f"{float(value):.5g}"


class Package:
    def __init__(self, source):
        self.source = source
        with ZipFile(source) as z:
            self.parts = {n: z.read(n) for n in z.namelist()}
        self.doc = E.fromstring(self.parts["word/document.xml"])
        self.body = self.doc.find(W + "body")
        self.ps = self.body.findall(W + "p")
        self.rels = E.fromstring(self.parts["word/_rels/document.xml.rels"])
        self.changed = []
        self.citation_props = E.Element(W + "rPr")
        E.SubElement(self.citation_props, W + "vertAlign", {W + "val": "superscript"})
        self.font(self.citation_props)
        self.next_id = max([int(x.get("id")) for x in self.doc.xpath(".//wp:docPr", namespaces=NS)] + [1]) + 1

    @staticmethod
    def font(props):
        for name, attrs in [("rFonts", {"ascii": "Times New Roman", "hAnsi": "Times New Roman", "cs": "Times New Roman"}), ("sz", {"val": "22"}), ("szCs", {"val": "22"})]:
            child = props.find(W + name)
            if child is None:
                child = E.SubElement(props, W + name)
            for k, v in attrs.items():
                child.set(W + k, v)

    def revise(self, p, value):
        value = value.replace(" degrees C", " \u00b0C")
        prior.revise(p, value, self.citation_props)
        for r in p.findall(W + "r"):
            rp = r.find(W + "rPr")
            if rp is None:
                rp = E.Element(W + "rPr")
                r.insert(0, rp)
            self.font(rp)
        self.changed.append(p)

    def paragraph(self, value, style="Normal", bold=False, page_break=False, keep=False):
        p = E.Element(W + "p")
        pp = E.SubElement(p, W + "pPr")
        E.SubElement(pp, W + "pStyle", {W + "val": style})
        E.SubElement(pp, W + "spacing", {W + "after": "120"})
        if page_break:
            E.SubElement(pp, W + "pageBreakBefore")
        if keep:
            E.SubElement(pp, W + "keepNext")
        self.revise(p, value)
        if bold:
            for r in p.findall(W + "r"):
                E.SubElement(r.find(W + "rPr"), W + "b")
        return p

    def before(self, target, nodes):
        for node in nodes:
            target.addprevious(node)

    def after(self, target, nodes):
        for node in nodes:
            target.addnext(node)
            target = node

    def image(self, path, width=6.5, page_break=False):
        fragment = Document()
        p = fragment.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.page_break_before = page_break
        p.paragraph_format.space_after = Pt(6)
        p.add_run().add_picture(str(path), width=Inches(width))
        node = E.fromstring(E.tostring(p._p))
        rid = "rIdRevision" + str(self.next_id)
        name = f"revision_{self.next_id}_" + path.name
        self.parts["word/media/" + name] = path.read_bytes()
        E.SubElement(self.rels, "{" + P + "}Relationship", Id=rid, Type=R + "/image", Target="media/" + name)
        for blip in node.xpath(".//a:blip", namespaces=NS):
            blip.set("{" + R + "}embed", rid)
        for dp in node.xpath(".//wp:docPr", namespaces=NS):
            dp.set("id", str(self.next_id))
            dp.set("name", path.stem)
            dp.set("descr", path.stem.replace("_", " "))
        self.next_id += 1
        return node

    def replace_image(self, p, path):
        new = self.image(path, 6.5)
        pp = new.find(W + "pPr")
        E.SubElement(pp, W + "pageBreakBefore")
        self.body.replace(p, new)

    def table(self, headers, data, widths=None):
        d = Document()
        t = d.add_table(rows=1, cols=len(headers))
        t.autofit = False
        if widths is None:
            widths = [6.5 / len(headers)] * len(headers)
        for col, width in zip(t.columns, widths):
            col.width = Inches(width)
        for cell, heading in zip(t.rows[0].cells, headers):
            cell.text = heading
        for values in data:
            for cell, value in zip(t.add_row().cells, values):
                cell.text = str(value)
        for row_index, row in enumerate(t.rows):
            trp = row._tr.get_or_add_trPr()
            E.SubElement(trp, W + "cantSplit")
            if row_index == 0:
                E.SubElement(trp, W + "tblHeader")
            for c, width in zip(row.cells, widths):
                c.width = Inches(width)
                cp = c._tc.get_or_add_tcPr()
                if row_index == 0:
                    E.SubElement(cp, W + "shd", {W + "fill": "E7EFF1"})
                for p in c.paragraphs:
                    p.alignment = 0
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(4)
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(11)
                        r.bold = row_index == 0
        tp = t._tbl.tblPr
        borders = E.SubElement(tp, W + "tblBorders")
        for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            E.SubElement(borders, W + edge, {W + "val": "single", W + "sz": "4", W + "color": "C2CCD0"})
        node = E.fromstring(E.tostring(t._tbl))
        self.changed.append(node)
        return node

    def save(self, path, marked=False):
        doc = deepcopy(self.doc)
        if marked:
            tree = self.doc.getroottree()
            changed_paths = {tree.getpath(p) for n in self.changed for p in ([n] if n.tag == W + "p" else n.findall(".//" + W + "p")) if p.getroottree().getroot() is self.doc}
            for original, cloned in zip(self.doc.iter(W + "p"), doc.iter(W + "p")):
                if tree.getpath(original) in changed_paths:
                    for r in cloned.findall(W + "r"):
                        if not text(r).strip():
                            continue
                        rp = r.find(W + "rPr")
                        if rp is None:
                            rp = E.Element(W + "rPr")
                            r.insert(0, rp)
                        old = rp.find(W + "highlight")
                        if old is not None:
                            rp.remove(old)
                        E.SubElement(rp, W + "highlight", {W + "val": "yellow"})
        parts = dict(self.parts)
        parts["word/document.xml"] = E.tostring(doc, xml_declaration=True, encoding="UTF-8", standalone=True)
        parts["word/_rels/document.xml.rels"] = E.tostring(self.rels, xml_declaration=True, encoding="UTF-8", standalone=True)
        settings = E.fromstring(parts["word/settings.xml"])
        update = settings.find(W + "updateFields")
        if update is None:
            update = E.SubElement(settings, W + "updateFields")
        update.set(W + "val", "true")
        parts["word/settings.xml"] = E.tostring(settings, xml_declaration=True, encoding="UTF-8", standalone=True)
        with ZipFile(path, "w", ZIP_DEFLATED) as z:
            for name, data in parts.items():
                z.writestr(name, data)


AUTHORS = "Amirreza Mottafegh[[\u2020]], Mincheol Park[[\u2020]], [Name to be added], Dr. Myeong, Professor Park, and Dr. Gwang-Noh Ahn"
EQUAL = "\u2020Amirreza Mottafegh and Mincheol Park contributed equally to this work. Author names, affiliations, and the additional author will be finalized before submission."

NEW_REFS = {
46: "Chauhan, R.; Rana, A.; Mottafegh, A.; Kim, D. P.; Singh, A. K. Manual to Auto-Optimization Platform of Multistep Apixaban Synthesis. Organic Process Research & Development 2025, 29 (3), 881-888. https://doi.org/10.1021/acs.oprd.4c00535.",
47: "Aand, D.; Rana, A.; Mottafegh, A.; Kim, D. P.; Singh, A. K. Autonomous closed-loop photochemical reaction optimization for the synthesis of various angiotensin II receptor blocker molecules. Reaction Chemistry & Engineering 2024, 9 (9), 2427-2435. https://doi.org/10.1039/D4RE00138A.",
48: "Chauhan, R.; Mottafegh, A.; Rana, A.; Singh, A. K. UV-visible Auto-Reactometry enabled parallel screening for accelerating chemical synthesis development with LLM. Chemical Engineering Journal 2026, 533, 174829. https://doi.org/10.1016/j.cej.2026.174829.",
49: "Mottafegh, A.; Ahn, G.-N. Adaptive human-in-the-loop optimization using language-guided priors for chemical experiments. Journal of Chemical Information and Modeling 2026, 66 (15), 8833-8847. https://doi.org/10.1021/acs.jcim.6c00976.",
50: "Lu, Y.; Yan, Z.; Jiang, Y.; He, H.; Zhou, J.; Gao, S.; Xing, D. \u03b1-Bromination of Aldehydes by Continuous Flow Chemistry and Its Application to the Flow Synthesis of 2-Aminothiazoles. Organic Process Research & Development 2025, 29, 2047-2055. https://doi.org/10.1021/acs.oprd.5c00076.",
}


def expanded_cites(value):
    values = []
    for bit in re.split(r"\s*,\s*", value):
        if re.fullmatch(r"\d+", bit):
            values.append(int(bit))
        elif re.fullmatch(r"\d+\s*[-\u2013]\s*\d+", bit):
            a, b = map(int, re.split(r"\s*[-\u2013]\s*", bit))
            values.extend(range(a, b + 1))
        else:
            return None
    return values


def compressed_cites(values):
    groups = []
    vals = sorted(set(values))
    i = 0
    while i < len(vals):
        j = i
        while j + 1 < len(vals) and vals[j + 1] == vals[j] + 1:
            j += 1
        groups.append(str(vals[i]) if j == i else f"{vals[i]}-{vals[j]}")
        i = j + 1
    return ",".join(groups)


def renumber_references(m):
    old_entries = {i + 1: p for i, p in enumerate(m.ps[86:131])}
    for key, value in NEW_REFS.items():
        p = deepcopy(m.ps[86])
        m.revise(p, value)
        journal = {46: "Organic Process Research & Development", 47: "Reaction Chemistry & Engineering", 48: "Chemical Engineering Journal", 49: "Journal of Chemical Information and Modeling", 50: "Organic Process Research & Development"}[key]
        author_title, tail = value.split(journal, 1)
        match = re.match(r" (\d{4}), (\d+)(.*)", tail)
        for child in list(p):
            if child.tag != W + "pPr":
                p.remove(child)
        for part, decoration in [(author_title, None), (journal, "i"), (" ", None), (match[1], "b"), (", ", None), (match[2], "i"), (match[3], None)]:
            props = E.Element(W + "rPr")
            m.font(props)
            if decoration:
                E.SubElement(props, W + decoration)
            prior.append_run(p, part, props)
        old_entries[key] = p
    order = []
    refs_heading = m.ps[85]
    for p in list(m.body)[:list(m.body).index(refs_heading)]:
        for r in p.findall(".//" + W + "r"):
            if not r.xpath("./w:rPr/w:vertAlign[@w:val='superscript']", namespaces=NS):
                continue
            ids = expanded_cites(text(r))
            if ids:
                for old in ids:
                    if old not in order:
                        order.append(old)
    assert set(order) == set(old_entries), (set(old_entries) - set(order), set(order) - set(old_entries))
    mapping = {old: new for new, old in enumerate(order, 1)}
    for p in list(m.body)[:list(m.body).index(refs_heading)]:
        for r in p.findall(".//" + W + "r"):
            if not r.xpath("./w:rPr/w:vertAlign[@w:val='superscript']", namespaces=NS):
                continue
            ids = expanded_cites(text(r))
            if ids:
                ts = r.findall(W + "t")
                ts[0].text = compressed_cites([mapping[x] for x in ids])
                for t in ts[1:]:
                    r.remove(t)
    for p in m.ps[86:131]:
        m.body.remove(p)
    m.after(refs_heading, [old_entries[i] for i in order])
    return mapping


def manuscript():
    m = Package(BASE / "manuscript_text_revised_round2.docx")
    p = m.ps
    m.before(p[1], [m.paragraph(AUTHORS, "BodyText")])
    m.after(p[1], [m.paragraph(EQUAL, "BodyText")])
    m.revise(p[77], "A.M. and M.P. contributed equally to this work. The complete contributor-role statement and approval of the final author list will be confirmed by all authors before submission.")
    m.revise(p[5], prior.marked_text(p[5]) + " Multistep apixaban synthesis[[46]] and integrated photochemical synthesis and separation of sartan intermediates[[47]] illustrate the practical importance of coordinating reaction development with downstream operations.")
    m.revise(p[7], "Recent large-language-model (LLM) copilots for chemistry[[33,34]] have expanded what AI can do for a chemist, from synthesis planning to autonomous experimentation[[35-38]]. Flow-specific developments include UV-visible auto-reactometry coupled to LLM-assisted Bayesian optimization for photochemical screening[[48]]. Such systems address condition optimization, whereas the task considered here is to construct and reconcile an equipment-constrained process proposal from a batch procedure. When a monolithic LLM is prompted directly for a flow reactor design, it must handle chemistry, kinetics, fluid mechanics, heat transfer, and safety within a single inference pass[[39,40]], without a separate calculation and reconciliation workflow. A fluent answer therefore does not by itself establish cross-domain consistency or show which constraints were checked.")
    m.revise(p[37], "The first connected-process case couples photoredox Giese addition with oxidation of the resulting sulfide to a sulfoxide (Figure 5). In the literature transformation, photoinduced single-electron transfer generates an alpha-thiomethyl radical from an alpha-silyl sulfide; addition to acrylonitrile constructs the C-C bond, and subsequent oxygen-dependent photochemistry changes the sulfur oxidation state.[[42]] This sequence tests more than residence-time selection: the oxygen-free first stage and oxidizing second stage must remain distinct within a continuously connected process. The collaborator's batch protocol uses 4 h under argon followed by 6 h exposed to air. These durations define source observations, not a kinetic law that can be scaled by a universal intensification factor.")
    m.after(p[37], [
        m.paragraph("The latest proposal binds each coil to a compatible irradiation module and introduces pure oxygen only at the Stage 2 inlet. For response Set 1, a 2 mL PFA coil in the Vapourtec UV-150 module is followed by a 20 mL FEP coil in the separately illuminated manual module. At a liquid feed of 0.020 mL min-1, Stage 1 has a nominal liquid residence time of 100 min. The 0.100 M substrate feed corresponds to 0.00200 mmol min-1; an oxygen inlet flow of 0.090 mL min-1 at 273.15 K and 1 atm supplies 0.00402 mmol min-1, or 2.01 equivalents. Both stages are proposed at 25 degrees C, with 450 and 448 nm irradiation, respectively. The Stage 2 inlet/STP index, V/(Q_liquid + Q_gas,STP), is 181.82 min; it is not a measurement of residence under reactor pressure.", "BodyText"),
        m.paragraph("This revision responds to the collaborator's observation of gas entering the upstream reactor during an earlier high-gas-flow trial. Requiring pure oxygen lowers the total gas feed needed for a given oxygen molar input relative to air, but does not by itself resolve pressure-driven backflow. The archived proposal contains a gas-line check valve and a 7 bar(g) back-pressure setting; liquid-branch protection, pressure compatibility, and startup behavior still require laboratory review. The later council backflow tools provide conditional diagnostics, not a validated probability of failure or permission to execute. Figure S20 and Tables S29-S31 document the inputs, component feeds, three response-set outputs, and measurement plan. Sets 1 and 3 produce the same reported settings, whereas Set 2 increases the first-stage volume to 5 mL; these are not three independent replicates. Final sulfoxide yields are reserved as XXX, with conversion, selectivity, analytical method, and actual operating conditions to be entered alongside them.", "BodyText"),
    ])
    m.replace_image(p[38], FIG / "figure5_prototype.png")
    m.revise(p[39], "Figure 5. Inventory-constrained photoredox Giese addition/oxidation. (a) Protocol and design context. (b) Archived Set 1 topology with compatible irradiation modules and oxygen introduced only before Stage 2. (c) Chemical sequence. (d) Proposed conditions for three response sets; XXX fields below reserve wet-lab yields, not measured or simulated data. Stage 1 time is V/Q_liquid; the Stage 2 value is an inlet/STP reporting index, not in-channel residence. Pressure/backflow checks remain unresolved. Full feeds, source inputs and other topologies are in ESI Section 8.1.")
    m.revise(p[41], "The second case is a thermal, two-stage amidation of 3-methyl-4-nitrobenzoic acid with benzylamine (Figure 6). DPDTC and DMAP first generate a 2-pyridyl thioester; downstream aminolysis forms the amide C(O)-N bond while displacing the sulfur-containing leaving group.[[43]] The input specifies 30 min at 95 degrees C for each batch stage, separated by cooling and amine addition. In a telescoped flow proposal, the thioester is neither isolated nor assumed to form quantitatively. Its formation in Stage 1 limits the reactive intermediate available in Stage 2, making final amide yield the appropriate combined-process response.")
    m.after(p[41], [
        m.paragraph("The proposed feed structure retains the amine as a separate stream until the interstage junction. Feed A contains acid (0.500 M), DPDTC (0.525 M), and DMAP (0.0500 M) in 2-MeTHF; Feed B contains benzylamine (2.10 M). In Set 1, Q_A = 0.160 and Q_B = 0.0400 mL min-1 give acid and amine molar flows of 0.0800 and 0.0840 mmol min-1, preserving 1.05 equivalents of amine. The second reactor receives the sum of the two liquid feeds. Two 5 mL ETFE coils therefore give 31.25 and 25.00 min, respectively, at 95 degrees C. The 7 bar(g) setting is a proposed means of maintaining liquid operation above the solvent's normal boiling temperature; pressure ratings and thermal behavior must still be confirmed for the assembled train. Omitting batch cooling is a design hypothesis, not evidence that interstage temperature is irrelevant.", "BodyText"),
        m.paragraph("Sets 2 and 3 select the same reported operating point: a 10 mL first reactor and 5 mL second reactor with liquid feeds of 0.280 and 0.0700 mL min-1, yielding stage times of 35.71 and 14.29 min. These proposals shorten the second stage while increasing the first-stage time, rather than applying a common speed-up factor to the complete sequence. The distinction can be tested by measuring residual acid, thioester, and final amide alongside the final yield. Feed concentrations are nominal preparation targets; premix stability, solubility, and the actual intermediate concentration require measurement. Figure S21 and Tables S32-S34 give the response sets, stage and feed calculations, and experimental reporting template. Final amide yields remain YYY pending insertion of the verified measurements and their analytical basis.", "BodyText"),
        m.paragraph("The alpha-bromination example from the original preprint[[50]] is retained unchanged as Figure S22 in ESI Section 9 for historical comparison. It is distinct from the current amidation case and is not presented as a new laboratory validation.", "BodyText"),
    ])
    m.replace_image(p[42], FIG / "figure6_prototype.png")
    m.revise(p[43], "Figure 6. Connected two-stage DPDTC-mediated amidation. (a) Protocol and design context. (b) Archived Set 1 topology with downstream amine addition and no intermediate isolation. (c) Acid activation through a 2-pyridyl thioester to N-benzyl-3-methyl-4-nitrobenzamide. (d) Proposed stage-specific conditions; YYY fields below reserve wet-lab yields. Times follow V/Q for the specified liquid-phase design, not measured conversion; downstream flow includes both feeds. Sets 2 and 3 coincide and are not independent replicates. Feed concentrations, molar flows and reporting fields are in ESI Section 8.2.")
    for caption in [p[39], p[43]]:
        E.SubElement(caption.find(W + "pPr"), W + "keepLines")
    m.revise(p[60], prior.marked_text(p[60]) + " Language-guided priors provide a related route for incorporating chemist hypotheses into experimental optimization while testing them against accumulating evidence[[49]]; that separate method is a prospective integration opportunity, not an algorithm evaluated within the present FlowPilot benchmark.")
    m.revise(p[69], prior.marked_text(p[69]) + " Major development milestones and their dated evidence are summarized in Table S35 (ESI Section 10). Later scientific-policy and backflow extensions are separated from the frozen software used for the earlier architecture benchmarks.")
    mapping = renumber_references(m)
    m.save(BASE / "manuscript_revised_20260922.docx")
    m.save(BASE / "manuscript_revised_20260922_marked.docx", True)
    return m, mapping


MILESTONES = [
    ("2026-03-31; snapshot 213ad6b3", "Batch parsing, chemistry planning, literature retrieval and Streamlit diagrams", "Ground batch-to-flow proposals in chemistry and literature."),
    ("2026-04-15 to 04-22; f08f178a, 491d436c, 3ecf7e44", "Deterministic engineering calculator and specialist council", "Separate dependent numerical calculations from model judgment; record conflicting assessments."),
    ("2026-04-24 onward; e596a073 and April 27 runs", "Model-pair and candidate-budget benchmark harness", "Persist model calls and compare council configurations, costs and candidate diversity."),
    ("2026-06-26 release note; June 29 snapshot 02a03f89", "Closed-loop experiment refinement and experiment-loop interface", "Use measured outcomes to revise unsupported design assumptions."),
    ("2026-06-29 run; June 30 snapshot 1d90c256", "Fixed-ID intake and frozen DesignInputPackage", "Record objectives, evidence, equipment and hypotheses before design."),
    ("2026-07-06; snapshot 98d92241", "Stronger inventory enforcement and gas/liquid reconciliation", "Retain hardware limits, specified equivalents and numerical closure after changes."),
    ("2026-07-28 runs; July 30 snapshot 0530af8e", "Architecture-ablation study with saved cells", "Separate contributions of retrieval, engineering, council and inventory."),
    ("2026-08-07 and 08-19 audits; ec9a48a7", "Canonical final-design contract", "Prevent disagreement among selected conditions, inventory realization, GUI and topology."),
    ("2026-08-19 to 08-25 campaigns; ec9a48a7, 6e49d9ac", "Repeated fixed-criterion architecture comparison", "Measure judge scores, critical flags, variability and resource use under shared inputs."),
    ("2026-09-04; a72e7c1f", "React/FastAPI GUI and inlet/STP reporting", "Expose consistent stage-resolved outputs and reproducible intake behavior."),
    ("2026-09-07 documents; September 14 snapshot 7a130a40", "Scientific design policy and inventory-bound candidate screening", "Keep stage order, objective, uncertain kinetics and selected conditions explicit."),
    ("2026-09-14 to 09-18 laboratory-inventory revisions; 7a130a40, 7c3998e7; v6/v7 profiles", "System-aware pump/module compatibility, feed stoichiometry and setting increments", "Translate collaborator clarifications into implementable equipment combinations and explicit component flows."),
    ("2026-09-18 trial records; additions beyond b48a5929", "Branch-specific backflow checks and experimental transient-pressure tools", "Expose unprotected reverse paths and conditional startup risks; not a calibrated failure-probability model."),
]


def esi():
    e = Package(BASE / "esi.docx")
    p = e.ps
    e.before(p[1], [e.paragraph(AUTHORS)])
    e.revise(p[5], EQUAL)
    nodes = []

    def para(t, **kwargs):
        node = e.paragraph(t, **kwargs)
        nodes.append(node)
        return node

    def heading(t, level=2, page_break=False):
        return para(t, style=f"Heading{level}", bold=True, keep=True, page_break=page_break)

    def table(n, caption, headers, data, widths=None, page_break=False):
        para(f"Table S{n}. {caption}", bold=True, keep=True, page_break=page_break)
        nodes.append(e.table(headers, data, widths))
        para("")

    heading("Connected-process case studies supporting main-text Figures 5 and 6", 1, True)
    para("This section distinguishes the supplied batch protocols, frozen chemist responses, archived model proposals, and experimental measurements. It adds no new wet-lab outcomes. All XXX (Figure 5) and YYY (Figure 6) fields are deliberate placeholders. Model-selected conditions are hypotheses for screening, not demonstrated optima. The September case-study exports are separate from the August architecture-comparison campaigns in Section 3.")
    heading("Photoredox Giese addition followed by sulfoxidation (main-text Figure 5)")
    para("The target sequence converts the alpha-silyl sulfide and acrylonitrile to a sulfide, then to sulfoxide 4a. The chemistry source is Park et al. [4]. Stage 1 is oxygen-free; pure oxygen is introduced only at the Stage 2 inlet in the revised flow proposal. The original batch protocol uses ambient air in the second step. The change to pure oxygen is a collaborator-requested constraint in the September 18 reruns, not an autonomous discovery attributed to the model.")
    source = (ROOT / "inventory_khu/source_review_20260915/selected_protocols_and_responses.md").read_text()
    chunks = re.split(r"^## Slide (\d+):[^\n]*\n", source, flags=re.M)
    slide = {int(chunks[i]): chunks[i + 1].strip() for i in range(1, len(chunks), 2)}

    def source_block(s):
        for t in re.split(r"\n\s*\n", s):
            t = t.strip()
            if t and not t.startswith("Figure ") and not t.startswith("Green Chem") and not t.startswith("ACS Sustainable"):
                para(t)

    para("Supplied batch protocol (verbatim collaborator transcription).", bold=True)
    source_block(slide[1])
    table(29, "Figure 5 response-set provenance and shared constraints.", ["Set", "Design emphasis", "Source / shared constraints"], [
        ["1", "Integrated performance", "KHU revised V2, slide 5"],
        ["2", "Conversion-focused", "KHU revised V2, slide 8"],
        ["3", "Throughput / processing time", "KHU revised V2, slide 11"],
        ["All", "Oxygen exclusion in Stage 1; at least 2.0 equivalents in Stage 2", "Original fixed-ID responses below; September 18 rerun additionally requires pure O2. Inventory v7 adds MFC setting-grid information to v6."],
    ], [0.5, 2.4, 3.6])
    para("Fixed-ID responses (verbatim; hypotheses are not measured facts).", bold=True)
    for k, n in enumerate([5, 8, 11], 1):
        para(f"Response Set {k}", bold=True, keep=True)
        source_block(slide[n])
    para("Archived execution: outputs/figure5_pure_oxygen_physics_20260918, attempt_01 for each set; upstream Claude Opus 4.6 and downstream Claude Sonnet 4.6; scientific_v2 policy with a candidate budget of 12. The run folders retain the input package, inventory snapshot, model-call records, council record, calculation checks and result JSON. This description does not imply bit-for-bit reproducibility from model names alone.")
    stages5 = rows(F5 / "all_stage_parameters.csv")
    stages6 = [r for r in rows(F6 / "all_stage_parameters.csv") if r["case"].startswith("figure6")]
    table(30, "Figure 5 nominal stage conditions and component-feed targets from the archived final exports. Panel I: stages.", ["Set / stage", "Coil / module", "Q liquid / gas (mL min-1)", "Time (min)", "T / BPR"], [
        [f"{r['case'][-1]} / {int(float(r['stage']))}", f"{num(r['volume_mL'])} mL {r['material']}; {r['module']}; {num(r['wavelength_nm'])} nm", f"{num(r['liquid_flow_mL_min'])} / {num(r['gas_inlet_STP_mL_min'])}", num(r['nominal_inlet_residence_min']), f"{num(r['temperature_C'])} C / {num(r['BPR_bar'])} bar(g)"] for r in stages5
    ], [0.6, 2.05, 1.5, 0.9, 1.45])
    para("Table S30 (continued). Panel II: identical nominal component feeds in all three sets. Liquid Feed A uses EtOH:pH 9 buffer (5:1, v/v). Its common volumetric flow is 0.0200 mL min-1; listing three components does not mean three liquid feeds. The gas feed is pure oxygen at inlet/STP.", bold=True, keep=True)
    nodes.append(e.table(["Component", "Feed / entry", "C (M)", "Molar flow (mmol min-1)", "Equiv."], [
        ["Alpha-silyl sulfide precursor", "A / Stage 1", "0.100", "0.00200", "1.00"],
        ["Acrylonitrile", "A / Stage 1", "0.200", "0.00400", "2.00"],
        ["Ir photocatalyst", "A / Stage 1", "0.000500", "0.0000100", "0.00500"],
        ["Oxygen", "Gas / Stage 2", "Not a liquid stock", "0.004016", "2.008"],
    ], [1.9, 1.1, 1.0, 1.7, 0.8]))
    para("Calculation basis: Q_gas,STP = 0.0900 mL min-1 and STP = 273.15 K and 1 atm. The archived calculator uses a gas molar volume of 22.41272243 mL mmol-1 (approximately 22.414 at STP). Thus n_dot(O2) = 0.0900/22.41272243 = 0.004016 mmol min-1 and equiv(O2) = 0.004016/(0.100 x 0.0200) = 2.008. The nominal Stage 2 inlet/STP index is 20/(0.0200 + 0.0900) = 181.82 min. This index is useful for consistent reporting but is not physical in-channel residence, gas holdup, oxygen uptake or a mass-transfer measurement. A dissolved-oxygen or conversion measurement is still needed to connect oxygen feed to reaction performance.")
    para("Equipment interpretation: Stage 1 uses an inventory-confirmed UV-150-compatible PFA coil and the UV-150 light; Stage 2 uses the Manual 2 module and its own 448 nm strip light. The UV-150 light is not detached for use with an unrelated manual module. The liquid pump is a Vapourtec E-series channel with BLUE tubing; gas delivery uses the FF-C00 MFC, a CV3301 gas-line check valve and the P713 mixing junction. No inline degasser is assigned. Offline deoxygenation and protection of the Stage 1 feed remain required by the oxygen-free protocol.")
    para("Unresolved experimental checks: the gas-line check valve addresses liquid entry into the gas branch, not gas migration into the upstream liquid reactor. The earlier collaborator incident is therefore not declared solved by its presence. Confirm valve orientation and pressure rating, liquid-branch isolation/protection, actual junction pressures, oxygen-compatible hardware, and startup/shutdown procedures before testing the archived 7 bar(g) proposal. The optional transient tool used illustrative dynamics and a 12 bar(g) simulated gas-source pressure; that value is not a laboratory setpoint and exceeds the recorded MFC inlet rating. Its outputs must not be treated as a calibrated incident prediction or safety clearance. Gas pressure traces and equipment response data were unavailable.")
    table(31, "Figure 5 experimental completion template. The three sets are design-response variants, not three replicate measurements.", ["Required record", "Set 1", "Set 2", "Set 3"], [
        ["Final sulfoxide yield (%)", "XXX", "XXX", "XXX"],
        ["Yield basis / internal standard / calibration", "XXX", "XXX", "XXX"],
        ["Substrate conversion and sulfoxide selectivity (%)", "XXX", "XXX", "XXX"],
        ["Residual sulfide / overoxidation products", "XXX", "XXX", "XXX"],
        ["Actual flow, temperature and pressure; steady-state sample window", "XXX", "XXX", "XXX"],
        ["Replicate count, individual values and SD if applicable", "XXX", "XXX", "XXX"],
        ["Startup/backflow observations and deviations", "XXX", "XXX", "XXX"],
    ], [3.5, 1, 1, 1])
    para("Report the observed operating point rather than copying proposed values when the experiment differs. Identical Set 1 and Set 3 settings do not create an independent repeat unless separate experiments are performed and recorded. Analytical yield, isolated yield, conversion and selectivity must not be interchanged. No wet-lab error bars should be added until the actual replicate structure is known.")
    add_topologies(e, nodes, 20, F5, "figure5", "Figure 5", "The Stage 2 time label in the archived topology is an inlet/STP index, not measured in-channel contact time. The branch-protection and pressure checks described above remain unresolved.")

    heading("Telescoped DPDTC-mediated amidation (main-text Figure 6)", page_break=True)
    para("The supplied protocol forms N-benzyl-3-methyl-4-nitrobenzamide through a 2-pyridyl thioester; the chemistry source is Saunders et al. [5]. Batch timing, nominal concentrations and stage order are protocol facts. Flow residence times, the omission of interstage cooling, and the use of a premixed activation feed are proposed choices requiring validation.")
    para("Supplied batch protocol (verbatim collaborator transcription).", bold=True)
    source_block(slide[12])
    table(32, "Figure 6 response-set provenance and common chemistry constraints.", ["Set", "Design emphasis", "Source / common requirements"], [
        ["1", "Integrated performance", "KHU revised V2, slide 16"],
        ["2", "Conversion-focused", "KHU revised V2, slide 19"],
        ["3", "Throughput / processing time", "KHU revised V2, slide 22"],
        ["All", "Sequential thioester formation then amidation", "Amine enters after Stage 1. Inventory v6; component concentration and molar-flow reporting; system-aware pump selection."],
    ], [0.5, 2.4, 3.6])
    para("Fixed-ID responses (verbatim).", bold=True)
    for k, n in enumerate([16, 19, 22], 1):
        para(f"Response Set {k}", bold=True, keep=True)
        source_block(slide[n])
    para("Archived numerical source: outputs/khu_revised_six_20260915/presentation/figure6_set1 through figure6_set3. The recorded routing is Claude Opus 4.6 upstream and Claude Sonnet 4.6 downstream, using scientific_v2 with a candidate budget of 12. The collaborator export dated 20260915_163845 provides the displayed topologies and component tables. These are preserved existing outputs, not newly generated designs in this document revision. The original model records and finalized JSON, not a later figure redraw, define the case provenance.")
    table(33, "Figure 6 stage conditions and component-feed targets. Panel I: final stage values; all coils are ETFE, 2.4 mm ID, operated at a proposed 95 C and 7 bar(g).", ["Set / stage", "Volume (mL)", "Total liquid Q (mL min-1)", "V/Q (min)", "Interpretation"], [
        [f"{r['case'][-1]} / {int(float(r['stage']))}", num(r['volume_mL']), num(r['liquid_flow_mL_min']), num(r['nominal_inlet_residence_min']), "Activation; Feed A" if int(float(r['stage'])) == 1 else "Amidation; A + B"] for r in stages6
    ], [0.8, 1, 1.7, 1, 2])
    para("Table S33 (continued). Panel II: nominal component concentrations and molar flows. Feed A contains acid, DPDTC and DMAP; Feed B contains benzylamine. Both use 2-MeTHF. Listed concentrations are preparation targets, not measured reactor concentrations.", bold=True, keep=True)
    feed6 = []
    for k in range(1, 4):
        feed6.extend(rows(F6 / f"figure6_set{k}/feed_parameters.csv"))
    nodes.append(e.table(["Set / feed", "Component", "C (M)", "Feed Q (mL min-1)", "Molar flow (mmol min-1)"], [
        [f"{r['case'][-1]} / {r['stream']}", r['component'], num(r['concentration_M']), num(r['flow_mL_min']), num(r['molar_flow_mmol_min'])] for r in feed6
    ], [0.7, 2.4, 0.8, 1.25, 1.35]))
    para("Stoichiometric closure: for Set 1, n_dot(acid) = 0.500 x 0.160 = 0.0800 mmol min-1 and n_dot(benzylamine) = 2.10 x 0.0400 = 0.0840 mmol min-1; their ratio is 1.05. Q_Stage2 = 0.160 + 0.0400 = 0.200 mL min-1, giving 5/0.200 = 25.00 min. Sets 2 and 3 preserve the same molar ratio at 0.280 and 0.0700 mL min-1. The fully mixed acid-equivalent concentration is 0.400 M in all sets. This is not a measured thioester concentration: at Stage 1 fractional conversion X1, the maximum thioester concentration available after mixing is 0.400 X1 M before considering side reactions.")
    para("Stock preparation and materials record: prepare and document separate Feed A and Feed B solutions at the target concentrations above, recording the actual final solution volumes, reagent assay/purity, preparation time and any precipitation or discoloration. The batch notation '0.4 M' refers to acid-equivalent concentration after the total solvent volume reaches 1.25 mL; it is not the concentration of the separate benzylamine stock. The proposed 2.10 M benzylamine stock and premixed activation feed require solubility and stability checks. If acid activation proceeds in the reservoir, the true reaction history extends beyond the stated coil residence time. Do not assume zero premix reaction or quantitative thioester formation. Feed vessels, pump channels and sampling must be documented so that this possible confound can be evaluated.")
    para("Process interpretation: choose compatible channels of the same Vapourtec E-series system where possible; the diagrams use the general Pump icon rather than implying that every pump is a syringe pump. No irradiation or gas feed is part of this thermal sequence. Downstream volumetric flow includes the benzylamine addition, so the two stage times cannot be calculated from a single unchanged inlet flow. Confirm temperature and pressure limits of every wetted component, operation above the solvent's normal boiling temperature, and whether direct hot interstage transfer is acceptable. Startup waste and offline workup are separate from the two nominal heated-coil residence times. Workup, isolation and analytical methods must be recorded from the actual experiment, not inferred from the topology.")
    table(34, "Figure 6 experimental completion template. YYY denotes data to be inserted, not zero yield.", ["Required record", "Set 1", "Set 2", "Set 3"], [
        ["Final amide yield (%) and analytical basis", "YYY", "YYY", "YYY"],
        ["Stage 1 acid / thioester composition", "YYY", "YYY", "YYY"],
        ["Final residual acid / thioester / side products", "YYY", "YYY", "YYY"],
        ["Actual stock concentrations, pump settings and reactor temperatures", "YYY", "YYY", "YYY"],
        ["Premix age, steady-state window and sample identity", "YYY", "YYY", "YYY"],
        ["Workup / isolated mass / identity confirmation", "YYY", "YYY", "YYY"],
        ["Replicate count, individual values and SD if applicable", "YYY", "YYY", "YYY"],
    ], [3.5, 1, 1, 1], page_break=True)
    para("Sets 2 and 3 have identical reported hardware and flows. Any comparison of their eventual yields must distinguish separate wet-lab repetitions from repeated presentation of one design. Report assay and isolated yields separately; include the NMR internal standard or chromatographic calibration, integration method, sample timing, and identity/purity evidence before replacing YYY.")
    add_topologies(e, nodes, 21, F6, "figure6", "Figure 6", "Stage residence times use the liquid flow entering each reactor; Stage 2 includes both liquid feeds. The diagrams are archived design proposals, not measured conversions or pressure validation.")

    heading("Archived preprint alpha-bromination case", 1, True)
    para("Figure S22 reproduces the original preprint Figure 6 artwork without modification, including its original internal labels and values. This earlier aldehyde-bromination example, based on Lu et al. [6], is retained as a historical computational illustration and is not the DPDTC amidation now shown in main-text Figure 6. The original artwork records unresolved mixing/design concerns; its model confidence, timings and predicted outcomes are not new measured results or an approved procedure. Reproduction here must not be taken to validate those archived statements.")
    nodes.append(e.image(FIG / "figureS22_preprint_figure6_exact.png"))
    para("Figure S22. Unmodified Figure 6 from the supplied FlowPilot preprint (preprint.pdf, page 14): historical alpha-bromination case and the associated model discussion and process comparison. The complete original artwork is preserved, not redrawn. Its chemistry source is Lu et al. [6]. Any predictions or unresolved qualifications within the image belong to that archived output and do not describe the new wet-lab amidation case. Source artwork extraction and image hashes are retained with this revision.")
    heading("Major pipeline evolution and version boundaries", 1, True)
    para("Table S35 records major workflow changes using repository snapshots, dated reports and archived runs. Dates identify available evidence, not independently verified release or deployment dates; no release tags were found. The purpose is to explain why the workflow evolved, not to count minor bug fixes or claim that each change increased chemical yield. Earlier benchmark numbers remain attached to their frozen campaigns. In particular, a recorded August 20 campaign used two candidates, whereas the later scientific-policy case studies used twelve; later intake, inventory and backflow behavior must not be retroactively attributed to the earlier benchmark.")
    table(35, "Major FlowPilot development milestones and their rationale. Snapshot identifiers and dated artifacts support chronology; they are not controlled before/after efficacy experiments.", ["Date / evidence", "Major addition", "Reason / interpretation"], MILESTONES, [1.9, 2.3, 2.3])
    para("Evidence anchors: outputs/FlowPilot_latest_release_2026-06-26.md; outputs/thq_intake_claude_gpt4o_20260629/run.log; outputs/gui_visual_audit_20260807/final_contract_overhaul/validation_summary.json; outputs/flowpilot_systemic_fix_20260819/SYSTEMIC_FIX_REPORT.md; deliverables/manuscript_benchmark_visualizations_20260825/METHODS_AND_SCOPE.md; outputs/validation/20260904_standardized_intake_stp_report.md; docs/flowpilot_2_development.md; inventory_khu/KHU_inventory_20260915_v6.json; outputs/council_backflow_ab_20260918_111811/REPORT.md; and docs/council_transient_physics.md. The transient-pressure extension was an experimental addition beyond snapshot b48a5929, using assumed dynamics without laboratory pressure traces.")
    para("A later rerun may evaluate a revised workflow, but it does not replace or improve a previously recorded benchmark result without a separately identified campaign. Computational replay, mocked integration checks, repeated model generations, and independent laboratory repeats are different evidence types and must remain separately labeled.")
    e.before(p[752], nodes)
    refs = [
        "4. Park, J.; Kim, S. H.; Cho, J.-Y.; Atriardi, S. R.; Kim, J.-Y.; Mardhiyah, H.; Park, B. Y.; Woo, S. K. Batch and Flow Synthesis of Sulfides and Sulfoxides Using Green Solvents and Oxidant through Visible-Light Photocatalysis. Green Chemistry 2025, 27, 3284-3292. https://doi.org/10.1039/D4GC05769D.",
        "5. Saunders, J. M.; Oceguera Nava, E.; Li, J.; Wong, M. J.; Freiberg, K. M.; Lipshutz, B. H. Flow-to-Flow Technology: Amide Formation in the Absence of Traditional Coupling Reagents Using DPDTC. ACS Sustainable Chemistry & Engineering 2025, 13, 6646-6655. https://doi.org/10.1021/acssuschemeng.5c00914.",
        "6. " + NEW_REFS[50],
    ]
    e.after(p[755], [e.paragraph(t) for t in refs])
    e.save(BASE / "esi_revised_20260922.docx")
    e.save(BASE / "esi_revised_20260922_marked.docx", True)
    return e


def add_topologies(e, nodes, fig_no, folder, caseprefix, mainfig, caveat):
    for i, letter in enumerate("abc", 1):
        nodes.append(e.paragraph(f"Figure S{fig_no}" + ("" if i == 1 else " (continued)") + f". Panel ({letter}): {mainfig}, response Set {i}.", bold=True, keep=True, page_break=True))
        nodes.append(e.image(folder / f"{caseprefix}_set{i}/topology.png"))
        nodes.append(e.paragraph(f"({letter}) Actual saved GUI topology for response Set {i}. " + caveat))


def audit(main, esi_doc, mapping):
    report = {}
    for stem, obj in [("manuscript", main), ("esi", esi_doc)]:
        with ZipFile(obj.source) as z:
            old = {n: sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith("word/media/")}
        dest = BASE / f"{stem}_revised_20260922.docx"
        with ZipFile(dest) as z:
            preserved = {n: sha256(z.read(n)).hexdigest() == h for n, h in old.items()}
            assert all(preserved.values())
            assert z.read("word/styles.xml") == obj.parts["word/styles.xml"]
            assert z.read("word/numbering.xml") == obj.parts["word/numbering.xml"]
        report[stem] = {"source": str(obj.source.relative_to(ROOT)), "source_sha256": sha256(obj.source.read_bytes()).hexdigest(), "output": str(dest.relative_to(ROOT)), "source_images_byte_preserved": preserved, "styles_and_numbering_unchanged": True}
    captions = {}
    for p in esi_doc.body.findall(W + "p"):
        match = re.match(r"(Figure|Table) S(\d+)\.", text(p))
        if match:
            captions[(match[1], int(match[2]))] = text(p)
    citations = []
    pattern = r"\b(Figures?|Tables?) S(\d+(?:\s*[-\u2013]\s*S?\d+)?(?:\s*(?:,|and)\s*S?\d+(?:\s*[-\u2013]\s*S?\d+)?)*)"
    for i, p in enumerate(main.body.findall(W + "p")):
        value = text(p)
        for match in re.finditer(pattern, value):
            kind = match[1].rstrip("s")
            for a, b in re.findall(r"(\d+)(?:\s*[-\u2013]\s*S?(\d+))?", match[2]):
                for n in range(int(a), int(b or a) + 1):
                    assert (kind, n) in captions, (kind, n)
                    citations.append({"paragraph": i, "reference": f"{kind} S{n}", "esi_caption": captions[(kind, n)], "manuscript_text": value})
    orders = {}
    for kind in ["Figure", "Table"]:
        seen = list(dict.fromkeys(int(r["reference"].split(" S")[1]) for r in citations if r["reference"].startswith(kind)))
        assert seen == sorted(seen), (kind, seen)
        assert set(seen) == {n for k, n in captions if k == kind}
        orders[kind] = seen
    with (OUT / "main_to_esi_citation_audit.csv").open("w", newline="") as f:
        if citations:
            w = csv.DictWriter(f, fieldnames=citations[0].keys())
            w.writeheader()
            w.writerows(citations)
    report["bibliography_old_to_new"] = mapping
    report["esi_first_citation_order"] = orders
    report["wet_lab_values_inserted"] = False
    report["case_figures"] = "New prototypes; saved model topologies; no new generation or laboratory run"
    (OUT / "document_audit.json").write_text(json.dumps(report, indent=2))
    (OUT / "bibliography_renumbering.json").write_text(json.dumps(mapping, indent=2))
    print(json.dumps({"main_references": len(mapping), "source_images_preserved": {k: len(report[k]["source_images_byte_preserved"]) for k in ["manuscript", "esi"]}, "citation_rows": len(citations)}, indent=2))


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    m, mapping = manuscript()
    e = esi()
    audit(m, e, mapping)
