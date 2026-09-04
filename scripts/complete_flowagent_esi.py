#!/usr/bin/env python3
"""Complete missing FlowAgent ESI material while preserving its Word template."""

from __future__ import annotations

from copy import deepcopy
import csv
from pathlib import Path
import shutil
import tempfile

import cairosvg
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DOCUMENT = ROOT / "FlowAgent_Esi.docx"
BACKUP = ROOT / "FlowAgent_Esi_before_completion_20260831.docx"
ASSET_ROOT = (
    ROOT
    / ".local_artifacts_backup_20260831/generated/deliverables/"
    "flowpilot_esi_figures_s1_s10_20260811"
)


def paragraph_with_text(document: Document, startswith: str) -> Paragraph:
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(startswith):
            return paragraph
    raise KeyError(f"Paragraph not found: {startswith}")


def table_with_header(document: Document, first_cell: str) -> Table:
    for table in document.tables:
        if table.rows and table.rows[0].cells[0].text.strip() == first_cell:
            return table
    raise KeyError(f"Table not found: {first_cell}")


def insert_paragraph_before(
    anchor,
    text: str = "",
    *,
    style: str = "Normal",
    template: Paragraph | None = None,
    bold_prefix: str | None = None,
    bold_title: bool = False,
) -> Paragraph:
    node = OxmlElement("w:p")
    anchor._element.addprevious(node)
    paragraph = Paragraph(node, anchor._parent)
    paragraph.style = style
    if template is not None and template._p.pPr is not None:
        old_ppr = paragraph._p.pPr
        if old_ppr is not None:
            paragraph._p.remove(old_ppr)
        paragraph._p.insert(0, deepcopy(template._p.pPr))
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        run = paragraph.add_run(text)
        run.bold = True if bold_title else None
    return paragraph


def insert_paragraph_after(
    anchor: Paragraph,
    text: str,
    *,
    style: str = "Normal",
    template: Paragraph | None = None,
) -> Paragraph:
    node = OxmlElement("w:p")
    anchor._p.addnext(node)
    paragraph = Paragraph(node, anchor._parent)
    paragraph.style = style
    if template is not None and template._p.pPr is not None:
        old_ppr = paragraph._p.pPr
        if old_ppr is not None:
            paragraph._p.remove(old_ppr)
        paragraph._p.insert(0, deepcopy(template._p.pPr))
    paragraph.add_run(text)
    return paragraph


def insert_image_before(
    anchor,
    image_path: Path,
    *,
    width_inches: float,
    page_break: bool = True,
) -> Paragraph:
    paragraph = insert_paragraph_before(anchor, style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.page_break_before = page_break
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))
    return paragraph


def insert_table_before(
    document: Document,
    anchor,
    headers: list[str],
    rows: list[list[str]],
    *,
    widths: list[float] | None = None,
) -> Table:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for col, value in enumerate(headers):
        table.rows[0].cells[col].text = value
    for values in rows:
        cells = table.add_row().cells
        for col, value in enumerate(values):
            cells[col].text = value
    for row_index, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        cannot_split = OxmlElement("w:cantSplit")
        row_properties.append(cannot_split)
        if row_index == 0:
            repeat_header = OxmlElement("w:tblHeader")
            repeat_header.set(qn("w:val"), "true")
            row_properties.append(repeat_header)
        for col, cell in enumerate(row.cells):
            if widths:
                cell.width = Inches(widths[col])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(12)
                    if row_index == 0:
                        run.bold = True
    anchor._element.addprevious(table._tbl)
    return table


def retitle_figure(source_number: int, target_number: int, directory: Path) -> Path:
    source = ASSET_ROOT / "figures" / f"Figure_S{source_number:02d}.svg"
    text = source.read_text(encoding="utf-8")
    text = text.replace(f"Figure S{source_number}", f"Figure S{target_number}")
    svg = directory / f"Figure_S{target_number:02d}.svg"
    png = directory / f"Figure_S{target_number:02d}.png"
    svg.write_text(text, encoding="utf-8")
    cairosvg.svg2png(bytestring=text.encode("utf-8"), write_to=str(png), dpi=300)
    return png


def rule_summary_rows() -> list[list[str]]:
    path = ASSET_ROOT / "source_data" / "fig2a_rule_landscape.csv"
    values: dict[str, dict[str, int | str]] = {}
    with path.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            item = values.setdefault(
                row["category_key"],
                {
                    "label": row["category_label"],
                    "total": int(row["category_total"]),
                    "hard_rule": 0,
                    "guideline": 0,
                    "tip": 0,
                    "safety": 0,
                },
            )
            item[row["severity"]] = int(row["count"])
    top = sorted(values.values(), key=lambda item: int(item["total"]), reverse=True)[:14]
    return [
        [
            str(item["label"]),
            str(item["total"]),
            str(item["hard_rule"]),
            str(item["guideline"]),
            str(item["tip"]),
            str(item["safety"]),
        ]
        for item in top
    ]


def set_caption(paragraph: Paragraph, label: str, caption: str) -> None:
    for run in paragraph.runs:
        for text_node in run._r.xpath(".//w:t"):
            text_node.text = ""
    first = paragraph.add_run()
    first.text = f"{label}. "
    first.bold = True
    second = paragraph.add_run(caption)
    second.bold = None


def add_table_caption_before(table: Table, template: Paragraph, label: str, caption: str) -> None:
    paragraph = insert_paragraph_before(
        table,
        f"{label}. {caption}",
        style="Normal",
        template=template,
        bold_prefix=f"{label}. ",
    )
    paragraph.paragraph_format.keep_with_next = True


def add_missing_knowledge_material(document: Document, images: dict[int, Path]) -> None:
    ablation = paragraph_with_text(document, "Ablation test and architecture comparison")
    body_template = paragraph_with_text(document, "The engineering knowledge store contains")
    subheading_template = paragraph_with_text(document, "Engineering rule-base construction and use")
    caption_template = paragraph_with_text(document, "Figure S3.")
    table_caption_template = paragraph_with_text(document, "Table S4.")

    caption = insert_paragraph_before(
        ablation,
        "Table S5. Largest engineering rule categories and severity composition.",
        style="Normal",
        template=table_caption_template,
        bold_prefix="Table S5. ",
    )
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.page_break_before = True
    insert_table_before(
        document,
        ablation,
        ["Category", "Total", "Hard rule", "Guideline", "Tip", "Safety"],
        rule_summary_rows(),
        widths=[1.70, 0.70, 1.00, 1.05, 0.75, 0.85],
    )
    insert_image_before(ablation, images[4], width_inches=6.25)
    figure_caption = insert_paragraph_before(
        ablation,
        (
            "Figure S4. Engineering rule-base structure. (A) Counts by category and severity "
            "for the 14 largest categories. (B) Fraction of rules containing a machine-detected "
            "quantitative expression; this indicates expression coverage, not independent "
            "equation verification. (C) Non-exclusive rule-category associations across chemistry "
            "classes. (D) Co-occurrence network of frequent engineering concepts."
        ),
        style="Normal",
        template=caption_template,
        bold_prefix="Figure S4. ",
    )
    figure_caption.paragraph_format.keep_with_next = False

    retrieval_heading = insert_paragraph_before(
        ablation,
        "Plan-aware retrieval and tier transitions",
        style="Body Text",
        template=subheading_template,
        bold_title=True,
    )
    retrieval_heading.paragraph_format.page_break_before = True
    insert_paragraph_before(
        ablation,
        (
            "The retrieval query is constructed from the ChemistryPlan rather than from protocol "
            "keywords alone. Reaction class, mechanism, catalyst or photocatalyst, solvent, phase "
            "regime, wavelength, temperature, and concentration are used when available. Retrieval "
            "begins with mechanism and phase filters over paired batch-flow records. If fewer than "
            "three candidates are found, the metadata filters are relaxed; if no paired records "
            "remain, the search expands to the complete corpus. Candidate records are reranked using "
            "0.60 semantic similarity and 0.40 field similarity. The field score combines "
            "photocatalyst, solvent, wavelength, temperature, and concentration terms with weights "
            "of 0.30, 0.20, 0.20, 0.15, and 0.15, respectively (Figure S5)."
        ),
        style="Body Text",
        template=body_template,
    )
    retrieval_caption = insert_paragraph_before(
        ablation,
        "Table S6. Frozen plan-aware retrieval tiers and reranking weights.",
        style="Normal",
        template=table_caption_template,
        bold_prefix="Table S6. ",
    )
    retrieval_caption.paragraph_format.keep_with_next = True
    insert_table_before(
        document,
        ablation,
        ["Component", "Setting", "Scope or transition"],
        [
            ["Tier 1", "Paired records; mechanism + phase filters", "Relax when fewer than three hits remain"],
            ["Tier 2", "Paired records; filters relaxed", "Expand when no hits remain"],
            ["Tier 3", "All records; no metadata filter", "Return ranked available records"],
            ["Semantic similarity", "0.60", "Contribution to final reranking score"],
            ["Field similarity", "0.40", "Contribution to final reranking score"],
            ["Photocatalyst", "0.30", "Contribution within field similarity"],
            ["Solvent / wavelength", "0.20 / 0.20", "Contributions within field similarity"],
            ["Temperature / concentration", "0.15 / 0.15", "Contributions within field similarity"],
        ],
        widths=[1.55, 2.45, 2.25],
    )
    insert_image_before(ablation, images[5], width_inches=6.25)
    insert_paragraph_before(
        ablation,
        (
            "Figure S5. Plan-aware retrieval workflow. (A) ChemistryPlan fields enrich the query "
            "before semantic or deterministic lexical retrieval. (B) Tiered retrieval starts with "
            "mechanism and phase filters, relaxes filters when coverage is insufficient, and finally "
            "searches all records. (C) Semantic and field similarities are combined during reranking. "
            "Held-out source identifiers are excluded before final top-k selection."
        ),
        style="Normal",
        template=caption_template,
        bold_prefix="Figure S5. ",
    )

    benchmark_heading = insert_paragraph_before(
        ablation,
        "Retrieval benchmark and leakage controls",
        style="Body Text",
        template=subheading_template,
        bold_title=True,
    )
    benchmark_heading.paragraph_format.page_break_before = True
    insert_paragraph_before(
        ablation,
        (
            "The frozen retrieval analysis contains 1,600 query-result pairs. Reranking changed "
            "25.1% of pairwise ranks. Top-five photocatalyst-family matching increased from 40.0% "
            "to 95.0% for iridium systems, 42.7% to 90.7% for organic dyes, 77.3% to 94.7% for "
            "ruthenium systems, 87.1% to 98.6% for TiO2, and 45.0% to 60.0% for ZnO (Figure S6). "
            "These measurements quantify metadata alignment and retrieval behavior; they do not "
            "measure reaction yield, process optimality, or experimental success. For benchmark "
            "cases, the hidden source record identifier is removed after candidate construction and "
            "before top-k selection. This leave-one-source-out control prevents direct retrieval of "
            "the designated reference record but cannot establish absence from model pretraining."
        ),
        style="Body Text",
        template=body_template,
    )
    benchmark_caption = insert_paragraph_before(
        ablation,
        "Table S7. Frozen retrieval benchmark summary.",
        style="Normal",
        template=table_caption_template,
        bold_prefix="Table S7. ",
    )
    benchmark_caption.paragraph_format.keep_with_next = True
    insert_table_before(
        document,
        ablation,
        ["Metric or family", "Semantic retrieval", "Plan-aware reranking", "Evaluation n"],
        [
            ["Query-result pairs reranked", "-", "25.1% changed rank", "1,600 pairs"],
            ["Iridium family match", "40.0%", "95.0%", "8 queries"],
            ["Organic-dye family match", "42.7%", "90.7%", "15 queries"],
            ["Ruthenium family match", "77.3%", "94.7%", "15 queries"],
            ["TiO2 family match", "87.1%", "98.6%", "14 queries"],
            ["ZnO family match", "45.0%", "60.0%", "4 queries"],
        ],
        widths=[1.75, 1.50, 1.70, 1.30],
    )
    insert_image_before(ablation, images[6], width_inches=6.25)
    insert_paragraph_before(
        ablation,
        (
            "Figure S6. Retrieval benchmark and leakage controls. (A) Top-five photocatalyst-family "
            "match rates for semantic retrieval and FlowPilot reranking. (B) Rank changes across "
            "1,600 frozen query-result pairs. (C) Field-score component availability. (D) Example "
            "iridium query before and after reranking. (E) Leave-one-source-out control and "
            "deterministic retrieval tests."
        ),
        style="Normal",
        template=caption_template,
        bold_prefix="Figure S6. ",
    )


def complete_benchmark_labels(document: Document) -> None:
    table_caption_template = paragraph_with_text(document, "Table S4.")
    captions = [
        ("Score", "Table S8", "Fixed 0-4 scoring anchors used by every independent judge."),
        ("ID", "Table S9", "Universal outcome criteria and applicability rules."),
        ("Condition", "Table S10", "Internal FlowPilot module-ablation conditions and outcomes."),
        ("Generator model", "Table S11", "Matched one-shot and FlowPilot architecture scores by generator model."),
        ("Architecture", "Table S12", "Aggregate architecture-level score and critical-error summary."),
    ]
    for header, label, caption in captions:
        add_table_caption_before(
            table_with_header(document, header),
            table_caption_template,
            label,
            caption,
        )

    replacements = [
        (
            "FigSX. Ranked one-shot",
            "Figure S7",
            "Ranked one-shot and FlowPilot fixed-criteria multi-LLM evaluation scores for CuAAC, "
            "photochemical oxidation, and hydrogenolysis. Points are means and error bars are sample "
            "standard deviations across three generation repeats.",
        ),
        (
            "FigSX. Mean paired",
            "Figure S8",
            "Mean paired FlowPilot-minus-one-shot score difference for each applicable universal "
            "criterion, separated by chemistry. Positive values favor FlowPilot. UO-08 and UO-09 "
            "are excluded where applicability is not uniform.",
        ),
        (
            "FigSX. Critical-error",
            "Figure S9",
            "Critical-error criteria identified by the independent judge panel for each model, "
            "chemistry, architecture, and repeat. Cell numbers report independent judge agreement "
            "for the same campaign-criterion error; totals count distinct affected criteria.",
        ),
        (
            "FigSX. FlowPilot-only",
            "Figure S10",
            "FlowPilot-only token use, measured generation cost, observed runtime, and "
            "chemistry-specific quality-per-cost index. Resource panels report mean and sample "
            "standard deviation across nine campaigns per model. Runtime is environment-dependent, "
            "and judge-evaluation costs are excluded.",
        ),
    ]
    for startswith, label, text in replacements:
        set_caption(paragraph_with_text(document, startswith), label, text)

    purpose = paragraph_with_text(document, "The NewGen 2.0 benchmark")
    purpose.text = purpose.text.replace(
        "The NewGen 2.0 benchmark",
        "The fixed-criteria multi-LLM evaluation",
    )
    typo = paragraph_with_text(document, "The retained manuscript comparison included")
    typo.text = typo.text.replace("the sfollowing", "the following")


def add_engineering_sections(document: Document) -> None:
    references = paragraph_with_text(document, "REFERENCES")
    heading_template = paragraph_with_text(document, "FlowPilot workflow, data contracts")
    subheading_template = paragraph_with_text(document, "System architecture and decision authority")
    body_template = paragraph_with_text(document, "FlowPilot separates chemical interpretation")
    caption_template = paragraph_with_text(document, "Table S4.")

    insert_paragraph_before(
        references,
        "Deterministic engineering, inventory realization, and refinement",
        style="Heading 1",
        template=heading_template,
    )
    insert_paragraph_before(
        references,
        (
            "FlowPilot separates model-proposed independent variables from quantities that can be "
            "recomputed. Flow rates, gas compression, molar equivalents, reactor volumes, stage "
            "times, pressure-drop indicators, heat-transfer indicators, and equipment assignments "
            "are recalculated after candidate selection and again after inventory realization. A "
            "numerical value becomes a run instruction only if the recomputed proposal, stage table, "
            "stream table, inventory allocation, process graph, and final contract agree within the "
            "implemented tolerances."
        ),
        style="Normal",
        template=body_template,
    )

    insert_paragraph_before(
        references,
        "Residence-time and gas-flow conventions",
        style="Normal",
        template=subheading_template,
        bold_title=True,
    )
    insert_paragraph_before(
        references,
        (
            "Liquid molar flow is obtained from liquid concentration and volumetric flow. For a "
            "single liquid phase, nominal residence time closes as tau = V/Q. Gas feeds are stored "
            "as MFC setpoints at the declared inlet/STP reference. The corresponding in-channel "
            "volumetric flow is calculated by ideal-gas scaling, Qg,ch = Qg,STP(Tch/TSTP)"
            "(PSTP/Pabs). FlowPilot reports both the inlet/STP apparent time, V/(QL + Qg,STP), and "
            "the pressure-corrected in-channel apparent time, V/(QL + Qg,ch), when those definitions "
            "are applicable. Gas equivalents are calculated from gas molar flow and limiting-liquid "
            "molar flow; they are not inferred from a gas-to-liquid volumetric ratio. In packed beds "
            "or poorly characterized multiphase systems, contact time is not equated automatically "
            "with empty-bed space time: measured phase holdup or residence-time-distribution data are "
            "requested where required."
        ),
        style="Normal",
        template=body_template,
    )
    table_caption = insert_paragraph_before(
        references,
        "Table S13. Deterministic quantities, definitions, and closure checks.",
        style="Normal",
        template=caption_template,
        bold_prefix="Table S13. ",
    )
    table_caption.paragraph_format.keep_with_next = True
    insert_table_before(
        document,
        references,
        ["Quantity", "Implemented definition or basis", "Required closure"],
        [
            ["Limiting molar flow", "n_lim = C_lim QL", "Concentration, liquid flow, and stream composition agree"],
            ["Liquid residence time", "tau_L = V_L/QL", "Reactor liquid volume equals QL tau_L"],
            ["In-channel gas flow", "Qg,ch = Qg,STP(Tch/TSTP)(PSTP/Pabs)", "Temperature and absolute-pressure bases are explicit"],
            ["Gas equivalents", "n_gas/n_lim from standard gas molar volume", "Gas identity and gas fraction are included"],
            ["Gas holdup estimate", "epsilon_g = Qg,ch/(QL + Qg,ch)", "Used only with the declared two-phase approximation"],
            ["Multistage total", "tau_total = sum(tau_i)", "Each stage uses cumulative local inlet flow and assigned volume"],
            ["Pressure drop", "Single- or two-phase model selected by phase regime", "Calculated drop remains within pump, tubing, reactor, and BPR limits"],
            ["Thermal conductance", "UA = U A_wall; Da_th = Qgen/Qrem", "Assumptions, geometry, and thermal margin are reported"],
        ],
        widths=[1.3, 2.55, 2.4],
    )

    insert_paragraph_before(
        references,
        "Multistage calculation and stream introduction",
        style="Normal",
        template=subheading_template,
        bold_title=True,
    )
    insert_paragraph_before(
        references,
        (
            "Each reactive stage receives a separate stage identifier, inlet composition, cumulative "
            "liquid and gas flow, assigned reactor volume, temperature, pressure, and residence-time "
            "basis. A feed introduced between reactors contributes only to downstream cumulative "
            "flow. Quench and work-up operations are represented in topology but are not counted as "
            "reactive stages. If only a total target time is available, allocation uses declared "
            "stage times or stage batch-time weights and otherwise applies an explicit equal-allocation "
            "fallback. Final validation checks both the stage sum and each local volume-flow-time "
            "identity."
        ),
        style="Normal",
        template=body_template,
    )

    insert_paragraph_before(
        references,
        "Heat-transfer screening",
        style="Normal",
        template=subheading_template,
        bold_title=True,
    )
    insert_paragraph_before(
        references,
        (
            "The heat-transfer module reports wall area, surface-to-volume ratio, overall thermal "
            "conductance UA, estimated heat generation and removal rates, a thermal Damkohler-like "
            "ratio, and a screening margin. The implemented relations are A_wall = pi d L, "
            "Qrem = UA DeltaTlm, and Da_th = Qgen/Qrem. Reaction enthalpy, overall heat-transfer "
            "coefficient, and driving-temperature difference can be estimated from reaction and "
            "reactor classes when measurements are unavailable. These values are screening "
            "assumptions rather than calorimetric evidence; an apparently favorable estimate does "
            "not replace reaction calorimetry for scale-up or strongly exothermic chemistry."
        ),
        style="Normal",
        template=body_template,
    )

    insert_paragraph_before(
        references,
        "Inventory normalization and executable topology",
        style="Normal",
        template=subheading_template,
        bold_title=True,
    )
    insert_paragraph_before(
        references,
        (
            "Laboratory inventories may be entered as structured JSON or extracted from PDF, Word, "
            "PowerPoint, spreadsheet, and text sources. Extracted content is normalized into typed "
            "equipment classes and reviewed before use. Each item carries a stable equipment_id and "
            "available capability fields. Free-text statements such as 'no inline degasser' become "
            "explicit unavailable capabilities or hard constraints rather than prompt-only notes. "
            "During design realization, every required operation is assigned to an inventory item. "
            "The allocator checks type, material, inner diameter, volume, flow range, temperature, "
            "pressure, wavelength, phase compatibility, and connectivity where declared. Under strict "
            "assignment, absent or invented hardware cannot be published in the executable topology."
        ),
        style="Normal",
        template=body_template,
    )
    inv_caption = insert_paragraph_before(
        references,
        "Table S14. Inventory classes and deterministically checked capabilities.",
        style="Normal",
        template=caption_template,
        bold_prefix="Table S14. ",
    )
    inv_caption.paragraph_format.keep_with_next = True
    insert_table_before(
        document,
        references,
        ["Inventory class", "Representative capability fields", "Design role"],
        [
            ["Pumps and MFCs", "Minimum/maximum flow, pressure, compatible fluids or gas identity", "Quantified feed delivery"],
            ["Mixers/connectors", "Inputs, material, pressure, supported ID, connectivity", "Stream joining and interstage connection"],
            ["Reactors/trains", "Type, volume, ID, material, temperature, pressure, light compatibility", "Reactive residence volume"],
            ["Light/temperature control", "Wavelength, power, allowed temperature setpoints", "Photochemical and thermal control"],
            ["BPRs/tubing", "Setpoints or range, pressure rating, ID, length, material", "Pressure and hydraulic realization"],
            ["Separators/collectors", "Supported phases, flow, pressure, volume, venting", "Depressurization and product collection"],
            ["Safety accessories", "Hazard compatibility and control capability", "Purge, check valve, shielding, containment, and venting"],
        ],
        widths=[1.35, 3.15, 1.75],
    )
    insert_paragraph_before(
        references,
        (
            "A process diagram is generated from the same inventory-assigned topology stored in the "
            "final contract. If assignment or numerical closure fails, FlowPilot may render a "
            "requirements topology for diagnosis, but it does not label that diagram executable and "
            "does not publish numerical run instructions. This distinction prevents a polished "
            "intermediate graph from being mistaken for a laboratory-ready process."
        ),
        style="Normal",
        template=body_template,
    )

    insert_paragraph_before(
        references,
        "Incremental experimental refinement",
        style="Normal",
        template=subheading_template,
        bold_title=True,
    )
    insert_paragraph_before(
        references,
        (
            "Closed-loop refinement accepts one or more ExperimentRecord objects containing actual "
            "operating conditions, measured response, and operational observations. Measured evidence "
            "has higher authority than model inference. The module diagnoses underconversion, "
            "selectivity loss, solubility or fouling, pressure instability, and gas-liquid instability; "
            "it then proposes bounded changes to residence time, liquid and gas flow, temperature, "
            "concentration, pressure, tubing, filtration, or safety controls. For multiple experiments, "
            "the campaign preserves every observation, identifies the best observed anchor, and may "
            "fit a simple apparent response-versus-residence-time model. Extrapolated targets are "
            "reported as screening estimates with acceptance criteria, not guaranteed yields. Every "
            "cycle creates a new version while retaining the preceding design and measurements."
        ),
        style="Normal",
        template=body_template,
    )
    exp_caption = insert_paragraph_before(
        references,
        "Table S15. Experimental feedback fields and their role in refinement.",
        style="Normal",
        template=caption_template,
        bold_prefix="Table S15. ",
    )
    exp_caption.paragraph_format.keep_with_next = True
    insert_table_before(
        document,
        references,
        ["Evidence group", "Recorded fields", "Refinement use"],
        [
            ["Timing and geometry", "Residence-time basis, inlet and channel time, flow, volume, tubing ID", "Reconcile actual exposure and throughput"],
            ["Gas delivery", "Gas identity, MFC/STP flow, channel flow, equivalents, pressure", "Recompute gas supply and apparent time bases"],
            ["Chemistry response", "Yield, conversion, selectivity, product and starting-material fractions", "Anchor response trend and target acceptance"],
            ["Operating state", "Temperature, concentration, wavelength, light power", "Separate kinetic, photonic, and concentration changes"],
            ["Run observations", "Pressure drift, clogging, precipitation, phase stability, impurity notes", "Trigger feasibility and safety interventions"],
        ],
        widths=[1.35, 3.2, 1.7],
    )

    insert_paragraph_before(
        references,
        "Computational provenance and interpretation boundaries",
        style="Heading 1",
        template=heading_template,
    )
    insert_paragraph_before(
        references,
        "Model routing and run records",
        style="Normal",
        template=subheading_template,
        bold_title=True,
    )
    insert_paragraph_before(
        references,
        (
            "The user interface exposes separate model routes for upstream chemistry interpretation "
            "and downstream proposal/council execution. The selected provider, model identifier, and "
            "local OpenAI-compatible endpoint mapping are applied for the duration of one serialized "
            "run and restored afterward. Model routing is written to pipeline-runtime provenance. "
            "Because hosted model aliases and local checkpoints can change, a reproducible report "
            "should retain the exact model identifier, endpoint or provider, generation parameters, "
            "software revision, dependency environment, frozen DesignInputPackage, inventory profile, "
            "raw model events, deterministic audit, and canonical final result."
        ),
        style="Normal",
        template=body_template,
    )
    repro_caption = insert_paragraph_before(
        references,
        "Table S16. Minimum records required for audit and computational replay.",
        style="Normal",
        template=caption_template,
        bold_prefix="Table S16. ",
    )
    repro_caption.paragraph_format.keep_with_next = True
    insert_table_before(
        document,
        references,
        ["Record", "Purpose", "Typical stored object"],
        [
            ["Frozen input", "Reconstruct chemist-provided facts and authority", "DesignInputPackage and inventory profile"],
            ["Model configuration", "Identify stochastic generators and routes", "Provider, model ID, endpoint, temperature, seed where supported"],
            ["Prompts/events", "Audit module inputs and raw responses", "Timestamped LLM event log and token telemetry"],
            ["Deterministic state", "Verify calculations and gates", "Calculation sheet, validation report, allocation, topology"],
            ["Canonical output", "Identify the published design", "FinalDesignContract, content hash, rendered artifacts"],
            ["Software environment", "Support replay of code-owned behavior", "Git revision, Python/dependency versions, configuration"],
        ],
        widths=[1.4, 2.35, 2.5],
    )

    insert_paragraph_before(
        references,
        "Statistical and evidential limitations",
        style="Normal",
        template=subheading_template,
        bold_title=True,
    )
    insert_paragraph_before(
        references,
        (
            "The architecture comparison uses three generation repeats for each retained "
            "model-case-architecture cell. Error bars therefore describe observed generation-repeat "
            "variation for this frozen benchmark, not population confidence intervals. LLM judges "
            "apply a fixed rubric and are identity-masked, but they are not substitutes for blinded "
            "flow-chemist assessment. Deterministic verification reduces arithmetic ambiguity but "
            "does not validate unknown kinetics, selectivity, catalyst lifetime, phase holdup, or "
            "scale-dependent transport. Benchmark scores measure the completeness and defensibility "
            "of proposed screening designs. Experimental execution and measured chemical performance "
            "remain the final validation."
        ),
        style="Normal",
        template=body_template,
    )

    note = insert_paragraph_after(
        references,
        "No references are cited exclusively in the Supporting Information; literature citations are provided in the main article.",
        style="Normal",
        template=body_template,
    )
    note.paragraph_format.space_before = Pt(6)


def request_field_update(document: Document) -> None:
    settings = document.settings.element
    element = settings.find(qn("w:updateFields"))
    if element is None:
        element = OxmlElement("w:updateFields")
        settings.append(element)
    element.set(qn("w:val"), "true")


def main() -> None:
    if not DOCUMENT.is_file():
        raise SystemExit(f"Missing document: {DOCUMENT}")
    if not BACKUP.exists():
        shutil.copy2(DOCUMENT, BACKUP)

    document = Document(BACKUP if BACKUP.exists() else DOCUMENT)

    with tempfile.TemporaryDirectory(prefix="flowpilot-esi-") as temporary:
        directory = Path(temporary)
        images = {
            4: retitle_figure(8, 4, directory),
            5: retitle_figure(9, 5, directory),
            6: retitle_figure(10, 6, directory),
        }
        add_missing_knowledge_material(document, images)
        complete_benchmark_labels(document)
        add_engineering_sections(document)
        request_field_update(document)
        output = directory / DOCUMENT.name
        document.save(output)
        shutil.copy2(output, DOCUMENT)

    print(f"Updated: {DOCUMENT}")
    print(f"Backup:  {BACKUP}")


if __name__ == "__main__":
    main()
