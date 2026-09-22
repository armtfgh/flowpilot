"""Read back saved PowerPoint tables and check them against canonical results."""
import argparse
import json
from pathlib import Path
import shutil
import sys
from PIL import Image, ImageStat
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from export_figure5_oxygen_slides import fmt, verify, alias


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('campaign', type=Path)
    args = parser.parse_args()
    out = args.campaign / 'slides'
    deck = Presentation(out / 'Figure5_three_sets_pure_oxygen.pptx')
    assert len(deck.slides) == 6
    audit = []
    for number in range(1, 4):
        name = f'figure5_set{number}'
        result, checked = verify(args.campaign / name / 'attempt_01')
        dest = out / name
        original = json.loads((dest / 'process_topology.json').read_text())
        display = json.loads((dest / 'display_topology.json').read_text())
        assert original == result['process_topology']
        for canonical, shown in zip(original['unit_operations'], display['unit_operations']):
            canonical, shown = json.loads(json.dumps(canonical)), json.loads(json.dumps(shown))
            if canonical['op_type'] == 'pump':
                canonical['parameters'].pop('contents', None)
                shown['parameters'].pop('contents', None)
            assert canonical == shown, 'Display changed a condition or equipment assignment'
        with Image.open(dest / 'topology.png') as image:
            assert image.width > 1000 and image.height > 500
            assert min(ImageStat.Stat(image.convert('RGB')).stddev) > 10
        tables = [shape.table for shape in deck.slides[number * 2 - 1].shapes if shape.has_table]
        assert len(tables) == 2
        stage_table, feed_table = tables
        rows = checked['stage_rows']
        for col, stage in enumerate(rows, 1):
            expected = {
                1: f'{fmt(stage["volume_mL"])} mL {stage["material"]}',
                2: fmt(stage['ID_mm']), 3: stage['module'], 4: fmt(stage['wavelength_nm']),
                5: fmt(stage['temperature_C']), 6: fmt(stage['liquid_flow_mL_min']),
                7: fmt(stage['gas_inlet_STP_mL_min']),
                8: f'{stage["nominal_inlet_residence_min"]:.2f}', 9: fmt(stage['BPR_bar']),
            }
            for row, value in expected.items():
                assert stage_table.cell(row, col).text == value
        for row, feed in enumerate(checked['component_rows'], 1):
            expected = [feed['stream'], alias(feed), fmt(feed['concentration_M']),
                        fmt(feed['flow_mL_min']), fmt(feed['molar_flow_mmol_min']), fmt(feed['equivalents'])]
            assert [feed_table.cell(row, col).text for col in range(6)] == expected
        audit.append(dict(case=name, independent_checks=len(checked['checks']),
                          final_data_preserved=True, display_changes='feed aliases only',
                          pptx_table_readback=True, topology_nonblank=True))
    render = json.loads((out / 'visual_review/render_audit.json').read_text())
    assert render['slides'] == 6 and not render['out_of_page_text']
    (out / 'export_verification.json').write_text(json.dumps(dict(
        slides=6, cases=audit, pdf_text_within_bounds=True,
        scope='Arithmetic, equipment and table readback checks are not laboratory validation.'), indent=2))
    shutil.copy2(Path(__file__), out / 'export_source' / Path(__file__).name)
    print(json.dumps(audit, indent=2))


if __name__ == '__main__':
    main()
