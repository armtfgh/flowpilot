"""Read-only source-based release review; never rewrite archived results."""
import csv
import hashlib
import json
import math
from collections import Counter
from datetime import datetime
from pathlib import Path
import re
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
PACKAGE=ROOT/'outputs/khu_revised_six_20260915/presentation'
WORKBOOK=ROOT/'inventory_khu/Inventory_final.xlsx'
NS={'s':'http://www.w3.org/2000/svg'}


def cell(book, sheet, address):
    tab=book[sheet]
    for merged in tab.merged_cells.ranges:
        if address in merged:
            return tab.cell(merged.min_row,merged.min_col).value
    return tab[address].value


def near(a,b,tolerance=4e-5):
    return a is not None and b is not None and math.isclose(float(a),float(b),rel_tol=tolerance,abs_tol=1e-8)


def table(path):
    with path.open(newline='') as f:
        return list(csv.DictReader(f))


def main():
    out=PACKAGE.parent/('pre_send_review_'+datetime.now().strftime('%Y%m%d_%H%M%S'))
    out.mkdir(exist_ok=False)
    workbook=load_workbook(WORKBOOK,data_only=True)
    hashes={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest()
            for p in [WORKBOOK,*PACKAGE.rglob('*')] if p.is_file()}
    checks=[];stages=[];feeds=[];findings=[];corrections=[];warnings=[]
    def add(case,criterion,passed,detail=''):
        checks.append(dict(case=case,criterion=criterion,passed=bool(passed),detail=detail))
    molar_volume=.08314*273.15/1.01325
    for case in json.loads((PACKAGE/'selected_results.json').read_text()):
        name=case['case'];folder=PACKAGE/name
        r=json.loads((folder/'result.json').read_text());p=r['proposal']
        inv=json.loads((folder/'inventory_profile.json').read_text())['lab_inventory']
        catalogs={cat:{x['equipment_id']:x for x in inv[cat]} for cat in
            ['pumps','reactors','tubing','light_sources','mixers','gas_hardware','pressure_controllers','safety_accessories']}
        bpr=p['BPR_bar'];ss=p['streams'];ops=r['process_topology']['unit_operations']
        opmap={o['op_id']:o for o in ops}
        svg=ET.parse(folder/'topology.svg')
        nodetext={n.find('s:title',NS).text:' '.join(t.text or '' for t in n.findall('s:text',NS))
                  for n in svg.findall('.//s:g[@class="node"]',NS)}
        csv_stages=table(folder/'stage_summary.csv');csv_feeds=table(folder/'component_feeds.csv')
        for stage in p['stage_parameters']:
            n=stage['stage_number'];key=f'stage{n}'
            q=sum(s['flow_rate_mL_min'] or 0 for s in ss if s['phase']!='gas' and s['introduction_stage']<=n)
            g=sum(s['gas_flow_sccm'] or 0 for s in ss if s['phase']=='gas' and s['introduction_stage']<=n)
            volume=stage['reactor_volume_mL'];tau=volume/(q+g)
            add(name,key+'_V_over_Q',near(tau,stage['residence_time_min']))
            add(name,key+'_cumulative_liquid',near(q,stage['Q_liquid_mL_min']))
            final=next(s for s in r['final_design']['stages'] if s['stage_number']==n)
            add(name,key+'_final_contract',all(near(final[k],stage[k]) for k in ['reactor_volume_mL','residence_time_min','temperature_C','Q_liquid_mL_min']))
            diagram=opmap[f'st{n}_reactor']['parameters']
            add(name,key+'_diagram_parameters',all(near(diagram[a],stage[b]) for a,b in
                [('volume_mL','reactor_volume_mL'),('temperature_C','temperature_C'),('residence_time_inlet_min','residence_time_min'),('ID_mm','d_mm')]))
            saved=next(s for s in csv_stages if int(s['stage'])==n)
            add(name,key+'_csv',all(near(saved[a],stage[b]) for a,b in
                [('volume_mL','reactor_volume_mL'),('nominal_inlet_residence_min','residence_time_min'),('temperature_C','temperature_C')]))
            text=nodetext[f'st{n}_reactor']
            add(name,key+'_svg_values',all(near(re.search(pattern,text)[1],value) for pattern,value in
                [(r'Time \([^)]*\): ([\d.]+) min',tau),(r'V=([\d.]+) mL',volume),(r'ID=([\d.]+) mm',stage['d_mm'])]))
            row=int(stage['reactor_equipment_id'].removeprefix('coil_row'))
            add(name,key+'_workbook_geometry',near(cell(workbook,'Reactor (tubing)',f'D{row}'),volume)
                and near(cell(workbook,'Reactor (tubing)',f'E{row}'),stage['d_mm'])
                and cell(workbook,'Reactor (tubing)',f'C{row}')==stage['material'])
            add(name,key+'_workbook_limits',stage['temperature_C']<=cell(workbook,'Reactor (tubing)',f'I{row}') and bpr<cell(workbook,'Reactor (tubing)',f'J{row}'))
            if stage.get('light_equipment_id'):
                light=catalogs['light_sources'][stage['light_equipment_id']]
                lr=int(re.search(r'row(\d+)',light['equipment_id'])[1])
                declared=cell(workbook,'Reactor (tubing)',f'K{row}')
                add(name,key+'_module',('UV-150' if light['module_id']=='uv150' else 'manual photoreactor') in declared
                    and volume<=cell(workbook,'Photoreactor module',f'J{lr}'))
                add(name,key+'_module_temperature',cell(workbook,'Photoreactor module',f'K{lr}')<=stage['temperature_C']<=cell(workbook,'Photoreactor module',f'L{lr}'))
                add(name,key+'_platform',light['module_id']!='uv150' or any(
                    catalogs['pumps'].get(s['pump_equipment_id'],{}).get('platform_id') in ['e_series','r_series'] for s in ss if s['phase']!='gas'))
            stages.append(dict(case=name,stage=n,volume_mL=volume,liquid_mL_min=q,gas_STP_mL_min=g,
                calculated_time_min=tau,reported_time_min=stage['residence_time_min'],temperature_C=stage['temperature_C'],BPR_gauge_bar=bpr))
        used=Counter()
        for assignment in r['inventory_allocation']['assignments']:
            for eid in assignment['equipment_item_ids']:
                used[(assignment['category'],eid)]+=1
        add(name,'physical_quantities',all(count<=catalogs[cat][eid]['quantity'] for (cat,eid),count in used.items()))
        add(name,'single_pump_platform',len({catalogs['pumps'][s['pump_equipment_id']]['platform_id'] for s in ss if s['phase']!='gas'})==1)
        for s in ss:
            if s['phase']=='gas':
                continue
            pump=catalogs['pumps'][s['pump_equipment_id']];q=s['flow_rate_mL_min']
            # These six saved runs all select the E-series: compare to source, not JSON defaults.
            add(name,'pump_'+s['stream_label']+'_source_settings',pump['platform_id']=='e_series'
                and cell(workbook,'Pump','E3')<=q<=cell(workbook,'Pump','F3')
                and near(q/cell(workbook,'Pump','G3'),round(q/cell(workbook,'Pump','G3'))))
            quantities=[c for c in csv_feeds if c['stream']==s['stream_label'] and c['phase']!='gas']
            add(name,'feed_'+s['stream_label']+'_quantified',bool(quantities) and all(c['concentration_M'] and c['molar_flow_mmol_min'] for c in quantities))
            for c in quantities:
                conc=float(c['concentration_M']);molar=q*conc
                add(name,'molar_flow_'+c['component'],near(molar,c['molar_flow_mmol_min']) and near(q,c['flow_mL_min']))
                feeds.append(dict(case=name,stream=s['stream_label'],component=c['component'],concentration_M=conc,flow_mL_min=q,molar_flow_mmol_min=molar))
        reference=next(c for c in csv_feeds if 'trimethylsilane' in c['component'].lower() or 'nitrobenzoic acid' in c['component'].lower())
        ref=float(reference['molar_flow_mmol_min'])
        ratios={'acrylonitrile':2,'ir(': .005} if name.startswith('figure5') else {'dpdtc':1.05,'dmap':.1,'benzylamine':1.05}
        for word,expected in ratios.items():
            c=next(c for c in csv_feeds if word in c['component'].lower())
            add(name,'stoichiometry_'+word,near(float(c['molar_flow_mmol_min'])/ref,expected))
        for gas in (s for s in ss if s['phase']=='gas'):
            q=gas['gas_flow_sccm'];fraction=gas['gas_reagent_mole_fraction'];eq=q/molar_volume*fraction/ref
            add(name,'oxygen_equivalents_from_STP',near(eq,2) and gas['introduction_stage']==2)
            step=cell(workbook,'MFC','D3')
            add(name,'MFC_source_increment',near(q/step,round(q/step)),f'MFC!D3={step}; saved={q}')
            rounded=math.ceil((q-1e-12)/step)*step
            st=p['stage_parameters'][1]
            corrections.append(dict(case=name,source_cell='MFC!D3',increment_mL_min=step,
                saved_gas_STP_mL_min=q,proposed_gas_STP_mL_min=rounded,
                proposed_O2_equiv=rounded/molar_volume*fraction/ref,
                saved_stage2_time_min=st['residence_time_min'],proposed_stage2_time_min=st['reactor_volume_mL']/(st['Q_liquid_mL_min']+rounded)))
            edges={(s['from_op'],s['to_op']) for s in r['process_topology']['streams']}
            cv=next(o for o in ops if o['op_type']=='check_valve')
            mfc=next(o for o in ops if o['op_type']=='mfc')
            add(name,'check_valve_position',cv['inventory_item_id']=='cv-3301' and (mfc['op_id'],cv['op_id']) in edges and (cv['op_id'],'st2_mixer') in edges)
        for eng in r['final_stage_engineering']['stages']:
            for step in eng['calculations']['steps']:
                v=step['values']
                if step['name']=='Back-Pressure Regulator':
                    add(name,f'engineering_stage{eng["stage_number"]}_selected_BPR',near(v['bpr_bar'],bpr),f'engineering={v["bpr_bar"]}, selected={bpr}; minimum-requirement calculation is not the applied setting')
                if step['status']=='WARNING':
                    warnings.append(dict(case=name,stage=eng['stage_number'],name=step['name'],summary=step['summary'],values=v))
        findings.append(dict(case=name,disposition='NOT CLEARED AS FINAL EXECUTION PACKAGE',
            failures=[c for c in checks if c['case']==name and not c['passed']]))
    unchanged=all(hashlib.sha256((ROOT/name).read_bytes()).hexdigest()==digest for name,digest in hashes.items())
    assert unchanged,'Review must not change input artifacts'
    for filename,rows in [('checks',checks),('verified_stage_arithmetic',stages),('verified_feed_concentrations',feeds),('proposed_gas_setting_corrections',corrections)]:
        with (out/(filename+'.csv')).open('w',newline='') as f:
            writer=csv.DictWriter(f,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)
    (out/'audit.json').write_text(json.dumps(dict(checks=checks,findings=findings,engineering_warnings=warnings,
        proposed_corrections=corrections,source_files_unchanged=unchanged,source_hashes=hashes),indent=2))
    lines=['# Pre-send review of the six KHU designs','',
        '**Decision: do not describe the existing package as fully resolved or ready for execution.** The stage arithmetic, installed reactor geometry, equipment counts, pump settings and feed stoichiometry agree, but this deeper review found gaps not covered by the previous 174 checks.','',
        '## Required corrections','',
        '1. MFC!D3 specifies a 0.01 mL/min STP setting increment. The inventory conversion omitted it, so all three Figure 5 gas setpoints are off-grid. The earlier statement that MFC resolution was unspecified was incorrect. Round upward to retain at least 2 equivalents, then revalidate and re-export the full final design. The table below is a proposed correction, not an already changed design.',
        '2. Detailed final-stage engineering JSON contains BPR=0 and gas_pressure_abs_bar=1.01325 in stage-local BPR calculations despite a selected shared 2.5 or 7 bar BPR. Its pressure margins are based on friction alone and should not be presented as whole-system headroom. Retain any minimum-pressure requirement as a separately labeled quantity; recompute applied pressure/headroom from the connected process.',
        '3. All Figure 6 stages retain axial-dispersion warnings (model Pe about 28.6 to 71.4). Their V/Q arithmetic is correct, but ideal plug flow and high conversion are not established. These warnings should be visible in the collaborator summary.', '',
        '| Figure 5 set | Saved air flow | Proposed air flow | Proposed O2 equiv | Proposed Stage 2 nominal time |',
        '|---|---:|---:|---:|---:|']
    lines += [f'| {c["case"]} | {c["saved_gas_STP_mL_min"]:.6f} | {c["proposed_gas_STP_mL_min"]:.2f} | {c["proposed_O2_equiv"]:.5f} | {c["proposed_stage2_time_min"]:.5f} |' for c in corrections]
    lines += ['', 'Gas flows are mL/min at the declared STP reference; times are minutes. The source supplies an adjustment increment, not an accuracy specification. No correction has been silently applied to the archived designs.', '',
        '## What checks out','',
        '- Each reactor volume and ID matches an installed coil after resolving merged workbook cells. No custom-fabricated volume is used.',
        '- Figure 5: Stage 1 uses the 2 mL PFA coil installed in UV-150 with its own 450 nm light and the E-series platform. Stage 2 uses Manual 2 with its own 448 nm light. The 20 mL coils are not assigned to UV-150.',
        '- Figure 6: two ETFE coils use thermal heating at 95 C, below the declared 100 C tubing limit. Set 1 uses two of the two available 5 mL coils. Sets 2/3 use one 10 mL and one 5 mL coil. No photoreactor-module 80 C limit applies to these thermal stages.',
        '- All liquid feeds use E-series BLUE pump tubing. Figure 6 uses two distinct channels of the same three-channel system, not a single channel supplying two rates. Rates meet the workbook minimum and 0.01 mL/min increment.',
        '- Gas enters Stage 2 only. The assigned CV-3301 sits between the MFC and Stage 2 mixer.',
        '- Final proposal, final-stage contract, topology operation values, SVG values and stage CSV agree on volume, flow-derived time and temperature within display precision.',
        '- The legacy topology residence_time_min field is rounded to two decimals. The renderer uses residence_time_inlet_min, which agrees with the precise final-stage time. This precision difference is not a different design.',
        '- Feed tables and preparation slides contain reagent concentrations. Molar flows recomputed from C times Q agree with the CSVs. Acrylonitrile is 2 equiv and photocatalyst 0.5 mol%; DPDTC and benzylamine are each 1.05 equiv and DMAP is 0.1 equiv.',
        '- For Figure 6, Feed A contains acid 0.5 M, DPDTC 0.525 M and DMAP 0.05 M; Feed B is benzylamine 2.1 M. Stage 2 cumulative liquid flow includes both feeds. The hypothetical intermediate concentration after mixing is not measured conversion.', '',
        '## Laboratory confirmations remain','',
        '- CV-3301 maximum working pressure is absent; 1 bar is the cracking differential, not its maximum rating.',
        '- Confirm pressure reference conventions, regulated pressurized air supply, MFC gas calibration and actual unit count. The specified inlet/outlet pressure limits do not by themselves establish operating differential pressure across the MFC and check valve.',
        '- Confirm buffer identity/strength, full-feed solubility, premix stability, pump-tubing compatibility and the heated/pressurized installation. Generic reagent-family compatibility is not a complete-mixture certification.',
        '- Demonstrate effective oxygen exclusion in Stage 1 and stage temperature under irradiation. Offline degassing instructions do not constitute a measured oxygen-free condition.',
        '- The six response sets produce four distinct numerical conditions: Figure 5 sets 1/3 match; Figure 6 sets 2/3 match. Do not present six unique experimental conditions.',
        '- Custom tubing stock is not explored in these six runs; these are not proven optima over the entire updated inventory.', '',
        '## Scope','',
        'No provider calls or new designs were generated. This review uses the workbook, saved final JSON, source-derived inventory, topology SVG and component/stage CSVs. It records existing engineering warnings, not a new kinetic or safety certification. Nominal gas-stage V/(Q_liquid+Q_gas,STP) is not actual hydrodynamic residence time. Original files and sealed archives were not modified.', '',
        f'Automated review: {sum(c["passed"] for c in checks)}/{len(checks)} checks passed. See checks.csv for the uncovered failures, and audit.json for source hashes.']
    (out/'REVIEW.md').write_text('\n'.join(lines)+'\n')
    print(out)
    print(f'{sum(c["passed"] for c in checks)}/{len(checks)} passed; original files unchanged={unchanged}')
    print(json.dumps(corrections,indent=2))


if __name__=='__main__':main()
