"""Create a readable Word copy of the second-pass audit report."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'manuscript/october1_postmeeting_round2_work'


def main():
    doc=Document()
    for style in ['Normal','Title','Heading 1','Heading 2','List Bullet','List Number']:
        s=doc.styles[style]
        s.font.name='Times New Roman'
        s.font.size=Pt(11)
        s.font.color.rgb=RGBColor(0,0,0)
        fonts=s.element.rPr.rFonts
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            fonts.attrib.pop(qn('w:'+attr),None)
    for sec in doc.sections:
        sec.top_margin=sec.bottom_margin=Inches(.7)
        sec.left_margin=sec.right_margin=Inches(.8)
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing=1.05
    for line in (OUT/'REVIEW_REPORT.md').read_text().splitlines():
        if not line:continue
        if line.startswith('# '):p=doc.add_paragraph(line[2:],'Title')
        elif line.startswith('## '):p=doc.add_paragraph(line[3:],'Heading 1')
        else:
            style='Normal'
            if line.startswith('- '):style='List Bullet';line=line[2:]
            elif re.match(r'^\d+\. ',line):style='List Number';line=re.sub(r'^\d+\. ','',line)
            p=doc.add_paragraph(style=style)
            for part in re.split(r'(\*\*.*?\*\*)',line.replace('`','')):
                r=p.add_run(part[2:-2] if part.startswith('**') else part)
                if part.startswith('**'):r.bold=True
        for r in p.runs:
            r.font.name='Times New Roman'
            r.font.size=Pt(11)
    doc.save(OUT/'FlowPilot_round2_review_report.docx')
    print(OUT/'FlowPilot_round2_review_report.docx')


if __name__=='__main__':main()
