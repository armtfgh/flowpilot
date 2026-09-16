#!/usr/bin/env python3
"""Build the manuscript-ready ESI corpus, rule-base, and retrieval section."""

from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.shared import Pt

from build_esi_workflow_section_docx import (
    BLUE,
    MID_GREY,
    add_caption,
    add_paragraph,
    add_prompt_box,
    add_table,
    configure_document,
    set_font,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = (
    ROOT
    / "deliverables"
    / "flowpilot_esi_figures_s1_s10_20260811"
    / "source_data"
)
RULES_PATH = ROOT / "flora_fundamentals" / "data" / "rules.json"
OUT_DIR = ROOT / "deliverables" / "manuscript_esi_sections_20260828"
OUT_PATH = OUT_DIR / "S2_Literature_Corpus_Rule_Base_and_Retrieval.docx"


FIGURE_S3_PROMPT = """Create a publication-ready, six-panel scientific data figure on a pure white background. Use a wide landscape layout, flat two-dimensional styling, large readable labels, restrained but distinct colors, and consistent panel spacing. Do not use gradients, shadows, decorative illustrations, or fabricated data. Build every plotted value from the supplied frozen source tables.

Title: "Figure S3 | Composition of the frozen flow-chemistry literature corpus"

Panel A - Reaction class, n = 464 classified records:
Use a sorted horizontal bar chart. Clearly show at least the leading categories: Thermal synthesis 96 (20.7%), Photoredox catalysis 85 (18.3%), Other 66 (14.2%), Heterogeneous photocatalysis 35 (7.5%), and Cross-coupling 32 (6.9%). Retain the remaining categories as smaller bars rather than silently removing them.

Panel B - Reactor type, n = 464:
Use a sorted horizontal bar chart. Highlight Capillary/coil reactor 122 (26.3%), Microreactor/chip 121 (26.1%), Continuous-flow reactor 71 (15.3%), Photoreactor 47 (10.1%), and Packed-bed reactor 27 (5.8%). Keep the long-tail classes visible in a grouped "remaining classified types" bar only if the aggregation is stated explicitly.

Panel C - Reactor material, n = 464 classified entries:
Use a sorted bar chart or compact treemap. Show Fluoropolymer (PTFE/PFA/FEP) 272 (58.6%), Glass/quartz 61 (13.1%), Stainless steel 58 (12.5%), and PDMS/silicon 28 (6.0%). Do not imply that every publication supplied equally detailed material information.

Panel D - Bond or transformation label, n = 464:
Use a sorted horizontal bar chart. Show C-C 123 (26.5%), Other/multiple 84 (18.1%), C-N 70 (15.1%), C-O 62 (13.4%), and C-S 22 (4.7%), followed by the remaining labels.

Panel E - Number of inlet streams, n = 154 records with explicit annotations:
Use discrete bars for one through higher stream counts. Show one stream 22 (14.3%), two streams 91 (59.1%), three streams 30 (19.5%), four streams 5 (3.2%), and five streams 3 (1.9%). Place "n = 154 annotated processes" directly in the panel so the denominator cannot be confused with 464.

Panel F - Reported yield distributions:
Use side-by-side violin plus box plots, or transparent histograms with box-plot summaries, for available batch yields (n = 135; median 80%; interquartile range 56-93%) and optimized-flow yields (n = 285; median 89%; interquartile range 75-96%). Label this panel "descriptive available-value distributions" and add a note that unequal metadata availability and study selection prevent a paired causal comparison.

Use counts on bars and percentages where space permits. Retain unknown, other, and long-tail categories in the denominator. Use a consistent sans-serif font and ensure all text remains readable at approximately 17 cm figure width. Add a small footer: "Frozen manuscript export; field-specific denominators shown in each panel.""" 


FIGURE_S4_PROMPT = """Create a publication-ready, four-panel scientific figure on a pure white background that summarizes a structured engineering rule base containing exactly 2,537 rules from six source documents. Use a wide landscape layout, flat vector styling, large labels, and a restrained palette. Do not use gradients, shadows, three-dimensional effects, decorative icons, or invented validation claims.

Title: "Figure S4 | Engineering rule-base coverage and severity structure"

Panel A - Rule category and severity:
Create sorted stacked horizontal bars for the 14 largest categories. Stack each category by severity using consistent colors: hard_rule, guideline, tip, and safety. Use these totals: Reactor design 403; General 347; Heat transfer 227; Catalyst 226; Residence time 204; Photochemistry 175; Materials 165; Scale-up 157; Pressure 134; Mixing 133; Safety 122; Solvent 72; Mass transfer 59; Temperature 38. Place each category total at the bar end. State that the complete store contains 29 categories.

Panel B - Global severity composition:
Show four horizontal bars or a compact 100% stacked bar with exact counts: guideline 1,772; hard_rule 717; tip 47; safety severity 1. Add a note: "Rule category and severity are separate labels; most records in the Safety category are classified as hard_rule rather than safety severity."

Panel C - Chemistry-class coverage:
Create a heatmap of rule-category association counts across representative chemistry classes: Photoredox catalysis, Heterogeneous photocatalysis, Photocycloaddition, Thermal synthesis, Cross-coupling, Hydrogenation, Oxidation/reduction, Electrochemistry, Biocatalysis, Polymer synthesis, Organocatalysis, and Precipitation/crystallization. Use a perceptually uniform light-to-dark scale and print values only where readable. State that associations are non-exclusive and can exceed the number of unique rules.

Panel D - Provenance and quantitative-content indicator:
Show six source-document bars with extracted-rule counts: Flow Chemistry Volume 1 - Fundamentals, 1,077; Flow Chemistry: Integrated Approaches for Practical Applications, 952; The Hitchhiker's Guide to Flow Chemistry, 238; Microreactors in Organic Synthesis and Catalysis, 184; Principles of Flow Chemistry, 61; d3sc00992k, 25. Beside the bars, show a compact indicator of the fraction of rules containing a machine-detected quantitative expression by category. Label it explicitly: "expression detected, not independently equation-validated."

At the bottom, add a methodological note: "Rules are retrieved as handbook-grounded prompt context. A hard_rule label does not by itself constitute an executable deterministic gate unless the corresponding check is implemented in code." Use a consistent sans-serif typeface and keep the figure readable at approximately 17 cm width."""


FIGURE_S5_PROMPT = """Create a publication-ready, three-panel scientific algorithm diagram on a pure white background. Use a wide landscape layout, flat vector styling, precise arrows, restrained colors, and large readable labels. Do not use gradients, shadows, three-dimensional effects, decorative illustrations, or fabricated retrieval results.

Title: "Figure S5 | Plan-aware literature retrieval and tier transitions"

Panel A - Query construction:
Show two validated inputs, BatchRecord and ChemistryPlan, merging into a plan-aware query. From BatchRecord include solvent, temperature, wavelength, concentration, and photocatalyst when available. From ChemistryPlan include reaction name, reaction class, mechanism type, catalyst or sensitizer names, key intermediate, bond formed, up to six retrieval keywords, up to four similar reaction classes, recommended wavelength, and oxygen-sensitivity context. Show two execution modes after query construction: semantic embedding search and deterministic lexical search. Add a dashed fallback arrow from embedding failure to lexical search.

Panel B - Three retrieval tiers:
Draw a clear top-to-bottom state machine. Tier 1: search paired batch-to-flow records with mechanism and, when available, phase filters; retrieve up to 20 candidates. If fewer than three hits are returned, move to Tier 2. Tier 2: search paired records with filters relaxed; retrieve up to 20 candidates. If zero hits are returned, move to Tier 3. Tier 3: search all records without metadata filters. If no records are returned, output an empty analogy list. Use amber transition labels for "<3 hits" and "0 hits." Do not imply that phase filtering is always populated.

Panel C - Field-aware reranking and leakage boundary:
Show the equation: final score = 0.60 x semantic score + 0.40 x field score. Below it show: field score = 0.30 x photocatalyst similarity + 0.20 x solvent match + 0.20 x wavelength proximity + 0.15 x temperature proximity + 0.15 x concentration proximity. Add the operational definitions: photocatalyst = 1.0 same class, 0.5 same metal family; solvent = exact normalized-name match; wavelength = 1.0 within 30 nm then linear decay to zero at 100 nm; temperature = 1.0 within 10 degrees C then linear decay to zero at 50 degrees C; concentration = 1.0 within a factor of two then linear decay to zero at a factor of five. State: "Missing metadata contributes zero; the field score is not renormalized over available fields." Show the final sequence: construct candidate list -> remove held-out record IDs when leave-one-source-out evaluation is active -> sort by final score -> return top three analogies. Make the held-out exclusion boundary visually explicit and label it "evaluation only; normal application calls may omit exclusions."

Use blue for structured chemistry inputs, teal for retrieval computation, orange for fallback decisions, purple for score components, green for returned analogies, and red only for excluded hidden-source records. Keep all labels horizontal and readable at approximately 17 cm figure width."""


def read_csv(name: str) -> list[dict[str, str]]:
    with (SOURCE_DIR / name).open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def leading_categories(name: str, n: int = 5) -> str:
    rows = read_csv(name)[:n]
    return "; ".join(
        f"{row['category']} {int(float(row['count']))} ({float(row['percent']):.1f}%)"
        for row in rows
    )


def percentile(values: list[float], fraction: float) -> float:
    values = sorted(values)
    position = (len(values) - 1) * fraction
    lower = int(position)
    upper = min(lower + 1, len(values) - 1)
    return values[lower] + (values[upper] - values[lower]) * (position - lower)


def corpus_rows() -> list[tuple[str, ...]]:
    inlet = read_csv("fig1f_inlet_streams.csv")
    inlet_summary = "; ".join(
        f"{row['inlet_streams']} streams: {int(row['count'])} ({float(row['percent']):.1f}%)"
        for row in inlet[:5]
    )
    batch = [float(row["yield_percent"]) for row in read_csv("fig1g_batch_yields_raw.csv")]
    flow = [float(row["yield_percent"]) for row in read_csv("fig1g_flow_yields_raw.csv")]
    return [
        ("Reaction class", "464", "100% classified export", leading_categories("fig1a_reaction_classes.csv"), "Other and long-tail classes retained."),
        ("Reactor type", "464", "100% classified export", leading_categories("fig1b_reactor_types.csv"), "Long-tail hardware labels retained."),
        ("Reactor material", "464", "100% classified export", leading_categories("fig1c_reactor_materials.csv", 4), "Classification granularity varies across records."),
        ("Bond/transformation label", "464", "100% classified export", leading_categories("fig1d_bond_types.csv"), "Other/multiple is an explicit category."),
        ("Explicit inlet-stream count", "154", "33.2% of 464", inlet_summary, "Percentages use n = 154, not the full corpus."),
        ("Reported batch yield", str(len(batch)), "29.1% of 464", f"Median {percentile(batch, 0.5):.0f}%; IQR {percentile(batch, 0.25):.0f}-{percentile(batch, 0.75):.0f}%", "Available numeric values only."),
        ("Reported optimized-flow yield", str(len(flow)), "61.4% of 464", f"Median {percentile(flow, 0.5):.0f}%; IQR {percentile(flow, 0.25):.0f}-{percentile(flow, 0.75):.0f}%", "Not a paired causal comparison with batch yield."),
    ]


def rule_data() -> tuple[list[tuple[str, ...]], list[tuple[str, ...]], Counter]:
    payload = json.loads(RULES_PATH.read_text(encoding="utf-8"))
    severity = Counter(str(rule.get("severity") or "unclassified") for rule in payload["rules"])
    category_rows = read_csv("fig2a_rule_landscape.csv")
    categories: dict[str, dict[str, int]] = {}
    for row in category_rows:
        item = categories.setdefault(
            row["category_label"],
            {"total": int(row["category_total"]), "hard_rule": 0, "guideline": 0, "tip": 0, "safety": 0},
        )
        item[row["severity"]] = int(row["count"])
    top = sorted(categories.items(), key=lambda item: -item[1]["total"])[:14]
    category_table = [
        (
            name,
            str(values["total"]),
            str(values["hard_rule"]),
            str(values["guideline"]),
            str(values["tip"]),
            str(values["safety"]),
        )
        for name, values in top
    ]
    def clean_source(source: dict) -> tuple[str, str, str]:
        raw = str(source["title"])
        if "Volume 1 Flow Chemistry" in raw:
            title, year = "Flow Chemistry, Volume 1: Flow Chemistry - Fundamentals", "2021 (stored title)"
        elif "Microreactors in Organic Synthesis" in raw:
            title, year = "Microreactors in Organic Synthesis and Catalysis", "2008 (stored title)"
        elif "integrated approaches for practical applications" in raw.lower():
            title, year = "Flow Chemistry: Integrated Approaches for Practical Applications", "2020 (stored title)"
        elif "Principles-of-Flow-Chemistry" in raw:
            title, year = "Principles of Flow Chemistry", "Not recorded"
        elif "hitchhiker" in raw.lower():
            title, year = "The Hitchhiker's Guide to Flow Chemistry", "Not recorded"
        else:
            title, year = raw, str(source.get("year") or "Not recorded")
        return title, str(source["n_rules_extracted"]), year

    source_table = sorted(
        (clean_source(source) for source in payload["handbooks"]),
        key=lambda row: -int(row[1]),
    )
    return category_table, source_table, severity


RETRIEVAL_TIER_ROWS = [
    ("Tier 1", "Paired batch-to-flow records", "Mechanism plus phase when available", "Up to 20", "If fewer than three hits, relax filters."),
    ("Tier 2", "Paired batch-to-flow records", "No hard metadata filters", "Up to 20", "If zero hits, search all records."),
    ("Tier 3", "All indexed records", "No hard metadata filters", "Up to 20", "If zero hits, return an empty analogy list."),
    ("Final selection", "Candidate list from active tier", "Optional held-out ID exclusion", "Top 3", "Sort descending by combined score."),
]


RETRIEVAL_SCORE_ROWS = [
    ("Semantic similarity", "0.60 of final score", "For normalized embeddings: max(0, 1 - L2^2/2). Lexical mode supplies the store's lexical ranking output."),
    ("Field similarity", "0.40 of final score", "Weighted sum of five metadata components; not renormalized when fields are missing."),
    ("Photocatalyst", "0.30 of field score", "1.0 for the same mapped class; 0.5 for the same mapped metal family; otherwise zero."),
    ("Solvent", "0.20 of field score", "1.0 for exact case-normalized solvent-name equality; otherwise zero."),
    ("Wavelength", "0.20 of field score", "1.0 within 30 nm; linear decay to zero at a 100 nm difference."),
    ("Temperature", "0.15 of field score", "1.0 within 10 degrees C; linear decay to zero at a 50 degrees C difference."),
    ("Concentration", "0.15 of field score", "1.0 within a factor of two; linear decay to zero at a factor of five."),
]


RETRIEVAL_EVIDENCE_ROWS = [
    ("Frozen query-result pairs", "1,600", "Pairs used to compare semantic rank with field-aware reranked position."),
    ("Pairs whose rank changed", "25.1%", "Demonstrates that metadata reranking materially changed ordering; not a measure of reaction yield."),
    ("TiO2 top-5 family match", "87.1% -> 98.6% (n = 14 queries)", "Absolute increase 11.4 percentage points."),
    ("Iridium top-5 family match", "40.0% -> 95.0% (n = 8 queries)", "Absolute increase 55.0 percentage points."),
    ("Ruthenium top-5 family match", "77.3% -> 94.7% (n = 15 queries)", "Absolute increase 17.3 percentage points."),
    ("Organic-dye top-5 family match", "42.7% -> 90.7% (n = 15 queries)", "Absolute increase 48.0 percentage points."),
    ("ZnO top-5 family match", "45.0% -> 60.0% (n = 4 queries)", "Absolute increase 15.0 percentage points; small query set."),
    ("Held-out source control", "Optional exclude_record_ids", "Candidate IDs are removed after candidate construction and before final sorting and top-k selection."),
]


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    category_rows, source_rows, severity = rule_data()

    document = Document()
    configure_document(document)
    document.core_properties.title = "Literature corpus, rule base, and retrieval"
    document.core_properties.subject = "Electronic Supporting Information section"
    document.core_properties.keywords = "FlowPilot, literature corpus, engineering rules, retrieval, RAG"

    document.add_heading("S2. Literature corpus, rule base, and retrieval", level=1)
    note = document.add_paragraph()
    note.paragraph_format.space_after = Pt(8)
    run = note.add_run(
        "Production note: Figure-generation prompts are intentionally placed at the proposed "
        "figure locations. Replace each prompt box with finalized artwork before submission. "
        "Add the final corpus repository identifier, access date, and source citations before publication."
    )
    set_font(run, "Arial", 8.5, color=MID_GREY)

    add_paragraph(
        document,
        "FlowPilot uses two complementary knowledge resources. A structured literature corpus provides "
        "empirical batch-to-flow precedents and process metadata, whereas a handbook-derived rule base "
        "provides general engineering principles. Plan-aware retrieval selects a small set of literature "
        "analogies for the chemistry under consideration, and the upstream chemistry layer receives a "
        "prioritized subset of relevant handbook rules. These resources provide evidence and context; they "
        "do not replace deterministic calculations, laboratory inventory constraints, or experimental "
        "validation.",
    )

    document.add_heading("S2.1 Frozen literature-corpus composition", level=2)
    add_paragraph(
        document,
        "The frozen manuscript export contains 464 classified flow-chemistry literature records. The "
        "classification spans thermal synthesis, photoredox and heterogeneous photocatalysis, cross-coupling, "
        "oxidation and reduction, electrochemistry, biocatalysis, polymer synthesis, crystallization, and "
        "other continuous-flow applications. For each record, available information was normalized into "
        "fields describing reaction class, reactor type, reactor material, bond or transformation label, "
        "operating conditions, number of inlet streams, and reported batch and flow outcomes.",
    )
    add_paragraph(
        document,
        "The primary categorical exports each contain 464 classified entries, but field completeness is "
        "not uniform. Explicit inlet-stream counts are available for 154 processes, batch yields for 135 "
        "records, and optimized-flow yields for 285 records (Table S4). Unknown, other, multiple, and "
        "long-tail classifications are retained rather than being excluded from denominators. This avoids "
        "inflating apparent coverage and makes the limitations of the extracted metadata visible.",
    )
    add_paragraph(
        document,
        "The available yield distributions are descriptive. They should not be interpreted as a paired "
        "estimate of the causal benefit of flow because batch and flow values are missing for different "
        "subsets, publication and optimization selection can affect the distributions, and individual "
        "records differ in chemistry and objective. Figure S3 therefore reports field-specific denominators "
        "directly in each panel.",
    )

    add_prompt_box(document, "FIGURE PLACEHOLDER - GENERATION PROMPT FOR FIGURE S3", FIGURE_S3_PROMPT)
    add_caption(
        document,
        "Figure S3",
        "Composition of the frozen flow-chemistry literature corpus. Distributions are shown for "
        "reaction class, reactor type, reactor material, bond or transformation label, number of inlet "
        "streams, and available batch and optimized-flow yields. The four principal categorical exports "
        "contain 464 classified records. Inlet-stream and yield panels use their own available-value "
        "denominators. Unknown, other, multiple, and long-tail categories are retained. Yield distributions "
        "are descriptive and are not treated as a paired causal comparison.",
    )

    add_table(
        document,
        "Table S4. Frozen literature-corpus fields and available denominators.",
        ["Metadata field", "Available n", "Coverage", "Leading categories or summary", "Interpretation note"],
        corpus_rows(),
        [2.8, 1.7, 2.3, 6.3, 3.9],
    )

    document.add_heading("S2.2 Engineering rule-base construction and use", level=2)
    add_paragraph(
        document,
        "The engineering knowledge store contains 2,537 structured rules extracted from six handbook, "
        "guide, and literature sources. Each stored rule includes a rule identifier, category, triggering "
        "condition, recommendation, reasoning, optional quantitative expression, exceptions, severity, "
        "confidence, source document, source page, and source context when available. The complete export "
        "contains 29 rule categories; the 14 largest are summarized in Table S5.",
    )
    add_paragraph(
        document,
        f"Across the complete store, {severity['guideline']:,} rules are labelled guideline, "
        f"{severity['hard_rule']:,} hard_rule, {severity['tip']:,} tip, and {severity['safety']:,} safety "
        "severity. Category and severity are distinct labels. For example, most entries in the Safety "
        "category are encoded as hard_rule rather than with the separate safety severity label. Figure S4 "
        "therefore displays both dimensions rather than treating the Safety category as a severity class.",
    )
    add_paragraph(
        document,
        "At run time, the knowledge store selects broad engineering categories from the reaction context "
        "and adds keyword matches from the mechanism, catalyst, solvent, phase regime, and operating "
        "conditions. Duplicate rule identifiers are removed, hard rules are prioritized, and a compact "
        "subset is formatted for the upstream chemistry prompt. The downstream pipeline receives the "
        "resulting ChemistryPlan, while executable numerical requirements remain the responsibility of "
        "deterministic calculators and validation gates.",
    )
    add_paragraph(
        document,
        "The rule-base labels should not be overinterpreted. A machine-detected quantitative expression "
        "indicates that a rule contains formula-like content; it does not demonstrate independent equation "
        "verification. Similarly, a hard_rule label controls prompt prioritization but is not automatically "
        "equivalent to an executable coded constraint. A rule becomes a deterministic gate only when the "
        "corresponding calculation or validation check is implemented in code.",
    )

    add_prompt_box(document, "FIGURE PLACEHOLDER - GENERATION PROMPT FOR FIGURE S4", FIGURE_S4_PROMPT)
    add_caption(
        document,
        "Figure S4",
        "Engineering rule-base coverage and severity structure. The 2,537-rule store is summarized by "
        "category, severity, chemistry-class association, source document, and machine-detected quantitative "
        "content. Severity counts are mutually exclusive, whereas chemistry-class associations are "
        "non-exclusive. Quantitative-expression detection is a coverage indicator and not evidence that "
        "every extracted expression has been independently validated. Rule labels provide prioritized "
        "knowledge context; deterministic enforcement requires a corresponding coded check.",
    )

    add_table(
        document,
        "Table S5. Counts and severity structure of the 14 largest rule categories.",
        ["Rule category", "Total", "Hard rule", "Guideline", "Tip", "Safety severity"],
        category_rows,
        [4.5, 2.0, 2.5, 2.5, 2.0, 3.5],
    )
    document.add_page_break()
    add_table(
        document,
        "Table S6. Source documents represented in the frozen 2,537-rule store.",
        ["Stored source title", "Extracted rules", "Stored year"],
        source_rows,
        [11.5, 2.8, 2.7],
    )

    document.add_heading("S2.3 Plan-aware retrieval and tier transitions", level=2)
    add_paragraph(
        document,
        "The retriever accepts a structured BatchRecord and, when available, a ChemistryPlan. The plan-aware "
        "query extends the batch summary with reaction name, reaction class, mechanism, catalyst or "
        "sensitizer names, solvent, temperature, measured or recommended wavelength, key intermediate, bond "
        "formed, retrieval keywords, similar reaction classes, and oxygen-sensitivity context. Retrieval can "
        "operate through semantic embeddings or deterministic lexical search. If embedding generation fails, "
        "the current implementation falls back to lexical search.",
    )
    add_paragraph(
        document,
        "Candidate scope is relaxed in three tiers (Table S7). The first tier searches paired batch-to-flow "
        "records using mechanism and, when present in the plan, phase filters. Fewer than three results "
        "trigger an unfiltered search over paired records. If that search returns no records, the final tier "
        "searches the complete record collection. Each tier retrieves up to 20 candidates, after which the "
        "combined score is calculated and the three highest-ranked analogies are returned.",
    )
    add_paragraph(
        document,
        "The final score combines 0.60 semantic similarity and 0.40 structured field similarity. Field "
        "similarity is a weighted sum of photocatalyst, solvent, wavelength, temperature, and concentration "
        "components (Table S8). Missing metadata contributes zero and the score is not renormalized over "
        "available fields. This conservative implementation avoids treating an unknown field as a match, but "
        "it can disadvantage incompletely annotated records and should be considered when interpreting ranks.",
    )
    add_paragraph(
        document,
        "For leave-one-source-out evaluation, hidden reference identifiers are supplied through the runtime "
        "configuration. Matching candidate IDs are removed after candidate construction and before final "
        "sorting and top-k selection. Normal application calls may omit this exclusion. This control prevents "
        "the held-out reference record from being returned directly, although exact benchmark replay also "
        "requires the frozen corpus, embedding or lexical configuration, query text, and exclusion manifest.",
    )

    add_prompt_box(document, "FIGURE PLACEHOLDER - GENERATION PROMPT FOR FIGURE S5", FIGURE_S5_PROMPT)
    add_caption(
        document,
        "Figure S5",
        "Plan-aware literature retrieval and tier transitions. (A) BatchRecord and ChemistryPlan fields "
        "are combined into an enriched semantic or lexical query. (B) Retrieval begins with filtered paired "
        "records, relaxes to unfiltered paired records when fewer than three hits are available, and searches "
        "all records only when the paired search is empty. (C) Up to 20 candidates are reranked using 0.60 "
        "semantic similarity and 0.40 field similarity before the top three analogies are returned. In "
        "leave-one-source-out evaluation, hidden record identifiers are removed before final ranking and "
        "selection.",
    )

    add_table(
        document,
        "Table S7. Retrieval tiers, scope, and fallback conditions.",
        ["Stage", "Search scope", "Hard filters", "Candidate limit", "Transition or output"],
        RETRIEVAL_TIER_ROWS,
        [2.0, 3.8, 3.7, 2.2, 5.3],
    )
    add_table(
        document,
        "Table S8. Field-aware retrieval score and operational definitions.",
        ["Score component", "Configured weight", "Operational definition"],
        RETRIEVAL_SCORE_ROWS,
        [4.0, 3.5, 9.5],
    )

    document.add_heading("S2.4 Frozen retrieval evaluation and interpretation", level=2)
    add_paragraph(
        document,
        "The frozen retrieval analysis contains 1,600 query-result pairs. Field-aware reranking changed the "
        "position of 25.1% of pairs relative to semantic ranking alone. In the annotated photochemical query "
        "sets, top-five photocatalyst-family agreement increased for TiO2, iridium, ruthenium, organic-dye, "
        "and ZnO systems (Table S9). These results demonstrate improved alignment with selected metadata "
        "attributes. They do not establish that a retrieved analogy is mechanistically complete, that its "
        "operating conditions transfer directly, or that it improves downstream experimental yield.",
    )
    add_paragraph(
        document,
        "The family-specific query counts are modest, particularly for ZnO and iridium. The evaluation is "
        "therefore reported as a retrieval diagnostic rather than a universal estimate of chemical relevance. "
        "Downstream modules retain responsibility for checking chemistry identity, engineering compatibility, "
        "inventory feasibility, and safety before any analogy can influence an executable design.",
    )
    add_table(
        document,
        "Table S9. Frozen retrieval diagnostics and leakage-control boundary.",
        ["Diagnostic", "Observed value", "Interpretation"],
        RETRIEVAL_EVIDENCE_ROWS,
        [4.3, 5.2, 7.5],
    )

    document.add_heading("S2.5 Scope and limitations", level=2)
    add_paragraph(
        document,
        "The literature corpus is a curated but incomplete representation of flow chemistry, and metadata "
        "quality depends on what source documents report and what the extraction pipeline can normalize. "
        "The rule base is likewise a structured knowledge resource rather than a substitute for primary "
        "source inspection. Retrieved analogies and handbook rules should be treated as traceable support for "
        "reasoning, while measured evidence, declared constraints, deterministic calculations, and wet-lab "
        "validation retain higher authority. Dataset and rule-store versions should be frozen and deposited "
        "with the final manuscript so that all figure values and retrieval claims can be reconstructed.",
    )

    document.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build())
