"""Generate the manuscript ESI figure and content plan as a Word document."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables" / "manuscript_esi_plan_20260811"
OUT_PATH = OUT_DIR / "FlowPilot_ESI_Figure_and_Content_Plan.docx"

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
TEAL = "0F6B6D"
LIGHT_TEAL = "DDEFEF"
ORANGE = "B45F06"
LIGHT_ORANGE = "FCE4D6"
RED = "9C0006"
LIGHT_RED = "FFC7CE"
GREY = "666666"
LIGHT_GREY = "F2F2F2"
WHITE = "FFFFFF"


SECTIONS = [
    {
        "heading": "S1. Workflow, data contracts, and reproducibility",
        "purpose": (
            "Document what enters the system, which component owns each decision, "
            "and how a complete run can be reconstructed from stored artifacts."
        ),
        "figures": [
            {
                "n": 1,
                "title": "Extended FlowPilot architecture and authority boundaries",
                "purpose": "Expand main Figure 1 into the complete executable architecture, including standardized intake, chemistry interpretation, retrieval, deterministic calculations, council, inventory reconciliation, final validation, topology rendering, and autosave.",
                "panels": "(a) end-to-end modules; (b) data objects passed between modules; (c) authority hierarchy: measured evidence > hard constraints > protocol facts > hypotheses > model inference; (d) executable versus blocked output paths.",
                "source": "Existing starting asset: ablation_results/figures/23_full_flowpilot_pipeline.pdf. Update from flora_translate/main.py, schemas.py, final_design_contract.py, and topology_compiler.py.",
                "status": "Regenerate",
            },
            {
                "n": 2,
                "title": "Component-level schema and data-flow contracts",
                "purpose": "Show that the LLM agents exchange typed objects rather than unconstrained prose and identify the authoritative source for each final field.",
                "panels": "BatchRecord -> ChemistryPlan -> FlowProposal -> DesignCalculations -> council candidate -> inventory-bound topology -> FinalDesignContract.",
                "source": "Generate from flora_translate/schemas.py and representative full_result.json artifacts.",
                "status": "New from code",
            },
            {
                "n": 3,
                "title": "Reproducible standardized intake workflow",
                "purpose": "Explain the conversational intake LLM while demonstrating fixed question identifiers and deterministic readiness rules.",
                "panels": "(a) protocol submission; (b) fixed IDs Q-BATCH-001 through Q-PREF-001; (c) answered/unavailable states; (d) frozen DesignInputPackage; (e) design gate.",
                "source": "components/intake_wizard.py, flora_translate/intake_agent.py, and intake tests.",
                "status": "New from code",
            },
            {
                "n": 4,
                "title": "Representative frozen DesignInputPackage and evidence hierarchy",
                "purpose": "Provide a readable example of how protocol facts, objectives, historical measurements, inventory constraints, hypotheses, and safety limits are stored before design.",
                "panels": "Annotated JSON excerpt plus a provenance map showing which package sections influence upstream chemistry, deterministic calculations, and council review.",
                "source": "Use a redacted intake_package.json from outputs/benchmarks/khu_three_protocols_canonical_20260810_161346.",
                "status": "Existing data; compose",
            },
            {
                "n": 5,
                "title": "GUI workflow and result provenance",
                "purpose": "Document the user-visible path from intake and inventory selection to final tabs while demonstrating that every tab reads the same FinalDesignContract.",
                "panels": "Intake, inventory selection, final summary, engineering design, stream assignment, council trace, process topology, and raw JSON.",
                "source": "Capture with Streamlit/Playwright from pages/flora_design_unified.py and components/*.py after the final GUI regression.",
                "status": "New screenshot set",
            },
            {
                "n": 6,
                "title": "Autosave artifact tree and run-level provenance",
                "purpose": "Demonstrate reproducibility through exact prompts, model events, inventory profile, raw result, canonical result, topology files, validation audit, logs, and checksums.",
                "panels": "(a) folder tree; (b) artifact-to-claim mapping; (c) checksum verification; (d) model/component event count.",
                "source": "outputs/gui_runs, outputs/benchmarks/khu_three_protocols_canonical_20260810_161346, and flora_translate/gui_autosave.py.",
                "status": "New from existing artifacts",
            },
        ],
    },
    {
        "heading": "S2. Literature corpus, rule base, and retrieval",
        "purpose": (
            "Provide the detailed distributions and retrieval controls underlying "
            "main Figures 2 and 3 without crowding the main manuscript."
        ),
        "figures": [
            {
                "n": 7,
                "title": "Full literature-corpus composition",
                "purpose": "Expand the corpus overview beyond the selected categories shown in main Figure 2.",
                "panels": "Reaction class, reactor type, tubing material, phase regime, number of streams, temperature, pressure, residence time, publication year, and source-paper counts.",
                "source": "Regenerate from the curated literature records used for the 464-paper corpus.",
                "status": "Regenerate",
            },
            {
                "n": 8,
                "title": "Engineering rule-base coverage and severity structure",
                "purpose": "Document all 2,537 rules, their categories, severity levels, source type, and chemistry-class associations.",
                "panels": "Category counts; severity distribution; category x chemistry heatmap; rule-source distribution; sparse-domain map.",
                "source": "flora_fundamentals knowledge store and rule export used for main Figure 2.",
                "status": "Regenerate",
            },
            {
                "n": 9,
                "title": "Plan-aware retrieval algorithm and tier transitions",
                "purpose": "Show the complete hard-filter, relaxed-filter, and semantic-fallback logic with field-aware reranking.",
                "panels": "Algorithm diagram, query fields, fallback triggers, weighted score, and one worked query.",
                "source": "Retrieval implementation plus main Figure 3 source data.",
                "status": "New detailed schematic",
            },
            {
                "n": 10,
                "title": "Retrieval performance, reranking shifts, and leakage controls",
                "purpose": "Support the retrieval claim with family-level results, rank changes, and proof that hidden benchmark references are excluded.",
                "panels": "Top-k family matching; rank-shift distribution; field-score contribution; similarity-to-reference heatmap; per-case hidden-source exclusion audit.",
                "source": "benchmark manuscript_comparison figures and ablation_test/tests/test_retrieval_exclusion.py.",
                "status": "Existing analyses; combine",
            },
        ],
    },
    {
        "heading": "S3. Deterministic engineering, multiphase calculations, and inventory",
        "purpose": (
            "Expose the equations and closure checks that distinguish FlowPilot from "
            "a prose-only LLM while documenting recent heat-transfer, gas-liquid, "
            "multistage, and inventory additions."
        ),
        "figures": [
            {
                "n": 11,
                "title": "Nine-step deterministic engineering calculator",
                "purpose": "Provide a detailed version of the calculator sequence and identify the inputs, equations, outputs, and failure gates at every step.",
                "panels": "Kinetics/residence time; reactor sizing; fluid dynamics; pressure drop; mixing/mass transfer; heat transfer; vapour pressure/BPR; process metrics; final validation.",
                "source": "flora_translate/design_calculator.py and components/design_steps.py.",
                "status": "New from code",
            },
            {
                "n": 12,
                "title": "Engineering arithmetic closure and unit-consistency validation",
                "purpose": "Show deterministic tests for V = Q tau, coil geometry, pressure drop, Re, Pe, heat-transfer area, and stage-total reconciliation.",
                "panels": "Expected-versus-computed parity plots and pass-rate bars across tests/benchmark outputs.",
                "source": "flora_translate/tests and canonical case audit.json files.",
                "status": "New from tests",
            },
            {
                "n": 13,
                "title": "Heat-transfer calculation workflow and sensitivity",
                "purpose": "Document the recent UA, area-to-volume, heat-generation/removal, and thermal-Damkohler calculations.",
                "panels": "Equation schematic; UA versus tubing ID/volume; heat-removal margin; representative thermal and photochemical cases; warning thresholds.",
                "source": "Design calculation heat_transfer_metrics and associated tests.",
                "status": "New from code",
            },
            {
                "n": 14,
                "title": "Gas-liquid bookkeeping: inlet/STP versus in-channel basis",
                "purpose": "Make the gas-flow convention unambiguous and show why pressure-corrected gas flow changes channel residence time but not inlet equivalents.",
                "panels": "STP feed; pressure/temperature correction; inlet residence; channel residence; O2-equivalent closure; example from the photoredox oxidation stage.",
                "source": "Existing ablation_results/figures/15_gas_bookkeeping.pdf plus current multiphase code and tests.",
                "status": "Update existing",
            },
            {
                "n": 15,
                "title": "Multistage residence-time and stream-flow accounting",
                "purpose": "Show how upstream effluent and newly added feeds are propagated without double counting and how per-stage residence times sum to the final total.",
                "panels": "Two-stage liquid example; gas-addition example; quench-only operation; serial-reactor example; stage arithmetic audit.",
                "source": "flora_translate/multistage_inventory.py, topology compiler, and three-protocol canonical audits.",
                "status": "New from code",
            },
            {
                "n": 16,
                "title": "Inventory ingestion and canonical profile generation",
                "purpose": "Explain how PDF/PPT/DOCX/free-text equipment descriptions become a validated, reusable LabInventory JSON profile.",
                "panels": "Document upload; LLM extraction; null normalization; Pydantic validation; human confirmation; saved profile; import into design.",
                "source": "Inventory Manager GUI and deliverables/flowpilot_inventory_management_20260804.",
                "status": "Existing schematic; adapt",
            },
            {
                "n": 17,
                "title": "Inventory capability coverage and deterministic allocation",
                "purpose": "Demonstrate selection among pumps, tubing, reactors, lights, MFCs, mixers, BPRs, connectors, and temperature limits, including blocked and assumed-accessory states.",
                "panels": "Capability matrix; allocation graph; accepted design; rejected unavailable reactor; temperature clamp; unresolved topology; final instrument manifest.",
                "source": "khu_inventory_updated_flowpilot.json, ablation_results/figures/12_inventory_feasibility.pdf, and topology artifacts.",
                "status": "Regenerate with final inventory logic",
            },
        ],
    },
    {
        "heading": "S4. Council behavior, model matrix, and reproducibility",
        "purpose": (
            "Provide the detailed evidence behind main Figure 4, including candidate "
            "trajectories, model dependence, repeated runs, and candidate-budget effects."
        ),
        "figures": [
            {
                "n": 18,
                "title": "Council candidate lifecycle and bounded revision",
                "purpose": "Show generation, domain scoring, deterministic rejection, specialist revision, skeptic audit, selection, and final recomputation for one complete run.",
                "panels": "Candidate Sankey/funnel; specialist decisions; rejected revisions; winning-candidate trajectory; final score components.",
                "source": "Council logs and llm_events.jsonl from a representative model-matrix run.",
                "status": "New from logs",
            },
            {
                "n": 19,
                "title": "Complete 4 x 4 model-matrix engineering outcomes",
                "purpose": "Move the detailed per-cell metrics supporting main Figure 4 into the ESI.",
                "panels": "Residence time, flow rate, reactor volume, tubing ID, BPR, validation status, screen-required status, and concern count heatmaps.",
                "source": "benchmark/data/model_matrix_benchmark_20260504_101732/visualizations.",
                "status": "Existing; assemble",
            },
            {
                "n": 20,
                "title": "Pre/post-council repeatability and metric covariance",
                "purpose": "Show all repeats behind the representative radar plot and quantify run-to-run variation rather than displaying only one aggregate polygon.",
                "panels": "Repeat overlays; paired deltas; coefficient-of-variation heatmap; metric-correlation matrix; design-regime scatter.",
                "source": "prepost_radar outputs and ablation_results/figures/17_metric_correlation.pdf.",
                "status": "Existing analyses; combine",
            },
            {
                "n": 21,
                "title": "Candidate-budget effects across B = 1, 6, 12, and 24",
                "purpose": "Provide the complete distributions behind main Figure 4e-f and document outlier/pathology rates.",
                "panels": "Design-family counts; metric distributions; revision/disqualification activity; CV heatmap; pathology rate; completion rate.",
                "source": "benchmark/data/protocol_budget_benchmark_20260427_181556/visualizations.",
                "status": "Existing; assemble",
            },
            {
                "n": 22,
                "title": "Runtime, token use, LLM calls, and stage-level cost",
                "purpose": "Separate computational cost from design quality and identify which stages dominate execution.",
                "panels": "Runtime; tokens; calls; stage breakdown; cost-quality scatter; failed/retried requests.",
                "source": "Model-matrix visualizations and ablation_results/figures/09, 10, 16, and 19.",
                "status": "Existing; update model names",
            },
        ],
    },
    {
        "heading": "S5. Matched ablation and architecture comparison",
        "purpose": (
            "Place the complete ablation in the ESI using the defensible matched-model "
            "design: the same model, protocol, objective, inventory, condition, and "
            "repeat input for one-shot and full FlowPilot."
        ),
        "figures": [
            {
                "n": 23,
                "title": "Ablation study design and frozen endpoint calculation",
                "purpose": "Define the matched comparisons, feasible/infeasible inventories, repeats, primary executability endpoint, secondary quality dimensions, and interpretation boundaries.",
                "panels": "Study matrix; input matching; endpoint decision tree; score calculation; frozen-weight statement.",
                "source": "ablation_results/presentation/flowpilot_vs_oneshot_qwen_gpt_20260810/figures/03_score_calculation.pdf and benchmark manifests.",
                "status": "Existing; adapt for manuscript",
            },
            {
                "n": 24,
                "title": "Qwen 27B one-shot versus Qwen 27B within FlowPilot",
                "purpose": "Isolate the architecture effect while holding the local base model fixed.",
                "panels": "Executable rate; correct-block rate; win/tie/loss; paired case scores; repeat dispersion.",
                "source": "qwen_gpt_architecture_comparison_v2 and presentation figures.",
                "status": "Existing publication candidate",
            },
            {
                "n": 25,
                "title": "GPT one-shot versus the same GPT model within FlowPilot",
                "purpose": "Repeat the architecture comparison with the commercial model while preserving exactly matched inputs and scoring.",
                "panels": "Executable rate; correct-block rate; win/tie/loss; paired case scores; repeat dispersion.",
                "source": "qwen_gpt_architecture_comparison_v2 and presentation figures. Replace provider/model labels with the exact endpoint reported in Methods.",
                "status": "Existing publication candidate",
            },
            {
                "n": 26,
                "title": "Paired architecture uplift by chemistry family",
                "purpose": "Show whether the architecture benefit is consistent across reaction families rather than driven by one protocol.",
                "panels": "Case-level paired deltas and chemistry-family cluster intervals for Qwen and GPT separately.",
                "source": "presentation/02_architecture_effects.pdf and comparison tables.",
                "status": "Existing publication candidate",
            },
            {
                "n": 27,
                "title": "Quality-dimension decomposition and schema-neutral sensitivity",
                "purpose": "Address the criticism that FlowPilot wins only because it emits its own schema by showing both the frozen dimensions and a schema-neutral sensitivity analysis.",
                "panels": "Engineering closure; safety; inventory; evidence; assurance; actionability; formal validity; schema-neutral paired effect.",
                "source": "reduced_frontier_final_package figures 03, 11, and 12 plus presentation/04_quality_score_breakdown.pdf.",
                "status": "Existing; combine carefully",
            },
            {
                "n": 28,
                "title": "Hard-check and deployment-gate rates",
                "purpose": "Keep scientific design quality separate from immediate deployment readiness and expose every failing gate.",
                "panels": "Geometry closure; gas bookkeeping; pump range; pressure; temperature; inventory assignment; topology; screen-required; deployment-ready rate.",
                "source": "reduced_frontier_final_package figures 06 and 07 and ablation_results/figures/29-30.",
                "status": "Existing; use current deterministic gates",
            },
            {
                "n": 29,
                "title": "Weight sensitivity, bootstrap uncertainty, and rank robustness",
                "purpose": "Demonstrate which conclusions are robust to reasonable weighting choices and which remain value-dependent.",
                "panels": "Six prespecified weight schemes; 20,000-weight-set winner frequency; paired bootstrap intervals; per-case win/tie/loss.",
                "source": "presentation/05_weight_sensitivity.pdf and ablation_results/figures/27-28.",
                "status": "Existing; report as sensitivity, not proof",
            },
            {
                "n": 30,
                "title": "Evidence-first safeguards and correction of fallback dependence",
                "purpose": "Show the before/after effect of removing forced intensification, preserving measured residence-time evidence, and deterministically validating final revisions.",
                "panels": "Council path before/after; residence-time evidence adherence; joint disposition/engineering success; safeguard-event counts.",
                "source": "stage3_evidence_first_validation_package figures 10-13.",
                "status": "Existing publication candidate",
            },
            {
                "n": 31,
                "title": "Agent call trace and proof of multi-agent execution",
                "purpose": "Provide direct evidence that upstream, deterministic engineering, council specialists, skeptic, revision, and orchestrator components actually executed.",
                "panels": "Chronological call trace; component call counts; prompt/response hashes; representative bounded messages; retry/fallback annotation.",
                "source": "ablation_results/tables/agent_call_events.csv, reports/agent_trace_summary.md, and representative llm_events.jsonl.",
                "status": "Existing logs; compose",
            },
        ],
    },
    {
        "heading": "S6. Case studies and wet-lab validation",
        "purpose": (
            "Preserve the displaced bromination case and provide the complete protocols, "
            "design histories, inventory constraints, experimental conditions, and raw "
            "analytical evidence behind revised main Figures 5 and 6."
        ),
        "figures": [
            {
                "n": 32,
                "title": "Thermal alpha-bromination with inline thiosulfate quench",
                "purpose": "Move the current main Figure 6 intact into the ESI so the original literature-comparison case remains part of the evidence record.",
                "panels": "Current panels: council deliberation, validated process flow diagram, and design-versus-literature comparison.",
                "source": "flent_fig_png/fig7.png and current main_manuscript.pdf Figure 6.",
                "status": "Existing; move intact",
            },
            {
                "n": 33,
                "title": "Bromination candidate space, mixing warning, and validation details",
                "purpose": "Supply the quantitative background omitted from the moved composite figure and preserve the unresolved micromixing limitation.",
                "panels": "Candidate geometry map; mixing criteria; conversion estimates; selected-versus-rejected candidates; scale comparison; safety/quench topology audit.",
                "source": "Bromination benchmark result and council logs; multiphase validation protocol.",
                "status": "New from existing run",
            },
            {
                "n": 34,
                "title": "Photoredox Giese/aerobic oxidation: complete protocol and inventory-bound topology",
                "purpose": "Provide the full chemistry, stage-separation logic, oxygen timing, no-inline-degasser constraint, instrument assignments, STP/in-channel gas flows, and residence-time arithmetic supporting revised main Figure 5.",
                "panels": "Batch protocol; chemistry mechanism/stream separation; initial topology; no-degasser inventory topology; instrument manifest; gas-basis calculation.",
                "source": "case_study1.png, scripts/run_case_study1_no_inline_degas.py, and canonical case-01 artifacts.",
                "status": "Existing data; regenerate after final experimental design is frozen",
            },
            {
                "n": 35,
                "title": "Photoredox case: full closed-loop experimental refinement history",
                "purpose": "Expand the wet-lab panel added to main Figure 5 into a complete cycle-by-cycle record, including rejected below-pump-limit designs and inventory changes.",
                "panels": "Cycle timeline; reactor volume; liquid/O2 flow; inlet/channel residence; temperature/pressure; measured yield; agent feedback; next design; uncertainty and stopping decision.",
                "source": "profpark_latest_july26.png, THQ/KRICT iterative run folders where applicable, collaborator records, and final wet-lab dataset. Use only the chemistry actually shown in main Figure 5.",
                "status": "Needs final wet-lab table and identity audit",
            },
            {
                "n": 36,
                "title": "DPDTC protocol 1: nitrobenzoic acid-benzylamine coupling",
                "purpose": "Provide the first replacement chemistry in revised main Figure 6 with enough detail to reproduce feed preparation and understand each refinement decision.",
                "panels": "Batch protocol; stream A/B recipes; stage topology; inventory assignment; temperature/residence-time design history; wet-lab conditions and yield; analytical trace.",
                "source": "2 protocols.pdf and canonical case_02 artifacts.",
                "status": "Blocked pending corrected rerun and wet-lab data",
            },
            {
                "n": 37,
                "title": "DPDTC protocol 2: benzoic acid-morpholine coupling",
                "purpose": "Provide the second replacement chemistry in revised main Figure 6, emphasizing aqueous interstage addition and phase/mixing considerations.",
                "panels": "Batch protocol; stream A/B recipes; stage topology; inventory assignment; temperature/residence-time design history; wet-lab conditions and yield; analytical trace.",
                "source": "2 protocols.pdf and canonical case_03 artifacts.",
                "status": "Blocked pending corrected rerun and wet-lab data",
            },
            {
                "n": 38,
                "title": "Cross-case experimental agreement, refinement efficiency, and failure modes",
                "purpose": "Close the ESI by summarizing what FlowPilot predicted, what was built, what was measured, and how many refinement cycles were needed across all wet-lab cases.",
                "panels": "Predicted-versus-measured yield or conversion; initial-to-final parameter shifts; cycles to target; inventory violations caught; unsuccessful screens retained; final uncertainty classification.",
                "source": "Final harmonized wet-lab table for main Figures 5 and 6.",
                "status": "Needs completed experiments",
            },
        ],
    },
]


TEXT_SECTIONS = [
    ("S1", "Software architecture and versioning", "Commit identifier, release date, Python environment, package versions, model endpoints, prompts, seeds where supported, retry policy, and deterministic/stochastic boundaries."),
    ("S2", "Standardized intake and authority policy", "Question bank, readiness requirements, unavailable responses, package freezing, and evidence/constraint authority order."),
    ("S3", "Literature corpus construction", "Inclusion criteria, paper deduplication, metadata extraction, manual/machine verification status, and hidden-reference exclusions."),
    ("S4", "Engineering rule-base construction", "Rule sources, schema, category/severity definitions, deduplication, conflict handling, and sparse-domain limitations."),
    ("S5", "Deterministic equations and assumptions", "All equations, units, correlations, property assumptions, thresholds, and validity ranges for the nine-step calculator."),
    ("S6", "Gas-liquid and multistage conventions", "Definitions of inlet/STP and in-channel flow, gas equivalents, holdup treatment, stage-flow propagation, total residence time, and quench-only operations."),
    ("S7", "Inventory profile and allocation rules", "Equipment schema, PDF/PPT/DOCX extraction, hard limits, accessory assumptions, quantity allocation, reactor trains, and blocked-design behavior."),
    ("S8", "Council prompts and bounded authority", "System prompts, specialist scopes, candidate generation, deterministic recomputation after revision, skeptic audit, and orchestrator scoring."),
    ("S9", "Benchmark and reproducibility protocol", "Case selection, model matrix, repeats, candidate budgets, endpoints, failed runs, exclusions, and statistical methods."),
    ("S10", "Ablation protocol and interpretation", "Matched-input design, feasible/infeasible controls, frozen scoring, schema-neutral analysis, bootstrap strategy, and limits on superiority claims."),
    ("S11", "Experimental procedures", "Full batch references, feed preparation, startup/steady-state definitions, reactor assembly, pressure/temperature control, sampling, workup, analytical methods, and safety controls."),
    ("S12", "Data and code availability", "Artifact manifests, raw prompts/completions, JSON schemas, logs, topology files, analysis scripts, checksums, and repository release plan."),
]


TABLES = [
    (1, "Software, model, and execution configuration", "Exact model IDs by component, endpoint/provider, sampling controls, retry policy, embedding model, environment, and commit."),
    (2, "Schema objects and field ownership", "Every major field, owning module, allowed modifiers, validation rule, and final source of truth."),
    (3, "Literature corpus summary", "Paper and record counts by reaction class, reactor, phase, material, and verification status."),
    (4, "Engineering rule-base summary", "Rule counts by category, severity, source, and chemistry class."),
    (5, "Equations, symbols, units, and validity ranges", "Complete deterministic calculator reference."),
    (6, "Model-matrix run manifest", "All upstream/council pairings, repeats, completion state, latency, tokens, and artifact path."),
    (7, "Candidate-budget results", "Per-repeat selected design, revisions, disqualifications, pathologies, and variability."),
    (8, "Ablation cases and matched inputs", "Protocol family, objective, inventory state, repeat hashes, feasibility oracle, and hidden-reference exclusion."),
    (9, "Ablation scoring specification", "Primary endpoint, quality dimensions, weights, hard gates, and schema-neutral sensitivity."),
    (10, "Per-case ablation outcomes", "Qwen and GPT one-shot versus full-pipeline results without averaging away failures."),
    (11, "KHU laboratory inventory", "Pumps, tubing, reactors, lights, MFCs, BPRs, connectors, temperature/pressure limits, quantities, and assumptions."),
    (12, "Photoredox wet-lab cycles", "Complete design and measured result for every cycle, including unsuccessful and rejected conditions."),
    (13, "DPDTC feed-preparation and operating conditions", "Masses, concentrations, reservoir volumes, flow ratios, stage flows, residence times, temperatures, pressures, yields, and analytical method."),
    (14, "Cross-case deviations and unresolved assumptions", "Difference between proposed and executed design, reason, consequence, and resolution."),
]


MANUSCRIPT_REVISIONS = [
    ("Abstract", "Replace the claim that the two end-to-end case studies are Giese oxidation and alpha-bromination. State the final wet-lab case set and distinguish literature comparison from experimental validation."),
    ("Introduction contributions", "Update the named validation cases and add matched architecture ablation, standardized intake, inventory-bound design, and closed-loop refinement only if these are claimed in Results."),
    ("Figure 5 Results", "Add wet-lab cycles and measured outcomes. Move full protocol, hardware, gas-basis arithmetic, unsuccessful cycles, and analytical data to Figures S34-S35 and Tables S12-S14."),
    ("Figure 6 Results", "Replace bromination with the two DPDTC protocols plus wet-lab data. Move the current bromination text and figure to ESI Figures S32-S33."),
    ("Discussion", "Remove bromination-specific evidence from the main narrative or cite ESI. Add what experimental feedback changed and clearly separate topology success from yield prediction."),
    ("Limitations", "Delete the statement that FlowPilot does not execute experiments once wet-lab validation is included. Retain that experiments are human-executed and that the system proposes screens. Update the obsolete statement that gas holdup is not modeled if the current implementation is reported."),
    ("Intensification claims", "The current code is evidence-first and no longer forces intensification without an explicit objective. Revise language claiming that every design is systematically forced toward shorter residence time."),
    ("Conclusion", "Replace the original two-case conclusion with the final experimental cases and report failed/iterative screens, not only the best yields."),
    ("Methods", "Add intake, inventory extraction/validation, final-design contract, gas-basis definitions, closed-loop update procedure, ablation endpoints, reproducibility, and wet-lab methods."),
    ("Model names and dates", "Use exact deployed model identifiers consistently. Keep architecture-effect and model-effect claims separate."),
]


CRITICAL_GATES = [
    "Correct DPDTC protocol 1 stream stoichiometry. The current topology uses equal A/B flow while describing B as 2.10 M benzylamine, which does not preserve 1.05 equivalents.",
    "Resolve the 95 °C batch condition versus the KHU inventory maximum of 80 °C. A deterministic clamp cannot be presented as chemically equivalent without a residence-time screen or verified 95 °C hardware rating.",
    "Regenerate final operating instructions after inventory reconciliation. Current retained text can say 95 °C/7 bar while the final contract says 80 °C/2.5 bar.",
    "Freeze one authoritative wet-lab data table and use it for main figures, ESI figures, captions, and JSON artifacts. Do not reconstruct values independently in multiple graphics.",
    "Retain failed, rejected, and low-yield cycles. Removing them would turn the iterative evaluation into a selected-success narrative.",
    "Use the matched Qwen and GPT architecture comparisons as the primary ablation. Label older composite-score analyses as exploratory or omit them where the metric favored no-council variants.",
    "Do not claim wet-lab yield superiority from the architecture ablation. The ablation supports constraint-compliant executability and repeatable design behavior.",
    "Confirm that Figure 5's wet-lab chemistry is the photoredox Giese/aerobic oxidation case. Do not mix THQ iterative records into that figure unless the manuscript explicitly introduces THQ as a separate chemistry.",
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)


def add_note_box(doc: Document, heading: str, text: str, *, color: str = LIGHT_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    set_cell_margins(cell, 120, 140, 120, 140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(heading)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE if color != LIGHT_RED else RED)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure_entry(doc: Document, item: dict) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"Figure S{item['n']}. {item['title']}")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    labels = ("Purpose", "Recommended panels", "Code/data source", "Production status")
    values = (item["purpose"], item["panels"], item["source"], item["status"])
    for row, label, value in zip(table.rows, labels, values):
        row.cells[0].width = Inches(1.25)
        row.cells[1].width = Inches(5.85)
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], LIGHT_GREY)
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
        row.cells[0].paragraphs[0].runs[0].bold = True
    status_cell = table.rows[3].cells[1]
    status = item["status"].lower()
    if "blocked" in status or "needs" in status:
        set_cell_shading(status_cell, LIGHT_RED)
    elif "existing" in status or "publication" in status or "move intact" in status:
        set_cell_shading(status_cell, LIGHT_TEAL)
    else:
        set_cell_shading(status_cell, LIGHT_ORANGE)


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    for name, size, color in (
        ("Title", 24, BLUE),
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, TEAL),
        ("Heading 3", 11, BLUE),
    ):
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(color)
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(5)
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = "FlowPilot | Supporting Information production plan"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GREY)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("FlowPilot Supporting Information\nFigure and Content Plan")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Prepared after review of main_manuscript.pdf and the current FlowPilot codebase\n").bold = True
    subtitle.add_run(f"Version date: {date.today().isoformat()} | Proposed scope: 38 figures, 14 tables, 12 text sections")

    add_note_box(
        doc,
        "Editorial decision captured in this plan",
        "Revised main Figure 5 will retain the photoredox Giese/aerobic-oxidation story and add wet-lab refinement data. Current main Figure 6 (alpha-bromination with inline thiosulfate quench) moves to the ESI. Revised main Figure 6 will present the two replacement DPDTC protocols with wet-lab results. The ESI preserves full protocols, calculations, logs, unsuccessful cycles, ablation evidence, and analytical details.",
        color=LIGHT_TEAL,
    )

    doc.add_heading("Executive assessment", level=1)
    p = doc.add_paragraph()
    p.add_run("The current manuscript tells a clear six-figure story: ").bold = True
    p.add_run("architecture (Figure 1), corpus/rules (Figure 2), retrieval (Figure 3), model/council behavior (Figure 4), photoredox case study (Figure 5), and bromination case study (Figure 6). The revised paper should keep the main text compact and shift detailed validation into a deliberately ordered ESI. The ESI should not merely collect unused plots; each section below supports a specific claim in the main manuscript.")

    doc.add_heading("Main-manuscript consequences", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "Location"
    hdr[1].text = "Required revision"
    for cell in hdr:
        set_cell_shading(cell, BLUE)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
    for location, revision in MANUSCRIPT_REVISIONS:
        row = table.add_row().cells
        row[0].text = location
        row[1].text = revision
        row[0].paragraphs[0].runs[0].bold = True
        for cell in row:
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)

    doc.add_heading("Recommended ESI structure", level=1)
    for section_item in SECTIONS:
        add_bullet(doc, f"{section_item['heading']}: {section_item['purpose']}")

    doc.add_page_break()
    doc.add_heading("Proposed supporting figures", level=1)
    intro = doc.add_paragraph()
    intro.add_run("Numbering principle. ").bold = True
    intro.add_run("Figures proceed in the same logical order as the manuscript: system inputs and provenance, knowledge grounding, deterministic engineering, council behavior, ablation, and finally chemical/experimental validation. Existing assets are retained only when their benchmark design and labels remain compatible with the final Methods.")

    for section_item in SECTIONS:
        doc.add_heading(section_item["heading"], level=2)
        doc.add_paragraph(section_item["purpose"])
        for figure in section_item["figures"]:
            add_figure_entry(doc, figure)

    doc.add_page_break()
    doc.add_heading("Supporting text sections", level=1)
    for number, title, content in TEXT_SECTIONS:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"{number}. {title}. ")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(BLUE)
        p.add_run(content)

    doc.add_heading("Proposed supporting tables", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, text in enumerate(("Table", "Title", "Required content")):
        cell = table.rows[0].cells[idx]
        cell.text = text
        set_cell_shading(cell, BLUE)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
    for n, title, content in TABLES:
        row = table.add_row().cells
        row[0].text = f"Table S{n}"
        row[1].text = title
        row[2].text = content
        row[0].paragraphs[0].runs[0].bold = True
        for cell in row:
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.2)

    doc.add_page_break()
    doc.add_heading("Publication gates before figure production", level=1)
    add_note_box(
        doc,
        "These are scientific consistency gates, not cosmetic edits",
        "The two DPDTC outputs presently contain conflicts that must be corrected before they are shown as experimental recommendations. Figures S36-S38 and revised main Figure 6 should be produced only from a new frozen, internally consistent dataset.",
        color=LIGHT_RED,
    )
    for gate in CRITICAL_GATES:
        add_bullet(doc, gate)

    doc.add_heading("Ablation figures to use and figures to avoid", level=1)
    p = doc.add_paragraph()
    p.add_run("Primary ESI evidence: ").bold = True
    p.add_run("the matched Qwen-one-shot versus Qwen-FlowPilot and GPT-one-shot versus GPT-FlowPilot study, the frozen endpoint/scoring specification, schema-neutral sensitivity, deployment gates, and the evidence-first validation package.")
    p = doc.add_paragraph()
    p.add_run("Exploratory only: ").bold = True
    p.add_run("the older architecture composite in which no-council or no-retrieval variants can outrank the full pipeline because formal validity, rejection behavior, and decision assurance are incompletely represented. These plots should not be used as the primary superiority evidence. They may be archived in the repository or described as development-stage diagnostics.")

    doc.add_heading("Recommended production order", level=1)
    ordered = [
        "Freeze the revised manuscript claims, exact model identifiers, and final list of wet-lab chemistries.",
        "Correct and rerun both DPDTC cases under a coherent 80 °C screen or a documented 95 °C-capable inventory; freeze feed recipes and final-design contracts.",
        "Create one master wet-lab CSV with all successful and unsuccessful cycles and generate main Figures 5-6 plus ESI Figures S34-S38 from that file.",
        "Freeze the matched ablation manifest and regenerate Figures S23-S31 using manuscript-consistent labels and vector output.",
        "Regenerate engineering and inventory Figures S11-S17 from the final code revision, then run arithmetic/topology audits.",
        "Regenerate corpus/retrieval Figures S7-S10 from the exact dataset version cited in Methods.",
        "Capture GUI/provenance Figures S3-S6 last, after the interface and artifact format are frozen.",
        "Perform a final cross-document audit: every number in the main text, ESI, tables, captions, JSON, and diagrams must derive from the same source row or result contract.",
    ]
    for i, item in enumerate(ordered, 1):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)

    doc.add_heading("Deliverable summary", level=1)
    summary = doc.add_table(rows=5, cols=2)
    summary.style = "Table Grid"
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = [
        ("Supporting figures", "38 proposed figures (S1-S38)"),
        ("Supporting tables", "14 proposed tables (S1-S14)"),
        ("Text sections", "12 method/data sections"),
        ("Main-figure change", "Figure 5 gains wet-lab refinement; Figure 6 becomes two DPDTC wet-lab protocols"),
        ("Moved figure", "Current alpha-bromination Figure 6 becomes Figure S32, with diagnostics in Figure S33"),
    ]
    for row, (left, right) in zip(summary.rows, values):
        row.cells[0].text = left
        row.cells[1].text = right
        set_cell_shading(row.cells[0], LIGHT_GREY)
        row.cells[0].paragraphs[0].runs[0].bold = True
        for cell in row.cells:
            set_cell_margins(cell)

    return doc


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.core_properties.title = "FlowPilot Supporting Information Figure and Content Plan"
    doc.core_properties.subject = "ESI plan following review of the current main manuscript and codebase"
    doc.core_properties.author = "FlowPilot project team"
    doc.core_properties.keywords = "FlowPilot, supporting information, ESI, figures, ablation, wet lab"
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
