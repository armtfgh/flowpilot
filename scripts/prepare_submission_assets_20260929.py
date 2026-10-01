"""Extract supplied scientific artwork without resampling the source records."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
import json
import struct
import requests
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt
from lxml import etree as E
import fitz

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'manuscript/Submission'
OUT=BASE/'revision_20260929'
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
R='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
A='{http://schemas.openxmlformats.org/drawingml/2006/main}'
CT='{http://schemas.openxmlformats.org/package/2006/content-types}'

def main():
    folder=OUT/'sources/khu'
    with ZipFile(BASE/'Supplementary Information (KRICT)_final (1).docx') as z:
        for name in z.namelist():
            if name.startswith('word/media/'):
                (folder/Path(name).name).write_bytes(z.read(name))
    names=['image5.emf','image10.emf','image12.emf','image16.emf','image17.emf','image18.emf','image19.emf','image20.emf']
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Inches(11.7);sec.page_height=Inches(8.3)
    sec.left_margin=sec.right_margin=Inches(.4)
    sec.top_margin=sec.bottom_margin=Inches(.4)
    replacements={}
    for i,name in enumerate(names):
        if i: doc.add_page_break()
        p=doc.add_paragraph(name);p.runs[0].font.size=Pt(10)
        # Distinct PNG placeholders ensure distinct package relationships.
        img=Image.new('RGB',(80+i,50),(255,255,255));b=BytesIO();img.save(b,format='PNG');b.seek(0)
        left,top,right,bottom=struct.unpack_from('4i',(folder/name).read_bytes(),24)
        aspect=(right-left)/(bottom-top)
        width=min(10.7,6.7*aspect); height=width/aspect
        pic=doc.add_paragraph().add_run().add_picture(b,width=Inches(width),height=Inches(height))
        rid=pic._inline.find('.//'+A+'blip').get(R+'embed')
        replacements[rid]=name
    stream=BytesIO();doc.save(stream);stream.seek(0)
    with ZipFile(stream) as z: parts={n:z.read(n) for n in z.namelist()}
    rels=E.fromstring(parts['word/_rels/document.xml.rels'])
    for rel in rels:
        if rel.get('Id') in replacements:
            name=replacements[rel.get('Id')]
            rel.set('Target','media/'+name);parts['word/media/'+name]=(folder/name).read_bytes()
    ct=E.fromstring(parts['[Content_Types].xml'])
    E.SubElement(ct,CT+'Default',Extension='emf',ContentType='image/x-emf')
    parts['[Content_Types].xml']=E.tostring(ct)
    parts['word/_rels/document.xml.rels']=E.tostring(rels)
    with ZipFile(OUT/'source_emf_gallery.docx','w',ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)
    (OUT/'emf_gallery_pages.json').write_text(json.dumps(names,indent=2)+'\n')
    # Preserve public access evidence without using credentials.
    evidence={'checked_on':'2026-09-29','repository':'https://github.com/armtfgh/flowpilot'}
    for suffix in ('','/releases','/tags','/contents/ablation_results','/contents/data'):
        url='https://api.github.com/repos/armtfgh/flowpilot'+suffix
        try:
            response=requests.get(url,timeout=35)
            evidence[suffix or 'repository_metadata']={'url':url,'status':response.status_code,'data':response.json()}
        except Exception as exc:evidence[suffix or 'repository_metadata']={'error':str(exc)}
    (OUT/'repository_access.json').write_text(json.dumps(evidence,indent=2)+'\n')
    p=fitz.open(OUT/'source_pdfs/Supplementary Information (KRICT)_final (1).pdf')
    for i in (1,3,5,8,9):p[i].get_pixmap(matrix=fitz.Matrix(1.6,1.6)).save(OUT/f'khu_page_{i+1:02d}.png')

if __name__=='__main__':main()
