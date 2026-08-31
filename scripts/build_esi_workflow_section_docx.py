#!/usr/bin/env python3
"""Build the manuscript-ready ESI workflow and provenance section."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables" / "manuscript_esi_sections_20260828"
OUT_PATH = OUT_DIR / "S1_FlowPilot_Workflow_Data_Contracts_and_Run_Provenance.docx"

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREY = "F2F2F2"
MID_GREY = "666666"
WHITE = "FFFFFF"


FIGURE_S1_PROMPT = """Create a publication-ready, two-panel scientific workflow diagram on a pure white background. Use a wide landscape layout, flat vector styling, precise alignment, restrained colors, and large readable labels. Do not use gradients, shadows, three-dimensional effects, decorative backgrounds, or fictional numerical data. The figure must look like a manually prepared scientific schematic rather than an AI illustration.

Title: "Figure S1 | FlowPilot architecture, authority, and typed contracts"

Panel A - Extended FlowPilot architecture and authority boundaries:
Show one continuous execution path with ten numbered stages: (1) Standardized intake: chemist + intake LLM; (2) Batch protocol parser; (3) Chemistry analysis: upstream LLM + rules; (4) Plan-aware literature and rule retrieval; (5) Deterministic engineering calculator + design-space search; (6) Flow-proposal LLM; (7) Multi-agent council + deterministic candidate gates; (8) Evidence and inventory reconciliation; (9) Final deterministic validation and contract compilation; (10) Executable topology, operating procedure, GUI result, and autosave. Use blue for human/intake boundaries, purple for LLM interpretation, teal for deterministic computation, orange for review gates, and green for validated published artifacts. Add a compact vertical authority ladder beside the pipeline, ordered from highest to lowest: measured experimental evidence; hard safety and inventory constraints; confirmed batch-protocol facts; chemist hypotheses; model inference. State below the ladder: "Higher-authority information constrains lower-authority proposals." At the final validation stage, show two branches: a green branch labelled "EXECUTABLE: parameters + inventory-assigned topology" and a red/amber branch labelled "BLOCKED OR INVENTORY CONFIRMATION REQUIRED: diagnostic requirements only." Add the statement: "Model confidence cannot override deterministic feasibility gates."

Panel B - Typed contracts and numerical closure:
Show the validated object chain: DesignInputPackage -> BatchRecord -> ChemistryPlan -> FlowProposal -> DesignCalculations -> inventory-assigned ProcessTopology -> FinalDesignContract. Put a small validation symbol at every object boundary and label the chain "schema validation at module boundaries." Under the chain, show a bounded revision loop: specialist or council proposes a candidate edit -> recompute dependent quantities -> reconcile exact inventory hardware -> validate arithmetic, gas basis, topology, chemistry identity, safety controls, and operating procedure -> publish or block. Make clear that a council candidate is not directly executable and that only FinalDesignContract supplies published run parameters. Do not include unstable implementation details such as field counts. Include a small legend for Human boundary, LLM interpretation, Deterministic computation, Review gate, and Published artifact.

Use lowercase and uppercase exactly as supplied for object and status names. Keep all labels horizontal and avoid crossing arrows. Use a consistent sans-serif typeface. The final figure should remain readable when printed at approximately 17 cm width."""


FIGURE_S2_PROMPT = """Create a publication-ready, three-panel scientific workflow diagram on a pure white background. Use a wide landscape layout, flat vector styling, restrained colors consistent with Figure S1, and large readable text. Do not use gradients, shadows, three-dimensional effects, decorative illustrations, or fabricated experimental results. The figure should resemble a carefully designed methods schematic.

Title: "Figure S2 | Reproducible standardized intake, GUI execution, and run provenance"

Panel A - Reproducible standardized intake:
Show a short conversation between a chemist and the intake LLM beginning with a free-text batch protocol. Beside the conversation, show the fixed question identifiers in this exact order: Q-BATCH-001 Batch protocol; Q-OBJ-001 Design objective; Q-CHEM-001 Chemistry identity; Q-HIST-001 Historical experiments; Q-INV-001 Laboratory inventory; Q-CONSTR-001 Operating limits; Q-HYP-001 Chemist hypotheses; Q-PREF-001 Output preference. Show answered and unavailable as the only resolution states for optional-information categories, while protocol, objective, and chemistry identity must be answered. Then show a deterministic readiness gate. The completed path produces a sealed object labelled "DesignInputPackage, flowpilot_intake_v1.0" containing protocol facts, objective, confirmed chemistry identity, measured evidence, inventory, operating limits, hypotheses, output preference, and question log. Add: "The LLM extracts content but cannot invent question IDs or bypass readiness rules."

Panel B - GUI execution and single-source rendering:
Show a clean GUI workflow, not a fictional screenshot: Paste protocol -> Analyze intake -> Resolve fixed questions -> Import or select inventory JSON -> Freeze package -> Run FlowPilot design -> Review final disposition. From one central green object labelled "FinalDesignContract, flowpilot_final_design_v2.0," draw arrows to the result views Summary, Engineering design, Process diagram, Stream assignments, Equipment and inventory, Safety and operating procedure, and Raw JSON. Place council deliberation and experiment history in a separate grey diagnostic/context group. Add: "Executable numerical views are derived from the same post-validation contract." Show that blocked runs display requirements topology and reasons but withhold executable parameters.

Panel C - Autosaved run provenance:
Show an ordered provenance chain: input.txt and intake_package.json -> result.json -> final_design.json -> topology.json and inventory_allocation.json -> instrument_manifest.json -> process.svg/process.png and render_manifest.json -> summary.json. Use dashed outlines for conditional diagnostic artifacts: process_requirements_topology.json, diagnostic_topology.json, diagnostic_process.svg/png, and diagnostic_render_manifest.json. Connect the final contract to the canonical topology and rendered diagram with a topology/hash consistency check. Add a small note: "Prompt logs, model snapshots, and software revision identifiers require separate campaign-level capture when exact computational replay is claimed." Do not show checksums as a default GUI artifact.

Keep the three panels visually balanced. Use blue for chemist/intake information, purple for LLM operations, teal for deterministic operations, orange for readiness or validation gates, green for executable canonical outputs, grey for diagnostic context, and red only for blocked status. All arrows must have a clear direction and all text must remain readable at approximately 17 cm figure width."""


COMPONENT_ROWS = [
    (
        "Standardized intake",
        "LLM extraction + deterministic state rules",
        "Free-text protocol and chemist answers",
        "DesignInputPackage",
        "Structures user information; fixed IDs and readiness rules remain deterministic.",
    ),
    (
        "Batch protocol parser",
        "Structured parser + schema validation",
        "Frozen protocol facts",
        "BatchRecord",
        "Normalizes batch conditions; cannot outrank the original confirmed protocol.",
    ),
    (
        "Chemistry analysis",
        "Upstream LLM + retrieved fundamentals",
        "BatchRecord, intake context",
        "ChemistryPlan",
        "Interprets mechanism, stages, sensitivities, and stream logic; does not publish hardware settings.",
    ),
    (
        "Plan-aware retrieval",
        "Deterministic filtering and ranking",
        "ChemistryPlan and retrieval query",
        "Ranked analogies and rules",
        "Supplies supporting precedents; retrieved analogies do not override measured evidence or constraints.",
    ),
    (
        "Engineering calculator and design-space search",
        "Deterministic",
        "BatchRecord, ChemistryPlan, inventory limits",
        "DesignCalculations and feasible seeds",
        "Owns arithmetic closure, candidate feasibility calculations, and bounded search outputs.",
    ),
    (
        "Flow-proposal agent",
        "Downstream LLM",
        "Chemistry, calculations, retrieval, inventory context",
        "FlowProposal",
        "Proposes a complete candidate and rationale; proposal remains non-executable until validation.",
    ),
    (
        "Multi-agent council",
        "Specialist LLMs + deterministic gates",
        "Candidate proposals and evidence",
        "Selected or revised candidate",
        "Critiques chemistry, engineering, safety, and practicality; revisions are bounded and recalculated.",
    ),
    (
        "Inventory reconciliation and topology compiler",
        "Deterministic",
        "Selected candidate and LabInventory",
        "Inventory allocation and ProcessTopology",
        "Binds required operations to declared equipment and rejects unresolved assignments.",
    ),
    (
        "Final validation and artifact compiler",
        "Deterministic",
        "Reconciled proposal, calculations, topology, safety context",
        "FinalDesignContract",
        "Owns executable/blocked disposition and canonical operating parameters.",
    ),
    (
        "GUI renderer and autosave",
        "Deterministic presentation and serialization",
        "FinalDesignContract plus diagnostic traces",
        "Result views and stored artifacts",
        "Publishes canonical outputs; diagnostic information is labelled and kept separate from run instructions.",
    ),
]


INTAKE_ROWS = [
    ("Q-BATCH-001", "Batch protocol", "Required", "Answered only", "Establishes confirmed protocol facts and batch conditions."),
    ("Q-OBJ-001", "Design objective", "Required", "Answered only", "Defines how candidates are ranked and what the first design should achieve."),
    ("Q-CHEM-001", "Chemistry identity", "Required", "Answered only", "Confirms transformation family, substrate/product identity, and bond changes."),
    ("Q-HIST-001", "Historical experiments", "Required resolution", "Answered or unavailable", "Provides measured evidence for closed-loop calibration and trend constraints."),
    ("Q-INV-001", "Laboratory inventory", "Required resolution", "Answered or unavailable", "Supplies hard equipment capabilities and discrete hardware options."),
    ("Q-CONSTR-001", "Operating limits", "Required resolution", "Answered or unavailable", "Defines pressure, temperature, flow, material, and safety boundaries."),
    ("Q-HYP-001", "Chemist hypotheses", "Required resolution", "Answered or unavailable", "Guides candidate screening but is explicitly treated below measured evidence and confirmed facts."),
    ("Q-PREF-001", "Output preference", "Optional", "Answered or omitted", "Controls presentation and requested screening format without changing physical feasibility rules."),
]


ARTIFACT_ROWS = [
    ("input.txt", "When raw user input is supplied to autosave", "Original GUI text submitted for the run."),
    ("intake_package.json", "When an intake package is used", "Frozen standardized context, answers, readiness state, and authority order."),
    ("result.json", "Every completed autosave", "Complete JSON-safe pipeline result, including available intermediate and diagnostic records."),
    ("final_design.json", "When the final contract is available", "Canonical executable or blocked FinalDesignContract used by numerical result views."),
    ("summary.json", "Every completed autosave", "Compact status and principal design fields for indexing and review."),
    ("topology.json", "When a process topology is compiled", "Stored unit-operation graph and stream connections."),
    ("inventory_allocation.json", "When inventory allocation is attempted", "Equipment assignments, unresolved requirements, and allocation status."),
    ("instrument_manifest.json", "When a manifest is produced", "Declared instruments assigned to the final design."),
    ("process.svg; process.png", "For a rendered executable topology", "Portable visual representation of the canonical process graph."),
    ("render_manifest.json", "When executable rendering is attempted", "Rendering status, artifact paths, and topology-hash metadata where available."),
    ("process_requirements_topology.json", "Conditional: unresolved inventory preflight", "Required process structure before complete equipment assignment."),
    ("diagnostic_topology.json; diagnostic_process.svg/png; diagnostic_render_manifest.json", "Conditional: blocked or diagnostic run", "Non-executable requirements view and rendering diagnostics; not operating instructions."),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def set_font(run, name: str, size: float, *, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.footer_distance = Cm(0.8)
    add_page_number(section.footer.paragraphs[0])

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(4)

    for name, size in (("Title", 16), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("1F1F1F")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)

    document.core_properties.title = "FlowPilot workflow, data contracts, and run provenance"
    document.core_properties.subject = "Electronic Supporting Information section"
    document.core_properties.author = "FlowPilot"
    document.core_properties.keywords = "FlowPilot, workflow, data contracts, standardized intake, provenance"


def add_paragraph(document: Document, text: str, *, bold_lead: str | None = None) -> None:
    paragraph = document.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)


def add_prompt_box(document: Document, label: str, prompt: str) -> None:
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(3)
    run = heading.add_run(label)
    set_font(run, "Arial", 10.5, bold=True, color=BLUE)

    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GREY)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(prompt)
    set_font(run, "Arial", 8.5)


def add_caption(document: Document, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    label_run = paragraph.add_run(label + ". ")
    set_font(label_run, "Times New Roman", 9, bold=True)
    body = paragraph.add_run(text)
    set_font(body, "Times New Roman", 9)


def add_table(document: Document, title: str, headers: list[str], rows: list[tuple[str, ...]], widths: list[float]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(title)
    set_font(run, "Times New Roman", 10, bold=True)

    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    set_repeat_header(table.rows[0])
    for index, (cell, header, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = Cm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, LIGHT_BLUE)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(header)
        set_font(run, "Arial", 8, bold=True, color=BLUE)

    for row_values in rows:
        cells = table.add_row().cells
        for cell, value, width in zip(cells, row_values, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(str(value))
            set_font(run, "Times New Roman", 7.8)


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)

    document.add_heading("S1. FlowPilot workflow, data contracts, and run provenance", level=1)
    note = document.add_paragraph()
    note.paragraph_format.space_after = Pt(8)
    note_run = note.add_run(
        "Production note: Figure-generation prompts are intentionally placed at the proposed "
        "figure locations. Replace each prompt box with the finalized artwork before submission."
    )
    set_font(note_run, "Arial", 8.5, color=MID_GREY)

    add_paragraph(
        document,
        "FlowPilot separates chemical interpretation, engineering calculation, equipment realization, "
        "and publication of executable instructions into explicit computational stages. This separation "
        "is intended to prevent an unconstrained language-model response from being presented directly "
        "as a laboratory procedure. The complete workflow begins with a standardized chemist-facing "
        "intake and ends with either a validated, inventory-assigned design or a diagnostic blocked state.",
    )
    add_paragraph(
        document,
        "The workflow is governed by an explicit information-authority order: measured experimental "
        "evidence > hard safety and inventory constraints > confirmed batch-protocol facts > chemist "
        "hypotheses > model inference. Consequently, historical measurements and declared laboratory "
        "limits constrain candidate generation and revision, whereas a hypothesis is retained as an idea "
        "to test rather than promoted to an observed fact.",
    )

    document.add_heading("S1.1 System architecture and decision authority", level=2)
    add_paragraph(
        document,
        "A FlowPilot run passes through standardized intake, batch parsing, upstream chemistry analysis, "
        "plan-aware retrieval, deterministic calculation and design-space exploration, downstream proposal "
        "generation, multi-agent council review, inventory reconciliation, final validation, and artifact "
        "publication (Figure S1A). The upstream chemistry layer identifies the transformation, mechanistic "
        "features, stage structure, stream incompatibilities, phase behavior, and candidate process "
        "bottlenecks. The retrieval layer supplies literature analogies and engineering rules relevant to "
        "that plan. These model-derived interpretations inform design, but they cannot supersede confirmed "
        "protocol facts, measurements, or hard operating limits.",
    )
    add_paragraph(
        document,
        "Deterministic components calculate and reconcile dependent quantities such as flow rate, reactor "
        "volume, residence-time basis, gas delivery, pressure settings, and inventory assignments. The "
        "council can criticize a candidate and propose bounded revisions, but a council response is not an "
        "executable output. Revised fields are recalculated and checked before publication. The final "
        "disposition therefore depends on deterministic closure rather than model confidence.",
    )
    add_paragraph(
        document,
        "Three user-facing outcomes are distinguished. An executable screening design contains canonical "
        "parameters and a fully assigned process topology. An inventory-confirmation state identifies the "
        "missing capability required before numerical execution. A blocked state records the failed "
        "requirements and diagnostic topology but withholds executable run instructions.",
    )

    document.add_heading("S1.2 Typed contracts and numerical closure", level=2)
    add_paragraph(
        document,
        "Major pipeline boundaries are represented by validated data objects rather than prose alone "
        "(Figure S1B). The principal chain is DesignInputPackage -> BatchRecord -> ChemistryPlan -> "
        "FlowProposal -> DesignCalculations -> inventory-assigned ProcessTopology -> FinalDesignContract. "
        "Natural-language rationales may accompany these objects, but downstream code consumes typed "
        "fields and performs boundary validation.",
    )
    add_paragraph(
        document,
        "The FinalDesignContract is rebuilt after engineering realization and inventory reconciliation. "
        "For an executable result, it contains canonical parameters, stage definitions, streams, instrument "
        "manifest, process graph, chemistry identity, component placement, safety controls, operating "
        "procedure, and validation experiments. If arithmetic, residence-time basis, chemistry identity, "
        "phase semantics, equipment allocation, topology, safety, procedure, or diagram consistency does "
        "not close, executable parameters remain absent from the published contract and preliminary values "
        "are retained only under a diagnostic section.",
    )

    add_prompt_box(document, "FIGURE PLACEHOLDER - GENERATION PROMPT FOR FIGURE S1", FIGURE_S1_PROMPT)
    add_caption(
        document,
        "Figure S1",
        "FlowPilot architecture, authority boundaries, and typed data contracts. (A) The end-to-end "
        "workflow separates chemist input, LLM interpretation, deterministic computation, review gates, "
        "and published artifacts. Information authority decreases from measured evidence and hard "
        "constraints to model inference. Final deterministic gates distinguish executable designs from "
        "inventory-confirmation and blocked diagnostic outputs. (B) Validated data objects are passed "
        "between modules. Candidate edits trigger recalculation, inventory reconciliation, and consistency "
        "checks before a FinalDesignContract can publish operating parameters.",
    )

    add_table(
        document,
        "Table S1. FlowPilot components, data contracts, and decision responsibilities.",
        ["Component", "Computational class", "Principal input", "Principal output", "Decision responsibility"],
        COMPONENT_ROWS,
        [2.8, 2.8, 3.4, 2.8, 5.2],
    )

    document.add_heading("S1.3 Reproducible standardized intake", level=2)
    add_paragraph(
        document,
        "The intake layer converts an initial free-text protocol and follow-up answers into a frozen "
        "DesignInputPackage. The LLM may extract batch fields and conversationally present missing "
        "questions, but it selects from a fixed bank of stable identifiers. It cannot invent a new "
        "question identifier or mark an incomplete package as ready. The current question bank is listed "
        "in Table S2.",
    )
    add_paragraph(
        document,
        "The batch protocol, design objective, and chemistry identity must be answered. Historical data, "
        "inventory, operating limits, and hypotheses must either be answered or explicitly marked "
        "unavailable. Output preference is optional. Readiness and the ordered list of pending identifiers "
        "are calculated from package state. The frozen package records the raw protocol, extracted batch "
        "fields, objective, historical evidence, inventory constraints, operating limits, hypotheses, "
        "output preferences, chemistry confirmation, question log, answers, readiness state, and authority "
        "order.",
    )
    add_paragraph(
        document,
        "This design makes the intake boundary reproducible without claiming that every generative phrase "
        "is identical across model calls. Reproducibility applies to the allowed question identifiers, "
        "resolution rules, serialized chemist responses, readiness decision, and frozen context delivered "
        "to downstream modules. Where exact computational replay is required, the frozen package must be "
        "combined with model, decoding, software-revision, and environment records.",
    )

    add_table(
        document,
        "Table S2. Fixed standardized-intake question bank and readiness rules.",
        ["Question ID", "Information category", "Readiness role", "Permitted resolution", "Downstream use"],
        INTAKE_ROWS,
        [2.6, 3.0, 2.7, 3.0, 5.7],
    )

    document.add_heading("S1.4 GUI execution and single-source result rendering", level=2)
    add_paragraph(
        document,
        "In the standardized-intake GUI, the user submits a protocol, resolves the fixed intake questions, "
        "and imports or selects a laboratory-inventory profile before initiating design. The run action is "
        "enabled only when the package is ready. The finalized package is passed to the same translation "
        "pipeline used by programmatic execution, and the resulting package is attached to the result.",
    )
    add_paragraph(
        document,
        "After final validation, executable numerical views are constructed from the post-realization "
        "FinalDesignContract. These include the summary, engineering design, stream assignments, equipment "
        "allocation, process graph, safety controls, and canonical operating procedure. Chemistry-plan, "
        "retrieval, council, and experiment-loop records remain available as supporting context or audit "
        "traces. They are not independently promoted as authoritative run parameters. For a blocked run, "
        "the GUI displays the failure reasons and a requirements or diagnostic topology while withholding "
        "an executable procedure.",
    )

    document.add_heading("S1.5 Run-level provenance and reconstruction", level=2)
    add_paragraph(
        document,
        "Completed GUI runs are automatically written to a uniquely timestamped directory. The artifact "
        "set preserves the standardized input, complete result, canonical final contract, concise summary, "
        "compiled topology, equipment allocation, instrument manifest, and available rendered diagrams "
        "(Table S3). Diagnostic topology and rendering artifacts are stored separately when the run is "
        "blocked or requires inventory confirmation.",
    )
    add_paragraph(
        document,
        "The render manifest can associate diagram files with topology-hash metadata, allowing a stored "
        "diagram to be checked against the corresponding graph where that metadata is available. The "
        "autosave package supports inspection of what was entered, calculated, validated, displayed, and "
        "assigned to hardware. It should not, by itself, be described as a complete computational-replay "
        "bundle because separate prompt logs, exact model snapshots, decoding parameters, dependency "
        "versions, and software revision identifiers are not guaranteed by the GUI autosave function.",
    )

    add_prompt_box(document, "FIGURE PLACEHOLDER - GENERATION PROMPT FOR FIGURE S2", FIGURE_S2_PROMPT)
    add_caption(
        document,
        "Figure S2",
        "Reproducible standardized intake, GUI execution, and run provenance. (A) A conversational intake "
        "LLM extracts content while fixed question identifiers and deterministic readiness rules control "
        "completion of the DesignInputPackage. (B) The frozen package and selected inventory enter the "
        "design pipeline. Executable numerical views are derived from the same post-validation "
        "FinalDesignContract, whereas blocked runs expose requirements and reasons without run parameters. "
        "(C) Autosaved artifacts link user input and intake context to the complete result, canonical final "
        "design, inventory-assigned topology, instrument manifest, rendered process scheme, and compact "
        "summary. Diagnostic files are stored separately from executable artifacts.",
    )

    add_table(
        document,
        "Table S3. GUI autosave artifacts and their provenance function.",
        ["Artifact", "Creation condition", "Provenance function"],
        ARTIFACT_ROWS,
        [4.7, 4.6, 7.7],
    )

    document.add_heading("S1.6 Scope of the reproducibility claim", level=2)
    add_paragraph(
        document,
        "The workflow provides structural reproducibility: standardized question identifiers, explicit "
        "authority labels, typed intermediate objects, deterministic feasibility gates, a single final "
        "design contract, and stored run artifacts. These controls reduce ambiguity and make inconsistencies "
        "auditable. They do not imply that stochastic language-model generations will be textually identical "
        "or that a design is experimentally validated before wet-laboratory testing. Final experimental "
        "performance remains an empirical outcome and can be returned to the historical-data field for a "
        "subsequent closed-loop refinement cycle.",
    )

    document.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    print(build())
