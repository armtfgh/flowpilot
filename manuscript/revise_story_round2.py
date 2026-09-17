"""Second prose-only pass on the accepted manuscript layout; all artwork is immutable."""
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile
import csv
import json
import re
import sys

from lxml import etree as E
from docx import Document

import revise_text_only as base

ROOT=Path(__file__).resolve().parent
SOURCE=ROOT/"manuscript_text_revised.docx"
DEST=ROOT/"manuscript_text_revised_round2.docx"
CHECK=ROOT/"text_revision_round2_checks"
W=base.W
NS=base.NS

UPDATES={
3: "Translating a batch protocol into continuous flow requires chemical reasoning, engineering calculations, and equipment choices to remain consistent. We present FlowPilot, a multi-agent AI system that connects these decisions through standardized chemist intake, literature retrieval, deterministic calculation, specialist review, and inventory reconciliation. A corpus of 464 classified literature records and 2,537 handbook-derived rules supports the workflow. The resulting design record links stage conditions, stream assignments, equipment, and the process diagram. We evaluate the architecture using matched one-shot and full-pipeline generation across three chemistry cases, five retained generator models, and three repeats, yielding 90 outcomes. Separate LLM judges apply 14 fixed criteria informed by deterministic checks. The mean benchmark score increases from 0.794 for one-shot generation to 0.917 for FlowPilot, while raw judge critical flags decrease from 98 to 4. Internal ablations examine the contribution of council review. Photoredox Giese addition followed by aerobic oxidation and telescoped DPDTC-mediated amidation are selected for prospective laboratory evaluation. The results support an auditable route from a batch protocol to a first flow experiment, while leaving reaction performance to experimental validation.",
6: "Designing a flow process is a chain of interlocking decisions across fluid dynamics[[16,17]], mass transfer[[18]], heat transfer[[19]], materials compatibility, kinetics, and safety, where each decision constrains the next[[20,21]]. A change in feed concentration affects stoichiometry and throughput; an interstage addition changes the flow entering the next reactor; and a reactor choice constrains achievable volume, pressure, and irradiation. Existing tools address different parts of this problem. Process simulators[[22]] support engineering analysis, including extensions toward automated flowsheet design[[23]]; Bayesian and active-learning approaches[[24\u201328]] optimize experimental conditions; and retrosynthesis tools[[29\u201331]] address route planning. Handbook knowledge[[20,21,32]] supplies practical flow-design principles, but is distributed across sources and often expressed qualitatively. The gap addressed here is their integration: translating a chemistry brief into a coherent first flow proposal that can be checked against the laboratory's actual equipment.",
8: "The structure of flow process design suggests a division of responsibility: LLMs interpret chemistry and propose operating choices, deterministic code calculates dependent quantities, and specialist agents review assumptions across domains. The organizing principle is therefore not simply to use more agents, but to make chemical reasoning and engineering decisions refer to the same process. This leads to the central question of the study: under the same chemistry and inventory constraints, does this coordinated workflow produce more consistent and complete flow designs than one-shot generation by the same model?",
10: "The evaluation follows that question from information quality to design quality. We first characterize the literature corpus and engineering rule base, then test whether chemistry-aware retrieval improves the relevance of selected analogies. We next compare the delivered one-shot and FlowPilot designs using a common fixed rubric across local and commercial models. Module ablations examine the role of council review, while an earlier isoxazole cycloaddition model-matrix study[[41]] characterizes configuration sensitivity. Finally, photoredox Giese addition followed by aerobic oxidation[[42]] and telescoped DPDTC-mediated amidation[[43]] provide complementary cases for prospective evaluation of connected laboratory processes. Retrieval relevance, computational design consistency, and measured chemical performance are thus treated as distinct levels of evidence.",
12: "FlowPilot organizes batch-to-flow translation around a shared process description (Figure 1). The chemist supplies the protocol, design objective, and laboratory inventory, then resolves standardized follow-up questions about missing chemistry, constraints, and experimental history. The confirmed answers form a stored input package for the run. The upstream module converts that context into a chemistry plan identifying species, reaction stages, stream incompatibilities, sensitivities, and possible flow-specific advantages. Retrieval supplies relevant precedents, and engineering calculations establish an initial proposal containing reactor geometry, stream assignments, residence-time assumptions, and operating conditions. Council review and equipment reconciliation then assess that proposal before it becomes the final reported design.",
16: "For a reconciled result, the output is a structured design package containing stage and stream parameters, equipment assignments, a process diagram, supporting analogies, and the review record. Numerical views and the diagram refer to the same final design contract, while unresolved requirements remain visible. This common record connects a chemical proposal to specific equipment and makes subsequent experimental feedback attributable to the conditions actually considered. The authority boundaries, input contracts, and provenance mechanisms are described in Figures S1-S2 and Tables S1-S3; computational closure does not itself establish reaction performance.",
23: "Coverage alone does not show that a relevant precedent will be selected. We therefore evaluated the retrieval system against a semantic-search baseline (Figure 3). The baseline embeds the batch protocol or query summary and ranks records by cosine similarity. FlowPilot enriches the query using the chemistry plan, including reaction class, mechanism, photocatalyst family, solvent, and wavelength. Filtering starts with paired records under mechanism and phase constraints, relaxes filters when fewer than three matches remain, and expands to all records if no paired match remains. A field-aware score then reranks the candidates. Figure 3a summarizes the workflow; Figure S5 and Table S6 specify its tier transitions and scoring weights.",
29: "Improved retrieval does not by itself establish that the resulting flow design is consistent. We therefore compared one-shot generation with the full FlowPilot pipeline while holding the chemistry, inventory, and requested output contract constant (Figure 4a). The three cases were CuAAC, photochemical oxidation, and hydrogenolysis, with three generation repeats per model-architecture-case combination. Separate Qwen, OpenAI, and Claude judge calls evaluated normalized final designs against 14 fixed criteria, using the same source context and deterministic verification evidence. Criteria address chemical fidelity, stream and reactor consistency, equipment feasibility, operating procedure, and uncertainty. Applicable integer ratings from 0 to 4 were averaged with equal criterion and judge weights and normalized to a score from 0 to 1. Critical flags were recorded separately. Tables S8-S10 provide the scoring anchors, criterion definitions, and module-screen summaries.",
32: "Across the five-model cohort in Table S11, full-pipeline mean scores ranged from 0.904 to 0.932, compared with 0.684 to 0.890 for matched one-shot generation (Figure 4b). The mean gain was largest for GPT-4o (+0.226), followed by Qwen3.8-27B (+0.168) and Qwen3.6-27B (+0.139); the Claude gains were smaller (+0.034 and +0.049). The judges recorded 98 critical flags across the 45 one-shot outcomes and 4 across the 45 FlowPilot outcomes; 18 and 43 outcomes, respectively, had no critical flag (Figure 4c; Table S12). The score and flag summaries therefore point in the same direction within this cohort: improved completeness and fewer flagged inconsistencies. They do not establish universal model superiority or measured reaction success. Figure 4d provides a complementary internal comparison, with council removal causing the largest score decrease among the configurations shown.",
33: "The aggregate comparison is complemented by chemistry-specific scores (Figure S7), criterion-level score differences (Figure S8), and the location of critical flags within individual runs (Figure S9). Figure S10 reports tokens, estimated generation costs, and observed runtime. The higher Qwen score-per-generation-cost values in Figure 4e therefore reflect the recorded token-price schedule, not a measurement of total deployment cost. Tables S13-S14 define numerical closure and equipment checks; Tables S15-S16 specify feedback and replay records. Figure S11 and Table S17 distinguish user-composed input from system-assembled prompts and returned artifacts for one matched run. This supports the practical input-handling argument without implying that FlowPilot uses fewer total model tokens. The legacy heuristic QA values in that worked example are not the fixed-criteria scores used in Figure 4.",
35: "The ESI also shows how the chemist supplies and inspects the information behind a design. Figures S14-S15 and Table S23 document protocol entry and fixed-ID follow-up questions. Figure S16 and Table S24 show inventory import, normalization, and equipment warnings; Figure S17 and Table S25 show stage-resolved topology and stream parameters from a reopened run. Concrete failures are retained in Table S26, and Figure S18 and Table S27 expose actual council disagreement and selection records. Figure S19 and Table S28 place these views in the complete GUI and saved-run context. These operational records connect the computational benchmark to an inspectable workflow. The remaining question, addressed by the planned cases below, is how selected proposals perform in connected laboratory reactions.",
46: "The central contribution is a coordinated design workflow, rather than a claim that a larger number of model calls is intrinsically better. The literature and rule base supply context, deterministic calculations relate dependent quantities, and the council evaluates chemical and operational choices. The evaluations address these roles at different levels: Figure 3 measures retrieval alignment, Figure 4 compares delivered designs, and the internal ablations probe the contribution of review. Together, they support the use of a shared process record and explicit checks for batch-to-flow translation. They do not imply that a plausible chemical interpretation becomes experimentally correct merely because its arithmetic closes.",
47: "The archived failures explain why this distinction matters. Table S26 includes a one-shot residence-time value that conflicts with its stated basis and a FlowPilot preparation instruction that assigns a catalyst to both a feed and a packed bed. Table S22 records a disagreement between a final structured design and an accompanying generated narrative. These are different failure modes, so a favorable aggregate score cannot substitute for inspecting the underlying records. FlowPilot's practical value is to bring chemical assumptions, calculations, equipment assignments, and review decisions into an auditable workflow in which such conflicts can be identified and corrected. The same auditability is needed for evaluator mistakes, including the checked false positive in Table S26.",
49: "Related AI systems address complementary parts of chemical research: ChemCrow[[34]] combines language-model reasoning with chemistry tools; Coscientist[[38]] investigates autonomous chemical research; El Agente[[35]] targets quantum-chemistry calculations; Zeng et al.[[36]] study multi-agent process optimization; and Materealize[[37]] addresses materials design and synthesis. Process simulation[[22]] and experimental optimization[[24\u201328]] provide further capabilities for evaluating and refining operating conditions. FlowPilot focuses on the transition from a batch procedure and laboratory constraints to a first flow-process proposal. Its distinguishing emphasis is the joint handling of chemical interpretation, engineering consistency, equipment assignment, and traceable review. These systems are not evaluated head-to-head here, so this positioning describes task scope rather than a performance ranking.",
54: "Several limitations define the scope of the evidence. First, the three benchmark chemistries cover only part of flow-process design space, and three generation repeats provide a limited estimate of variability. The retained model cohort and archived software configurations also limit extrapolation to other models or releases. The matched comparison holds the task information constant, not the total inference budget, so it evaluates the delivered workflows rather than equal-compute performance. Figures 5 and 6 are reserved for testing selected proposals against laboratory observations; no reaction yield or experimental safety claim follows from the computational scores. Integration with experimental platforms[[44]] could extend that evaluation, but chemist review remains essential.",
62: "FlowPilot connects a batch protocol and chemist-supplied constraints to a traceable first flow-process proposal. Standardized intake establishes the design context, chemistry-aware retrieval supplies precedents, deterministic calculation links dependent parameters, and specialist review assesses candidate operating choices. Equipment reconciliation and a common final record connect the selected conditions to the process diagram and supporting report.",
63: "For the five-model, three-chemistry cohort, the mean fixed-criteria score increased from 0.794 for one-shot generation to 0.917 for FlowPilot, with fewer raw judge critical flags. Module ablations support a contribution from council review, while historical model-matrix and budget studies show that this contribution depends on configuration. These findings support the coordinated workflow as a practical basis for batch-to-flow design, within the limits of the evaluated cases and rubric. The Giese/oxidation and DPDTC-amidation studies will examine the separate question of experimental performance.",
64: "For the chemist, the intended benefit is a shorter path from an initial procedure to an inspectable proposal: enter the protocol, resolve the standardized questions, select an inventory profile, and review the linked design records. Measurements can then be returned with the actual conditions to inform a subsequent design cycle. This connects proposal generation with experimental development while preserving the distinction between what was assumed, calculated, proposed, and observed.",
}


def get_package(path):
    with ZipFile(path) as z:
        return z.infolist(),{name:z.read(name) for name in z.namelist()}


def semantic_review_rows():
    return [
        ("Figure 1; Figures S1-S2; Tables S1-S3", "Architecture, intake and provenance", "Describes design contracts and recorded workflow; not wet-lab validation."),
        ("Figure 2; Figures S3-S4; Tables S4-S5", "Knowledge coverage", "Counts and associations describe metadata coverage. Existing Figure 2(c) denominator discrepancy remains flagged."),
        ("Figure 3; Figures S5-S6; Tables S6-S7", "Retrieval relevance", "Family matching and rank changes are not measurements of reaction or design success."),
        ("Figure 4; Figures S7-S10; Tables S8-S12", "Outcome and architecture evaluation", "Mean scores, flag counts and distinct repeat/between-case SD definitions checked against ESI tables. The extra GPT-5.4 bars remain disclosed."),
        ("Tables S13-S16", "Numerical closure, inventory, feedback and replay", "Reference definitions without claiming every proposed operating point is experimentally validated."),
        ("Figure S11; Table S17", "Practical input and output example", "User-facing input differs from internal prompt volume; legacy QA is not the primary benchmark score."),
        ("Figure S12; Tables S18-S21", "Historical council matrix and radar study", "Separate five-repeat engineering profile; not the independent-judge outcome score."),
        ("Figure S13; Table S22", "Structured design and narrative discrepancy", "The caption and table concern a historical SCREEN record, not measured execution."),
        ("Figures S14-S17; Tables S23-S25", "Actual GUI and inventory examples", "Recorded intake and reopened-run views; not a newly executed benchmark or wet-lab run."),
        ("Table S26", "Concrete failure mechanisms", "One-shot timing inconsistency, FlowPilot solids assignment, and evaluator false positive remain explicitly differentiated."),
        ("Figure S18; Table S27", "Council disagreement", "Actual saved model assessments; no claims about hidden reasoning or experimental correctness."),
        ("Figure S19; Table S28", "Full application and companion records", "Navigation and provenance support, not performance evidence."),
        ("Figures 5-6", "Planned two-stage laboratory cases", "Unchanged placeholders and short prospective paragraphs; no fabricated measured data."),
    ]


def audit_citations(doc, esi_doc):
    base.CHECK=CHECK
    supplemental=base.references(doc,esi_doc)
    ps=doc.xpath("/w:document/w:body/w:p",namespaces=NS)
    captions={}
    for p in ps:
        match=re.match(r"Figure (\d+)\.",base.text(p))
        if match:captions[int(match[1])]=base.text(p)
    rows=[]
    pattern=r"\bFigures?\s+((?:\d+[a-f]?)(?:\s*[-\u2013]\s*\d+[a-f]?)?(?:\s*(?:,|and)\s*\d+[a-f]?(?:\s*[-\u2013]\s*\d+[a-f]?)?)*)"
    for i,p in enumerate(ps):
        value=base.text(p)
        if re.match(r"Figure \d+\.",value):continue
        for match in re.finditer(pattern,value):
            for a,b in re.findall(r"(\d+)[a-f]?(?:\s*[-\u2013]\s*(\d+)[a-f]?)?",match[1]):
                for n in range(int(a),int(b or a)+1):
                    assert n in captions,(n,value)
                    rows.append({"paragraph":i,"reference":f"Figure {n}","caption":captions[n],"context":value})
    with (CHECK/"main_figure_cross_references.csv").open("w",newline="") as f:
        writer=csv.DictWriter(f,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)
    seen=[]
    for row in rows:
        n=int(row["reference"].split()[1])
        if n not in seen:seen.append(n)
    assert seen==list(range(1,7)),seen
    assert len(captions)==6
    with (CHECK/"citation_support_review.csv").open("w",newline="") as f:
        writer=csv.writer(f);writer.writerow(["references","supports","interpretation_checked"]);writer.writerows(semantic_review_rows())
    superscripts=doc.xpath("//w:r[w:rPr/w:vertAlign[@w:val='superscript']]/w:t/text()",namespaces=NS)
    first=[]
    for item in superscripts:
        assert re.fullmatch(r"[\d,\s\u2013-]+",item),item
        for a,b in re.findall(r"(\d+)(?:[-\u2013](\d+))?",item):
            for n in range(int(a),int(b or a)+1):
                if n not in first:first.append(n)
    assert first==list(range(1,46)),first
    esi_headings=[p.text for p in Document(base.ESI).paragraphs if p.style.name.startswith("Heading")]
    assert "Ablation test and architecture comparison" in esi_headings
    assert "Council model matrix and candidate-budget sensitivity" in esi_headings
    return {**supplemental,"main_figures_cited_in_body_in_order":seen,
            "bibliography_references_cited_in_first_appearance_order":first,
            "section_3_and_6_1_targets_checked":True}


def check_numbers():
    tables=Document(base.ESI).tables
    scores=[[c.text for c in r.cells] for r in tables[10].rows]
    aggregate=[[c.text for c in r.cells] for r in tables[11].rows]
    expected={"Qwen3.6-27B":"+0.139","Qwen3.8-27B":"+0.168","GPT-4o":"+0.226", "Claude Sonnet 4.6":"+0.034","Claude Opus 4.6":"+0.049"}
    assert {r[0]:r[3] for r in scores[1:]}==expected
    assert aggregate[1][1:]==["0.794","18/45","98"]
    assert aggregate[2][1:]==["0.917","43/45","4"]
    return {"checked_against":"unchanged ESI Tables S11-S12","candidate_outcomes":90,
            "per_architecture":45,"mean_one_shot":.794,"mean_flowpilot":.917,
            "raw_judge_flags_one_shot":98,"raw_judge_flags_flowpilot":4,
            "model_score_gains":expected,"new_scores_or_experiments":False}


def main():
    CHECK.mkdir(exist_ok=True)
    source_hash=sha256(SOURCE.read_bytes()).hexdigest()
    esi_hash=sha256(base.ESI.read_bytes()).hexdigest()
    infos,parts=get_package(SOURCE)
    original=E.fromstring(parts["word/document.xml"])
    doc=deepcopy(original)
    ps=doc.xpath("/w:document/w:body/w:p",namespaces=NS)
    oldps=original.xpath("/w:document/w:body/w:p",namespaces=NS)
    citations=next(deepcopy(r) for r in ps[5].xpath(".//w:rPr",namespaces=NS) if r.find(W+"vertAlign") is not None)
    updates=dict(UPDATES)
    updates[15]=base.marked_text(ps[15]).replace("The downstream council provides the main error-control mechanism.","The downstream council supplies chemical and cross-domain review alongside deterministic numerical checks (Figure 1d).")
    updates[26]=base.marked_text(ps[26]).replace("materially changed which analogies were presented to the design engine","changed the ordering of the candidate analogies presented to the design engine")
    changes=[]
    for i,value in sorted(updates.items()):
        before=base.text(ps[i]);base.revise(ps[i],value,citations)
        assert E.tostring(ps[i].find(W+"pPr"))==E.tostring(oldps[i].find(W+"pPr")),i
        changes.append({"paragraph":i,"before":before,"after":base.text(ps[i])})
    assert len(ps)==len(oldps)
    for i in range(len(ps)):
        if i not in updates:assert E.tostring(ps[i])==E.tostring(oldps[i]),i
    old_draw=original.xpath("//w:drawing",namespaces=NS)
    new_draw=doc.xpath("//w:drawing",namespaces=NS)
    assert len(old_draw)==len(new_draw)==6
    assert all(E.tostring(a)==E.tostring(b) for a,b in zip(old_draw,new_draw))
    for tag in ["rFonts","sz","szCs"]:
        before={tuple(sorted(e.attrib.items())) for e in original.xpath("//w:"+tag,namespaces=NS)}
        after={tuple(sorted(e.attrib.items())) for e in doc.xpath("//w:"+tag,namespaces=NS)}
        assert after.issubset(before),(tag,after-before)
    assert E.tostring(doc.find(".//"+W+"sectPr"))==E.tostring(original.find(".//"+W+"sectPr"))
    assert "[[" not in base.text(doc)
    _,esi_parts=get_package(base.ESI)
    refs=audit_citations(doc,E.fromstring(esi_parts["word/document.xml"]))
    numbers=check_numbers()
    xml=E.tostring(doc,xml_declaration=True,encoding="UTF-8",standalone=True)
    with ZipFile(DEST,"w") as z:
        for info in infos:z.writestr(info,xml if info.filename=="word/document.xml" else parts[info.filename])
    _,newparts=get_package(DEST)
    changed=[name for name in parts if parts[name]!=newparts[name]]
    assert changed==["word/document.xml"],changed
    assert source_hash==sha256(SOURCE.read_bytes()).hexdigest()
    assert esi_hash==sha256(base.ESI.read_bytes()).hexdigest()
    # Confirm the preserved published artwork still matches the original, not a regeneration.
    _,original_parts=get_package(ROOT/"manuscript.docx")
    for n in range(1,5):
        key=f"word/media/image{n}.png";assert newparts[key]==original_parts[key]
    checks={"source_round1_unchanged":True,"esi_unchanged":True,"all_six_images_and_drawing_properties_unchanged":True,
            "figures_1_to_4_match_original_manuscript":True,"fonts_sizes_styles_page_settings_unchanged":True,
            "paragraph_order_and_count_unchanged":True,"figure_5_6_sections_and_placeholders_unchanged":True,
            "changed_docx_parts":changed,"text_paragraphs_revised":len(updates),"source_sha256":source_hash,
            "output_sha256":sha256(DEST.read_bytes()).hexdigest(),"esi_sha256":esi_hash,"citations":refs,"numeric_check":numbers}
    (CHECK/"verification.json").write_text(json.dumps(checks,indent=2)+"\n")
    (CHECK/"paragraph_changes.json").write_text(json.dumps(changes,indent=2,ensure_ascii=False)+"\n")
    print(json.dumps({"output":str(DEST),"text_paragraphs_revised":len(updates),"changed_docx_parts":changed,
                      "main_figures":6,"esi_figures":19,"esi_tables":28,"bibliography_references":45,
                      "all_artwork_and_formatting_preserved":True},indent=2))


def check_pdf():
    import fitz
    from PIL import Image,ImageDraw
    pdf=fitz.open(CHECK/"manuscript_text_revised_round2.pdf")
    images=[];outside=[];caption_pages={}
    for i,p in enumerate(pdf):
        path=CHECK/f"page_{i+1:02d}.png"
        p.get_pixmap(matrix=fitz.Matrix(1.3,1.3),alpha=False).save(path);images.append(path)
        for block in p.get_text("blocks"):
            m=re.match(r"Figure\s+([1-6])\.\s+\S",block[4].lstrip())
            if m:caption_pages.setdefault(m[1],i+1)
        for word in p.get_text("words"):
            if word[0]<-1 or word[1]<-1 or word[2]>p.rect.width+1 or word[3]>p.rect.height+1:outside.append({"page":i+1,"word":word[4]})
    for start in range(0,len(images),6):
        sheet=Image.new("RGB",(1250,1800),"#d9dde0");draw=ImageDraw.Draw(sheet)
        for j,path in enumerate(images[start:start+6]):
            tile=Image.open(path);tile.thumbnail((615,565));x=(j%2)*625;y=(j//2)*600
            draw.text((x+10,y+8),f"Round 2 - page {start+j+1}",fill="black");sheet.paste(tile,(x+(625-tile.width)//2,y+30))
        sheet.save(CHECK/f"contact_{start+1:02d}.png")
    assert not outside,outside
    assert set(caption_pages)==set("123456"),caption_pages
    result={"pages":len(pdf),"out_of_page_words":outside,"caption_pages":caption_pages}
    (CHECK/"layout_check.json").write_text(json.dumps(result,indent=2)+"\n");print(json.dumps(result,indent=2))


if __name__=="__main__":
    check_pdf() if "--check-pdf" in sys.argv else main()
