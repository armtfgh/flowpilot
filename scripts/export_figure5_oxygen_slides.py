"""Six review slides generated without changing the selected design numbers."""
import argparse
import copy
import csv
from datetime import date
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from PIL import Image
from pptx.util import Inches
from export_khu_collaborator_deck import Deck, fmt, write_csv, dump
from check_khu_six_revised import check
from flora_translate.diagram_artifacts import render_topology_artifacts
from flora_translate.schemas import ProcessTopology

VM_STP = .08314 * 273.15 / 1.01325


class ReviewDeck(Deck):
    def slide(self, title, subtitle):
        slide = self.p.slides.add_slide(self.p.slide_layouts[6])
        self.text(slide, title, .55, .25, 14.9, .6, 28, True)
        self.text(slide, subtitle, .55, .94, 14.9, .5, 16, color='51636C')
        self.text(slide, f'FlowPilot | KHU inventory | {date.today().isoformat()} | Proposed conditions: laboratory review required',
                  .55, 8.62, 14.1, .26, 11, color='51636C')
        self.text(slide, str(len(self.p.slides)), 15, 8.62, .45, .26, 11)
        return slide


def alias(row):
    name = row['component'].lower()
    if row['phase'] == 'gas':
        return 'Pure O2'
    if 'trimethylsilane' in name:
        return 'Substrate 1'
    if '[ir(' in name:
        return 'Ir catalyst'
    return 'Acrylonitrile' if 'acrylonitrile' in name else row['component']


def verify(folder):
    report = check(folder / 'result.json')
    result = json.loads((folder / 'result.json').read_text())
    profile = json.loads((folder / 'inventory_profile.json').read_text())
    gas = [s for s in result['proposal']['streams'] if s['phase'] == 'gas']
    report['checks']['pure_oxygen'] = bool(gas) and all(
        re.search(r'\b(?:O2|oxygen)\b', ' '.join(s['contents']), re.I)
        and not re.search(r'\bair\b', ' '.join(s['contents']), re.I)
        and s.get('gas_reagent_mole_fraction') == 1.0 for s in gas)
    devices = {d['equipment_id']: d for d in profile['lab_inventory']['gas_hardware']}
    report['checks']['declared_mfc_grid'] = all(math.isclose(
        s['gas_flow_sccm'] / devices[s['pump_equipment_id']]['flow_rate_increment_sccm'],
        round(s['gas_flow_sccm'] / devices[s['pump_equipment_id']]['flow_rate_increment_sccm']),
        abs_tol=1e-7, rel_tol=0) for s in gas)
    for s in gas:
        limiting = sum(r['molar_flow_mmol_min'] for r in report['component_rows']
                       if 'trimethylsilane' in r['component'].lower())
        actual = s['gas_flow_sccm'] / VM_STP / limiting
        report['checks']['reported_oxygen_equivalents_close'] = math.isclose(actual, s['molar_equiv'], abs_tol=1e-4)
        for r in report['component_rows']:
            if r['phase'] == 'gas':
                r.update(molar_flow_mmol_min=s['gas_flow_sccm'] / VM_STP,
                         equivalents=s['molar_equiv'],
                         concentration_basis=f'Pure O2; inlet/STP 273.15 K, 1.01325 bar; molar volume {VM_STP:.8f} mL/mmol.')
    physics = result['scientific_assessment'].get('physics_review', {})
    calls = physics.get('tool_calls', [])
    report['checks']['both_existing_agents_used_physics'] = {c['role'] for c in calls} == {'DrFluidics', 'DrSafety'}
    fluidics = next((c for c in calls if c['role'] == 'DrFluidics'), {})
    report['checks']['all_twelve_candidates_screened'] = len(fluidics.get('request', {}).get('candidate_ids', [])) == 12
    report['checks']['assumptions_not_measurements'] = physics.get('profile', {}).get('provenance') == 'illustrative_assumptions'
    report['checks']['no_probability_claim'] = physics.get('backflow_probability') is None
    report['checks']['lab_review_retained'] = result['final_design'].get('flow_operability', {}).get('laboratory_execution_status') == 'review_required'
    report['checks']['no_automatic_hardware_installation'] = physics.get('design_modification_applied') is False
    report.update(all_passed=all(report['checks'].values()), failed=[k for k, v in report['checks'].items() if not v])
    dump(folder / 'independent_checks.json', report)
    assert report['all_passed'], (str(folder), report['failed'])
    return result, report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('campaign', type=Path)
    parser.add_argument('--attempt', default='attempt_01')
    parser.add_argument('--output-name', default='slides')
    parser.add_argument('--sets', type=int, nargs='+', default=[1, 2, 3], choices=[1, 2, 3])
    args = parser.parse_args()
    out = args.campaign / args.output_name
    out.mkdir(parents=True, exist_ok=False)
    deck = ReviewDeck()
    summaries, all_stages, all_feeds, selected_files, comparisons = [], [], [], [], []
    report_text = ['# Figure 5: pure-oxygen reruns', '',
        'Three fresh Claude runs, preserving the three KHU response sets. Upstream: Claude Opus 4.6. Council/downstream: Claude Sonnet 4.6. Each council screens 12 candidates.', '',
        '**These are proposals for laboratory engineering review, not validated or safety-cleared experimental instructions.**', '',
        '## Changes and evidence', '',
        '- Pure oxygen is required at Stage 2; Stage 1 remains oxygen-free. Original batch air exposure is preserved as a protocol fact, not used as the flow gas specification.',
        '- At least 2.0 O2 equivalents are required; the actual rounded MFC setting determines the reported dose and stage time.',
        '- MFC increment 0.01 mL/min at inlet/STP comes from Inventory_final.xlsx, MFC!D3. The unchanged v6 profile is preserved; campaign profile v7 adds this typed field.',
        '- The reported KHU incident used oxygen at 0.43 mL/min and solution at 0.02 mL/min. Reverse flow was observed. This is qualitative evidence; no yield or pressure trace was invented.',
        '- DrFluidics and DrSafety use the existing experimental transient tool; no new council agent was added.',
        '- The intake parser now recognizes excluded alternatives such as "do not use air". Regression tests cover this and mol% parsing.', '',
        '## Interpretation', '',
        'Stage 1 nominal time is V/Q_liquid. Stage 2 time is the requested inlet/STP index V/(Q_liquid + Q_O2,STP); it is not the actual pressurized gas-liquid contact time. Gas values in slides are mL/min at STP (0 C, 1.01325 bar).', '',
        'The physics tool uses an explicitly illustrative, uncalibrated assumption profile, including assumed gas supply 12 bar(g). That assumption is NOT a recommended supply setting and is above the workbook-listed MFC inlet limit of 10 bar, whose pressure reference must also be confirmed. The simulation cannot establish equipment compliance or safety. No startup timing, pressure traces, pump dynamics, valve service verification, or gas-line dead volume were measured. Backflow probability remains unspecified.', '',
        'The installed inventory check valve is on the gas branch. It does not protect the liquid branch from gas entering toward Stage 1. A simulated liquid-branch valve is hypothetical, not installed or confirmed available. Gas service compatibility, supply/outlet ratings and a validated startup/interlock and liquid-branch protection strategy need laboratory review before execution.', '']
    for number in args.sets:
        name = f'figure5_set{number}'
        folder = args.campaign / name / args.attempt
        result, checks = verify(folder)
        dest = out / name
        dest.mkdir()
        stages, feeds = checks['stage_rows'], checks['component_rows']
        topology = result['process_topology']
        dump(dest / 'process_topology.json', topology)
        write_csv(dest / 'stage_parameters.csv', stages)
        write_csv(dest / 'feed_parameters.csv', feeds)
        display = copy.deepcopy(topology)
        for op in display['unit_operations']:
            if op['op_type'] == 'pump':
                rows = [f for f in feeds if f['stream'] == op['parameters']['stream']]
                op['parameters']['contents'] = [f'{alias(r)}: {fmt(r["concentration_M"])} M' for r in rows]
        artifacts = render_topology_artifacts(ProcessTopology.model_validate(display), title='', base_dir=dest / 'render')
        assert artifacts['manifest']['renderer'] == 'graphviz'
        for suffix in ['png', 'svg']:
            shutil.copy2(artifacts[f'{suffix}_path'], dest / f'topology.{suffix}')
        dump(dest / 'display_topology.json', display)
        selected_files.append(dict(case=name, source_result=str(folder / 'result.json'),
                              sha256=hashlib.sha256((folder / 'result.json').read_bytes()).hexdigest(),
                              display_changes='Short feed aliases only; all numbers preserved.'))
        audit = result['scientific_assessment']
        physics = audit['physics_review']
        simulations = [dict(candidate_id=c['candidate_id'], **s)
                       for c in physics['results'] for s in c.get('simulations', [])]
        if simulations:
            keys = ['candidate_id', 'variant', 'sensitivity', 'scenario', 'status', 'reverse_flow_predicted',
                    'reverse_displacement_uL', 'peak_reverse_liquid_mL_min', 'peak_junction_bar_g',
                    'peak_gas_plenum_bar_g', 'evidence_id']
            write_csv(dest / 'physics_scenarios.csv', [{k: s.get(k) for k in keys} for s in simulations])
        operability = result['final_design']['flow_operability']
        dump(dest / 'operability_review.json', operability)
        dump(dest / 'reviewer_assessments.json', physics.get('reviewer_assessments', {}))
        gas = next(r for r in feeds if r['phase'] == 'gas')
        events = [json.loads(line) for line in (folder / 'llm_calls.jsonl').read_text().splitlines() if line.strip()]
        token_totals = {k: sum((e.get('usage') or {}).get(k, 0) or 0 for e in events)
                        for k in ('input_tokens', 'output_tokens')}
        run_record = json.loads((folder / 'summary.json').read_text())
        summary = dict(case=name, selected_candidate=audit['selected_candidate_id'],
                       stage1_min=stages[0]['nominal_inlet_residence_min'],
                       stage2_inlet_STP_index_min=stages[1]['nominal_inlet_residence_min'],
                       liquid_mL_min=stages[0]['liquid_flow_mL_min'], oxygen_STP_mL_min=gas['flow_mL_min'],
                       oxygen_equivalents=gas['equivalents'], simulations=len(simulations),
                       final_contract=result['final_design']['status'], lab_status='review_required',
                       independent_checks=len(checks['checks']), all_checks_passed=True,
                       observed_model_calls=len(events), elapsed_seconds=run_record['elapsed_seconds'], **token_totals)
        summaries.append(summary)
        previous = ROOT / 'outputs/khu_revised_six_20260915/collaborator_slides_20260915_163845' / name
        if (previous / 'stage_parameters.csv').exists():
            with (previous / 'stage_parameters.csv').open(newline='') as f:
                old_stages = list(csv.DictReader(f))
            for old, new in zip(old_stages, stages):
                comparisons.append(dict(case=name, stage=new['stage'],
                    previous_source=str(previous / 'stage_parameters.csv'),
                    previous_gas='air (recommended, not the gas actually used in the incident)', new_gas='pure O2',
                    previous_volume_mL=old['volume_mL'], new_volume_mL=new['volume_mL'],
                    previous_liquid_mL_min=old['liquid_flow_mL_min'], new_liquid_mL_min=new['liquid_flow_mL_min'],
                    previous_gas_STP_mL_min=old['gas_inlet_STP_mL_min'], new_gas_STP_mL_min=new['gas_inlet_STP_mL_min'],
                    previous_nominal_time_min=old['nominal_inlet_residence_min'], new_nominal_time_min=new['nominal_inlet_residence_min'],
                    previous_BPR_bar=old['BPR_bar'], new_BPR_bar=new['BPR_bar']))
        report_text += [f'## Set {number}', '',
                       '| Parameter | Stage 1 | Stage 2 |', '|---|---:|---:|',
                       f'| Volume (mL) | {fmt(stages[0]["volume_mL"])} | {fmt(stages[1]["volume_mL"])} |',
                       f'| Liquid flow (mL/min) | {fmt(stages[0]["liquid_flow_mL_min"])} | {fmt(stages[1]["liquid_flow_mL_min"])} |',
                       f'| O2 at inlet/STP (mL/min) | 0 | {fmt(gas["flow_mL_min"])} |',
                       f'| Nominal/index time (min) | {summary["stage1_min"]:.3f} | {summary["stage2_inlet_STP_index_min"]:.3f} |',
                       f'| Temperature (C) | {fmt(stages[0]["temperature_C"])} | {fmt(stages[1]["temperature_C"])} |', '',
                       f'Delivered O2: {gas["equivalents"]:.4f} equiv. Selected candidate {audit["selected_candidate_id"]}. '
                       f'{len(checks["checks"])} independent checks pass; {len(simulations)} conditional simulations archived.', '',
                       f'Elapsed time: {run_record["elapsed_seconds"] / 60:.2f} min. Observed model calls: {len(events)}. '
                       f'Provider-reported model tokens (including retries): {token_totals["input_tokens"]:,} input, '
                       f'{token_totals["output_tokens"]:,} output. These totals exclude retrieval-embedding tokens.', '',
                       '### Selected-design warnings', '']
        report_text += [f'- {f.get("message", f)}' for f in operability.get('findings', [])]
        report_text += ['', '### Reviewer proposals (not approved hardware changes)', '']
        for role, assessment in physics.get('reviewer_assessments', {}).items():
            report_text += [f'**{role}**'] + [f'- {a}' for a in assessment.get('proposed_alternatives', [])]
        report_text += ['', f'Full source: `{name}/{args.attempt}/result.json`. Full prompts/responses: `llm_calls.jsonl`; physics traces are under `model_artifacts/scientific_council/`.', '']

        slide = deck.slide(f'Figure 5 | Set {number} | Flow topology',
                           'Giese addition followed by sulfoxide oxidation | Pure oxygen added only at Stage 2')
        with Image.open(dest / 'topology.png') as image:
            w, h = image.size
        scale = min(14.9 / w, 5.9 / h)
        slide.shapes.add_picture(str(dest / 'topology.png'), Inches((16 - w * scale) / 2),
                                Inches(1.45 + (5.9 - h * scale) / 2), width=Inches(w * scale), height=Inches(h * scale))
        deck.text(slide, 'Stage 1: oxygen-free, pre-degassed feed. Gas branch: MFC > inventory check valve > Stage 2 mixer.\n'
                  'Module / coil / light assignments are bound to KHU inventory; complete feed identities on the next slide.',
                  .55, 7.4, 14.9, .66, 16)
        deck.text(slide, 'Backflow remains under review: the gas-line valve does not protect the liquid branch. Not cleared for execution.',
                  .55, 8.1, 14.9, .38, 15, True, 'A13F31')
        slide = deck.slide(f'Figure 5 | Set {number} | Process parameters',
                           'Stock concentrations before mixing | Gas flow referenced to inlet/STP')
        rows = [
            ['Reactor', *[f'{fmt(s["volume_mL"])} mL {s["material"]}' for s in stages]],
            ['Tubing ID (mm)', *[fmt(s['ID_mm']) for s in stages]],
            ['Module', *[s['module'] for s in stages]],
            ['Light (nm)', *[fmt(s['wavelength_nm']) for s in stages]],
            ['Temperature (C)', *[fmt(s['temperature_C']) for s in stages]],
            ['Liquid flow (mL/min)', *[fmt(s['liquid_flow_mL_min']) for s in stages]],
            ['O2, STP (mL/min)', *[fmt(s['gas_inlet_STP_mL_min']) for s in stages]],
            ['Nominal / index time (min)', *[f'{s["nominal_inlet_residence_min"]:.2f}' for s in stages]],
            ['BPR setting (bar g)', *[fmt(s['BPR_bar']) for s in stages]],
        ]
        deck.table(slide, ['Stage conditions', 'Stage 1', 'Stage 2'], rows, .55, 1.6,
                   [3.05, 1.87, 1.87], [.52, .52, .46, .74, .46, .46, .58, .58, .70, .52], 16)
        feed_rows = [[r['stream'], alias(r), fmt(r['concentration_M']), fmt(r['flow_mL_min']),
                      fmt(r['molar_flow_mmol_min']), fmt(r['equivalents'])] for r in feeds]
        deck.table(slide, ['Feed', 'Component', 'C\n(M)', 'Q\n(mL/min)', 'n\n(mmol/min)', 'Equiv.'],
                   feed_rows, 7.65, 1.6, [.70, 1.95, .95, 1.2, 1.5, 1.1], [.7] + [.64] * len(feed_rows), 15)
        liquid = next(r for r in feeds if r['phase'] != 'gas')
        pumps = {d['equipment_id']: d['name'] for d in json.loads((folder / 'inventory_profile.json').read_text())['lab_inventory']['pumps']}
        identity = ('Substrate 1: (((4-methoxyphenyl)thio)methyl)trimethylsilane.\n'
                    'Ir catalyst: [Ir(dF(CF3)ppy)2(dtbpy)]PF6.\n'
                    'Feed solvent: EtOH:pH 9 buffer = 5:1 (v/v).\n'
                    f'Gas feed {gas["stream"]}: pure O2; Q, molar flow and equivalents all refer to O2.\n'
                    f'Liquid pump: {pumps.get(liquid["pump"], liquid["pump"])}.\n'
                    'Gas: FF-C00 MFC, CV-3301 check valve, T-mixer.')
        deck.text(slide, identity, 7.65, 5.0, 7.8, 2.34, 15)
        deck.text(slide, f'O2 delivery: {gas["equivalents"]:.4f} equiv\nMFC setting increment: 0.01 mL/min (STP)',
                  .55, 7.22, 6.8, .56, 15, True)
        deck.text(slide, 'Time: Stage 1 = V/Q liquid; Stage 2 = V/(Q liquid + Q O2 at STP). STP: 0 C, 1.01325 bar.\n'
                  'Stage 2 is an inlet/STP index, not actual pressurized contact time. No measured yield is predicted.',
                  .55, 7.88, 14.9, .63, 15, color='51636C')
        all_stages.extend(stages)
        all_feeds.extend(feeds)
    deck.p.save(out / 'Figure5_three_sets_pure_oxygen.pptx')
    write_csv(out / 'all_stage_parameters.csv', all_stages)
    write_csv(out / 'all_feed_parameters.csv', all_feeds)
    write_csv(out / 'run_summary.csv', summaries)
    if comparisons:
        write_csv(out / 'previous_vs_new.csv', comparisons)
        report_text += ['## Previous vs new conditions', '',
                       '`previous_vs_new.csv` compares the archived 15 September collaborator slide conditions with these fresh runs. '
                       'The earlier recommended gas was air; the collaborator reported actually using oxygen in the incident. '
                       'These reruns change the confirmed gas requirement, qualitative feedback and physics review together; '
                       'the comparison is not an isolated causal ablation or evidence of improved experimental yield.', '']
    dump(out / 'source_manifest.json', selected_files)
    dump(out / 'summary.json', summaries)
    source_dir = out / 'export_source'
    source_dir.mkdir()
    for script in (Path(__file__), ROOT / 'scripts/run_figure5_oxygen_rerun.py',
                   ROOT / 'scripts/export_khu_collaborator_deck.py', ROOT / 'scripts/check_khu_six_revised.py',
                   ROOT / 'scripts/run_khu_six_revised.py', ROOT / 'scripts/check_khu_slide_render.py'):
        shutil.copy2(script, source_dir / script.name)
    (out / 'REPORT.md').write_text('\n'.join(report_text), encoding='utf-8')
    if args.sets == [1, 2, 3]:
        shutil.copy2(out / 'REPORT.md', args.campaign / 'REPORT.md')
    print(json.dumps(summaries, indent=2))
    print('PRESENTATION', out)


if __name__ == '__main__':
    main()
