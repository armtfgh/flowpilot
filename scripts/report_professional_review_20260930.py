"""Write a concise author-facing review record without changing the paper."""
from collections import Counter
from hashlib import sha256
import csv
import json

from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import professional_review_20260930 as rev


def paragraph(doc, value, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(value)
    r.bold = bold
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    return p


def heading(doc, value):
    p = paragraph(doc, value, True)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8)


def main():
    audit = json.loads((rev.OUT / 'verification.json').read_text())
    edits = json.loads((rev.OUT / 'paragraph_changes.json').read_text())
    assert audit['passed'] == audit['total']
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2.1)
    d.styles['Normal'].font.name = 'Times New Roman'
    d.styles['Normal'].font.size = Pt(11)
    heading(d, 'FlowPilot manuscript and ESI: professional review record')
    paragraph(d, '30 September 2026 | Second editorial and source-consistency review')
    paragraph(d, 'This revision strengthens the scientific narrative while preserving the accepted artwork, source evidence, numerical outcomes, and document styling. It is a reviewed draft, not a declaration that all submission requirements or experimental metadata are complete.')

    heading(d, 'Delivered documents')
    for stem in ('manuscript', 'esi'):
        paragraph(d, f'{stem}_submission_reviewed_20260930.docx')
        paragraph(d, f'{stem}_submission_reviewed_20260930_marked.docx (yellow revisions)')
    paragraph(d, 'Clean PDFs accompany the Word files in manuscript/Submission. Original input files were not overwritten. The full change log, citation audit, checks, and rendered-page review are in professional_review_20260930/.')

    heading(d, 'Principal corrections')
    items = [
        ('Story and structure', 'Tightened the abstract, results interpretation, discussion, and conclusion. Kept the natural, prose-based Introduction with no main-text comparison table. The Discussion now comprises four connected paragraphs; Methods retains eight detailed subsections, referring to ESI for full prompts, inventories, recipes, and records.'),
        ('Retrieval methods', 'Distinguished the operational retrieval workflow from the offline analyses used for Figures 3 and S6. The archived rank analysis uses inverse Euclidean-distance similarity and only photocatalyst, solvent, and wavelength fields, not all five operational fields. Its 80 queries produce 1,600 pairs and 401 rank changes. The separate family analysis has 56 queries. Both are now described as metadata-alignment analyses, not independent proof of chemical correctness.'),
        ('Experimental interpretation', 'Corrected the backflow history: the reported failed trial already used pure oxygen at 0.43 mL/min. The revised point retained oxygen at 0.090 mL/min. Clarified that a gas-line check valve does not itself establish protection of the upstream liquid branch. Kept nominal STP-based arithmetic separate from physical residence time and unconfirmed gas-controller reference conditions.'),
        ('Chemical and statistical claims', 'Defined CuAAC and DPDTC. Removed unsupported causal interpretation of the amidation yield difference, universal superiority claims, and suggestions that sample-loop experiments establish sustained steady-state operation. Separate judge calls are not represented as statistically independent validation; score differences and module effects remain descriptive.'),
        ('References and terminology', 'Moved reference 56 to the actual isoxazole case, corrected the attribution of reference 59, completed the Yasukawa citation, and labeled the Materealize source as a preprint. All 62 references remain cited in first-appearance order. Removed an empty Notes heading without inventing a competing-interests declaration.'),
        ('ESI and presentation', 'Updated the corresponding methodological and experimental passages, replaced document-editing instructions with scientific prose, clarified the source label 181 min versus the calculated index 181.82 min, refreshed 45 contents entries, and kept Tables S23 and S24 together rather than stranding their final rows.'),
    ]
    for i, (label, detail) in enumerate(items, 1):
        paragraph(d, f'{i}. {label}. {detail}')

    heading(d, 'Verification completed')
    paragraph(d, f'{audit["passed"]}/{audit["total"]} automated checks passed. The {len(edits)} recorded edits include text revisions and two pagination adjustments. All original media are byte-identical; Figure 4 matches manuscript_revised_20260928.docx, and Figures 5–6 and their captions are unchanged. Styles, page geometry, table data, and archived transcripts are retained. Revised text uses Times New Roman, 11 pt; clean and marked copies have identical text and pagination.')
    paragraph(d, 'All parsed figure/table citations resolve, including main-text citations to all 31 ESI figures and 39 ESI tables. The benchmark means (0.794 one-shot; 0.917 FlowPilot), critical-flag totals (98; 4), and model-level SDs were independently recomputed from the 90 archived outcomes. Stage times, oxygen and amine equivalents, and rounded isolated-yield arithmetic were checked. No new model calls, experiments, benchmark selection, or rescoring were performed.')
    paragraph(d, 'The clean manuscript renders to 31 pages and the ESI to 118 pages. Page-by-page contact sheets and selected enlarged pages were inspected. Automated rendering checks found no out-of-page text or empty pages. These checks do not substitute for verification of every label embedded in preserved artwork.')

    heading(d, 'Items still requiring resolution or author confirmation')
    pending = [
        ('Retrieval-family audit', 'The family helper zeros the query self-similarity before adding metadata, but does not explicitly mask the query afterward. The archived family aggregates cannot certify exclusion for every query. This is a verification gap, not evidence that a self-hit actually occurred. An audited rerun with strict masking and saved query-level rankings is needed before claiming leakage-free independent retrieval validation. The 1,600-pair rank export does independently show no query-ID self-hit. The limitation is now explicit in ESI Section 2.4; no numerical result was silently changed.'),
        ('Preserved Figure 3 labels', 'The existing artwork repeats Tier 2 for its orange block and uses workflow/quality labels that can be mistaken for the actual offline evaluation. Percentage-point gains are shown with percent signs, and some chemically recognizable catalysts have unassigned classifier labels. The revised caption and methods clarify interpretation, but final artwork should be corrected from its editable source. It was not redrawn in this round because figure preservation remains the agreed scope.'),
        ('Laboratory record', 'Obtain or confirm gas-controller reference state/calibration; pressure reference, sensor position and calibration; oxygen-free feed preparation; check-valve branch/orientation and startup conditions; sample-loop collection windows and second-feed timing; raw quantitative-NMR calculation sheets and internal-standard details; original analytical exports; and independent repeat identifiers, if experiments were repeated. Confirm tubing dimensions and the reported 95 C bath arrangement. Do not infer missing precision or experimental SD.'),
        ('Final author and release approvals', 'The two corresponding authors and affiliations are consistent. The complete contribution/competing-interests declarations, final analytical-data release, benchmark archive and reproducible software/data deposition still require author approval or release completion. The current availability statement remains explicitly dated to the prior access check.'),
        ('Figure 4 cohort choice', 'The requested restored figure displays six generators, including archived GPT-5.4 results; the headline means and ESI tables use five. This is now explicit in the caption and Methods. Retaining that distinction is acceptable as a disclosed cohort choice; using exactly one cohort everywhere would require an author-approved figure change, which was not made.'),
    ]
    for i, (label, detail) in enumerate(pending, 1):
        paragraph(d, f'{i}. {label}. {detail}')

    heading(d, 'Evidence and audit files')
    paragraph(d, 'changes.csv and paragraph_changes.json record the reasons and before/after text. verification.json records the checks; cross_reference_audit.csv maps citations to captions; contents_page_audit.json records the refreshed contents. visual_review/ contains rendered contact sheets and enlarged pages. source_manifest.json and evidence_manifest.json record source hashes.')
    paragraph(d, 'Primary bibliographic checks included Yasukawa and Kobayashi (https://pmc.ncbi.nlm.nih.gov/articles/PMC8323109/) and the Materealize preprint record (https://arxiv.org/abs/2601.15743). Chemistry terminology was checked against the Saunders amidation publication record (https://pubmed.ncbi.nlm.nih.gov/40376597/).')
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    r = footer.add_run()
    r.font.name, r.font.size = 'Times New Roman', Pt(11)
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); r._r.addnext(fld)
    path = rev.BASE / 'Review_notes_20260930.docx'
    d.save(path)
    sources = [
        rev.ROOT / 'visualization/fig3c_score_decomposition.py',
        rev.ROOT / 'visualization/fig3d_rag_quality.py',
        rev.ROOT / 'visualization/panel_data_exports/fig3c_retrieval_pairs_raw.csv',
        rev.ROOT / 'visualization/panel_data_exports/fig3d_family_match_rates.csv',
        rev.ROOT / 'deliverables/manuscript_benchmark_visualizations_20260825/source_data_revised/fig08_campaign_error_summary.csv',
        rev.BASE / 'revision_20260929/source_data/khu_reported_experimental_results.csv',
        rev.BASE / 'Supplementary Information (KRICT)_final (1).docx',
        rev.ROOT / 'outputs/figure5_pure_oxygen_physics_20260918/REPORT.md',
    ]
    manifest = {str(p.relative_to(rev.ROOT)): sha256(p.read_bytes()).hexdigest() for p in sources}
    (rev.OUT / 'evidence_manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')
    (rev.OUT / 'review_summary.json').write_text(json.dumps({'edits_by_document': dict(Counter(e['document'] for e in edits)), 'checks_passed': audit['passed'], 'checks_total': audit['total'], 'outstanding': [{'item': title, 'detail': detail} for title, detail in pending]}, indent=2) + '\n')
    print(path)


if __name__ == '__main__':
    main()
