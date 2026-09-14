"""Add real GUI captures and traceable operational examples to the revised ESI."""

import hashlib
import json
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont
from matplotlib import font_manager

from capture_esi_gui_20260909 import OUT, ROOT, RUN
from revise_current_esi_20260909 import font, text, before, populate_table

SOURCE = ROOT / "manuscript/esi_revised_20260909.docx"
OUTPUT = ROOT / "manuscript/esi_revised_GUI_additions_20260909.docx"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    shots = OUT / "screenshots"
    evidence = OUT / "evidence"
    figures = OUT / "figures"
    figures.mkdir(exist_ok=True)
    r = json.loads((evidence / "gui_archive_result.json").read_text())
    profile = json.loads((evidence / "khu_inventory_v4.json").read_text())
    intake_check = json.loads((evidence / "intake_reproducibility_check.json").read_text())
    original_hash = sha(SOURCE)
    doc = Document(SOURCE)
    refs = next(p for p in doc.paragraphs if p.text == "References")
    first_new = None
    heading_records = []
    figure_records = []

    def para(value, bold=False):
        p = before(refs, value)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08
        p.paragraph_format.keep_with_next = False
        if bold:
            for run in p.runs:
                run.bold = True
        return p

    def heading(value, level, number):
        nonlocal first_new
        p = para(value, True)
        p.style = "Heading " + str(level)
        p.paragraph_format.keep_with_next = True
        if first_new is None:
            first_new = p
            p.paragraph_format.page_break_before = True
        mark = OxmlElement("w:bookmarkStart")
        identity = str(31000 + len(heading_records))
        mark.set(qn("w:id"), identity)
        mark.set(qn("w:name"), "ESI_GUI_" + number.replace(".", "_"))
        p._p.insert(1, mark)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), identity)
        p._p.append(end)
        heading_records.append({"text": value, "number": number, "level": level, "bookmark": mark.get(qn("w:name"))})
        return p

    def table(number, caption, headers, rows, widths):
        p = para(f"Table S{number}. {caption}", True)
        p.paragraph_format.keep_with_next = True
        t = doc.add_table(rows=2, cols=len(headers))
        t.style = doc.tables[16].style
        p._p.addnext(t._tbl)
        populate_table(t, headers, rows, widths)
        if number in (23, 24, 25, 28):
            for row in list(t.rows)[:-1]:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True
        return t

    def figure(number, files, caption, widths, labels=None, suffix=""):
        first = True
        for i, (filename, width) in enumerate(zip(files, widths)):
            if labels:
                label = para(labels[i], True)
                label.paragraph_format.page_break_before = first
                label.paragraph_format.keep_with_next = True
            p = para("")
            p.paragraph_format.page_break_before = first and not labels
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.keep_with_next = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(shots / filename), width=Inches(width))
            first = False
        cap = para(f"Figure S{number}{' (continued)' if suffix == 'b' else ''}. {caption}")
        cap.paragraph_format.line_spacing = 1
        cap.paragraph_format.keep_together = True
        figure_records.append({"figure": f"S{number}{suffix}", "screenshots": files, "widths_inches": widths})
        # Companion PNG sheets contain only screenshot pixels and external panel labels.
        ims = [Image.open(shots / f).convert("RGB") for f in files]
        target = 1800
        scaled = [im.resize((round(target * width / max(widths)), round(im.height * target * width / max(widths) / im.width))) for im, width in zip(ims, widths)]
        label_height = 50 if labels else 0
        canvas = Image.new("RGB", (target + 40, sum(im.height + label_height + 20 for im in scaled) + 20), "white")
        draw = ImageDraw.Draw(canvas)
        y = 20
        label_font = ImageFont.truetype(font_manager.findfont(font_manager.FontProperties(family="DejaVu Sans", weight="bold")), 30)
        for i, im in enumerate(scaled):
            if labels:
                draw.text((20, y), labels[i], fill="black", font=label_font)
                y += label_height
            canvas.paste(im, ((canvas.width - im.width) // 2, y))
            y += im.height + 20
        canvas.save(figures / f"Figure_S{number}{suffix}.png", dpi=(300, 300))

    heading("GUI operation, inventory management, and auditable examples", 1, "7")
    para("This section complements the architecture schematics with actual browser captures and archived operational evidence. The interface was exercised on 9 September 2026 using the existing web application; no screen was generated by an image model, no API response was mocked, and no design-generation request was submitted for this documentation exercise. Protocol entry, answer saving and inventory validation were performed live. Result and council views reopen the recorded run identified below. These are demonstrations of software interaction, not new chemical experiments or additions to the benchmark sample size.")
    para(f"The GUI example is the two-stage DPDTC-mediated amidation of 3-methyl-4-nitrobenzoic acid with benzylamine, with a thioester intermediate. The archived run is {RUN}. Its input specifies 30 min at 95 C for activation, 5-10 min cooling, and another 30 min at 95 C after amine addition. This example is separate from the three benchmark chemistries in Section 3. The original run used claude-opus-4-6 upstream and claude-sonnet-4-6 downstream with 12 candidates. The captured current interface and the archived scientific-design policy are identified in the companion records; they must not be assumed to be identical to every historical benchmark configuration.")
    heading("Batch entry, follow-up questions, and saved answers", 2, "7.1")
    para("Figures S14 and S15 show the user-facing input process. The chemist enters the protocol, binds the saved laboratory profile, and reviews the questions selected by the versioned question bank. The screenshot session used deterministic question selection with LLM-assisted extraction switched off, allowing the interface sequence to be reproduced without a new provider call. This check establishes repeatability of this question-selection mode; it does not establish reproducibility of model-generated chemistry or numerical designs.")
    ids = ", ".join(intake_check["observations"][0]["question_ids"])
    para(f"Three fresh requests with the same protocol, profile and extraction mode returned identical pending question IDs and question-set hashes. The pending IDs were {ids}. The user supplied the archived objective and chemistry confirmation, supplied the archived hypotheses, and marked absent experimental history unavailable. The answer-saving endpoint returned ready_for_design = true. Complete questions, answers and intermediate packages are retained in the companion evidence, including items not simultaneously visible inside the scrolling GUI panel.")
    table(23, "Observed GUI steps and the records that preserve them.", ["User action", "Observed behavior", "Saved evidence"], [
        ["Enter batch protocol", "The protocol remains editable before analysis; model and laboratory selectors are separate controls.", "intake_initial_api.json; Figure S14a"],
        ["Analyze and answer", "Stable question IDs are displayed; absent history can be explicitly unavailable.", "intake_recheck_1-3.json; Figure S15"],
        ["Save answers", "The demonstration reaches a frozen, ready input package. Readiness concerns input completeness, not chemical validation.", "intake_answered_api.json; Figure S14b"],
        ["Run or reopen", "A new run requires an explicit design command. This documentation session instead reopened an existing run.", "capture_manifest.json; archived result.json"],
        ["Inspect and export", "Stage conditions, topology, response bindings, council records and JSON are separate views of the recorded result.", "Figures S17-S18; Table S25"],
    ], [1.1, 3.15, 2.25])
    figure(14, ["01_batch_input.png", "03_intake_saved.png"],
           "Actual FlowPilot interface: (a) entry of the complete two-stage batch protocol before analysis; (b) the completion state after saving the required answers. The panels are different moments in the same live intake demonstration. LLM-assisted extraction was disabled. Frozen means that the answered input package is ready for submission, not that a chemical design or yield has been validated. Full-screen originals are retained alongside these interface crops.", [5.2, 4.1], ["(a) Batch protocol input", "(b) Answered package"])
    figure(15, ["02_followup_questions.png"],
           "Actual follow-up question panel with stable IDs, question-bank version, question-set hash and entered answers. The visible examples concern objective and chemistry identity; other questions remain within the scrollable panel. Scrollbars and partially visible textarea contents are native interface behavior, not edited text. Complete input and answer records accompany the ESI. Three identical deterministic intake requests produced the same pending IDs and hash.", [5.1])

    heading("Inventory input, normalization, and laboratory constraints", 2, "7.2")
    para("The inventory workspace separates source material from a normalized, editable profile (Figure S16). It accepts uploaded laboratory documents and free-text equipment or constraint descriptions; the JSON editor provides a structured validation route. Users review equipment identifiers, quantities, units, capability limits and unavailable functions before saving a version or selecting Use in design. Source extraction is not itself laboratory confirmation. A schema-valid profile may still contain warnings or missing capabilities, and design-specific compatibility is checked separately.")
    para("The demonstrated structured route uses KHU Laboratory Inventory, version 4, schema flowpilot_inventory_profile_v3.0. Its seven reactor entries and three pump entries are inventory records, not a claim of seven reactor families or three physical pumps: one pump entry describes three Chemyx units. The profile combines source-derived data with subsequent chemist confirmations. The BPR named Noname has confirmed setpoints 2, 8 and 9 bar even though the original PDF did not list BPR settings. The retained source-warning text refers to that original omission and is not a complete description of the later pressure-controller entry.")
    table(24, "Examples of the KHU version-4 inventory data and their interpretation.", ["Inventory item", "Stored information", "Design consequence or qualification"], [
        ["Inline degassing", "Unavailable; inline degasser equipment explicitly forbidden. Offline pre-degassing with a sealed reservoir is an allowed alternative.", "Omit the inline device; do not equate its absence with oxygen exclusion being unnecessary."],
        ["Reactor records", "PFA: 2 mL; two 5 mL units at 1.016 mm ID; 10 mL; 5 mL at 0.762 mm ID. FEP: 10 and 20 mL.", "Assign actual volume/material/ID combinations rather than inventing an intermediate coil size."],
        ["Temperature limits", "Listed PFA coils: maximum 80 C; listed FEP coils: maximum 50 C.", "These profile limits restrict this inventory selection; they are not universal material-property limits."],
        ["Pump records", "Chemyx Fusion 100 (three units), Vapourtec SF-10, Vapourtec R2C+; separate pressure and flow limits.", "Use the chosen pump and syringe configuration. The profile minimum is not automatically a globally suitable experimental feed rate."],
        ["Pressure controller", "Chemist-confirmed BPR: setpoints 2, 8, 9 bar; maximum 10 bar.", "Preserve pressure-basis declarations and check the entire flow path, not just this device."],
        ["Mixers and connectors", "No dedicated mixer entry; one connector entry. The displayed design uses a generic T-mixer requiring verification.", "An assumed passive accessory is not evidence that a suitable three-port mixer exists. A two-port union alone cannot merge two feeds."],
        ["Source caveats", "Conflicting inch/metric tubing labels; missing radiant power; original BPR omission.", "Preserve caveats and provenance. A numerical unknown-value sentinel must not be treated as measured zero power."],
    ], [1.1, 2.8, 2.6])
    para("During the interface check, uploading this JSON through the generic Extract inventory operation with LLM assistance disabled returned an empty inventory and validation warnings. Pasting the existing JSON into the editor and selecting Validate edits retained all seven reactor and three pump records. Figure S16 therefore documents the structured validation route, not a successful PDF extraction experiment. The failed extraction response and screenshot are retained as a current input-path limitation; the source inventory was not overwritten or resaved during the capture session.")
    figure(16, ["04_inventory_source.png"], "Inventory workspace, panel (a): actual upload controls and the separate field for additional equipment or constraints. The KHU JSON file is selected here. This capture shows available controls, not evidence that every supported document format has been validated. The structured JSON-validation result is shown on the following page.", [6.0], ["(a) Inventory input"] , "a")
    figure(16, ["05_inventory_preview.png"], "Inventory workspace, panel (b): validated KHU profile after loading the structured JSON through the editor and selecting Validate edits. Equipment record counts, assigned names, warnings, and save/use controls remain visible. VALID denotes schema/profile validation, not complete hardware availability or laboratory safety. Warnings were not removed for presentation.", [6.0], ["(b) Normalized profile and warnings"], "b")

    heading("Stage-resolved output and run provenance", 2, "7.3")
    para("Figure S17 reopens the archived DPDTC run and shows its actual topology and process-summary views. Table S25 transcribes the final stage values at readable manuscript text size. The second liquid feed enters between the reactors, so the Stage 2 flow is the sum of the upstream liquid and the new amine feed. Both stages are liquid-only in this example. The generic gas-basis note visible in the GUI is not evidence of a gas feed here. For a gas-containing design, inlet/STP apparent time is a reporting convention and must not be presented as measured in-channel contact time.")
    table(25, "Archived DPDTC screening parameters shown by the GUI; not a validated synthesis procedure.", ["Quantity", "Stage 1", "Stage 2"], [
        ["Reaction", "Thioester formation", "Amide formation"],
        ["Reactor volume / material", "5 mL / PFA", "5 mL / PFA"],
        ["Tubing ID", "0.762 mm", "1.016 mm"],
        ["Cumulative liquid flow", "0.298142 mL/min", "0.372678 mL/min"],
        ["New feed at this stage", "Stream A: 0.298142 mL/min", "Stream B: 0.074536 mL/min"],
        ["Nominal liquid time", "16.7705 min", "13.4164 min"],
        ["Temperature", "80 C", "80 C"],
        ["Recorded pressure", "2 bar gauge", "2 bar gauge"],
        ["Gas feed", "None", "None"],
    ], [2.1, 2.2, 2.2])
    para("The approximately 30.19 min combined nominal time is a proposed screen, not a measured conversion time. The source batch used 95 C; the selected profile restricts the chosen PFA coils to 80 C. The UI reports LOW confidence and requires verification of the generic interstage mixer. Neither a passing contract nor council agreement establishes that reduced temperature and shorter holds preserve yield. The pump pictograms are generic: equipment names and identifiers specify the peristaltic and HPLC pumps and take precedence over the icon's internal syringe-pump label.")
    figure(17, ["08_process_topology.png", "07_process_summary.png"],
           "Actual archived-run output: (a) icon-based topology with equipment names and stage-specific values; (b) the process-summary table, including the interstage feed and cumulative Stage 2 flow. These two tabs show the same reopened result. Numerical values are also transcribed in Table S25. The generic T-mixer remains a disclosed verification requirement. The screenshot is a software screening output, not a wet-lab validated protocol. Full-resolution interface captures and original graph files accompany the ESI.", [6.25, 6.25], ["(a) Process topology", "(b) Stage and feed tables"])

    heading("Examples of critical findings and evaluator limitations", 2, "7.4")
    para("Table S26 gives traceable examples from retained one-shot and FlowPilot outcomes. They were selected to illustrate different failure mechanisms, not sampled to estimate a failure rate. A judge flag, a source-supported design inconsistency, and a deterministic software block are different observations. Multiple criteria can flag the same underlying defect. The archived benchmark scores and Figure S9 flag counts are unchanged; the arithmetic recheck below is an explicit post hoc audit, not a silent rescore.")
    table(26, "Concrete archived issues and one checked evaluator false positive.", ["Record / criterion", "Observed issue", "Interpretation"], [
        ["One-shot Opus 4.6; hydrogenolysis repeat 01; N2-D1FC65D8; UO-05", "The procedure places a vented gas-liquid separator before the sole pressure-defining BPR while describing a pressurized reactor. The original post-reactor instructions confirm this order.", "One judge flagged UO-05. The stated pressure-control and venting arrangement is incomplete as written; component pressure ratings alone do not resolve the topology."],
        ["One-shot Opus 4.6; hydrogenolysis repeat 02; N2-B6923FF7; UO-07", "The main parameter reports 20.27 min, while its own basis text gives 3.0 mL / 0.1 mL/min = 30.0 min empty-bed time and a different holdup time.", "Two judges flagged UO-07. The claimed basis and serialized value disagree. Distinct time definitions must not be collapsed into one unlabeled value."],
        ["FlowPilot Qwen3.6-27B; CuAAC repeat 03; N2-110E4388; UO-02/03/11", "The final preparation instruction charges 0.25 mmol Cu/C into a 10 mL feed even though the catalyst is also assigned to a packed-bed cartridge. The saved operating_procedure confirms the instruction.", "One judge raised three criterion flags for one solids-assignment defect. The feed-preparation compiler and heterogeneous-catalyst placement must agree; passing geometry checks is insufficient."],
        ["FlowPilot Opus 4.6; oxidation repeat 01; worked example in Section 6", "The archived council summary reports no usable candidates and a calculator center-point fallback. The raw narrative and final structured conditions then differ (Table S22).", "Operational fallback and artifact synchronization are separate limitations. A completed pipeline does not establish successful council selection or consistency of every report artifact."],
        ["FlowPilot GPT-4o; oxidation repeat 03; N2-E85D424E; UO-08", "One Qwen judge called 3.3158 sccm versus 3.3158 mL/min at STP a factor-of-1000 defect.", "This specific allegation is incorrect. These volumetric units are equal at the same standard conditions. The arithmetic gives approximately 1 O2 equivalent, not 0.001. The historical flag remains disclosed as a judge false positive."],
    ], [1.5, 2.55, 2.45])
    para("Unit audit for N2-E85D424E: oxygen feed = (3.3158 mL/min air x 0.21) / (22.414 mL/mmol) = 0.0310662 mmol/min O2. Dividing by the recorded substrate flow of 0.031068 mmol/min gives 0.99994 equivalent. The same reference temperature and pressure must be used for both volume units. This checks the stated thousand-fold allegation only; it does not validate all other aspects of that design. It also explains why the heatmap should be described as judge-reported critical flags rather than confirmed error counts.")

    heading("Recorded council discussion and disagreement", 2, "7.5")
    para("Figure S18 shows the actual Council tab for the same DPDTC run. Four domain reviewers evaluate the candidate set, after which a Skeptic and Chief record cross-domain assessment and selection. Table S27 reproduces exact saved excerpts showing a disagreement about candidate 6 and the eventual selection of candidate 1. These are explicit model-authored review outputs, not a reconstruction of hidden model reasoning and not a fabricated dialogue. The GUI groups domain reviews and synthesis into two display rounds; this must not be interpreted as two complete candidate-generation/revision cycles.")
    figure(18, ["12_council_collapsed.png"], "Actual Council tab with expandable domain, Skeptic and Chief records. WARNING states are retained. Display rounds organize the recorded contributions and are not evidence of independent physical validation or two full optimization cycles. Exact source excerpts and the complete machine-readable discussion accompany Table S27.", [6.2])
    excerpts = json.loads((evidence / "council_exact_excerpts.json").read_text())
    table(27, "Exact archived council excerpts; statements are model assessments, not endorsed chemical facts.", ["Speaker and decision", "Verbatim recorded assessment"], [[x["speaker"] + "\n" + x["subject"] + "\n" + x["recommendation"], x["exact_text"]] for x in excerpts], [1.25, 5.25])
    para("The discussion exposes uncertainty rather than proving the selected timing is correct. The Chief's phrase '60-minute batch cycle (including cooling)' is itself inconsistent with the source's two 30 min heated holds plus 5-10 min cooling. The quoted statement is preserved as an audit example. The saved Skeptic record also assumes a connector can serve an interstage T-junction; the actual port count and compatible mixing hardware must be verified rather than inferred from agreement among agents. No laboratory validation follows from these recorded assessments.")

    heading("Evidence package and capture reproducibility", 2, "7.6")
    para("The companion evidence package separates fresh GUI actions from archived generation artifacts. Full-screen originals and cropped screenshots are saved at device scale factor 2. Cropping and page arrangement are editorial; no UI text, warnings, numerical values or image contents were altered. The capture manifest records the API/build identity, request URLs and screenshot hashes. The live intake rechecks used the same saved protocol and KHU profile three times. No design job was submitted, no saved inventory version was overwritten, and no benchmark score was recalculated.")
    para("Figure S19 retains the full application context: navigation between Design studio, Inventory and Saved runs; the build/run identifier; intake, evidence and inventory state; the screening disposition and equipment warning; and the result tabs. It locates the enlarged interface crops in the overall workflow. The server address shown is a local capture endpoint, not a public service URL.")
    figure(19, ["07_process_summary_full.png"], "Full application screenshot for the reopened DPDTC run. The sidebar exposes the input, inventory and archived-run workspaces, while the header identifies the build and run. The result retains LOW confidence and the equipment-verification warning despite showing a completed screening design. This image provides navigation context; detailed numerical values are in Table S25 and full-resolution screenshots in the evidence package.", [6.3])
    table(28, "Companion records for the added operational examples.", ["Record", "Purpose"], [
        ["capture_manifest.json; screenshots/", "Actual UI capture sequence, build identity and full-resolution images."],
        ["intake_initial_api.json; intake_answered_api.json", "Live question-selection and answer-saving responses."],
        ["intake_reproducibility_check.json", "Three identical deterministic question sets and hashes."],
        ["khu_inventory_v4.json; inventory_validated_api.json", "Saved inventory source and structured validation result."],
        ["inventory_import_api.json", "Preserved empty-profile result from the generic extraction path; not presented as successful parsing."],
        ["gui_archive_result.json; gui_archive_final_design.json", "Unmodified run result and final structured conditions."],
        ["council_exact_excerpts.json; council_full_record.json", "Verbatim excerpt provenance and complete discussion records."],
        ["benchmark_examples/; selected_judge_findings.csv", "Candidate results, selected judge outputs and identities for Table S26."],
        ["gas_flag_arithmetic_check.json; SHA256SUMS.txt", "Transparent post hoc unit check and evidence-file checksums."],
    ], [3.25, 3.25])

    # Extend the existing contents control without replacing unrelated document formatting.
    sdt = doc.element.xpath(".//w:sdtContent")[0]
    toc_entries = [p for p in sdt if p.tag == qn("w:p") and p.xpath(".//w:hyperlink")]
    ref_entry = toc_entries[-1]
    major_template = deepcopy(ref_entry)
    minor_template = deepcopy(toc_entries[-2])
    for item in heading_records:
        entry = deepcopy(major_template if item["level"] == 1 else minor_template)
        link = entry.xpath(".//w:hyperlink")[0]
        link.set(qn("w:anchor"), item["bookmark"])
        texts = entry.xpath(".//w:t")
        texts[0].text = item["number"] + " " + item["text"]
        texts[-1].text = "1"
        for instr in entry.xpath(".//w:instrText"):
            instr.text = f" PAGEREF {item['bookmark']} \\h "
        ref_entry.addprevious(entry)
    ref_entry.xpath(".//w:t")[0].text = "8 References"
    inventory_note = next(p for p in doc.paragraphs if p.text.startswith("During the interface check,"))
    inventory_table_caption = next(p for p in doc.paragraphs if p.text.startswith("Table S24."))
    inventory_table_caption._p.addprevious(inventory_note._p)
    for p in doc.paragraphs:
        if p.text.startswith("Unit audit for N2-E85D424E:"):
            p.paragraph_format.keep_together = True
        if p.text == "Evidence package and capture reproducibility":
            p.paragraph_format.page_break_before = True
    for i, p in enumerate(doc.paragraphs):
        if any(p.text.startswith(f"Figure S{n}.") or p.text.startswith(f"Figure S{n} (continued).") for n in range(14, 20)):
            for following in doc.paragraphs[i + 1:]:
                if following.text.strip() or following._p.xpath(".//w:drawing"):
                    following.paragraph_format.page_break_before = True
                    break
    # Prevent inherited page-break placeholders from adding a final blank page.
    refs.paragraph_format.page_break_before = True
    doc.save(OUTPUT)
    assert sha(SOURCE) == original_hash
    checked = Document(OUTPUT)
    assert len(checked.tables) == 28
    assert len(checked.element.xpath(".//m:oMath")) == len(Document(SOURCE).element.xpath(".//m:oMath"))
    # Previously restored raw benchmark blocks remain bit-for-bit identical as text.
    old_checks = json.loads((ROOT / "deliverables/esi_revision_20260909/revision_checks.json").read_text())
    for name, expected in old_checks["raw_blocks_verified"].items():
        parts, active, identity = [], False, None
        for p in checked.paragraphs:
            start = p._p.xpath(f'.//w:bookmarkStart[@w:name="{name}"]')
            if start:
                active, identity = True, start[0].get(qn("w:id"))
            if active:
                parts.append(p.text)
                if p._p.xpath(f'.//w:bookmarkEnd[@w:id="{identity}"]'):
                    break
        assert hashlib.sha256("\n".join(parts).encode()).hexdigest() == expected["sha256"]
    checks = {"source": str(SOURCE.relative_to(ROOT)), "source_sha256": original_hash, "output_sha256": sha(OUTPUT),
              "tables": 28, "new_figures": figure_records, "new_headings": heading_records,
              "raw_blocks_preserved": True, "original_unchanged": True, "new_benchmark_generations": 0}
    (OUT / "document_checks.json").write_text(json.dumps(checks, indent=2))
    print(OUTPUT)


if __name__ == "__main__":
    main()
