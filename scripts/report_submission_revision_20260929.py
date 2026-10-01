"""Create the collaborator-facing revision checklist and pending-materials note."""
from pathlib import Path
from hashlib import sha256
import json
import csv
import shutil
from io import BytesIO
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt, Inches
from lxml import etree as E
from revise_layout_discussion_20260928 import style_tables

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'manuscript/Submission';OUT=BASE/'revision_20260929'

CHANGES=[
 ('Experimental evidence','Addressed with remaining metadata','Replaced prospective yield placeholders with KHU-reported measurements. Added ESI Section 12, Figures S23–S31, and Tables S37–S39: lamp photographs/spectra, actual setups, original comparison schemes, procedures, product characterization, four NMR spectra, and implementation/feed reconciliation. No new experiment or model run was represented as having been performed.'),
 ('Author information','Addressed; roles await author approval','Confirmed author list and affiliation assignments used consistently. Corresponding authors: Boyoung Y. Park and Gwang-Noh Ahn. Asterisks moved after names; both email addresses included. Equal contribution of Amirreza Mottafegh and Mincheol Park retained. KHU NRF grant RS-2026-25594673 added.'),
 ('Numbers, labels, and cohorts','Addressed','Figure 1 BPR unit corrected from min to bar; illustrative value retained. Figure 4b/c rebuilt from the same five-model cohort as Tables S11–S12. Table S10 restored to all 15 archived conditions (45 outcomes), rather than labeling an eight-row subset as complete. Nine underlying campaigns per model/architecture were used to independently check score means and repeat-level SDs.'),
 ('Figures 5 and 6','Addressed','New readable chemistry/process/conditions/yield layouts. Proposed configuration is labeled separately from reported operation. NMR and isolated yields are distinguished. Shared-set entries are not called independent repeats. Full archived GUI topologies are retained in Figures S20–S21; original KHU artwork and actual photographs are included separately.'),
 ('Gas time and pressure conventions','Addressed; metrology awaits KHU','Figure 5 R2 uses the explicit nominal inlet-reference index 20/(0.020 + 0.090) = 181.82 min. This is not operating-pressure residence time. Proposal STP is 273.15 K and 1 atm; actual MFC reference conditions are unconfirmed. The 7 bar cartridge BPR is distinguished from measured system-pressure ranges.'),
 ('Literature comparison','Addressed','Added main-text Table 1, condensed from Table S36, covering design task/output, equipment and stage scope, and reported validation for six related systems and FlowPilot. The comparison is descriptive, not a head-to-head performance ranking. Full Table S36 retained and updated with laboratory results.'),
 ('Repetition and traceability','Addressed','Replaced the old empty reporting templates, consolidated experimental procedures in Section 12, preserved source records, and clarified finite sample-loop operation. Existing figures and tables were retained except the explicitly targeted main-text revisions; all original embedded media remain archived in the packages.'),
 ('Code and data links','Addressed; data release incomplete','Added the confirmed public GitHub repository, benchmark-code link, and data-access index. Public repository/API access was checked on 29 September 2026. There are currently no release tags; ablation_results contains an index README rather than the complete outcome archive. No unprovided data DOI or release was invented.'),
]
PENDING=[
 ('Quantitative analytical records','KHU','Supply raw quantitative-NMR files and integration worksheets, the amount and purity of 1,3-benzodioxole, dilution/calculation details, and sample identifiers. The supplied spectra are product-characterization spectra, not the complete yield-calculation records. Original NMR/LRMS exports and final approval of reported assignments are also needed.'),
 ('Independent repeats','KHU','Confirm independent run counts and individual yields. Figure 5 Sets 1/3 and Figure 6 Sets 2/3 are grouped in the supplied result tables. No experimental SD or significance claim can be made from those grouped entries alone.'),
 ('Sample-loop and collection timing','KHU','Give sample injection and collection windows, carrier/stock transitions, dispersion handling, and benzylamine-feed start/stop timing and total volume. The 2 mL loop corresponds to 100 min injection for Figure 5, and 12.50 or 7.14 min for Figure 6; these are not the reactor residence times.'),
 ('Gas/pressure and backflow records','KHU','Confirm the MFC reference temperature/pressure and calibration; pressure-sensor position and gauge/absolute convention; check-valve branch/orientation; oxygen-free feed preparation; and startup/backflow observations. The proposal STP calculation must not be confused with a measured physical residence time or safety validation.'),
 ('Protocol and dimensional confirmation','KHU','Confirm the 95 °C water-bath description and actual tubing IDs (0.04 in = 1.016 mm; 0.093 in = 2.3622 mm). The original literature comparison uses 180 min for Giese Stage 1 versus 240 min in the supplied design prompt; amidation literature batch DPDTC is 1.10 equiv versus 1.05 in the input/current experiments. These differences are explicitly retained, not silently reconciled.'),
 ('Submission administration','All authors / corresponding authors','Approve the final CRediT roles, author spelling, affiliations, correspondence emails, funding, and any conflict-of-interest statement. Publish a fixed software release and complete data archive (or documented reviewer-access package), with a persistent identifier and updated availability text. These items cannot be completed by inference.'),
]

def main():
    d=Document();s=d.sections[0];s.top_margin=s.bottom_margin=Inches(.7);s.left_margin=s.right_margin=Inches(.8)
    for st in ('Normal','Heading 1','Heading 2'):
        d.styles[st].font.name='Times New Roman';d.styles[st].font.size=Pt(11)
    d.add_heading('Submission revision and outstanding confirmations',0)
    for r in d.paragraphs[-1].runs:r.font.name='Times New Roman';r.font.size=Pt(14)
    d.add_paragraph('29 September 2026. Revised from the three files supplied in manuscript/Submission. The user-revised source documents were not overwritten. Clean copies and yellow-highlighted revision copies are provided. The revision incorporates the KHU results but is not described as unconditionally submission-complete while the confirmations below remain open.')
    d.add_heading('Files to review',1)
    for name in ['manuscript_submission_revised_20260929.docx','esi_submission_revised_20260929.docx']:
        d.add_paragraph(name+' (clean); companion filename ending _marked.docx highlights revised text in yellow.')
    d.add_heading('Response to the collaborator',1)
    for point,status,description in CHANGES:
        p=d.add_paragraph();r=p.add_run(point+': '+status+'. ');r.bold=True;p.add_run(description)
    d.add_heading('Reported laboratory outcomes',1)
    t=d.add_table(rows=1,cols=4)
    for c,v in zip(t.rows[0].cells,['Case','Reported response sets','NMR yield','Isolated yield']):c.text=v
    for row in [('Sulfoxide 4a','Sets 1/3, grouped','97%','94%; 42.0 mg'),('Sulfoxide 4a','Set 2','98%','Not reported'),('Amide 3b','Set 1','86%','84%; 227.0 mg'),('Amide 3b','Sets 2/3, grouped','68%','Not reported')]:
        for c,v in zip(t.add_row().cells,row):c.text=v
    d.add_paragraph('The data show experimentally productive implementations, not proof that the proposals were optimal. A one-percentage-point sulfoxide difference has no established significance without repeat data. The amidation comparison changes multiple process variables and does not independently identify the rate-limiting stage.')
    d.add_heading('Materials or confirmations still needed',1)
    for i,(title,owner,detail) in enumerate(PENDING,1):
        p=d.add_paragraph();p.add_run(f'{i}. {title} ({owner}). ').bold=True;p.add_run(detail)
    d.add_heading('Verification and provenance',1)
    check=json.loads((OUT/'verification.json').read_text()) if (OUT/'verification.json').exists() else {}
    d.add_paragraph(f"Automated verification: {check.get('passed','pending')}/{check.get('total','pending')} checks passed. Checks cover unchanged originals, package structure, author consistency, figure/table references, table typography, numerical closure, score/SD recomputation, critical-flag totals, and rendered page bounds. Visual-review contact sheets and high-resolution figure previews accompany the audit. Automated checks do not establish the correctness of unprovided laboratory metadata.")
    d.add_paragraph('revision_20260929/ contains the source-document hashes, extracted original KHU artwork, unchanged source PDFs, revised figure PNG/PDF/SVG files, source-data CSVs, the complete cross-reference audit, repository-access evidence, verification results, and rendered revised documents. No missing measurement, error bar, uncertainty estimate, reference-state calibration, or independent experiment was invented.')
    d.add_paragraph('Correspondence-email source: https://pharm.khu.ac.kr/pharm_eng/user/bbs/BMSR00047/list.do?menuNo=9000008. Public code: https://github.com/armtfgh/flowpilot. Data-access index: https://github.com/armtfgh/flowpilot/tree/master/ablation_results. New characterization reference: https://doi.org/10.1039/D1CC05417A.')
    xml=E.fromstring(E.tostring(d._element))
    style_tables(xml)
    buffer=BytesIO()
    d.save(buffer)
    with ZipFile(buffer) as source, ZipFile(BASE/'Submission_revision_notes_20260929.docx','w') as target:
        for item in source.infolist():
            data=E.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True) if item.filename=='word/document.xml' else source.read(item.filename)
            target.writestr(item,data)
    md=['# Submission revision, 29 September 2026','', 'Original Submission files are unchanged. Review the clean DOCX files; use *_marked.docx for yellow-highlighted text revisions.','', '## Collaborator comments']
    for title,status,detail in CHANGES:md.extend([f'### {title}: {status}',detail,''])
    md.append('## Outstanding confirmations')
    for i,(title,owner,detail) in enumerate(PENDING,1):md.extend([f'{i}. **{title} ({owner})**: {detail}',''])
    md.extend(['## Audit',f"Automated checks: {check.get('passed')}/{check.get('total')}. See verification.json, cross_reference_audit.csv, rendered_page_map.json, and visual_review/.",'No new model benchmark or wet-lab experiment was performed during this revision.'])
    (OUT/'REVISION_REPORT.md').write_text('\n'.join(md)+'\n')
    mapping=[('KHU Figure S1A','Figure S23a; Table S37','MR16 lamp photograph, spectrum and specifications'),('KHU Figure S1B','Figure S23b; Table S37','Vapourtec lamp photograph, spectrum and specifications'),('KHU Figure S1C','Figure S23c; Table S37','Strip LED photograph, spectrum and specifications'),('KHU Figure S2','Figure S24; main Figure 5; Table S31','Photochemical setup, chemical scheme and result table'),('KHU Figure S3','Figure S25','Actual photochemical apparatus'),('KHU Figure S4','Figure S26; main Figure 6; Table S34','Amidation setup, chemical scheme and result table'),('KHU Figure S5','Figure S27','Actual amidation apparatus'),('KHU 4a proton/carbon spectra','Figures S28-S29','Native spectrum traces retained'),('KHU 3b proton/carbon spectra','Figures S30-S31','Native spectrum traces retained'),('KHU characterization structures','Section 12.5','Both product structures, NMR and LRMS text'),('KHU experimental text','Sections 12.1-12.5','All methods and analytical information incorporated; typographical fixes and missing metadata identified')]
    with (OUT/'khu_content_mapping.csv').open('w',newline='') as f:
        w=csv.writer(f);w.writerow(['Supplied content','Revised location','Description']);w.writerows(mapping)
    for stem in ('manuscript','esi'):
        path=OUT/f'rendered/{stem}_submission_revised_20260929.pdf'
        if path.exists():shutil.copy2(path,BASE/path.name)
    outputs={}
    for f in BASE.glob('*20260929*'):
        if f.is_file():outputs[f.name]={'bytes':f.stat().st_size,'sha256':sha256(f.read_bytes()).hexdigest()}
    (OUT/'delivery_manifest.json').write_text(json.dumps(outputs,indent=2)+'\n')
    print('Saved revision notes, pending-items list, content mapping and review PDFs.')

if __name__=='__main__':main()
