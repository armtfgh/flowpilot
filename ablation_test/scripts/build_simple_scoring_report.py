"""Build a concise, plain-language report for the architecture scoring system."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = (
    ROOT
    / "ablation_results"
    / "presentation"
    / "flowpilot_vs_oneshot_qwen_gpt_20260810"
)
QWEN = (
    ROOT
    / "ablation_results"
    / "studies"
    / "qwen_architecture_outcome_benchmark_v1_20260731"
)
SPEC_PATH = OUT / "frozen" / "scoring_spec.json"
FIGURE_PATH = OUT / "figures" / "03_score_calculation.png"
MD_PATH = OUT / "SCORING_SYSTEM_SIMPLE_REPORT.md"
DOCX_PATH = OUT / "SCORING_SYSTEM_SIMPLE_REPORT.docx"

INK = "17212B"
BLUE = "2468A2"
TEAL = "16847A"
ORANGE = "D17A22"
RED = "C44E52"
LIGHT = "F3F6F8"
WHITE = "FFFFFF"


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_text(cell, text: object, *, bold: bool = False, color: str = INK) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[object]], widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[index], BLUE)
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=WHITE)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(row):
            if row_index % 2:
                set_cell_shading(cells[column_index], LIGHT)
            set_cell_text(cells[column_index], value)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5 if level == 1 else 2)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(14 if level == 1 else 11)
    run.font.color.rgb = RGBColor.from_string(BLUE if level == 1 else INK)


def add_body(document: Document, text: str, *, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.05
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        first.font.name = "Arial"
        first.font.size = Pt(9)
        rest = paragraph.add_run(text[len(bold_prefix) :])
        rest.font.name = "Arial"
        rest.font.size = Pt(9)
    else:
        run = paragraph.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(9)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def load_example():
    cells = pd.read_csv(QWEN / "tables" / "cell_level_scores.csv")
    checks = pd.read_csv(QWEN / "tables" / "check_level_results.csv")
    selector = (cells["scenario_id"] == "photo_oxidation_feasible") & (cells["repeat"] == 1)
    examples = {}
    for architecture in ("one_shot", "flowpilot"):
        row = cells[selector & (cells["architecture"] == architecture)].iloc[0]
        detail = checks[
            (checks["scenario_id"] == "photo_oxidation_feasible")
            & (checks["repeat"] == 1)
            & (checks["architecture"] == architecture)
        ].copy()
        examples[architecture] = (row, detail)
    return examples


def check_summary(detail: pd.DataFrame, category: str) -> tuple[int, int, float]:
    subset = detail[detail["category"] == category]
    passed = int(subset["passed"].sum())
    total = len(subset)
    return passed, total, 100.0 * passed / total


def build_markdown(spec: dict, examples) -> None:
    one, one_checks = examples["one_shot"]
    flow, flow_checks = examples["flowpilot"]
    categories = [
        ("Inventory compliance", "inventory_compliance", "30%"),
        ("Numerical closure", "numerical_closure", "30%"),
        ("Chemistry fidelity", "chemistry_fidelity", "20%"),
        ("Process completeness", "process_completeness", "20%"),
    ]
    lines = [
        "# FlowPilot Ablation Benchmark: Simple Scoring Guide",
        "",
        "## The scoring idea in one sentence",
        "",
        "Every output is checked using deterministic pass/fail tests. Safety and executability are decided first; the weighted quality score is secondary.",
        "",
        "![Scoring workflow](figures/03_score_calculation.png)",
        "",
        "## 1. Dimension scores",
        "",
        "For each dimension:",
        "",
        "`Dimension score = 100 x passed applicable checks / all applicable checks`",
        "",
        "| Dimension | Weight | What is checked |",
        "|---|---:|---|",
        "| Inventory compliance | 30% | Listed reactors, tubing, pumps, BPR, ratings, light source, and no assumed equipment |",
        "| Numerical closure | 30% | V = Q x tau, stream sums, gas pressure correction, gas equivalents, and residence times |",
        "| Chemistry fidelity | 20% | Correct disposition, chemistry requirements, operating window, and gas identity |",
        "| Process completeness | 20% | Structured proposal, required fields and topology, feed preparation, work-up, and safety |",
        "",
        "`Quality = 0.30 x Inventory + 0.30 x Numerical + 0.20 x Chemistry + 0.20 x Process`",
        "",
        "Checks have equal value inside their dimension. Critical status does not change the dimension arithmetic; it controls executability.",
        "",
        "## 2. Executability gate",
        "",
        "A feasible design is executable only if it returns SCREEN, contains a structured proposal, passes every applicable critical check, and assumes no unlisted equipment. A high quality score cannot rescue a critical failure.",
        "",
        "Deliberately infeasible cases are evaluated by correct BLOCK rate and receive no artificial zero quality score.",
        "",
        "## 3. Exact worked example",
        "",
        "The saved Qwen photo-oxidation feasible case, repeat 1, gives:",
        "",
        "| Dimension | One-shot calculation | One-shot score | FlowPilot calculation | FlowPilot score |",
        "|---|---:|---:|---:|---:|",
    ]
    for label, category, _ in categories:
        op, ot, os = check_summary(one_checks, category)
        fp, ft, fs = check_summary(flow_checks, category)
        lines.append(f"| {label} | {op}/{ot} passed | {os:.3f} | {fp}/{ft} passed | {fs:.3f} |")
    lines.extend(
        [
            "",
            f"One-shot quality = 0.30 x {one['inventory_compliance_score']:.3f} + 0.30 x {one['numerical_closure_score']:.3f} + 0.20 x {one['chemistry_fidelity_score']:.3f} + 0.20 x {one['process_completeness_score']:.3f} = **{one['quality_score']:.3f}**.",
            "",
            f"It was still non-executable because it had **{int(one['critical_failure_count'])} critical failures**. FlowPilot passed every check and scored **{flow['quality_score']:.3f}** with no critical failures.",
            "",
            "### Why the one-shot inventory score was 71.429",
            "",
            "It passed reactor matching, temperature rating, pump-flow range, pressure rating, and light-source availability. It failed the available-BPR check and the no-assumed-equipment check. Therefore: `100 x 5/7 = 71.429`.",
            "",
            "### Why the one-shot numerical score was 42.857",
            "",
            "It passed reactor-volume closure, liquid-stream summation, and presence of a positive STP gas flow. It failed gas pressure correction, gas-equivalent closure, inlet/STP residence-time closure, and in-channel residence-time closure. Therefore: `100 x 3/7 = 42.857`.",
            "",
            "### Why chemistry fidelity was 100",
            "",
            "The example passed all 5 chemistry checks: correct SCREEN disposition, wavelength within 410-430 nm, liquid flow within 0.01-5.0 mL/min, reactor volume equal to an allowed 5 or 10 mL value, and air/oxygen as the aerobic gas. Therefore: `100 x 5/5 = 100`.",
            "",
            "### Why process completeness was 100",
            "",
            "The example passed all 16 process checks: structured proposal; six positive core fields; five required topology elements; explicit liquid-feed contents; pre-reactor preparation; post-reactor handling; and a safety/hazard statement. Therefore: `100 x 16/16 = 100`.",
            "",
            "Applicable chemistry constraints and topology requirements change with the protocol, so the denominator can differ between chemistry families.",
            "",
            "## 4. Pairwise comparison",
            "",
            "1. Executable beats non-executable.",
            "2. If both execute, a quality difference greater than 5 points determines the winner.",
            "3. If neither executes, fewer critical failures wins.",
            "4. Otherwise, the result is a tie.",
            "",
            "## Interpretation boundary",
            "",
            "This benchmark measures inventory-constrained engineering executability, numerical consistency, chemistry fidelity, and process completeness. It does not measure experimental yield, novelty, or global process optimality. It is a retrospective paired pilot.",
            "",
            f"Frozen specification: `{SPEC_PATH.relative_to(ROOT)}`",
        ]
    )
    MD_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def build_docx(spec: dict, examples) -> None:
    one, one_checks = examples["one_shot"]
    flow, flow_checks = examples["flowpilot"]
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(9)

    header = section.header.paragraphs[0]
    header.text = "FlowPilot | Simple Scoring Guide"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("FlowPilot Ablation Benchmark")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor.from_string(INK)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    run = subtitle.add_run("Simple Guide to the Scoring System")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    add_body(
        document,
        "Every model output is evaluated with the same deterministic pass/fail checks. Executability is decided first; the weighted quality score is secondary.",
    )
    image = document.add_paragraph()
    image.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image.add_run().add_picture(str(FIGURE_PATH), width=Inches(6.9))
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Primary claim: constraint-compliant executability, not wet-lab yield.")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(RED)

    document.add_page_break()
    add_heading(document, "1. How a score is calculated")
    add_body(document, "Each applicable check is pass or fail. Inside a dimension, all checks have equal arithmetic value:")
    formula = document.add_paragraph()
    formula.alignment = WD_ALIGN_PARAGRAPH.CENTER
    formula.paragraph_format.space_after = Pt(8)
    run = formula.add_run("Dimension score = 100 x passed checks / applicable checks")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(ORANGE)
    add_table(
        document,
        ["Dimension", "Weight", "Main criteria"],
        [
            ["Inventory compliance", "30%", "Listed hardware, operating ratings, no assumed equipment"],
            ["Numerical closure", "30%", "V = Q x tau, stream sums, gas correction, equivalents, residence times"],
            ["Chemistry fidelity", "20%", "Correct disposition, reaction requirements, operating window, gas identity"],
            ["Process completeness", "20%", "Required fields and topology, feeds, preparation, work-up, safety"],
        ],
        widths=[1.65, 0.7, 4.6],
    )
    add_body(document, "Final quality = 0.30 x Inventory + 0.30 x Numerical + 0.20 x Chemistry + 0.20 x Process", bold_prefix="Final quality")
    add_heading(document, "2. Executability is a separate hard gate")
    for item in spec["executable_definition"]:
        add_bullet(document, item)
    add_body(document, "A critical failure makes a feasible proposal non-executable even when its partial quality score is high.", bold_prefix="A critical failure")
    add_body(document, "Infeasible cases are evaluated separately by correct BLOCK rate. They are not assigned artificial zero quality scores.")

    document.add_page_break()
    add_heading(document, "3. Exact worked example")
    add_body(document, "Saved case: Qwen photo-oxidation feasible condition, repeat 1.")
    rows = []
    for label, category in (
        ("Inventory", "inventory_compliance"),
        ("Numerical", "numerical_closure"),
        ("Chemistry", "chemistry_fidelity"),
        ("Process", "process_completeness"),
    ):
        op, ot, os = check_summary(one_checks, category)
        fp, ft, fs = check_summary(flow_checks, category)
        rows.append([label, f"{op}/{ot}", f"{os:.3f}", f"{fp}/{ft}", f"{fs:.3f}"])
    rows.append(["Weighted quality", "-", f"{one['quality_score']:.3f}", "-", f"{flow['quality_score']:.3f}"])
    add_table(
        document,
        ["Dimension", "One-shot passed", "One-shot score", "FlowPilot passed", "FlowPilot score"],
        rows,
        widths=[1.2, 1.15, 1.15, 1.15, 1.15],
    )
    add_body(
        document,
        f"One-shot quality = 0.30 x {one['inventory_compliance_score']:.3f} + 0.30 x {one['numerical_closure_score']:.3f} + 0.20 x {one['chemistry_fidelity_score']:.3f} + 0.20 x {one['process_completeness_score']:.3f} = {one['quality_score']:.3f}.",
    )
    add_body(document, f"The one-shot proposal was non-executable because it had {int(one['critical_failure_count'])} critical failures. FlowPilot had zero critical failures.", bold_prefix="The one-shot proposal was non-executable")
    add_heading(document, "Why inventory compliance was 71.429", level=2)
    add_table(
        document,
        ["Inventory check", "One-shot result"],
        [
            ["Reactor volume and tubing ID match inventory", "Pass"],
            ["Temperature is within reactor and tubing ratings", "Pass"],
            ["Liquid flow is within an available pump range", "Pass"],
            ["BPR is one of the available settings", "Fail"],
            ["Pressure is within equipment ratings", "Pass"],
            ["Wavelength matches an available light source", "Pass"],
            ["No unlisted equipment is assumed", "Fail"],
        ],
        widths=[5.2, 1.2],
    )
    add_body(document, "Inventory score = 100 x 5/7 = 71.429.", bold_prefix="Inventory score")
    add_heading(document, "Why numerical closure was 42.857", level=2)
    add_table(
        document,
        ["Numerical check", "One-shot result"],
        [
            ["Reactor volume closes against flow and residence time", "Pass"],
            ["Liquid streams sum to total liquid flow", "Pass"],
            ["Positive inlet/STP gas flow is reported", "Pass"],
            ["Pressure-corrected in-channel gas flow closes", "Fail"],
            ["Declared gas equivalents close", "Fail"],
            ["Inlet/STP residence time closes", "Fail"],
            ["In-channel residence time closes", "Fail"],
        ],
        widths=[5.2, 1.2],
    )
    add_body(document, "Numerical score = 100 x 3/7 = 42.857.", bold_prefix="Numerical score")

    document.add_page_break()
    add_heading(document, "4. Chemistry and process criteria")
    add_body(document, "The same pass-count rule is used. Critical status affects executability, not the arithmetic value of a check inside its dimension.")
    add_heading(document, "Chemistry fidelity: 5/5 = 100", level=2)
    add_table(
        document,
        ["Chemistry check", "Required value", "Critical"],
        [
            ["Disposition", "SCREEN for the feasible case", "Yes"],
            ["Photochemical wavelength", "410-430 nm", "Yes"],
            ["Liquid flow operating window", "0.01-5.0 mL/min", "Yes"],
            ["Allowed reactor volume", "5 or 10 mL", "Yes"],
            ["Aerobic gas identity", "Air or oxygen", "Yes"],
        ],
        widths=[2.8, 2.8, 0.8],
    )
    add_body(document, "Both one-shot and FlowPilot passed all five checks, so each received 100 for chemistry fidelity.")
    add_heading(document, "Process completeness: 16/16 = 100", level=2)
    add_table(
        document,
        ["Process group", "Checks included", "Count"],
        [
            ["Proposal and core fields", "Structured proposal plus positive temperature, residence time, flow, reactor volume, tubing ID, and BPR", "7"],
            ["Required topology", "Liquid pump, oxygen MFC, gas-liquid mixer, photoreactor, and BPR", "5"],
            ["Operational description", "Explicit liquid-feed contents, pre-reactor preparation, post-reactor handling, and safety/hazard statement", "4"],
        ],
        widths=[1.7, 4.8, 0.6],
    )
    add_body(document, "Both outputs passed all sixteen checks, so each received 100 for process completeness.")
    add_body(document, "The applicable chemistry constraints and topology checks change with the protocol. Therefore, other chemistry families may have different denominators.", bold_prefix="The applicable chemistry constraints")

    document.add_page_break()
    add_heading(document, "5. How matched outputs are compared")
    ranking = [
        "An executable design beats a non-executable design.",
        "If both execute, a quality difference greater than 5 points determines the winner.",
        "If neither executes, the design with fewer critical failures wins.",
        "All remaining pairs are ties.",
    ]
    for item in ranking:
        add_bullet(document, item)
    add_heading(document, "Tolerances", level=2)
    add_table(
        document,
        ["Calculation", "Tolerance"],
        [
            ["General relative closure", "5%"],
            ["Gas calculations", "10%"],
            ["Temperature", "0.1 deg C"],
            ["Wavelength", "1 nm"],
            ["Quality-score tie interval", "+/- 5 points"],
        ],
        widths=[4.5, 1.5],
    )
    add_heading(document, "What the result means", level=2)
    add_bullet(document, "It measures inventory-constrained executability and deterministic engineering consistency.")
    add_bullet(document, "It measures whether required chemistry and process information was preserved.")
    add_bullet(document, "It does not measure experimental yield, novelty, or global process optimality.")
    add_bullet(document, "This study is a retrospective paired pilot, not a prospective preregistered trial.")
    add_body(document, f"Frozen scoring specification: {SPEC_PATH.relative_to(ROOT)}")

    document.save(DOCX_PATH)


def main() -> None:
    spec = json.loads(SPEC_PATH.read_text(encoding="utf-8"))
    examples = load_example()
    build_markdown(spec, examples)
    build_docx(spec, examples)
    print(MD_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
