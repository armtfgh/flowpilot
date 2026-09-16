"""Check original-source changes and the selected diagrams independently."""
import base64
import hashlib
import json
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from flora_translate.schemas import ProcessTopology
from flora_design.visualizer.flowsheet_builder import _main_process_path

CAMPAIGN=ROOT/'outputs/khu_revised_six_20260915'
NS={'s':'http://www.w3.org/2000/svg'}


def main():
    old_path=ROOT/'outputs/khu_revised_six_20260914/presentation/sources/Inventory_final.xlsx'
    new_path=ROOT/'inventory_khu/Inventory_final.xlsx'
    old,new=[load_workbook(p,data_only=True) for p in [old_path,new_path]]
    differences=[]
    assert old.sheetnames==new.sheetnames
    for a,b in zip(old,new):
        for row in range(1,max(a.max_row,b.max_row)+1):
            for col in range(1,max(a.max_column,b.max_column)+1):
                x,y=a.cell(row,col),b.cell(row,col)
                if x.value != y.value:
                    differences.append(dict(sheet=a.title,cell=x.coordinate,previous=x.value,updated=y.value))
    report=dict(sheets_checked=new.sheetnames,changes=differences,
        source_sha256={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [old_path,new_path]})
    (CAMPAIGN/'inventory_workbook_diff.json').write_text(json.dumps(report,indent=2,ensure_ascii=False))
    out=CAMPAIGN/'presentation'
    selected=json.loads((out/'selected_results.json').read_text())
    rows=[]
    original_icon=(ROOT/'flora_design/visualizer/icons/check_valve.png').read_bytes()
    for case in selected:
        result=json.loads((out/case['case']/'result.json').read_text())
        topology=ProcessTopology.model_validate(result['process_topology'])
        nodes=ET.parse(out/case['case']/'topology.svg').findall('.//s:g[@class="node"]',NS)
        positions={}
        for node in nodes:
            image=node.find('s:image',NS)
            if image is not None:
                positions[node.find('s:title',NS).text]=(float(image.get('x')),float(image.get('y'))+float(image.get('height').removesuffix('px'))/2)
        main=[positions[key.replace('-','_').replace(' ','_')] for key in _main_process_path(topology.unit_operations,topology.streams)
              if key.replace('-','_').replace(' ','_') in positions]
        spread=max(y for x,y in main)-min(y for x,y in main)
        assert len(main)>=4 and spread<=2
        assert all(a[0]<b[0] for a,b in zip(main,main[1:]))
        valve=None
        if case['case'].startswith('figure5'):
            valve=any(base64.b64decode(image.get('{http://www.w3.org/1999/xlink}href').split(',',1)[1])==original_icon
                for node in nodes for image in node.findall('s:image',NS)
                if image.get('{http://www.w3.org/1999/xlink}href','').startswith('data:image/png;base64,'))
            assert valve, 'Uploaded check-valve PNG missing from embedded SVG'
        rows.append(dict(case=case['case'],main_icon_vertical_spread_pt=spread,
            left_to_right=True,uploaded_check_valve_embedded=valve))
    (out/'topology_layout_audit.json').write_text(json.dumps(rows,indent=2))
    print(json.dumps(rows,indent=2))


if __name__=='__main__':main()
