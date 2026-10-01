"""Two-page paragraph map of the revised manuscript, with one bullet per paragraph."""
from pathlib import Path
import csv
import json

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'manuscript/Submission'
OUT = BASE / 'narrative_revision_20261001'
SOURCE = BASE / 'manuscript_submission_narrative_20261001.docx'

# Each entry identifies the actual revised paragraph, not a proposed section.
GROUPS = [
    ('Abstract', [
        ('Designing a continuous-flow', 'Overview: research assistant, architecture, benchmark and experimental evidence.'),
    ]),
    ('Introduction', [
        ('Continuous-flow chemistry', 'Motivation: flow opportunities require laboratory-specific process design.'),
        ('The difficulty lies', 'Problem: chemistry, stage conditions and equipment decisions are coupled.'),
        ('Digital methods address', 'Context: planning, simulation and optimization support different development tasks.'),
        ('Large language models', 'Context: scientific agents connect researcher intent to tools and operations.'),
        ('Flow chemistry and process', 'Prior work: existing flow and process agents define the comparison.'),
        ('The relevant question', 'Design requirement: unite chemistry, quantitative conditions and available hardware.'),
        ('Here we introduce', 'Contribution: FlowPilot supports the researcher through end-to-end process design.'),
        ('The organizing principle', 'Mechanism: retrieval, council and calculations share one inspectable design record.'),
        ('We evaluate FlowPilot', 'Evidence roadmap: retrieval, architecture comparison and two connected reactions.'),
    ]),
    ('Results: Architecture and Knowledge (Figures 1-3)', [
        ('FlowPilot begins by', 'Define the problem: intake connects chemistry, objectives and laboratory constraints.'),
        ('The council evaluates', 'Coordinate decisions: specialists review alternatives and recalculate revised conditions.'),
        ('Equipment assignment is', 'Realize the proposal: compatible equipment, stage data and topology agree.'),
        ('Candidate design requires', 'Ground chemical choices: corpus coverage and annotated process precedents.'),
        ('The complementary handbook', 'Ground engineering choices: rule coverage, applicability and calculation boundaries.'),
        ('A broad corpus', 'Retrieve evidence: chemistry-aware queries and progressively relaxed search filters.'),
        ('In the two query', 'Measure search changes: enriched queries and observed reranking activity.'),
        ('In a separate family', 'Interpret retrieval: better metadata alignment is not transferable kinetic evidence.'),
    ]),
    ('Results: Architecture Comparison and Experimental Transition (Figure 4)', [
        ('We next asked', 'Test architecture: matched one-shot comparisons across three complementary chemistries.'),
        ('The five-model comparison', 'Define assessment: model diversity, repeats, fixed criteria and three judges.'),
        ('The comparison holds', 'Explain fairness: matched information, unequal computation and repeat-level statistics.'),
        ('FlowPilot increased', 'Interpret gains: model-dependent improvement and competitive local-model designs.'),
        ('The critical-flag analysis', 'Assess defects: fewer judge flags, distinguished from confirmed physical failures.'),
        ('The case-specific analysis', 'Locate benefits: process integration, chemistry-specific effects and generation costs.'),
        ('To examine which parts', 'Interpret ablations: council matters; additional refinement is not uniformly beneficial.'),
        ('A separate historical', 'Assess exploration: candidate budgets, engineering profiles and selection variability.'),
        ('These assessments also', 'Connect to practice: GUI provenance, inspectable disagreements and remaining defects.'),
        ('We next examined', 'Introduce experiments: three priority sets, two realized configurations per chemistry.'),
    ]),
    ('Results: Connected-Process Experiments (Figures 5-6)', [
        ('The first case couples', 'Giese purpose: separate radical addition from oxygen-dependent sulfoxide formation.'),
        ('We implemented this', 'Giese realization: compatible irradiation, feeds and explicitly defined time indices.'),
        ('Implementation revealed', 'Giese learning: backflow feedback, revised oxygen delivery and measured yields.'),
        ('The second case tests', 'Amidation purpose: couple thioester formation, amine addition and downstream conversion.'),
        ('We implemented separate', 'Amidation realization: preserve stoichiometry and account for cumulative downstream flow.'),
        ('Set 1 gave', 'Amidation learning: relate yield differences to stage allocation and flow.'),
    ]),
    ('Discussion', [
        ('The central contribution', 'Synthesize novelty: the connected, laboratory-specific process is the design target.'),
        ('The benchmark helps', 'Interpret architecture: coordinated review matters more than agent count alone.'),
        ('The experiments address', 'Interpret experiments: apparatus behavior and complete-sequence performance require measurement.'),
        ('A useful design assistant', 'Define progression: traceable decisions and measured feedback support prospective refinement.'),
    ]),
    ('Conclusion', [
        ('FlowPilot provides an', 'Restate contribution: one researcher-facing environment links knowledge, design and hardware.'),
        ('Across the five-model architecture', 'Consolidate evidence: benchmark gains and experimentally implemented connected reactions.'),
        ('These results position', 'Take-home message: support experimental development rather than replace experimentation.'),
    ]),
    ('Methods', [
        ('The standardized workflow', 'Inputs: fixed questions, frozen packages, evidence authority and campaign separation.'),
        ('The literature resource', 'Knowledge: corpus annotation, missing metadata and handbook-derived rules.'),
        ('For operational retrieval', 'Live retrieval: filters, weighting, relaxation and source exclusion.'),
        ('The retrospective figure', 'Offline analysis: embedding comparisons, query sampling and alignment definitions.'),
        ('The engineering layer', 'Liquid calculations: molar feeds, stage balances and residence-time definitions.'),
        ('Gas equivalents were', 'Gas/thermal calculations: reference states, time bases and screening assumptions.'),
        ('The downstream Designer', 'Council procedure: candidate generation, specialist review, revisions and selection records.'),
        ('Inventory profiles were', 'Inventory procedure: form/document input, validation and equipment-capability checks.'),
        ('The final design contract', 'Output procedure: synchronized rendering, autosave, feedback and reconstruction requirements.'),
        ('The primary comparison', 'Benchmark generation: matched inputs, models, repeats and computation differences.'),
        ('For judging, final', 'Judging: masked packets, fixed ratings, applicability and equal-weight aggregation.'),
        ('For each generator', 'Statistics: repeat-level means/SD, flag aggregation and cost definitions.'),
        ('The internal ablation', 'Ablation procedure: module conditions, case variation and separate historical study.'),
        ('We implemented the two', 'Experimental procedure: sample loops, staged feeds, equipment and operating conditions.'),
        ('NMR yields were', 'Analysis: NMR basis, characterization, grouped results and outstanding metadata.'),
    ]),
]


def main():
    source = Document(SOURCE)
    eligible = []
    for i, p in enumerate(source.paragraphs):
        if p.text == 'AUTHOR INFORMATION':
            break
        if i >= 6 and p.text and not p.style.name.startswith('Heading') and not p.text.startswith('Figure '):
            eligible.append((i,p.text))
    entries = [entry for _,group in GROUPS for entry in group]
    assert len(eligible) == len(entries) == 56
    for (_, actual), (prefix, _) in zip(eligible, entries):
        assert actual.startswith(prefix), (prefix, actual[:140])

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.left_margin = section.right_margin = Inches(.70)
    section.top_margin = section.bottom_margin = Inches(.62)
    section.footer_distance = Inches(.25)
    for style in ['Normal','List Bullet']:
        s=doc.styles[style]
        s.font.name='Times New Roman';s.font.size=Pt(11)
        s.paragraph_format.space_before=Pt(0);s.paragraph_format.space_after=Pt(1.5)
        s.paragraph_format.line_spacing=Pt(13)
    p=doc.add_paragraph()
    p.add_run('FlowPilot manuscript: paragraph-by-paragraph storyline').bold=True
    p.paragraph_format.space_after=Pt(4)
    p=doc.add_paragraph('One bullet per scientific paragraph, in manuscript order. P01-P56 exclude captions, author information and references.')
    p.paragraph_format.space_after=Pt(5)
    count=0;mapping=[]
    for heading,group in GROUPS:
        if heading.startswith('Results: Connected'):
            doc.add_page_break()
        p=doc.add_paragraph()
        p.add_run(heading).bold=True
        p.paragraph_format.space_before=Pt(5)
        p.paragraph_format.space_after=Pt(2)
        p.paragraph_format.keep_with_next=True
        for prefix, value in group:
            count+=1
            p=doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent=Inches(.16)
            p.paragraph_format.first_line_indent=Inches(-.13)
            p.paragraph_format.keep_together=True
            p.add_run(f'P{count:02d} ').bold=True
            p.add_run(value)
            mapping.append({'paragraph_id':f'P{count:02d}', 'section':heading,
                            'docx_paragraph_index':eligible[count-1][0], 'outline':value,
                            'source_text':eligible[count-1][1]})
    p=section.footer.paragraphs[0]
    p.alignment=2
    p.add_run('Paragraph map | ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    doc.core_properties.title='FlowPilot manuscript paragraph map'
    doc.core_properties.subject='Purpose and content of every scientific paragraph in the revised manuscript'
    dest=BASE/'FlowPilot_paragraph_map_20261001.docx'
    doc.save(dest)
    with (OUT/'paragraph_map.csv').open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(mapping[0]));w.writeheader();w.writerows(mapping)
    (OUT/'paragraph_map.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'paragraphs':count,'file':str(dest)},indent=2))


if __name__=='__main__':
    main()
