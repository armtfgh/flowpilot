#!/usr/bin/env python3
"""Document archived case-study council reports; do not run or alter the pipeline."""
from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import shutil
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Inches, Pt
from lxml import etree as E

from revise_manuscript_cases_20260922 import Package, W, NS, text

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "manuscript"
SOURCE = BASE / "FlowPilot_ESI_october2_corrected.docx"
TARGET = BASE / "FlowPilot_ESI_october2_council_discussions.docx"
OUT = BASE / "council_discussions_figures5_6_20261002"
APPENDIX = OUT / "FlowPilot_Figures5_6_full_council_reports.docx"
FIGTEXT = BASE / "FlowPilot_Figures5_6_council_figure_text.docx"
CAMPAIGNS = {
    5: ROOT / "outputs/figure5_pure_oxygen_physics_20260918",
    6: ROOT / "outputs/khu_revised_six_20260915",
}
ROLES = ["DrChemistry", "DrKinetics", "DrFluidics", "DrSafety"]


INTRO = [
    "This section documents the recorded council reports underlying the three proposed configurations for each connected-process case (Figures S12-S17; main-text Figures 5 and 6). The Figure 5 records are the pure-oxygen reruns of 18 September 2026; the Figure 6 records are the amidation runs of 15 September 2026. These are the campaigns already cited in Sections 7.1 and 7.2, not new generations. Reported choices describe proposals at the time of computation. Actual implementation and analytical results remain separately documented in Section 8.",
    "Each run used Claude Opus 4.6 for upstream chemistry interpretation and Claude Sonnet 4.6 for downstream generation and every council role. The council therefore comprises separate role-conditioned calls to the same downstream model, not independent model families or human reviewers. Twelve inventory-constrained connected-process candidates were supplied for review. Candidate screening times were generated using fixed exploratory multipliers relative to the batch holds and reconciled to available equipment; the multipliers are not fitted rate constants, yield predictions or experimentally established intensification factors.",
    "DrChemistry examined identity, feed composition and order of addition; DrKinetics examined stage exposure and uncertainty; DrFluidics examined flows, equipment and transport; and DrSafety examined operating limits and unresolved hazards. Their structured reports used candidate identifiers, recommendations, hard-violation flags, justifications and uncertainties. The Skeptic request included the specialist reviews, and the Chief request included both those reviews and the Skeptic response. Thus the recorded exchange is a sequence of structured reports passed between roles, rather than a free-form simultaneous conversation. The Chief selected one existing candidate and recorded alternatives, answer impacts, limitations and follow-up measurements. No additional iterative design-revision round should be inferred from a response-format retry.",
    "For each selected candidate, the complete four specialist verdicts, available candidate-specific Skeptic assessment and complete Chief response are reproduced below. Shared issues are introduced before the case records. The companion file FlowPilot_Figures5_6_full_council_reports.docx contains every recorded council response for all 12 candidates in all six runs, including tool requests and unsuccessful truncated responses. The source archive retains exact response text, full prompts and requests, candidate pools, final results and source hashes. The detailed ESI account is not a substitute for that unabridged archive. Source identifiers F5-S1 to F6-S3 map to their original paths in source_manifest.json.",
    "Quoted or transcribed model statements are historical outputs, not independently verified scientific conclusions or instructions to operate. Numerical discrepancies and interpretation limits are annotated explicitly rather than silently editing the reports. Final values are checked against the archived stage/feed records. A recommendation of acceptable or prefer means the model retained a candidate for screening; it does not establish chemical success, resolve all engineering uncertainties or authorize laboratory execution.",
    "The objective-policy record is also important for interpreting the three sets. In both chemistries, Set 2 was classified as yield_priority; Sets 1 and 3 were classified as balanced. Set 3 included higher-throughput wording in the human brief but was not classified as a distinct throughput_priority policy. Consequently, the three sets are different response-conditioned selections, not a controlled ablation of three reliably distinct policy modes. Candidate identifiers are local to each run, and convergent choices are retained rather than treated as independent physical designs.",
]

COMMON = {
    5: [
        "Chemistry and exposure. The reviews kept the oxygen-free Giese step separate from oxidation and assessed the complete connected sequence. The recorded liquid feed contained substrate at 0.1 M, acrylonitrile at 0.2 M and iridium photocatalyst at 0.0005 M. Offline feed deoxygenation was retained because no inline degasser was assigned. The council considered the available UV-150-compatible Stage 1 coils and the separate Manual 2 oxidation module with its own light. It did not establish photon-dose equivalence, reaction kinetics or a predicted sulfoxide yield. A larger coil was considered as a way to change Stage 1 exposure, not as proof of better conversion.",
        "Gas identity and dose. Pure oxygen at the Stage 2 inlet was already imposed by the researcher in the frozen input; the council did not autonomously discover the air-to-oxygen change. The 0.09 mL/min inlet/STP setting at 0.02 mL/min liquid feed delivers approximately 2.008 mol O2 per mol substrate, using 22.414 mL/mmol at 273.15 K and 1.01325 bar. It is a delivered feed ratio, not measured gas uptake. References in the reports to an unapproved oxygen change mean that laboratory oxygen-service approval remained outstanding, not that the researcher had failed to specify oxygen.",
        "Fluidics and safety. DrFluidics and DrSafety each requested the transient-flow tool for all 12 candidates. Their request identifiers point to the same deterministic result bundle within a run; these are two reviews of shared tool evidence, not independent simulation replications. The archive contains 192 conditional simulations per Figure 5 run. The gas-line check valve does not close the connected reverse path from the gas-injection junction into the Stage 1 liquid branch. Proposed liquid-branch protection, pressure-decoupled transfer and startup controls remained recommendations for engineering review; the tool record states design_modification_applied = false and laboratory_execution_status = review_required.",
        "Conditional-model boundary. Gas-source pressure, dead volume, compliance and device response were illustrative assumptions because measured startup traces were unavailable. The assumed 12 bar(g) gas-source pressure was not a laboratory setpoint and exceeded the recorded 10 bar MFC inlet limit. The pressure reference and device rating require confirmation. Tool outputs therefore cannot establish equipment compliance, predict incident probability or demonstrate absence of backflow. The qualitative earlier backflow observation supports examining the pathway, but does not calibrate the simulations or prove what would happen at the lower proposed gas flow.",
        "Outcome across response sets. Sets 1 and 3 selected Candidate 1: a 2 mL Stage 1 coil with a 100 min liquid V/Q time. Set 2 selected Candidate 6: a 5 mL Stage 1 coil with a 250 min liquid V/Q time. All three retained a 20 mL Stage 2 coil, 0.09 mL/min oxygen at STP and 0.02 mL/min liquid, giving an inlet/STP time index of 181.82 min. This Stage 2 index is not pressure-dependent liquid residence time, photon exposure or a kinetic equivalent of the 360 min batch hold. All candidates had the same substrate feed throughput (0.002 mmol/min), so selection of the compact design did not itself increase feed throughput.",
    ],
    6: [
        "Chemistry and stage coupling. The reports treated thioester formation and subsequent benzylamine aminolysis as one telescoped process. Stage 1 determines the amount and composition of intermediate entering Stage 2. Benzylamine was kept out of the activation feed and introduced between reactors. The council compared shorter, near-batch and extended holds against the 30 min + 30 min batch procedure. It discussed incomplete activation, incomplete aminolysis and possible prolonged-heating side reactions as uncertainties, not established rate laws or measured degradation.",
        "Feed and equipment review. The fixed activation feed contained acid (0.5 M), DPDTC (0.525 M) and DMAP (0.05 M); the separate benzylamine feed was 2.1 M, all in 2-MeTHF. The volumetric feed ratio Q_B/Q_A = 0.25 gives 1.05 equivalents of benzylamine. DrFluidics noted that Stage 2 time must use Q_A + Q_B, not Q_A alone. The reviewers addressed BLUE pump-tubing compatibility, independent pump channels, a real mixing junction and thermal/pressure limits. Final allocation identifies Pump A and Pump B as separate operations and assigns the P-713 T-mixer. Earlier narrative claims that the mixer was unassigned are not descriptions of that final topology; physical availability and assembly still require confirmation.",
        "Kinetic and operating uncertainty. The proposals retained 95 deg C and a 7 bar(g) BPR. Their pressure rationale was maintaining liquid-phase operation, not proof of stability for an uncharacterized mixture. Omission of the batch cooling step was treated as a direct-transfer hypothesis to evaluate. Solubility of the concentrated feeds, precipitation on mixing, pump accuracy and intermediate/product stability were listed as follow-up checks. No gas transient-tool call is present in these Figure 6 records; that later feature must not be retroactively attributed to them.",
        "Outcome across response sets. Set 1 selected Candidate 4 as a near-batch screening reference: two 5 mL coils, Q_A = 0.16 and Q_B = 0.04 mL/min, with 31.25 and 25.00 min stage times. Sets 2 and 3 selected Candidate 3: 10 mL then 5 mL, Q_A = 0.28 and Q_B = 0.07 mL/min, with 35.71 and 14.29 min stage times. The latter raises acid feed throughput from 0.08 to 0.14 mmol/min but leaves shorter aminolysis exposure as an explicit experimental question. It does not demonstrate higher product throughput or establish a yield-maximizing design.",
        "Interpretation of the decision. In Set 1, the Chief deferred Candidate 3 because its shorter second stage increased uncertainty relative to the near-batch reference. In Set 2, the same asymmetric allocation was preferred as an informative test of a longer activation stage combined with shorter aminolysis. In Set 3, substrate feed throughput was emphasized. This is an inspectable difference in decision priorities, not proof that one rationale is chemically optimal. A final-yield measurement alone cannot identify the rate-limiting stage; stage-resolved measurements would be needed. Later experimental results were not present as yield evidence in these council requests and are not retrospectively inserted into their reasoning.",
    ],
}

NOTES = {
    5: [
        "Time interpretation: some reports call Stage 2 liquid exposure approximately 1000 min, obtained from 20 mL/0.02 mL/min. This is a liquid-only space-time calculation using the entire reactor volume, not measured two-phase liquid residence. Likewise, the 181.82 min inlet/STP index must not be interpreted as a measured fraction of batch reaction exposure.",
        "Numerical transcription: the Figure 5 Set 3 Chief response gives the photocatalyst molar flow as 0.001 mmol/min. The final feed record and concentration-times-flow calculation give 0.0005 mol/L x 0.02 mL/min = 0.000010 mmol/min. The original statement is retained in the transcript as a 100-fold numerical error, not adopted in the figure text or experimental recipe.",
        "Geometry and transport claims: the Set 1 DrFluidics shared assessment groups Candidate 6 with large-bore ETFE options even though Candidate 6 uses 1 mm-ID PFA. The Set 3 Chief attributes a smaller in-channel gas volume fraction to the selected reactor relative to other 7 bar candidates with smaller reactors. Their gas and liquid feed rates, pressure and temperature are the same; changing reactor volume alone does not establish a different gas/liquid volumetric flow ratio or measured gas holdup. A known bore defines geometry but does not establish photon-dose adequacy or a measured transport advantage.",
        "Protocol interpretation: several reports describe the Stage 2 batch temperature as unspecified, whereas the supplied procedure continues under otherwise identical irradiation conditions after Stage 1 at 25 deg C. That wording does not establish an independently measured Stage 2 temperature, but it does not justify presenting 25 deg C as an unrelated new condition. Similarly, the qualitative backflow incident does not prove that every lower-flow alternative will reverse; the defensible finding is that the unprotected pathway remains.",
    ],
    6: [
        "Numerical transcription: in the Set 1 Chief answer_impacts entry for Q-PREF-001, benzylamine molar flow is written as 0.02 mmol/min. The final feed and independent arithmetic give 2.1 mol/L x 0.04 mL/min = 0.084 mmol/min, which is 1.05 times the acid feed of 0.08 mmol/min. For Sets 2/3, the corresponding amine rate is 0.147 mmol/min and acid rate is 0.14 mmol/min. The copied council text remains unedited; these corrected values govern the reporting.",
        "Inference limits: wording such as adequate exposure, confirmed photon geometry, or a single run distinguishing the rate-limiting step is not accepted as an experimentally demonstrated conclusion. A nominal hold near the batch duration does not establish equal conversion in flow. In particular, the Set 2 claim that incomplete Stage 1 conversion must have another cause because its hold exceeds batch is not justified. Changes in mixing, temperature distribution, feed state and residence-time distribution remain possible. The selected designs are therefore described as screening hypotheses.",
        "Inventory reconciliation: references to a missing mixer or one pump label belong to the council's view of the candidate record. The final allocation explicitly assigns the interstage P-713 mixer and separate stream-pump operations. This resolves the reported assignment at the software-record level, not by claiming that the council physically installed equipment or verified the laboratory configuration.",
    ],
}

FIGURE_COPY = {
    5: {
        "title": "Council review: oxygen placement and connected-stage exposure",
        "boxes": [
            ("Chemistry", "Protect the oxygen-free Giese stage; introduce O2 only before oxidation."),
            ("Kinetics", "Compare stage exposures without assuming flow acceleration or predicting yield."),
            ("Fluidics", "Check inlet O2 delivery and the reverse-flow path toward Reactor 1."),
            ("Safety", "Gas-line protection does not protect the liquid branch; retain the unresolved warning."),
            ("Skeptic", "Do not equate numerical closure with safe operation or proven conversion."),
            ("Chief", "Select a compact baseline or extend Stage 1 exposure under the conversion-focused brief."),
        ],
        "outcome": "Sets 1/3: R1 = 100 min. Set 2: R1 = 250 min. All sets: R2 inlet/STP index = 181.82 min; O2 = 0.09 mL/min at STP (2.008 equiv).",
        "compact": "Stage-specific oxygen placement | Exposure comparison | Reverse-path review | Candidate selection",
        "footnote": "Condensed from archived reports, not verbatim dialogue. Pure O2 was user-specified. Proposed conditions; backflow remained unresolved. The R2 STP index is not actual operating residence time.",
        "caption": "Council review of the connected Giese addition/oxidation process. Specialist reports compare chemical order, stage exposure, equipment and gas-delivery constraints; the Skeptic and Chief retain unresolved operability issues alongside the selected screening candidate. Full reports and set-specific decisions are documented in ESI Section 7.3. Pure oxygen was specified by the researcher rather than independently discovered by the council.",
    },
    6: {
        "title": "Council review: activation and amidation as one process",
        "boxes": [
            ("Chemistry", "Form the thioester first; add benzylamine between reactors without intermediate isolation."),
            ("Kinetics", "Compare near-batch holds with longer activation and shorter aminolysis."),
            ("Fluidics", "Recalculate Stage 2 time after amine addition; close feed concentrations and molar ratios."),
            ("Safety", "Review heated pressurized operation, feed solubility and pump compatibility."),
            ("Skeptic", "Flag uncertain Stage 2 conversion; no residence-time choice proves the final yield."),
            ("Chief", "Select a near-batch reference or a higher-feed-throughput, asymmetric screening design."),
        ],
        "outcome": "Set 1: R1/R2 = 31.25/25.00 min. Sets 2/3: R1/R2 = 35.71/14.29 min. Acid feed throughput: 0.08 versus 0.14 mmol/min.",
        "compact": "Preserve addition order | Close stoichiometry | Compare stage-time allocation | Select the first screen",
        "footnote": "Condensed from archived reports, not verbatim dialogue. Proposed conditions, not optimized yields. Shorter Stage 2 exposure remained an experimental question; feed throughput is not product throughput.",
        "caption": "Council review of telescoped thioester formation and amidation. Specialist reports connect addition order and stoichiometry to the residence time in each reactor, while comparing a near-batch reference with an asymmetric higher-feed-throughput alternative. The selected candidate is reported with its remaining chemical and operating uncertainties rather than as an established optimum (ESI Section 7.3).",
    },
}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def dump(path, value):
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def decoded(raw):
    raw = re.sub(r"^```(?:json)?\s*", "", raw.strip())
    raw = re.sub(r"\s*```$", "", raw)
    return json.loads(raw)


def scalar_leaves(value, prefix=""):
    if isinstance(value, dict):
        for key, child in value.items():
            yield from scalar_leaves(child, f"{prefix}.{key}".lstrip("."))
    elif isinstance(value, list):
        if not value:
            yield prefix, "[]"
        for i, child in enumerate(value):
            yield from scalar_leaves(child, f"{prefix}[{i}]")
    else:
        yield prefix, value if isinstance(value, str) else json.dumps(value)


def new_doc(title):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    sec.left_margin = sec.right_margin = Inches(0.8)
    for style in doc.styles:
        if style.type == 1:
            style.font.name = "Times New Roman"
            style.font.size = Pt(11)
            style.font.color.rgb = __import__('docx').shared.RGBColor(0, 0, 0)
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.05
    doc.add_heading(title, 0)
    return doc


def set_font(p):
    for run in p.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)


def field_label(path):
    names = {
        "candidate_id": "Candidate", "recommendation": "Recommendation",
        "hard_violation": "Hard violation reported", "justification": "Justification",
        "uncertainties": "Uncertainty", "alternatives": "Alternative",
        "reason_not_selected": "Reason not selected", "answer_impacts": "Answer impact",
        "question_id": "Question", "next_measurements": "Suggested measurement",
        "required_measurements": "Requested measurement", "limitations": "Limitation",
        "shared_assessment": "Shared assessment", "objective_alignment": "Objective alignment",
    }
    labels = []
    for bit in path.split("."):
        match = re.fullmatch(r"(.+?)(?:\[(\d+)\])?", bit)
        key, index = match.groups()
        label = names.get(key, key.replace("_", " ").capitalize())
        if index is not None:
            label += f" {int(index) + 1}"
        labels.append(label)
    return " / ".join(labels)


def report_value(doc, value, prefix=""):
    for path, content in scalar_leaves(value, prefix):
        p = doc.add_paragraph()
        p.add_run(field_label(path) + ": ").bold = True
        p.add_run(content)
        set_font(p)


def main():
    OUT.mkdir(exist_ok=True)
    original_digest = digest(SOURCE)
    snapshot = OUT / "source_ESI_snapshot.docx"
    if snapshot.exists():
        assert digest(snapshot) == original_digest, "Source ESI changed; do not overwrite the prior snapshot."
    else:
        shutil.copy2(SOURCE, snapshot)
    pkg = Package(SOURCE)
    original_body = list(pkg.body)
    before_images = {k: hashlib.sha256(v).hexdigest() for k, v in pkg.parts.items() if k.startswith("word/media/")}
    original_tables = [E.tostring(t) for t in pkg.body.findall(W + "tbl")]
    template = next(p for p in pkg.body.findall(W + "p") if text(p).startswith("The supplied protocol forms"))
    h2 = next(p for p in pkg.body.findall(W + "p") if text(p) == "Telescoped DPDTC-mediated amidation (main-text Figure 6)")
    target = next(p for p in pkg.body.findall(W + "p") if text(p) == "Laboratory implementation, experimental procedures, and analytical data")
    section = []
    manifest, audit_rows, candidate_rows, source_hashes = [], [], [], {}
    source_hashes[str(SOURCE.relative_to(ROOT))] = original_digest
    full_doc = new_doc("Supplementary council reports for Figures 5 and 6")
    full_doc.add_paragraph("Unabridged recorded council responses, including unsuccessful truncated responses. These reports are historical model outputs, not independently validated findings or laboratory instructions. Full prompts, machine-readable records and original source hashes accompany this document. See the interpretation notes in the revised ESI, Section 7.3. No new model calls were made.")

    def paragraph(value, bold=False, heading=False):
        p = E.Element(W + "p")
        pp = deepcopy((h2 if heading else template).find(W + "pPr"))
        if pp is None:
            pp = E.Element(W + "pPr")
        for node in list(pp):
            if node.tag in {W + "pageBreakBefore", W + "keepNext", W + "keepLines"}:
                pp.remove(node)
        p.append(pp)
        pkg.revise(p, value)
        if bold or heading:
            E.SubElement(pp, W + "keepNext")
            for r in p.findall(W + "r"):
                E.SubElement(r.find(W + "rPr"), W + "b")
        else:
            for r in p.findall(W + "r"):
                rp = r.find(W + "rPr")
                E.SubElement(rp, W + "b", {W + "val": "0"})
        section.append(p)
        return p

    def fields(value, prefix=""):
        for path, value in scalar_leaves(value, prefix):
            label = field_label(path)
            paragraph(f"{label}: {value}")

    paragraph("Recorded council deliberation for the two connected-process cases", heading=True)
    for value in INTRO:
        paragraph(value)
    records = []
    for figure in (5, 6):
        paragraph(f"Figure {figure}: {'Giese addition and oxidation' if figure == 5 else 'thioester formation and amidation'}", bold=True)
        for value in COMMON[figure]:
            paragraph(value)
        paragraph("Interpretation notes for the unedited reports", bold=True)
        for note in NOTES[figure]:
            paragraph(note)
        for set_number in (1, 2, 3):
            key = f"F{figure}-S{set_number}"
            source = CAMPAIGNS[figure] / f"figure{figure}_set{set_number}/attempt_01"
            raw_dir = OUT / "raw" / key
            raw_dir.mkdir(parents=True, exist_ok=True)
            for name in ["result.json", "llm_calls.jsonl", "intake_package.json", "inventory_profile.json", "request.json", "summary.json", "run.log", "source_hashes.json", "stage_summary.csv", "component_feeds.csv"]:
                p = source / name
                if p.exists():
                    shutil.copy2(p, raw_dir / name)
                    source_hashes[str(p.relative_to(ROOT))] = digest(p)
            data = json.loads((source / "result.json").read_text())
            assessment = data["scientific_assessment"]
            calls = assessment["calls"]
            chosen = assessment["selected_candidate_id"]
            assert assessment["candidate_count"] == len(assessment["candidates"]) == 12
            assert assessment["selected_design_preserved"]
            assert chosen == assessment["chief"]["candidate_id"]
            assert set(assessment["reviews"]) == set(ROLES)
            for role in ROLES:
                assert sorted(r["candidate_id"] for r in assessment["reviews"][role]) == list(range(1, 13))
            dump(raw_dir / "scientific_assessment.json", assessment)
            stages = data["result_report"]["stages"]
            for stage in stages:
                q = stage["liquid_flow_mL_min"] + stage["gas_flow_stp_mL_min"]
                assert math.isclose(stage["volume_mL"] / q, stage["residence_time_min"], abs_tol=0.0001)
            errors = [(c["role"], c["attempt"]) for c in calls if c.get("parse_error")]
            role_calls = [{"call": i + 1, "role": c["role"], "attempt": c["attempt"], "parse_error": c.get("parse_error"), "request_keys": list(c["request"])} for i, c in enumerate(calls)]
            science_fields = {
                "id": key, "source": str(source.relative_to(ROOT)),
                "council_archive": assessment["archive_path"],
                "selected_candidate": chosen, "candidate_count": 12,
                "objective": assessment["objective"], "priority": assessment["objective_policy"]["priority"],
                "upstream": "claude-opus-4-6", "downstream": assessment["model"],
                "calls": role_calls, "final_stages": stages,
                "all_four_reviewers_cover_all_12_candidates": True,
                "vetoed_ids": assessment["vetoed_ids"],
            }
            manifest.append(science_fields)
            records.append((key, data, source))
            selected = next(c for c in assessment["candidates"] if c["candidate_id"] == chosen)
            for candidate in assessment["candidates"]:
                ss = candidate["proposal"]["stage_parameters"]
                candidate_rows.append({
                    "run": key, "candidate": candidate["candidate_id"], "selected": candidate["candidate_id"] == chosen,
                    "volume_R1_mL": ss[0]["reactor_volume_mL"], "volume_R2_mL": ss[1]["reactor_volume_mL"],
                    "time_R1_min": ss[0]["residence_time_min"], "time_R2_min": ss[1]["residence_time_min"],
                    "time_R2_basis": ss[1]["residence_time_basis"],
                    "BPR_bar_g": candidate["proposal"]["BPR_bar"],
                    **{r: next(v["recommendation"] for v in assessment["reviews"][r] if v["candidate_id"] == candidate["candidate_id"]) for r in ROLES},
                })
            paragraph(f"Figure {figure}, response Set {set_number} ({key}): selected Candidate {chosen}", bold=True)
            paragraph("Recorded objective: " + assessment["objective"])
            times = "; ".join(f"R{s['number']}: {s['volume_mL']:g} mL, {s['residence_time_min']:.2f} min ({s['residence_basis']})" for s in stages)
            paragraph(f"Final-record cross-check: {times}. The model priority was {assessment['objective_policy']['priority']}. All four specialists reviewed all 12 candidates; the recorded veto list is empty. The selected candidate was preserved by final numerical reconciliation.")
            if errors:
                paragraph("Response-format retries: " + "; ".join(f"{r}, attempt {n}" for r, n in errors) + ". These responses were truncated or invalid JSON and were retried, not treated as extra deliberation rounds. The unsuccessful text remains in the companion archive.")
            else:
                paragraph("All six council report calls completed without a recorded response-format retry.")
            for role in ROLES:
                call_index = next(i for i in range(len(calls) - 1, -1, -1) if calls[i]["role"] == role and not calls[i].get("parse_error"))
                review = next(r for r in assessment["reviews"][role] if r["candidate_id"] == chosen)
                paragraph(f"{role}: complete selected-candidate verdict (recorded text; {key}, call {call_index + 1})", bold=True)
                fields(review)
                audit_rows.append({"run": key, "section": role, "json_pointer": f"scientific_assessment.reviews.{role}[candidate_id={chosen}]", "verbatim": True})
            paragraph("Skeptic: selected-candidate assessment and shared qualifications", bold=True)
            skeptic = assessment["skeptic"]
            # Retain shared assessments in full, but avoid repeating every rejected
            # candidate here; the companion contains the entire raw response.
            for k, v in skeptic.items():
                if isinstance(v, list) and any(isinstance(x, dict) and "candidate_id" in x for x in v):
                    v = [x for x in v if not isinstance(x, dict) or "candidate_id" not in x or x["candidate_id"] == chosen]
                fields(v, k)
            paragraph("Chief: complete recorded selection report", bold=True)
            fields(assessment["chief"])
            audit_rows.append({"run": key, "section": "Chief", "json_pointer": "scientific_assessment.chief", "verbatim": True})
            paragraph(f"Source: {key}, result.json, scientific_assessment; corresponding prompts and exact response text are retained in llm_calls.jsonl and the numbered response files in the companion archive. The original result and final design data have not been edited.")
            full_doc.add_heading(f"{key}: Figure {figure}, Set {set_number}", 1)
            full_doc.add_paragraph(f"Source: {source.relative_to(ROOT)}. Selected Candidate {chosen}; 12-candidate pool; downstream model {assessment['model']}. The objective policy is {assessment['objective_policy']['priority']}.")
            for i, call in enumerate(calls, 1):
                fn = f"call_{i:02d}_{call['role']}_attempt{call['attempt']}.txt"
                (raw_dir / fn).write_text(call["raw_response"], encoding="utf-8")
                full_doc.add_heading(f"Call {i}: {call['role']}, attempt {call['attempt']}", 2)
                full_doc.add_paragraph("Status: " + ("unsuccessful response; retained for provenance. " + call["parse_error"] if call.get("parse_error") else "recorded response, parsed successfully."))
                try:
                    value = decoded(call["raw_response"])
                except json.JSONDecodeError:
                    assert call.get("parse_error"), "Unrecorded parse failure"
                    for line in call["raw_response"].splitlines():
                        full_doc.add_paragraph(line)
                else:
                    report_value(full_doc, value)
                audit_rows.append({"run": key, "section": f"call_{i}", "json_pointer": f"scientific_assessment.calls[{i-1}].raw_response", "verbatim": True})
                assert (raw_dir / fn).read_text(encoding="utf-8") == call["raw_response"]
            full_doc.add_heading("Editorial cross-checks (not model statements)", 2)
            for note in NOTES[figure]:
                full_doc.add_paragraph(note)
    paragraph("What these records demonstrate", bold=True)
    paragraph("The reports show how chemical interpretation, stage exposure, equipment constraints and operating uncertainties were communicated to a final selector. They also expose limits: repeated warnings were not equivalent to installed mitigations; some narrative numbers were wrong; and similar briefs could produce the same selection. The records support an auditable design-assistance workflow, not a claim that discussion guarantees the highest yield or eliminates the need for experimental and engineering review. The six runs contain four distinct selected physical configurations: two for each chemistry.")
    pkg.before(target, section)
    assert [n for n in pkg.body if n not in section] == original_body
    # Change only the main document XML. Preserve every original image, style,
    # relationship, table and page-setup part in the supplied Word package.
    parts = dict(pkg.parts)
    parts["word/document.xml"] = E.tostring(pkg.doc, xml_declaration=True, encoding="UTF-8", standalone=True)
    with ZipFile(TARGET, "w", ZIP_DEFLATED) as z:
        for k, v in parts.items():
            z.writestr(k, v)
    for p in full_doc.paragraphs:
        set_font(p)
    full_doc.save(APPENDIX)
    figure_doc = new_doc("Council discussion text for main-text Figures 5 and 6")
    md = ["# Council text for Figures 5 and 6", "", "Condensed summaries of archived reports, not verbatim dialogue. No figures were modified.", ""]
    for figure, copy in FIGURE_COPY.items():
        if figure == 6:
            figure_doc.add_page_break()
        figure_doc.add_heading(f"Figure {figure}", 1)
        figure_doc.add_heading(copy["title"], 2)
        md.extend([f"## Figure {figure}: {copy['title']}", ""])
        for role, value in copy["boxes"]:
            p = figure_doc.add_paragraph()
            p.add_run(role + ": ").bold = True
            p.add_run(value)
            md.append(f"- **{role}:** {value}")
        for label, key in [("Selected proposals", "outcome"), ("Compact alternative", "compact"), ("Figure note", "footnote"), ("Caption insert", "caption")]:
            p = figure_doc.add_paragraph()
            p.add_run(label + ": ").bold = True
            p.add_run(copy[key])
            md.extend(["", f"**{label}:** {copy[key]}"])
        md.append("")
    for p in figure_doc.paragraphs:
        set_font(p)
    figure_doc.save(FIGTEXT)
    (OUT / "Figure5_6_council_figure_text.md").write_text("\n".join(md), encoding="utf-8")
    dump(OUT / "source_manifest.json", {"runs": manifest, "source_sha256": source_hashes})
    dump(OUT / "text_provenance.json", audit_rows)
    dump(OUT / "figure_text.json", FIGURE_COPY)
    for name, rows in [("candidate_decisions.csv", candidate_rows)]:
        with (OUT / name).open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=list(rows[0]))
            writer.writeheader(); writer.writerows(rows)
    checks = {
        "source_ESI_unchanged": digest(SOURCE) == original_digest,
        "original_body_nodes_unchanged_and_in_order": [n for n in pkg.body if n not in section] == original_body,
        "all_original_tables_byte_preserved": original_tables == [E.tostring(t) for t in pkg.body.findall(W + "tbl")],
        "all_images_byte_preserved": before_images == {k: hashlib.sha256(v).hexdigest() for k, v in parts.items() if k.startswith("word/media/")},
        "all_other_package_parts_unchanged": all(parts[k] == v for k, v in pkg.parts.items() if k != "word/document.xml"),
        "six_runs_documented": len(manifest) == 6,
        "72_candidates_archived": len(candidate_rows) == 72,
        "all_selected_stage_times_independently_recomputed": True,
        "all_response_files_exact": True,
        "new_paragraphs_Times_New_Roman_11": all(r.xpath('./w:rPr/w:rFonts[@w:ascii="Times New Roman"]', namespaces=NS) and r.xpath('./w:rPr/w:sz[@w:val="22"]', namespaces=NS) for p in section for r in p.findall(W + "r") if text(r)),
        "new_model_calls": 0,
        "new_figures": 0,
        "added_ESI_paragraphs": len(section),
        "archived_council_calls": sum(len(r["calls"]) for r in manifest),
    }
    assert all(v for k, v in checks.items() if isinstance(v, bool)), checks
    dump(OUT / "structural_verification.json", checks)
    readme = """# Figure 5/6 council documentation

The revised ESI adds Section 7.3. Existing text, figures, tables and page setup
are preserved. The original, user-edited October 2 ESI is not overwritten.

The ESI contains detailed case-level interpretation, all selected-candidate
specialist verdicts, relevant Skeptic assessments and complete Chief reports.
The full-council Word companion contains every recorded council response for
all candidates, including unsuccessful response-format attempts and tool requests.
This is an archive of model reports, not validated claims or operating instructions.

`raw/F5-S1` through `raw/F6-S3` contain exact prompts/responses and original
results. Numbered text files exactly reproduce `calls[*].raw_response`.
`source_manifest.json` maps all runs and hashes original source files.
`candidate_decisions.csv` contains all 72 candidate rows and specialist verdicts.
The separate figure-text document contains short, explicitly paraphrased labels.
The original figures and main manuscript are not modified.

No paid model calls or new experimental claims were introduced. Pure oxygen
was required by the user, not autonomously proposed by this rerun. Backflow
remained unresolved. Model arithmetic errors are retained and annotated rather
than replaced with apparently flawless dialogue. Raw transient traces remain at
the original paths named in the archived result; the result retains their summaries.
"""
    (OUT / "README.md").write_text(readme, encoding="utf-8")
    print(json.dumps(checks, indent=2))
    print(TARGET)


if __name__ == "__main__":
    main()
