from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/home/amirreza/Documents/codes/Flow Agent')
OUT = ROOT / 'manuscript' / 'Advanced_Science_cover_letter_20261006.docx'
TITLE = 'Agentic design of continuous flow processes with chemical knowledge and laboratory constraints'
AUTHORS = 'Amirreza Mottafegh, Mincheol Park, Jun-young Jo, Jin Suk Myung, Boyoung Y. Park, and Gwang-Noh Ahn'

BODY = [
    'We submit the above manuscript for consideration as a Research Article in Advanced Science. '
    'The study investigates how large language models and agentic reasoning can advance flow chemistry by translating a batch procedure into a connected experimental design. '
    'This task requires chemical knowledge to be reconciled with feed preparation, reactor configuration and the capabilities of the available equipment. '
    'We develop a multi-agent workflow that starts from a batch protocol, laboratory inventory and researcher objectives, then combines literature precedents, handbook-derived engineering guidance, quantitative calculations and specialist review to propose a complete flow process. '
    'Implemented as FlowPilot, the workflow connects feed compositions and stage conditions to specific equipment and a process diagram through a graphical interface. '
    'The broader advance is a practical way to organize language-model knowledge around the coupled decisions needed to turn established chemistry into a testable flow experiment.',
    'Our evaluation spans a matched five-model benchmark, module-ablation and retrieval analyses, computational dry runs, and laboratory implementation of two connected reaction sequences. '
    'Across 90 benchmark designs, the mean assessed score increased from 0.794 for one-shot generation to 0.917 for the agentic workflow, while judge-reported critical flags decreased from 98 to 4. '
    'For each experimental response set, computational review of 12 candidates produced one selected design. '
    'The photoredox addition–oxidation sequence achieved 97–98% NMR yield and 94% isolated yield, compared with reported literature flow yields of 90% and 88%, respectively, and a batch isolated yield of 95%, under different conditions. '
    'A second, telescoped amidation afforded up to 86% NMR yield and 84% isolated yield. '
    'These results show that agent-selected starting configurations can be experimentally productive when assessed against literature procedures developed through experimental optimization. '
    'By linking language-model evaluation to apparatus design and chemical outcomes, the study should interest Advanced Science readers working across artificial intelligence, chemistry and engineering.'
]

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.2677)
sec.page_height = Inches(11.6929)
sec.top_margin = sec.bottom_margin = Inches(0.9)
sec.left_margin = sec.right_margin = Inches(0.95)
sec.header_distance = sec.footer_distance = Inches(0.3)

for name in ('Normal', 'Title'):
    s = doc.styles[name]
    s.font.name = 'Times New Roman'
    s.font.size = Pt(11)
    s.font.color.rgb = RGBColor(0, 0, 0)
    rpr = s.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + attr), 'Times New Roman')
    for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:' + attr), None)
    s.paragraph_format.space_before = Pt(0)
    s.paragraph_format.space_after = Pt(8)
    s.paragraph_format.line_spacing = 1.05

doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.keep_with_next = True
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)

def para(text, *, style=None, after=8, bold=False, italic=False, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.widow_control = True
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0, 0, 0)
    r.bold = bold or style == 'Title'
    r.italic = italic
    fonts = r._element.get_or_add_rPr().rFonts
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + attr), 'Times New Roman')
    return p

para('6 October 2026', after=10)
para('The Editors\nAdvanced Science', after=12)
para('Agentic design of continuous flow processes\nwith chemical knowledge and laboratory constraints', style='Title', after=5)
para('Amirreza Mottafegh, Mincheol Park, Jun-young Jo, Jin Suk Myung,\nBoyoung Y. Park, and Gwang-Noh Ahn', after=12)
para('Dear Editors,', after=8)
for text in BODY:
    para(text, after=10)
para('Sincerely,', after=6)
para('Boyoung Y. Park and Gwang-Noh Ahn\nCorresponding authors', after=5)
para('Kyung Hee University | boyoungy.park@khu.ac.kr\nKorea Research Institute of Chemical Technology | gnahn@krict.re.kr', after=0)

doc.core_properties.title = 'Cover letter for Advanced Science'
doc.core_properties.subject = TITLE
doc.core_properties.author = 'Boyoung Y. Park; Gwang-Noh Ahn'
doc.core_properties.keywords = 'Advanced Science; cover letter; agentic flow process design'
doc.core_properties.comments = ''
doc.save(OUT)
print(OUT)
print('Body words:', sum(len(p.split()) for p in BODY))
