"""Export six archived selections as a new, traceable 12-slide review deck.

This is a deterministic reporting revision, not another council run or a claim
of experimental validation. Source results and their approval state stay intact.
"""
import copy
import csv
import hashlib
import json
import math
from datetime import datetime
from pathlib import Path
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt
from PIL import Image
from flora_translate.diagram_artifacts import render_topology_artifacts
from flora_translate.schemas import ProcessTopology

SOURCE = ROOT / 'outputs/khu_revised_six_20260915/presentation'
WORKBOOK = ROOT / 'inventory_khu/Inventory_final.xlsx'
VM = .08314 * 273.15 / 1.01325  # mL/mmol at the archived STP convention.


def dump(path, data):
    path.write_text(json.dumps(data, indent=2, ensure_ascii=True))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_csv(path):
    with path.open(newline='') as handle:
        return list(csv.DictReader(handle))


def write_csv(path, rows):
    with path.open('w', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def fmt(value):
    return '--' if value is None or value == '' else f'{float(value):.5g}'


def alias(row):
    name = row['component'].lower()
    if 'trimethylsilane' in name:
        return 'Substrate 1'
    if 'nitrobenzoic acid' in name:
        return 'Acid'
    if '[ir(' in name:
        return 'Ir catalyst'
    if row['phase'] == 'gas':
        return 'O2 in air'
    return {'acrylonitrile': 'Acrylonitrile', 'benzylamine': 'Benzylamine',
            'dpdtc': 'DPDTC', 'dmap': 'DMAP'}[name]


class Deck:
    def __init__(self):
        self.p = Presentation()
        self.p.slide_width, self.p.slide_height = Inches(16), Inches(9)

    def text(self, slide, text, x, y, w, h, size=17, bold=False, color='23333B'):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        for i, line in enumerate(text.split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name, p.font.size, p.font.bold = 'Arial', Pt(size), bold
            p.font.color.rgb = RGBColor.from_string(color)
            p.space_after = Pt(4)
        return shape

    def slide(self, title, subtitle):
        s = self.p.slides.add_slide(self.p.slide_layouts[6])
        self.text(s, title, .55, .25, 14.9, .6, 28, True)
        self.text(s, subtitle, .55, .94, 14.9, .5, 16, color='51636C')
        self.text(s, 'FlowPilot | KHU updated inventory | 15 September 2026 | Proposed conditions for laboratory review',
                  .55, 8.62, 14.1, .26, 11, color='51636C')
        self.text(s, str(len(self.p.slides)), 15, 8.62, .45, .26, 11)
        return s

    def table(self, slide, headers, rows, x, y, widths, row_heights, size=16):
        table = slide.shapes.add_table(len(rows)+1, len(headers), Inches(x), Inches(y),
                                      Inches(sum(widths)), Inches(sum(row_heights))).table
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
        for row, height in zip(table.rows, row_heights):
            row.height = Inches(height)
        for i, row in enumerate([headers, *rows]):
            for j, value in enumerate(row):
                cell = table.cell(i, j)
                cell.text = str(value)
                cell.margin_left = cell.margin_right = Inches(.08)
                cell.margin_top = cell.margin_bottom = Inches(.04)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string('215F6B' if i == 0 else 'EDF3F4' if i % 2 else 'FFFFFF')
                for p in cell.text_frame.paragraphs:
                    p.font.name, p.font.size = 'Arial', Pt(size)
                    p.font.bold = i == 0
                    p.font.color.rgb = RGBColor.from_string('FFFFFF' if i == 0 else '23333B')


def main():
    out = SOURCE.parent / ('collaborator_slides_' + datetime.now().strftime('%Y%m%d_%H%M%S'))
    out.mkdir()
    book = load_workbook(WORKBOOK, data_only=True)
    increment = float(book['MFC']['D3'].value)
    assert increment == .01, 'Review changed workbook before exporting.'
    source_hashes = {str(p.relative_to(ROOT)): sha(p) for p in SOURCE.rglob('*') if p.is_file()}
    deck = Deck()
    checks, changes, all_stages, all_feeds = [], [], [], []
    for selection in json.loads((SOURCE/'selected_results.json').read_text()):
        name = selection['case']
        src, target = SOURCE/name, out/name
        target.mkdir()
        result = json.loads((src/'result.json').read_text())
        topology = copy.deepcopy(result['process_topology'])
        stages, feeds = read_csv(src/'stage_summary.csv'), read_csv(src/'component_feeds.csv')
        for row in stages:
            for key in ('stage', 'volume_mL', 'ID_mm', 'temperature_C', 'liquid_flow_mL_min',
                        'gas_inlet_STP_mL_min', 'nominal_inlet_residence_min', 'BPR_bar', 'wavelength_nm'):
                row[key] = float(row[key]) if row[key] else None
        for row in feeds:
            for key in ('concentration_M', 'flow_mL_min', 'molar_flow_mmol_min', 'equivalents'):
                row[key] = float(row[key]) if row[key] else None
        photo = name.startswith('figure5')
        limiting = feeds[0]['molar_flow_mmol_min']
        gas = next((r for r in feeds if r['phase'] == 'gas'), None)
        correction = None
        if gas:
            old = gas['flow_mL_min']
            new = round(math.ceil((old-1e-12)/increment)*increment, 2)
            molar = new / VM * .21
            gas.update(flow_mL_min=new, molar_flow_mmol_min=molar, equivalents=molar/limiting,
                       concentration_basis='O2 molar flow from 21 mol% air, STP 273.15 K / 1.01325 bar')
            stages[1]['gas_inlet_STP_mL_min'] = new
            correction = dict(case=name, old_air_STP_mL_min=old, new_air_STP_mL_min=new,
                              actual_O2_equivalents=molar/limiting, source_cell='MFC!D3', increment=increment)
            changes.append(correction)
        for stage in stages:
            stage['nominal_inlet_residence_min'] = stage['volume_mL'] / (
                stage['liquid_flow_mL_min'] + stage['gas_inlet_STP_mL_min'])
        ops = {o['op_id']: o for o in topology['unit_operations']}
        bpr = ops['bpr_final']['parameters']['pressure_bar']
        for stage in stages:
            op = ops[f'st{int(stage["stage"])}_reactor']
            p = op['parameters']
            p.update(residence_time_min=stage['nominal_inlet_residence_min'],
                     residence_time_inlet_min=stage['nominal_inlet_residence_min'],
                     Q_inlet_mL_min=stage['liquid_flow_mL_min']+stage['gas_inlet_STP_mL_min'],
                     Q_liquid_mL_min=stage['liquid_flow_mL_min'],
                     Q_gas_sccm=stage['gas_inlet_STP_mL_min'] or None)
            # The export reports inlet/STP apparent time, not unverified in-channel estimates.
            for field in ('residence_time_in_channel_min', 'Q_gas_actual_mL_min',
                          'liquid_holdup_volume_mL', 'gas_holdup'):
                p.pop(field, None)
            op['rationale'] = 'Reported time calculated from the corrected export stage table; no kinetic claim.'
            assert math.isclose(stage['BPR_bar'], bpr)
            checks.append(dict(case=name, check=f'Stage {int(stage["stage"])} V/Q and selected BPR', passed=True))
        if gas:
            gas_op = next(o for o in ops.values() if o['op_type'] == 'mfc')
            valve_op = next(o for o in ops.values() if o['op_type'] == 'check_valve')
            p = gas_op['parameters']
            p.update(gas_flow_sccm=gas['flow_mL_min'], molar_equiv=gas['equivalents'])
            p.pop('gas_flow_actual_mL_min', None)
            for item in topology['instrument_manifest']:
                for setting in item.get('settings', []):
                    if 'flow_sccm' in setting:
                        setting['flow_sccm'] = gas['flow_mL_min']
            assert math.isclose(gas['flow_mL_min']/increment, round(gas['flow_mL_min']/increment))
            assert gas['equivalents'] >= 2
            edges = {(e['from_op'], e['to_op']) for e in topology['streams']}
            assert (gas_op['op_id'], valve_op['op_id']) in edges
            assert (valve_op['op_id'], 'st2_mixer') in edges
            checks.append(dict(case=name, check='MFC increment, >=2 O2 equivalents, Stage 2 addition through check valve', passed=True))
        for feed in feeds:
            if feed['phase'] != 'gas':
                assert math.isclose(feed['concentration_M']*feed['flow_mL_min'], feed['molar_flow_mmol_min'], rel_tol=1e-6)
                assert math.isclose(feed['molar_flow_mmol_min']/limiting, feed['equivalents'], rel_tol=1e-6)
                checks.append(dict(case=name, check=f'{feed["component"]}: C x Q and stoichiometry', passed=True))
        topology['residence_time_min'] = sum(s['nominal_inlet_residence_min'] for s in stages)
        topology['total_flow_rate_mL_min'] = stages[-1]['liquid_flow_mL_min'] + stages[-1]['gas_inlet_STP_mL_min']
        dump(target/'process_topology.json', topology)
        write_csv(target/'stage_parameters.csv', stages)
        write_csv(target/'feed_parameters.csv', feeds)
        dump(target/'export_provenance.json', dict(source_result=str(src/'result.json'),
             source_result_sha256=sha(src/'result.json'), workbook_sha256=sha(WORKBOOK),
             status='deterministically_revised_proposal_for_laboratory_review', new_model_calls=0,
             correction=correction, pressure_basis='Selected connected-process BPR setting, gauge basis in original proposal; laboratory to confirm.',
             excluded='Legacy stage-local BPR minimum calculations and in-channel estimates are not exported as operating settings.',
             experimental_validation=False))
        # Short display aliases retain full chemical identities in the next-slide table and CSV.
        display = copy.deepcopy(topology)
        for op in display['unit_operations']:
            if op['op_type'] == 'pump':
                stream = op['parameters']['stream']
                component_rows = [r for r in feeds if r['stream'] == stream]
                op['parameters']['contents'] = [f'{alias(row)}: {fmt(row["concentration_M"])} M'
                    for row in component_rows]
        artifacts = render_topology_artifacts(ProcessTopology.model_validate(display), title='', base_dir=target/'render')
        assert artifacts['manifest']['renderer'] == 'graphviz'
        shutil.copy2(artifacts['png_path'], target/'topology.png')
        shutil.copy2(artifacts['svg_path'], target/'topology.svg')
        figure, setno = name[6], name[-1]
        title = f'Figure {figure} | Set {setno}'
        chemistry = 'Giese addition followed by aerobic oxidation' if photo else 'DPDTC-mediated, telescoped amide formation'
        slide = deck.slide(title+' | Flow topology', chemistry)
        with Image.open(target/'topology.png') as im:
            w, h = im.size
        scale = min(14.9/w, 6.25/h)
        slide.shapes.add_picture(str(target/'topology.png'), Inches((16-w*scale)/2),
                                Inches(1.48+(6.25-h*scale)/2), width=Inches(w*scale), height=Inches(h*scale))
        topology_note = ('Stage 1 remains oxygen-free; offline solvent degassing. Air enters only before Stage 2.\n'
                         'PFA coil mounted in UV-150; Stage 2 uses Manual 2 and its own light. Full feed identities on next slide.' if photo else
                         'Feed B joins the Stage 1 effluent at the T-mixer; the intermediate is not isolated.\n'
                         'Two E-series pump channels; thermally heated ETFE reactors. Full feed identities on next slide.')
        deck.text(slide, topology_note, .55, 7.83, 14.9, .67, 16)
        slide = deck.slide(title+' | Process parameters', chemistry+' | Feed concentrations are stock concentrations, before mixing')
        stage_rows = [
            ['Reactor', *[f'{fmt(s["volume_mL"])} mL {s["material"]}' for s in stages]],
            ['Tubing ID (mm)', *[fmt(s['ID_mm']) for s in stages]],
            ['Module / heating', *([stages[0]['module'], stages[1]['module']] if photo else ['Thermal / oil bath']*2)],
            ['Light (nm)', *[fmt(s['wavelength_nm']) for s in stages]],
            ['Temperature (C)', *[fmt(s['temperature_C']) for s in stages]],
            ['Liquid flow (mL/min)', *[fmt(s['liquid_flow_mL_min']) for s in stages]],
            ['Air, STP (mL/min)', *[fmt(s['gas_inlet_STP_mL_min']) for s in stages]],
            ['Nominal time (min)', *[f'{s["nominal_inlet_residence_min"]:.2f}' for s in stages]],
            ['BPR setting (bar g)', *[fmt(s['BPR_bar']) for s in stages]],
        ]
        deck.table(slide, ['Stage conditions', 'Stage 1', 'Stage 2'], stage_rows,
                   .55, 1.6, [3.05, 1.87, 1.87], [.52, .52, .46, .70, .46, .46, .58, .58, .58, .52], 17)
        feed_rows = [[f'{r["stream"]}', alias(r), fmt(r['concentration_M']), fmt(r['flow_mL_min']),
                      fmt(r['molar_flow_mmol_min']), fmt(r['equivalents'])] for r in feeds]
        deck.table(slide, ['Feed', 'Component', 'C\n(M)', 'Q\n(mL/min)', 'n\n(mmol/min)', 'Equiv.'], feed_rows,
                   7.65, 1.6, [.75, 1.95, .95, 1.2, 1.5, 1.05], [.7]+[.64]*4, 15)
        identity = ('Substrate 1: (((4-methoxyphenyl)thio)methyl)trimethylsilane.\n'
                    'Ir catalyst: [Ir(dF(CF3)ppy)2(dtbpy)]PF6.\n'
                    'Feed A solvent: EtOH:pH 9 buffer = 5:1 (v/v).\n'
                    f'Feed {gas["stream"]}: air (21 mol% O2); Q is total air flow; n and equivalents refer to O2.\n'
                    f'Delivery: E-series / BLUE pump tubing (A); FF-C00 MFC ({gas["stream"]}), CV-3301, P-713 T-mixer.' if photo else
                    'Acid: 3-methyl-4-nitrobenzoic acid.\n'
                    'Feed A: acid + DPDTC + DMAP in 2-MeTHF.\n'
                    'Feed B: benzylamine in 2-MeTHF.\n'
                    'Delivery: E-series / BLUE pump tubing, two channels; P-713 T-mixer.\n'
                    'Feed B adds 1.05 equiv benzylamine relative to the acid feed.')
        deck.text(slide, identity, 7.65, 5.02, 7.8, 2.48, 16)
        total = sum(s['nominal_inlet_residence_min'] for s in stages)
        deck.text(slide, f'Total nominal sequence time: {total:.2f} min', .55, 7.03, 6.8, .4, 17, True)
        basis = ('Time = V / (liquid + air at STP); STP = 0 C, 1.01325 bar. This is not actual gas-liquid contact time.\n'
                 'Before use: confirm pressurized-air delivery, MFC calibration/pressure basis, check-valve rating and buffer/feed compatibility.' if photo else
                 'Time = V / cumulative liquid flow. Nominal times do not establish conversion or plug-flow behavior.\n'
                 'Before use: confirm feed solubility/stability, pump compatibility, pressure ratings and the heating arrangement.')
        deck.text(slide, basis, .55, 7.65, 14.9, .88, 15, color='51636C')
        all_stages.extend(stages)
        all_feeds.extend(feeds)
    assert len(deck.p.slides) == 12
    deck.p.save(out/'KHU_six_designs_for_review.pptx')
    write_csv(out/'all_stage_parameters.csv', all_stages)
    write_csv(out/'all_feed_parameters.csv', all_feeds)
    write_csv(out/'validation_checks.csv', checks)
    dump(out/'corrections.json', changes)
    assert all(sha(ROOT/p) == digest for p, digest in source_hashes.items())
    dump(out/'source_hashes.json', source_hashes)
    (out/'README.md').write_text(
        '# KHU collaborator slides\n\n12 slides: topology then parameters for each of six response sets.\n\n'
        'The MFC adjustment increment was applied deterministically to the selected archived designs: '
        'air 0.426909 -> 0.43 or 0.853818 -> 0.86 mL/min at STP. Oxygen equivalents and nominal times '
        'were recalculated. The tables use the assigned connected-process BPR, not the legacy stage-local minimum. '
        'No new model calls or council approval occurred. Old result files and presentations are untouched.\n\n'
        'These are proposals for laboratory review, not experimentally validated or safety-certified instructions. '
        'Outstanding laboratory confirmations are printed on every parameter slide. '
        'Six response sets correspond to four distinct numerical conditions; Figure 5 sets 1/3 and Figure 6 sets 2/3 match.\n\n'
        'Topology figures use the existing GUI icon renderer, including its generic pump and check-valve icons. '
        'Short display aliases are defined in the tables; full identities are retained in CSV. '
        'Production pipeline code and saved GUI inventory profiles were not changed by this export.\n')
    shutil.copy2(__file__, out/Path(__file__).name)
    print(out)
    print(f'12 slides; {len(checks)} arithmetic / reporting checks passed; original source hashes unchanged.')


if __name__ == '__main__':
    main()
