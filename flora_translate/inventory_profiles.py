"""Versioned inventory profiles for the FlowPilot Inventory Manager."""

from __future__ import annotations

import hashlib
import json
import os
import re
import zipfile
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Literal
from xml.etree import ElementTree

from pydantic import BaseModel, Field

import flora_translate.config as cfg
from flora_translate.engine.llm_agents import call_model_text
from flora_translate.schemas import (
    CollectorSpec,
    ConnectorSpec,
    DegasserSpec,
    FilterSpec,
    GasHardwareSpec,
    LAB_INVENTORY_SCHEMA_VERSION,
    LabInventory,
    LightSourceSpec,
    MixerSpec,
    PressureControllerSpec,
    PumpSpec,
    ReactorSpec,
    ReactorTrainSpec,
    SafetyAccessorySpec,
    SeparatorSpec,
    TemperatureControllerSpec,
    TubingSpec,
)


PROFILE_SCHEMA_VERSION = "flowpilot_inventory_profile_v3.0"
PROFILE_ROOT = Path(
    os.getenv(
        "FLOWPILOT_INVENTORY_PROFILE_DIR",
        "flora_translate/data/inventory_profiles",
    )
)


class EquipmentCapability(BaseModel):
    available: bool
    service_status: str = "available"
    allowed_alternatives: list[str] = Field(default_factory=list)
    notes: str = ""


class InventorySource(BaseModel):
    source_id: str
    filename: str
    media_type: str = ""
    sha256: str
    extracted_characters: int = 0
    notes: str = ""


class InventoryValidationReport(BaseModel):
    valid: bool = False
    errors: list[str] = Field(default_factory=list)
    warnings: list[str] = Field(default_factory=list)
    unresolved_fields: list[str] = Field(default_factory=list)
    checked_at_utc: str = ""


class InventoryProfile(BaseModel):
    schema_version: str = PROFILE_SCHEMA_VERSION
    profile_id: str
    name: str
    laboratory: str = ""
    version: int = 0
    status: Literal["draft", "validated", "archived"] = "draft"
    created_at_utc: str = Field(default_factory=lambda: _now())
    updated_at_utc: str = Field(default_factory=lambda: _now())
    lab_inventory: LabInventory = Field(
        default_factory=lambda: LabInventory(
            schema_version=LAB_INVENTORY_SCHEMA_VERSION,
            strict_assignment=True,
        )
    )
    equipment_capabilities: dict[str, EquipmentCapability] = Field(default_factory=dict)
    operating_constraints: dict[str, Any] = Field(default_factory=dict)
    provenance: list[InventorySource] = Field(default_factory=list)
    extraction_metadata: dict[str, Any] = Field(default_factory=dict)
    validation: InventoryValidationReport = Field(
        default_factory=InventoryValidationReport
    )

    def design_operating_limits(self) -> dict[str, Any]:
        """Return constraints in the form consumed by DesignInputPackage."""

        limits = dict(self.operating_constraints)
        inline = self.equipment_capabilities.get("inline_degassing")
        if inline is not None:
            limits["inline_degasser_available"] = inline.available
            if not inline.available:
                limits["forbidden_equipment"] = _ordered_unique(
                    _string_list(limits.get("forbidden_equipment"))
                    + [
                        "inline membrane degasser",
                        "inline vacuum degasser",
                        "dedicated inline degassing unit",
                    ]
                )
                if inline.allowed_alternatives:
                    limits["allowed_oxygen_exclusion_methods"] = _ordered_unique(
                        _string_list(limits.get("allowed_oxygen_exclusion_methods"))
                        + list(inline.allowed_alternatives)
                    )
                limits.setdefault(
                    "design_rule",
                    "Do not include an inline degasser unit operation in the final process topology.",
                )
        limits["inventory_profile"] = {
            "profile_id": self.profile_id,
            "name": self.name,
            "version": self.version,
            "schema_version": self.schema_version,
        }
        return limits


def empty_inventory_profile(name: str = "New laboratory inventory") -> InventoryProfile:
    return InventoryProfile(
        profile_id=_slug(name),
        name=name,
        lab_inventory=LabInventory(
            schema_version=LAB_INVENTORY_SCHEMA_VERSION,
            strict_assignment=True,
        ),
    )


def inventory_profile_from_payload(
    payload: dict[str, Any],
    *,
    default_name: str = "Imported laboratory inventory",
) -> InventoryProfile:
    """Load either an InventoryProfile wrapper or legacy LabInventory JSON."""

    if "lab_inventory" in payload:
        data = _migrate_profile_payload(payload)
        data.setdefault("profile_id", _slug(str(data.get("name") or default_name)))
        data.setdefault("name", default_name)
        profile = InventoryProfile.model_validate(data)
    else:
        inventory = LabInventory.model_validate(_migrate_lab_inventory_payload(payload))
        profile = InventoryProfile(
            profile_id=_slug(default_name),
            name=default_name,
            lab_inventory=inventory,
        )
    profile.validation = validate_inventory_profile(profile)
    if profile.status == "validated" and not profile.validation.valid:
        profile.status = "draft"
    return profile


def _migrate_profile_payload(payload: dict[str, Any]) -> dict[str, Any]:
    """Upgrade v1/v2 profile wrappers without discarding source fields."""

    data = json.loads(json.dumps(payload, default=str))
    source_version = str(data.get("schema_version") or "unversioned")
    data["schema_version"] = PROFILE_SCHEMA_VERSION
    data["lab_inventory"] = _migrate_lab_inventory_payload(
        data.get("lab_inventory") or {}
    )
    metadata = data.setdefault("extraction_metadata", {})
    if source_version != PROFILE_SCHEMA_VERSION:
        metadata["migrated_from_schema_version"] = source_version
        metadata["migration_target_schema_version"] = PROFILE_SCHEMA_VERSION
    return data


def _migrate_lab_inventory_payload(payload: dict[str, Any]) -> dict[str, Any]:
    """Normalize old inventory JSON into the v3 equipment categories."""

    data = json.loads(json.dumps(payload or {}, default=str))
    data["schema_version"] = LAB_INVENTORY_SCHEMA_VERSION
    data["strict_assignment"] = True
    categories = (
        "pumps",
        "tubing",
        "BPR_available",
        "light_sources",
        "gas_hardware",
        "reactors",
        "mixers",
        "pressure_controllers",
        "degassers",
        "filters",
        "separators",
        "connectors",
        "temperature_controllers",
        "collectors",
        "reactor_trains",
        "safety_accessories",
    )
    for category in categories:
        data.setdefault(category, [])

    for category in categories:
        if category == "BPR_available":
            continue
        values = data.get(category)
        if not isinstance(values, list):
            continue
        for item in values:
            if not isinstance(item, dict):
                continue
            item.setdefault("quantity", _quantity_from_name(item.get("name")))
            item.setdefault("service_status", "available")

    if data["BPR_available"] and not data["pressure_controllers"]:
        settings = sorted(set(_numbers(data["BPR_available"])))
        if settings:
            data["pressure_controllers"] = [
                {
                    "equipment_id": "pressurecontroller_legacy_bpr_settings",
                    "name": "Legacy declared BPR",
                    "type": "BPR",
                    "quantity": 1,
                    "service_status": "available",
                    "setpoints_bar": settings,
                    "min_pressure_bar": min(settings),
                    "max_pressure_bar": max(settings),
                    "notes": "Migrated from BPR_available settings.",
                }
            ]
    return data


def _quantity_from_name(value: Any) -> int:
    text = str(value or "")
    match = re.search(r"\b(\d+)\s*(?:units?|pcs?|pieces?)\b", text, re.IGNORECASE)
    return max(1, int(match.group(1))) if match else 1


def extract_inventory_profile(
    text: str,
    *,
    name: str,
    laboratory: str = "",
    use_llm: bool = True,
) -> InventoryProfile:
    """Extract an inventory profile from source text with deterministic fallback."""

    payload: dict[str, Any] = {}
    extraction_error = ""
    if use_llm and text.strip():
        try:
            system = (cfg.PROMPTS_DIR / "inventory_extract_system.txt").read_text()
            user_template = (cfg.PROMPTS_DIR / "inventory_extract_user.txt").read_text()
            response = call_model_text(
                model=cfg.MODEL_INPUT_PARSER,
                api_name="inventory_extraction_agent",
                max_tokens=5000,
                system=system,
                user_content=user_template.replace("{inventory_source_text}", text),
            )
            payload = _parse_json_object(response.text)
        except Exception as exc:
            extraction_error = str(exc)

    fallback = _fallback_inventory_extraction(text)
    payload = _merge_extraction(fallback, payload)
    payload.setdefault("name", name)
    payload.setdefault("laboratory", laboratory)
    payload.setdefault("profile_id", _slug(name))
    profile, normalization_warnings = _profile_from_extraction(payload)
    profile.extraction_metadata = {
        "method": "llm+deterministic" if use_llm and not extraction_error else "deterministic",
        "model": cfg.MODEL_INPUT_PARSER if use_llm else None,
        "source_characters": len(text),
        "llm_error": extraction_error or None,
        "normalization_warnings": normalization_warnings,
        "unresolved_fields": _string_list(payload.get("unresolved_fields")),
    }
    profile.validation = validate_inventory_profile(profile)
    return profile


def extract_uploaded_documents(
    files: list[tuple[str, bytes, str]],
) -> tuple[str, list[InventorySource], list[str]]:
    """Extract text from uploaded documents without requiring office software."""

    sections: list[str] = []
    sources: list[InventorySource] = []
    errors: list[str] = []
    for index, (filename, data, media_type) in enumerate(files, 1):
        try:
            extracted = extract_document_text(filename, data)
        except Exception as exc:
            errors.append(f"{filename}: {exc}")
            extracted = ""
        source_id = f"SRC-{index:03d}"
        sources.append(
            InventorySource(
                source_id=source_id,
                filename=filename,
                media_type=media_type,
                sha256=hashlib.sha256(data).hexdigest(),
                extracted_characters=len(extracted),
            )
        )
        if extracted:
            sections.append(f"## SOURCE {source_id}: {filename}\n{extracted}")
    return "\n\n".join(sections), sources, errors


def extract_document_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".tsv"}:
        return data.decode("utf-8", errors="replace")
    if suffix == ".json":
        parsed = json.loads(data.decode("utf-8"))
        return json.dumps(parsed, indent=2, ensure_ascii=False)
    if suffix == ".pdf":
        import fitz

        document = fitz.open(stream=data, filetype="pdf")
        return "\n\n".join(
            f"[Page {index + 1}]\n{page.get_text('text')}"
            for index, page in enumerate(document)
        )
    if suffix == ".docx":
        return _extract_ooxml_text(data, "word/document.xml")
    if suffix == ".pptx":
        with zipfile.ZipFile(BytesIO(data)) as archive:
            names = sorted(
                (
                    name
                    for name in archive.namelist()
                    if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
                ),
                key=_natural_key,
            )
            return "\n\n".join(
                f"[Slide {index + 1}]\n{_xml_text(archive.read(name))}"
                for index, name in enumerate(names)
            )
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        sections = []
        for sheet in workbook.worksheets:
            rows = [
                "\t".join("" if value is None else str(value) for value in row)
                for row in sheet.iter_rows(values_only=True)
            ]
            sections.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
        return "\n\n".join(sections)
    raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")


def validate_inventory_profile(profile: InventoryProfile) -> InventoryValidationReport:
    inventory = profile.lab_inventory
    errors: list[str] = []
    warnings: list[str] = []
    unresolved: list[str] = []

    if not inventory.reactors:
        errors.append("At least one reactor must be defined.")
    if not inventory.pumps:
        warnings.append("No pumps are defined; liquid-flow feasibility cannot be enforced.")
    if not inventory.tubing:
        warnings.append("No tubing is defined; material and pressure validation is incomplete.")
    if inventory.strict_assignment and not inventory.mixers:
        warnings.append(
            "No mixers are defined; a passive standard mixing fitting may be used "
            "as a disclosed screening assumption and must be verified before a run."
        )
    if inventory.strict_assignment and not inventory.pressure_controllers:
        warnings.append(
            "No pressure controllers are defined; pressurized designs will be blocked."
        )

    for index, pump in enumerate(inventory.pumps, 1):
        if pump.min_flow_rate_mL_min > pump.max_flow_rate_mL_min:
            errors.append(f"Pump {index} minimum flow exceeds maximum flow.")
        if pump.max_pressure_bar <= 0:
            errors.append(f"Pump {index} must have a positive pressure rating.")
    for index, reactor in enumerate(inventory.reactors, 1):
        if reactor.volume_mL <= 0 or reactor.ID_mm <= 0:
            errors.append(f"Reactor {index} requires positive volume and ID.")
        if reactor.component_volumes_mL:
            component_total = sum(reactor.component_volumes_mL)
            if abs(component_total - reactor.volume_mL) > 0.02:
                errors.append(
                    f"Reactor {index} component volumes sum to {component_total:g} mL, "
                    f"not {reactor.volume_mL:g} mL."
                )
        if reactor.allowed_temperatures_C and (
            reactor.min_temperature_C is not None
            or reactor.max_temperature_C is not None
        ):
            warnings.append(
                f"Reactor {index} defines both discrete and ranged temperatures; "
                "discrete values take precedence."
            )

    duplicate_reactors = _duplicates(
        (
            reactor.system.lower(),
            reactor.name.lower(),
            reactor.volume_mL,
            reactor.ID_mm,
        )
        for reactor in inventory.reactors
    )
    if duplicate_reactors:
        warnings.append("Duplicate reactor entries were detected.")
    if len(set(inventory.BPR_available)) != len(inventory.BPR_available):
        warnings.append("Duplicate BPR settings were detected.")
    if any(value < 0 for value in inventory.BPR_available):
        errors.append("BPR settings cannot be negative.")

    equipment_ids = [item.equipment_id for item in inventory.all_equipment()]
    if len(equipment_ids) != len(set(equipment_ids)):
        errors.append("Equipment IDs must be unique within one inventory profile.")
    reactor_by_id = {item.equipment_id: item for item in inventory.reactors}
    connector_ids = {item.equipment_id for item in inventory.connectors}
    for train in inventory.reactor_trains:
        missing_reactors = [
            item_id for item_id in train.component_reactor_ids if item_id not in reactor_by_id
        ]
        missing_connectors = [
            item_id for item_id in train.connector_ids if item_id not in connector_ids
        ]
        if missing_reactors:
            errors.append(
                f"Reactor train {train.equipment_id} references unknown reactors: "
                + ", ".join(missing_reactors)
            )
        if missing_connectors:
            errors.append(
                f"Reactor train {train.equipment_id} references unknown connectors: "
                + ", ".join(missing_connectors)
            )
        if train.total_volume_mL is not None and not missing_reactors:
            actual = sum(reactor_by_id[item_id].volume_mL for item_id in train.component_reactor_ids)
            if abs(actual - train.total_volume_mL) > 0.02:
                errors.append(
                    f"Reactor train {train.equipment_id} volume is {train.total_volume_mL:g} mL, "
                    f"but its components total {actual:g} mL."
                )

    unresolved.extend(
        str(item)
        for item in profile.extraction_metadata.get("unresolved_fields", [])
        if str(item).strip()
    )
    warnings.extend(
        str(item)
        for item in profile.extraction_metadata.get("normalization_warnings", [])
        if str(item).strip()
    )
    return InventoryValidationReport(
        valid=not errors and not unresolved,
        errors=errors,
        warnings=warnings,
        unresolved_fields=unresolved,
        checked_at_utc=_now(),
    )


def save_inventory_profile(
    profile: InventoryProfile,
    *,
    source_files: list[tuple[str, bytes, str]] | None = None,
    root: Path = PROFILE_ROOT,
) -> tuple[InventoryProfile, Path]:
    """Save an immutable new profile version and its supporting artifacts."""

    validation = validate_inventory_profile(profile)
    saved = profile.model_copy(deep=True)
    saved.validation = validation
    saved.status = "validated" if validation.valid else "draft"
    saved.updated_at_utc = _now()
    profile_root = root / _slug(saved.profile_id)
    existing = [
        int(match.group(1))
        for child in profile_root.glob("v*")
        if (match := re.fullmatch(r"v(\d+)", child.name))
    ]
    saved.version = max(existing, default=0) + 1
    version_dir = profile_root / f"v{saved.version:03d}"
    version_dir.mkdir(parents=True, exist_ok=False)

    if source_files:
        source_dir = version_dir / "sources"
        source_dir.mkdir()
        for filename, data, _ in source_files:
            (source_dir / Path(filename).name).write_bytes(data)

    _write_json(version_dir / "inventory_profile.json", saved.model_dump())
    _write_json(version_dir / "lab_inventory.json", saved.lab_inventory.model_dump())
    _write_json(version_dir / "operating_constraints.json", saved.design_operating_limits())
    _write_json(version_dir / "source_manifest.json", [item.model_dump() for item in saved.provenance])
    _write_json(version_dir / "validation_report.json", validation.model_dump())
    (version_dir / "extraction_log.jsonl").write_text(
        json.dumps(
            {
                "timestamp_utc": _now(),
                "profile_id": saved.profile_id,
                "version": saved.version,
                "extraction_metadata": saved.extraction_metadata,
            },
            ensure_ascii=False,
            default=str,
        )
        + "\n"
    )
    return saved, version_dir


def list_inventory_profiles(root: Path = PROFILE_ROOT) -> list[dict[str, Any]]:
    profiles: list[dict[str, Any]] = []
    if not root.exists():
        return profiles
    for profile_dir in sorted(path for path in root.iterdir() if path.is_dir()):
        versions = sorted(profile_dir.glob("v*/inventory_profile.json"))
        if not versions:
            continue
        path = versions[-1]
        try:
            profile = InventoryProfile.model_validate_json(path.read_text())
        except Exception:
            continue
        profiles.append(
            {
                "profile_id": profile.profile_id,
                "name": profile.name,
                "laboratory": profile.laboratory,
                "version": profile.version,
                "status": profile.status,
                "path": str(path),
                "reactor_count": len(profile.lab_inventory.reactors),
                "pump_count": len(profile.lab_inventory.pumps),
            }
        )
    return profiles


def load_inventory_profile(reference: str | Path, root: Path = PROFILE_ROOT) -> InventoryProfile:
    path = Path(reference)
    if path.is_file():
        return inventory_profile_from_payload(json.loads(path.read_text()))
    profile_dir = root / _slug(str(reference))
    versions = sorted(profile_dir.glob("v*/inventory_profile.json"))
    if not versions:
        raise FileNotFoundError(f"Inventory profile not found: {reference}")
    return inventory_profile_from_payload(json.loads(versions[-1].read_text()))


def _profile_from_extraction(payload: dict[str, Any]) -> tuple[InventoryProfile, list[str]]:
    warnings: list[str] = []
    inventory_candidate = payload.get("lab_inventory")
    if inventory_candidate and not isinstance(inventory_candidate, dict):
        warnings.append(
            "lab_inventory was not a JSON object; review the extracted equipment tables."
        )
        inventory_candidate = {}
    inventory_payload = inventory_candidate or {
        key: payload.get(key, [])
        for key in (
            "pumps",
            "tubing",
            "BPR_available",
            "light_sources",
            "gas_hardware",
            "reactors",
            "mixers",
            "pressure_controllers",
            "degassers",
            "filters",
            "separators",
            "connectors",
            "temperature_controllers",
            "collectors",
            "reactor_trains",
            "safety_accessories",
        )
    }
    inventory_payload = _migrate_lab_inventory_payload(inventory_payload)
    inventory = LabInventory(
        schema_version=LAB_INVENTORY_SCHEMA_VERSION,
        strict_assignment=True,
        pumps=_validated_items(inventory_payload.get("pumps"), PumpSpec, "pump", warnings),
        tubing=_validated_items(inventory_payload.get("tubing"), TubingSpec, "tubing", warnings),
        BPR_available=_numbers(inventory_payload.get("BPR_available")),
        light_sources=_validated_items(
            inventory_payload.get("light_sources"), LightSourceSpec, "light source", warnings
        ),
        gas_hardware=_validated_items(
            inventory_payload.get("gas_hardware"), GasHardwareSpec, "gas hardware", warnings
        ),
        reactors=_validated_items(
            inventory_payload.get("reactors"), ReactorSpec, "reactor", warnings
        ),
        mixers=_validated_items(
            inventory_payload.get("mixers"), MixerSpec, "mixer", warnings
        ),
        pressure_controllers=_validated_items(
            inventory_payload.get("pressure_controllers"),
            PressureControllerSpec,
            "pressure controller",
            warnings,
        ),
        degassers=_validated_items(
            inventory_payload.get("degassers"), DegasserSpec, "degasser", warnings
        ),
        filters=_validated_items(
            inventory_payload.get("filters"), FilterSpec, "filter", warnings
        ),
        separators=_validated_items(
            inventory_payload.get("separators"), SeparatorSpec, "separator", warnings
        ),
        connectors=_validated_items(
            inventory_payload.get("connectors"), ConnectorSpec, "connector", warnings
        ),
        temperature_controllers=_validated_items(
            inventory_payload.get("temperature_controllers"),
            TemperatureControllerSpec,
            "temperature controller",
            warnings,
        ),
        collectors=_validated_items(
            inventory_payload.get("collectors"), CollectorSpec, "collector", warnings
        ),
        reactor_trains=_validated_items(
            inventory_payload.get("reactor_trains"),
            ReactorTrainSpec,
            "reactor train",
            warnings,
        ),
        safety_accessories=_validated_items(
            inventory_payload.get("safety_accessories"),
            SafetyAccessorySpec,
            "safety accessory",
            warnings,
        ),
    )
    capabilities = _normalize_capabilities(
        payload.get("equipment_capabilities"), warnings
    )
    operating_constraints = _normalize_constraints(
        payload.get("operating_constraints"), warnings
    )
    profile = InventoryProfile(
        profile_id=_slug(str(payload.get("profile_id") or payload.get("name") or "inventory")),
        name=str(payload.get("name") or "Imported laboratory inventory"),
        laboratory=str(payload.get("laboratory") or ""),
        lab_inventory=inventory,
        equipment_capabilities=capabilities,
        operating_constraints=operating_constraints,
        extraction_metadata={
            "unresolved_fields": _string_list(payload.get("unresolved_fields")),
        },
    )
    return profile, warnings


def _fallback_inventory_extraction(text: str) -> dict[str, Any]:
    lower = text.lower()
    payload: dict[str, Any] = {
        "lab_inventory": {},
        "equipment_capabilities": {},
        "operating_constraints": {},
        "unresolved_fields": [],
    }
    no_degas = bool(
        re.search(
            r"\b(?:no|without|do not have|don't have|dont have|unavailable)\b[^.\n]{0,45}"
            r"\b(?:degas|degasser|degassing)\b",
            lower,
        )
        or re.search(r"\b(?:degas|degasser|degassing)\b[^.\n]{0,30}\bunavailable\b", lower)
    )
    if no_degas:
        alternatives = []
        if "argon" in lower or " ar " in f" {lower} ":
            alternatives.append("offline argon sparging")
        if "nitrogen" in lower or " n2 " in f" {lower} ":
            alternatives.append("offline nitrogen sparging")
        payload["equipment_capabilities"]["inline_degassing"] = {
            "available": False,
            "service_status": "unavailable",
            "allowed_alternatives": alternatives,
            "notes": "Normalized from source statement that no laboratory degasser is available.",
        }
        payload["operating_constraints"].update(
            {
                "inline_degasser_available": False,
                "forbidden_equipment": [
                    "inline membrane degasser",
                    "inline vacuum degasser",
                    "dedicated inline degassing unit",
                ],
                "design_rule": "Do not include an inline degasser unit operation in the final process topology.",
            }
        )
    return payload


def _merge_extraction(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    merged = json.loads(json.dumps(base))
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(merged.get(key), dict):
            merged[key] = _merge_extraction(merged[key], value)
        elif value not in (None, "", [], {}):
            merged[key] = value
    return merged


def _validated_items(values: Any, model, label: str, warnings: list[str]) -> list:
    if isinstance(values, dict):
        if values and all(isinstance(value, dict) for value in values.values()):
            normalized = []
            for key, value in values.items():
                item = dict(value)
                if "name" in model.model_fields:
                    item.setdefault("name", str(key))
                normalized.append(item)
            values = normalized
        else:
            values = [values]
    elif values is None:
        values = []
    elif not isinstance(values, (list, tuple)):
        warnings.append(f"Ignored non-list {label} inventory value: {values!r}")
        values = []

    output = []
    for index, value in enumerate(values, 1):
        try:
            output.append(model.model_validate(value))
        except Exception as exc:
            warnings.append(f"Skipped incomplete {label} {index}: {exc}")
    return output


def _numbers(values: Any) -> list[float]:
    if values is None:
        values = []
    elif isinstance(values, (str, int, float)):
        values = [values]
    elif not isinstance(values, (list, tuple, set)):
        values = []
    output = []
    for value in values:
        try:
            output.append(float(value))
        except (TypeError, ValueError):
            continue
    return output


def _normalize_capabilities(
    value: Any,
    warnings: list[str],
) -> dict[str, EquipmentCapability]:
    if value is None:
        return {}
    if isinstance(value, list):
        normalized: dict[str, Any] = {}
        for index, item in enumerate(value, 1):
            if not isinstance(item, dict):
                warnings.append(f"Ignored invalid equipment capability {index}: {item!r}")
                continue
            item = dict(item)
            key = item.pop("capability", None) or item.pop("name", None) or item.pop("id", None)
            if not key:
                warnings.append(f"Equipment capability {index} has no capability name.")
                continue
            normalized[str(key)] = item
        value = normalized
    if not isinstance(value, dict):
        warnings.append("equipment_capabilities was not a JSON object and was ignored.")
        return {}

    output: dict[str, EquipmentCapability] = {}
    for key, item in value.items():
        try:
            if isinstance(item, bool):
                item = {"available": item}
            elif isinstance(item, str) and item.lower() in {"available", "unavailable"}:
                item = {
                    "available": item.lower() == "available",
                    "service_status": item.lower(),
                }
            output[str(key)] = EquipmentCapability.model_validate(item)
        except Exception as exc:
            warnings.append(f"Ignored invalid equipment capability {key}: {exc}")
    return output


def _normalize_constraints(value: Any, warnings: list[str]) -> dict[str, Any]:
    if value is None:
        return {}
    if isinstance(value, dict):
        return dict(value)
    if isinstance(value, str):
        warnings.append("Converted string-form operating_constraints to a review list.")
        return {"additional_constraints": [value]}
    if isinstance(value, list):
        merged: dict[str, Any] = {}
        remainder: list[Any] = []
        for item in value:
            if isinstance(item, dict) and len(item) == 1:
                key, item_value = next(iter(item.items()))
                if key not in merged:
                    merged[str(key)] = item_value
                    continue
            remainder.append(item)
        if remainder:
            merged["additional_constraints"] = remainder
        warnings.append("Converted list-form operating_constraints to a JSON object.")
        return merged
    warnings.append("operating_constraints had an unsupported shape and was ignored.")
    return {}


def _extract_ooxml_text(data: bytes, member: str) -> str:
    with zipfile.ZipFile(BytesIO(data)) as archive:
        return _xml_text(archive.read(member))


def _xml_text(data: bytes) -> str:
    root = ElementTree.fromstring(data)
    chunks = [node.text for node in root.iter() if node.text and node.tag.endswith("}t")]
    return "\n".join(chunks)


def _parse_json_object(text: str) -> dict[str, Any]:
    stripped = text.strip()
    if stripped.startswith("```"):
        stripped = stripped.split("```", 2)[1]
        if stripped.lstrip().startswith("json"):
            stripped = stripped.lstrip()[4:]
    start = stripped.find("{")
    end = stripped.rfind("}")
    if start < 0 or end < start:
        raise ValueError("Inventory extraction response did not contain JSON")
    payload = json.loads(stripped[start : end + 1])
    if not isinstance(payload, dict):
        raise ValueError("Inventory extraction response must be a JSON object")
    return payload


def _natural_key(value: str) -> list[Any]:
    return [int(part) if part.isdigit() else part for part in re.split(r"(\d+)", value)]


def _duplicates(values) -> bool:
    seen = set()
    for value in values:
        if value in seen:
            return True
        seen.add(value)
    return False


def _string_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, str):
        return [value] if value.strip() else []
    if isinstance(value, (list, tuple, set)):
        return [str(item) for item in value if str(item).strip()]
    return [str(value)]


def _ordered_unique(values: list[str]) -> list[str]:
    return list(dict.fromkeys(values))


def _slug(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", value.lower()).strip("_")
    return slug or "inventory"


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _write_json(path: Path, payload: Any) -> None:
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False, default=str))
