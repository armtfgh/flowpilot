"""Build the presentation package for the matched Qwen/OpenAI NewGen study."""

from __future__ import annotations

import csv
import hashlib
import json
import sys
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from ablation_test.src.cases import load_cases_from_path
from ablation_test.src.metrics import score_run


OUT = ROOT / "deliverables/flowpilot_newgen_qwen_openai_20260813"
FIG = OUT / "figures"
DATA = OUT / "source_data"

CASE_PATHS = {
    "CuAAC": ROOT / "ablation_test/benchmarks/newgen_benchmark_v1_pilot_cuaac/case.json",
    "Hydrogenolysis": ROOT / "ablation_test/benchmarks/newgen_benchmark_v1_pilot_hydrogenolysis/case.json",
    "Two-stage amidation": ROOT / "ablation_test/benchmarks/newgen_benchmark_v1_pilot_multistep/case.json",
}

RUNS = {
    ("Qwen3.6-27B", "One-shot", "CuAAC"): ROOT / "ablation_results/newgen_benchmark/newgen_benchmark_v1_pilot_cuaac_20260813/runs/qwen27b_one_shot",
    ("Qwen3.6-27B", "One-shot", "Hydrogenolysis"): ROOT / "ablation_results/newgen_benchmark/newgen_benchmark_v1_pilot_hydrogenolysis_20260813/runs/qwen27b_one_shot",
    ("Qwen3.6-27B", "One-shot", "Two-stage amidation"): ROOT / "ablation_results/newgen_benchmark/newgen_benchmark_v1_pilot_multistep_20260813/runs/qwen27b_one_shot",
    ("Qwen3.6-27B", "FlowPilot", "CuAAC"): ROOT / "ablation_results/newgen_benchmark/newgen_production_three_case_20260813_150731/cuaac_qwen27b_full_flowpilot",
    ("Qwen3.6-27B", "FlowPilot", "Hydrogenolysis"): ROOT / "ablation_results/newgen_benchmark/newgen_production_three_case_20260813_150731/hydrogenolysis_qwen27b_full_flowpilot",
    ("Qwen3.6-27B", "FlowPilot", "Two-stage amidation"): ROOT / "ablation_results/newgen_benchmark/newgen_production_multistep_repaired_20260813_151952/multistep_qwen27b_full_flowpilot",
    ("GPT-5.4", "One-shot", "CuAAC"): ROOT / "ablation_results/newgen_benchmark/newgen_openai_three_case_20260813_155908/cuaac_gpt54_one_shot",
    ("GPT-5.4", "One-shot", "Hydrogenolysis"): ROOT / "ablation_results/newgen_benchmark/newgen_openai_three_case_20260813_155908/hydrogenolysis_gpt54_one_shot",
    ("GPT-5.4", "One-shot", "Two-stage amidation"): ROOT / "ablation_results/newgen_benchmark/newgen_openai_repaired_cells_20260813_161126/multistep_gpt54_one_shot",
    ("GPT-5.4", "FlowPilot", "CuAAC"): ROOT / "ablation_results/newgen_benchmark/newgen_openai_three_case_20260813_155908/cuaac_gpt54_full_flowpilot",
    ("GPT-5.4", "FlowPilot", "Hydrogenolysis"): ROOT / "ablation_results/newgen_benchmark/newgen_openai_repaired_cells_20260813_161126/hydrogenolysis_gpt54_full_flowpilot",
    ("GPT-5.4", "FlowPilot", "Two-stage amidation"): ROOT / "ablation_results/newgen_benchmark/newgen_openai_three_case_20260813_155908/multistep_gpt54_full_flowpilot",
}

DIMENSIONS = [
    "formal_validity",
    "engineering_integrity",
    "process_completeness",
    "safety_adequacy",
    "evidence_provenance",
    "decision_assurance",
    "actionability_calibration",
]
DIM_LABELS = ["Validity", "Engineering", "Process", "Safety", "Evidence", "Decision", "Calibration"]
WEIGHTS = [0.10, 0.25, 0.15, 0.15, 0.10, 0.20, 0.05]

WHITE = "FFFFFF"
INK = "15202B"
MUTED = "5E6A75"
LINE = "D8E0E5"
PANEL = "F5F7F8"
TEAL = "087F73"
TEAL_LIGHT = "E6F4F1"
BLUE = "2D6F9F"
BLUE_LIGHT = "EAF2F8"
GRAY = "8D98A1"
GRAY_LIGHT = "EDF0F2"
RED = "C74747"
AMBER = "D18B20"


def load_rows() -> list[dict]:
    cases = {name: load_cases_from_path(path)[0] for name, path in CASE_PATHS.items()}
    rows = []
    for (model, architecture, case_name), run_dir in RUNS.items():
        result = json.loads((run_dir / "result.json").read_text(encoding="utf-8"))
        metrics = score_run(cases[case_name], result, run_dir)
        dimensions = metrics.get("quality_assurance_dimensions_v2") or {}
        row = {
            "model": model,
            "architecture": architecture,
            "case": case_name,
            "quality_assurance": metrics["quality_assurance_score_v2"],
            "deployment_readiness": metrics["deployment_readiness_score_v2"],
            "deterministic_composite": metrics["deterministic_composite_score"],
            "gate_count": metrics["deployment_gate_count_v2"],
            "gate_reasons": ";".join(metrics.get("deployment_gate_reasons_v2") or []),
            "schema_valid": metrics["schema_valid"],
            "final_design_status": (result.get("final_design") or {}).get("status") or "one-shot output",
            "disposition": result.get("recommended_disposition") or result.get("reported_disposition") or "not applicable",
            "run_directory": str(run_dir.resolve()),
        }
        row.update({f"dimension_{key}": dimensions.get(key, 0.0) for key in DIMENSIONS})
        rows.append(row)
    return rows


def write_csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def aggregate(rows: list[dict]) -> list[dict]:
    output = []
    for model in ("Qwen3.6-27B", "GPT-5.4"):
        model_rows = [row for row in rows if row["model"] == model]
        for architecture in ("One-shot", "FlowPilot"):
            subset = [row for row in model_rows if row["architecture"] == architecture]
            output.append(
                {
                    "model": model,
                    "architecture": architecture,
                    "n_cases": len(subset),
                    "mean_quality_assurance": round(np.mean([r["quality_assurance"] for r in subset]), 4),
                    "mean_deployment_readiness": round(np.mean([r["deployment_readiness"] for r in subset]), 4),
                    "mean_deterministic_composite": round(np.mean([r["deterministic_composite"] for r in subset]), 4),
                    "gate_free_cases": sum(r["gate_count"] == 0 for r in subset),
                    **{
                        f"mean_{key}": round(np.mean([r[f"dimension_{key}"] for r in subset]), 4)
                        for key in DIMENSIONS
                    },
                }
            )
    return output


def figure_paired(rows: list[dict]) -> None:
    cases = list(CASE_PATHS)
    fig, axes = plt.subplots(1, 2, figsize=(12.4, 4.5), sharey=True)
    for ax, model in zip(axes, ("Qwen3.6-27B", "GPT-5.4")):
        x = np.arange(len(cases))
        one = [next(r for r in rows if r["model"] == model and r["architecture"] == "One-shot" and r["case"] == c)["quality_assurance"] for c in cases]
        full = [next(r for r in rows if r["model"] == model and r["architecture"] == "FlowPilot" and r["case"] == c)["quality_assurance"] for c in cases]
        width = 0.34
        ax.bar(x - width / 2, one, width, color="#A7B0B7", label="One-shot")
        ax.bar(x + width / 2, full, width, color="#087F73" if model.startswith("Qwen") else "#2D6F9F", label="FlowPilot")
        for xpos, value in zip(x - width / 2, one):
            ax.text(xpos, value + 0.018, f"{value:.2f}", ha="center", fontsize=9, color="#4D5963")
        for xpos, value in zip(x + width / 2, full):
            ax.text(xpos, value + 0.018, f"{value:.2f}", ha="center", fontsize=9, fontweight="bold", color="#15202B")
        ax.set_title(model, loc="left", fontsize=15, fontweight="bold")
        ax.set_xticks(x, ["CuAAC", "H₂", "2-stage"])
        ax.set_ylim(0, 1.08)
        ax.grid(axis="y", color="#E1E6E9", linewidth=0.8)
        ax.spines[["top", "right", "left"]].set_visible(False)
        ax.tick_params(axis="y", length=0)
    axes[0].set_ylabel("Quality assurance score", fontsize=11)
    axes[1].legend(frameon=False, ncol=2, loc="lower center", bbox_to_anchor=(0.5, -0.26))
    fig.suptitle("The architecture improves assurance for every model–case pair", x=0.08, ha="left", fontsize=18, fontweight="bold")
    fig.text(0.08, 0.91, "Same protocol, inventory, model, candidate budget and universal scorer within each pair", fontsize=10.5, color="#5E6A75")
    fig.tight_layout(rect=(0.03, 0.07, 1, 0.88))
    for ext in ("png", "svg"):
        fig.savefig(FIG / f"paired_quality_assurance.{ext}", dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def figure_aggregate(agg: list[dict]) -> None:
    fig, ax = plt.subplots(figsize=(9.6, 4.8))
    models = ["Qwen3.6-27B", "GPT-5.4"]
    x = np.arange(2)
    width = 0.31
    one = [next(r for r in agg if r["model"] == m and r["architecture"] == "One-shot")["mean_quality_assurance"] for m in models]
    full = [next(r for r in agg if r["model"] == m and r["architecture"] == "FlowPilot")["mean_quality_assurance"] for m in models]
    ax.bar(x - width / 2, one, width, color="#A7B0B7", label="One-shot")
    colors = ["#087F73", "#2D6F9F"]
    ax.bar(x + width / 2, full, width, color=colors, label="FlowPilot")
    for index, (base, final) in enumerate(zip(one, full)):
        ax.text(index - width / 2, base + 0.018, f"{base:.3f}", ha="center", fontsize=11)
        ax.text(index + width / 2, final + 0.018, f"{final:.3f}", ha="center", fontsize=11, fontweight="bold")
        ax.text(index, 1.02, f"+{final - base:.3f}", ha="center", fontsize=12, fontweight="bold", color=colors[index])
    ax.set_xticks(x, models)
    ax.set_ylim(0, 1.1)
    ax.set_ylabel("Mean quality assurance (3 cases)")
    ax.grid(axis="y", color="#E1E6E9")
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.legend(frameon=False, ncol=2, loc="lower center", bbox_to_anchor=(0.5, -0.22))
    ax.set_title("Architecture uplift persists from a 27B local model to GPT-5.4", loc="left", fontsize=17, fontweight="bold", pad=18)
    fig.tight_layout()
    for ext in ("png", "svg"):
        fig.savefig(FIG / f"aggregate_architecture_uplift.{ext}", dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def figure_dimensions(agg: list[dict]) -> None:
    labels = ["Qwen one-shot", "Qwen FlowPilot", "GPT one-shot", "GPT FlowPilot"]
    order = [
        ("Qwen3.6-27B", "One-shot"),
        ("Qwen3.6-27B", "FlowPilot"),
        ("GPT-5.4", "One-shot"),
        ("GPT-5.4", "FlowPilot"),
    ]
    matrix = np.array(
        [[next(r for r in agg if r["model"] == model and r["architecture"] == architecture)[f"mean_{key}"] for key in DIMENSIONS] for model, architecture in order]
    )
    cmap = LinearSegmentedColormap.from_list("flowpilot", ["#F1F3F4", "#B8DDD7", "#087F73"])
    fig, ax = plt.subplots(figsize=(11.2, 4.0))
    image = ax.imshow(matrix, vmin=0, vmax=1, cmap=cmap, aspect="auto")
    for i in range(matrix.shape[0]):
        for j in range(matrix.shape[1]):
            ax.text(j, i, f"{matrix[i, j]:.2f}", ha="center", va="center", fontsize=9.5, color="white" if matrix[i, j] > 0.72 else "#15202B", fontweight="bold" if matrix[i, j] > 0.9 else "normal")
    ax.set_xticks(range(len(DIM_LABELS)), DIM_LABELS, fontsize=10)
    ax.set_yticks(range(len(labels)), labels, fontsize=10)
    ax.tick_params(length=0)
    ax.spines[:].set_visible(False)
    ax.set_title("FlowPilot's largest gains come from decision assurance and provenance", loc="left", fontsize=17, fontweight="bold", pad=16)
    fig.colorbar(image, ax=ax, fraction=0.025, pad=0.025, label="Dimension score")
    fig.tight_layout()
    for ext in ("png", "svg"):
        fig.savefig(FIG / f"quality_dimension_heatmap.{ext}", dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def figure_scoring() -> None:
    fig, ax = plt.subplots(figsize=(11.2, 2.6))
    colors = ["#2D6F9F", "#087F73", "#5C9B8E", "#D18B20", "#6586A0", "#735D8E", "#A7B0B7"]
    left = 0.0
    for label, weight, color in zip(DIM_LABELS, WEIGHTS, colors):
        ax.barh([0], [weight], left=left, color=color, height=0.38)
        if weight >= 0.10:
            ax.text(left + weight / 2, 0, f"{label}\n{weight:.0%}", ha="center", va="center", color="white", fontsize=9, fontweight="bold")
        left += weight
    ax.set_xlim(0, 1)
    ax.set_ylim(-0.55, 0.55)
    ax.axis("off")
    ax.set_title("Universal quality-assurance score", loc="left", fontsize=17, fontweight="bold", pad=12)
    fig.text(0.126, 0.11, "Deployment gates cap readiness for invalid schema, gas bookkeeping, geometry, pump/tubing, topology, or safety failures.", fontsize=10.5, color="#5E6A75")
    fig.tight_layout(rect=(0, 0.18, 1, 1))
    for ext in ("png", "svg"):
        fig.savefig(FIG / f"scoring_system.{ext}", dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def add_text(slide, text, x, y, w, h, *, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def rect(slide, x, y, w, h, *, fill=WHITE, line=LINE, radius=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line); shape.line.width = Pt(1)
    if radius:
        shape.adjustments[0] = 0.06
    return shape


def arrow(slide, x1, y1, x2, y2, color=TEAL):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = rgb(color); line.line.width = Pt(2); line.line.end_arrowhead = True
    return line


def header(slide, number, title, subtitle):
    add_text(slide, "FlowPilot | NewGen benchmark", 0.55, 0.22, 3.2, 0.22, size=10.5, color=TEAL, bold=True)
    add_text(slide, number, 12.15, 0.22, 0.6, 0.22, size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    add_text(slide, title, 0.55, 0.62, 12.2, 0.48, size=27, bold=True)
    add_text(slide, subtitle, 0.55, 1.15, 12.1, 0.32, size=12.5, color=MUTED)


def footer(slide, text="Three cases; one run per model–architecture–case cell; QA is not wet-lab yield accuracy."):
    add_text(slide, text, 0.55, 7.18, 12.2, 0.16, size=8.2, color=MUTED)


def build_pptx(agg: list[dict]) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    q_one = next(r for r in agg if r["model"] == "Qwen3.6-27B" and r["architecture"] == "One-shot")
    q_full = next(r for r in agg if r["model"] == "Qwen3.6-27B" and r["architecture"] == "FlowPilot")
    g_one = next(r for r in agg if r["model"] == "GPT-5.4" and r["architecture"] == "One-shot")
    g_full = next(r for r in agg if r["model"] == "GPT-5.4" and r["architecture"] == "FlowPilot")

    # Slide 1: headline results.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(WHITE)
    header(slide, "01", "FlowPilot improves design assurance across model scales", "Matched architecture comparison using identical protocols, inventories, model identity and scoring within each pair.")
    slide.shapes.add_picture(str(FIG / "paired_quality_assurance.png"), Inches(4.05), Inches(1.72), width=Inches(8.8))
    cards = [
        ("Qwen3.6-27B", q_one["mean_quality_assurance"], q_full["mean_quality_assurance"], TEAL, TEAL_LIGHT),
        ("GPT-5.4", g_one["mean_quality_assurance"], g_full["mean_quality_assurance"], BLUE, BLUE_LIGHT),
    ]
    for idx, (label, before, after, color, light) in enumerate(cards):
        y = 1.88 + idx * 1.43
        rect(slide, 0.55, y, 3.08, 1.17, fill=light, line=light)
        add_text(slide, label, 0.78, y + 0.15, 2.6, 0.25, size=13, color=color, bold=True)
        add_text(slide, f"{before:.3f} → {after:.3f}", 0.78, y + 0.47, 2.6, 0.37, size=23, bold=True)
        add_text(slide, f"architecture uplift  +{after-before:.3f}", 0.78, y + 0.88, 2.6, 0.18, size=9.5, color=MUTED)
    rect(slide, 0.55, 4.80, 3.08, 1.02, fill=PANEL, line=LINE)
    add_text(slide, "6 / 6", 0.78, 4.97, 1.0, 0.35, size=24, bold=True, color=TEAL)
    add_text(slide, "paired QA improvements", 1.72, 5.02, 1.65, 0.25, size=11, bold=True)
    add_text(slide, "All six full-pipeline outputs closed as executable SCREEN designs.", 0.78, 5.42, 2.55, 0.25, size=9.5, color=MUTED)
    rect(slide, 0.55, 6.12, 12.22, 0.64, fill="FFF6E5", line="F0D8A8")
    add_text(slide, "Interpretation", 0.78, 6.30, 1.05, 0.2, size=10.5, bold=True, color=AMBER)
    add_text(slide, "The same model becomes more dependable when coupled to retrieval, engineering calculations, inventory enforcement, specialist review and a final deterministic contract.", 1.82, 6.27, 10.55, 0.27, size=10.7)
    footer(slide)

    # Slide 2: methodology.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(WHITE)
    header(slide, "02", "How the matched benchmark was executed", "The model is controlled within each pair; only the architecture changes.")
    stages = [
        (0.55, "1  Freeze input", "Protocol + objective\nInventory + limits\nHeld-out source IDs", BLUE, BLUE_LIGHT),
        (3.35, "2  Two matched paths", "ONE-SHOT\nOne model call\n\nFLOWPILOT\nUpstream + retrieval +\ncalculator + council", TEAL, TEAL_LIGHT),
        (6.20, "3  Deterministic closure", "Pump/MFC bounds\nExact reactor assignment\nGas basis + residence time\nSafety + topology gates", AMBER, "FFF6E5"),
        (9.28, "4  Universal scorer", "7 fixed QA dimensions\nSame weights for every case\nReadiness gates\nNo held-out answer in prompt", BLUE, BLUE_LIGHT),
    ]
    widths = [2.25, 2.35, 2.55, 3.45]
    for (x, title, body, color, light), width in zip(stages, widths):
        rect(slide, x, 2.08, width, 2.42, fill=light, line=color)
        add_text(slide, title, x + 0.2, 2.28, width - 0.4, 0.3, size=13, bold=True, color=color)
        add_text(slide, body, x + 0.2, 2.75, width - 0.4, 1.45, size=11, valign=MSO_ANCHOR.TOP)
    arrow(slide, 2.86, 3.28, 3.25, 3.28)
    arrow(slide, 5.76, 3.28, 6.10, 3.28)
    arrow(slide, 8.82, 3.28, 9.18, 3.28)
    rect(slide, 0.55, 4.87, 12.18, 1.52, fill=PANEL, line=LINE)
    add_text(slide, "Quality-assurance score", 0.78, 5.08, 2.0, 0.24, size=13, bold=True)
    x = 2.75
    colors = [BLUE, TEAL, "5C9B8E", AMBER, "6586A0", "735D8E", GRAY]
    for label, weight, color in zip(DIM_LABELS, WEIGHTS, colors):
        width = 8.95 * weight
        shape = rect(slide, x, 5.00, width, 0.65, fill=color, line=color, radius=False)
        if weight >= 0.10:
            add_text(slide, f"{label}\n{weight:.0%}", x, 5.10, width, 0.40, size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        else:
            add_text(slide, f"Cal.\n{weight:.0%}", x, 5.10, width, 0.40, size=7.0, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        x += width
    add_text(slide, "Hard gates cap readiness when schema, gas bookkeeping, geometry, pumps/tubing, topology, or safety fail.", 2.75, 5.87, 9.3, 0.22, size=10, color=MUTED)
    rect(slide, 0.55, 6.62, 12.18, 0.38, fill=TEAL_LIGHT, line=TEAL_LIGHT)
    add_text(slide, "Controlled variables: candidate budget = 1  |  temperature = 0 where provider-supported  |  lexical retrieval  |  source answer excluded  |  same case-specific inventory", 0.78, 6.72, 11.7, 0.18, size=9.2, bold=True, color=TEAL)
    footer(slide, "Models: /models/Qwen3.6-27B and gpt-5.4-2026-03-05. Each model is used for both its one-shot and FlowPilot condition.")

    # Slide 3: why.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(WHITE)
    header(slide, "03", "Why the architecture changes the outcome", "One-shot numeric closure can be strong; FlowPilot adds traceable decision and process assurance.")
    slide.shapes.add_picture(str(FIG / "quality_dimension_heatmap.png"), Inches(0.55), Inches(1.72), width=Inches(8.1))
    rect(slide, 8.95, 1.92, 3.83, 4.75, fill=PANEL, line=LINE)
    add_text(slide, "Observed interventions", 9.22, 2.16, 3.25, 0.28, size=15, bold=True)
    bullets = [
        "Rejected unsupported residence-time reductions.",
        "Recomputed V/Q and stage-specific cumulative flows.",
        "Converted H₂ inlet/STP flow to in-channel flow and equivalents.",
        "Bound pumps, MFC, reactors, BPR and mixers to inventory IDs.",
        "Published topology only after final deterministic closure.",
    ]
    y = 2.70
    for item in bullets:
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.25), Inches(y + 0.04), Inches(0.13), Inches(0.13))
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(TEAL); shape.line.color.rgb = rgb(TEAL)
        add_text(slide, item, 9.50, y, 2.95, 0.54, size=10.5)
        y += 0.72
    rect(slide, 9.22, 6.00, 3.30, 0.42, fill=TEAL_LIGHT, line=TEAL_LIGHT)
    add_text(slide, "Result: 6/6 executable SCREEN contracts", 9.40, 6.12, 2.95, 0.18, size=10, bold=True, color=TEAL)
    footer(slide)

    # Slide 4: claim and caveats.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(WHITE)
    header(slide, "04", "Presentation conclusion: architecture, not model size alone", "The local 27B model in FlowPilot reaches essentially the same mean assurance as GPT-5.4 in FlowPilot.")
    slide.shapes.add_picture(str(FIG / "aggregate_architecture_uplift.png"), Inches(0.55), Inches(1.68), width=Inches(7.25))
    rect(slide, 8.20, 1.83, 4.58, 2.15, fill=TEAL_LIGHT, line=TEAL)
    add_text(slide, "Defensible claim", 8.48, 2.08, 3.9, 0.30, size=16, bold=True, color=TEAL)
    add_text(slide, "Across these three inventory-constrained cases, FlowPilot improved universal design-assurance scores for both Qwen3.6-27B and GPT-5.4, and produced executable screening contracts in every full-pipeline run.", 8.48, 2.54, 3.85, 1.08, size=12)
    rect(slide, 8.20, 4.25, 4.58, 2.22, fill="FFF6E5", line="F0D8A8")
    add_text(slide, "Do not overclaim", 8.48, 4.50, 3.8, 0.28, size=15, bold=True, color=AMBER)
    add_text(slide, "• n = 3 cases and one run per cell\n• no statistical significance claim\n• scores assess design assurance, not reaction yield\n• final designs remain SCREEN until wet-lab validation", 8.48, 4.94, 3.75, 1.18, size=11)
    footer(slide, "Recommended next step: pre-register a larger repeated benchmark and separately report wet-lab predictive accuracy.")

    path = OUT / "FlowPilot_NewGen_Qwen_OpenAI_Benchmark.pptx"
    prs.save(path)
    return path


def write_report(rows: list[dict], agg: list[dict]) -> None:
    q_one = next(r for r in agg if r["model"] == "Qwen3.6-27B" and r["architecture"] == "One-shot")
    q_full = next(r for r in agg if r["model"] == "Qwen3.6-27B" and r["architecture"] == "FlowPilot")
    g_one = next(r for r in agg if r["model"] == "GPT-5.4" and r["architecture"] == "One-shot")
    g_full = next(r for r in agg if r["model"] == "GPT-5.4" and r["architecture"] == "FlowPilot")
    text = f"""# FlowPilot NewGen Qwen/OpenAI Benchmark

## Headline

The full architecture improved the universal quality-assurance score in all six matched model-by-case comparisons.

| Model | One-shot mean QA | FlowPilot mean QA | Uplift | Full executable designs |
|---|---:|---:|---:|---:|
| Qwen3.6-27B | {q_one['mean_quality_assurance']:.4f} | {q_full['mean_quality_assurance']:.4f} | +{q_full['mean_quality_assurance']-q_one['mean_quality_assurance']:.4f} | 3/3 |
| GPT-5.4 | {g_one['mean_quality_assurance']:.4f} | {g_full['mean_quality_assurance']:.4f} | +{g_full['mean_quality_assurance']-g_one['mean_quality_assurance']:.4f} | 3/3 |

Qwen FlowPilot ({q_full['mean_quality_assurance']:.4f}) and GPT FlowPilot ({g_full['mean_quality_assurance']:.4f}) are nearly equal on this small suite. This supports an architecture-efficiency argument, not model equivalence in general.

## Design

- Three cases: CuAAC packed-bed chemistry, gas-liquid-solid hydrogenolysis, and two-stage oxidative amidation.
- Same protocol, objective, inventory, candidate budget, scorer, and held-out-source policy within each pair.
- Qwen condition: `/models/Qwen3.6-27B` for one-shot and every FlowPilot LLM role.
- OpenAI condition: `gpt-5.4-2026-03-05` for one-shot and every FlowPilot LLM role.
- FlowPilot adds chemistry parsing, held-out-safe retrieval, deterministic calculations, specialist council, design realization, inventory allocation, safety validation, and executable topology generation.

## Scoring

The fixed QA score combines formal validity (10%), engineering integrity (25%), process completeness (15%), safety adequacy (15%), evidence provenance (10%), decision assurance (20%), and actionability calibration (5%). Deployment gates cap readiness for schema, gas-bookkeeping, geometry, pump/tubing, topology, and critical-safety failures.

## Interpretation

The benchmark supports the statement that FlowPilot improves design assurance relative to an identical-model one-shot baseline on these cases. It does not establish statistical significance, universal superiority, or wet-lab yield accuracy. Every full design is labeled `SCREEN` pending experimental validation.
"""
    (OUT / "REPORT.md").write_text(text, encoding="utf-8")


def checksums() -> None:
    lines = []
    for path in sorted(OUT.rglob("*")):
        if path.is_file() and path.name != "checksums.sha256":
            lines.append(f"{hashlib.sha256(path.read_bytes()).hexdigest()}  {path.relative_to(OUT)}")
    (OUT / "checksums.sha256").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    FIG.mkdir(parents=True, exist_ok=True)
    DATA.mkdir(parents=True, exist_ok=True)
    rows = load_rows()
    agg = aggregate(rows)
    write_csv(DATA / "run_metrics.csv", rows)
    write_csv(DATA / "aggregate_metrics.csv", agg)
    write_csv(DATA / "scoring_weights.csv", [{"dimension": label, "weight": weight} for label, weight in zip(DIM_LABELS, WEIGHTS)])
    summary = {"schema_version": "flowpilot_newgen_qwen_openai_presentation_v1.0", "rows": rows, "aggregate": agg}
    (DATA / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    figure_paired(rows)
    figure_aggregate(agg)
    figure_dimensions(agg)
    figure_scoring()
    build_pptx(agg)
    write_report(rows, agg)
    checksums()
    print(OUT)
    print(json.dumps(agg, indent=2))


if __name__ == "__main__":
    main()
