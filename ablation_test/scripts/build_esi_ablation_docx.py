#!/usr/bin/env python3
"""Build the manuscript-ready Word version of the NewGen 2.0 ESI section."""

from __future__ import annotations

import argparse
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
PACKAGE_DIR = ROOT / "deliverables" / "manuscript_benchmark_visualizations_20260825"
DEFAULT_INPUT = PACKAGE_DIR / "ESI_SECTION_1_ABLATION_TEST_AND_ARCHITECTURE_COMPARISON.md"
DEFAULT_OUTPUT = PACKAGE_DIR / "ESI_SECTION_1_ABLATION_TEST_AND_ARCHITECTURE_COMPARISON.docx"


FORMULAS = {
    "`C(i,k) = mean_j[s(i,k,j)]`": r"$$C_{i,k}=\frac{1}{J}\sum_{j=1}^{J}s_{i,k,j}$$",
    "`S(i) = mean_k in A(i)[C(i,k)] / 4`": (
        r"$$S_i=\frac{1}{|A_i|}\sum_{k\in A_i}\frac{C_{i,k}}{4}$$"
    ),
    "`R(model, architecture, repeat) = mean_case[S(i)]`": (
        r"$$R_{m,a,r}=\frac{1}{3}\sum_{q=1}^{3}S_{m,a,r,q}$$"
    ),
    "`Delta(i) = S(FlowPilot, model, case, repeat) - S(one-shot, model, case, repeat)`": (
        r"$$\Delta_{m,q,r}=S_{\mathrm{FlowPilot},m,q,r}"
        r"-S_{\mathrm{one\text{-}shot},m,q,r}$$"
    ),
}


def docx_markdown(markdown: str) -> str:
    """Remove browser-only disclosure markup while retaining all manuscript content."""
    markdown = re.sub(r"<!--.*?-->", "", markdown, flags=re.DOTALL)
    markdown = re.sub(
        r"<summary><strong>(.*?)</strong></summary>",
        r"#### \1",
        markdown,
        flags=re.DOTALL,
    )
    markdown = markdown.replace("<details>", "").replace("</details>", "")
    for source, replacement in FORMULAS.items():
        markdown = markdown.replace(source, replacement)
    return markdown


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def set_run_font(run, name: str, size: float, *, bold: bool | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def configure_styles(document: Document) -> None:
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08

    heading_sizes = {"Title": 16, "Heading 1": 15, "Heading 2": 12.5, "Heading 3": 11}
    for style_name, size in heading_sizes.items():
        if style_name not in styles:
            continue
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)

    if "Source Code" not in styles:
        source = styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
    else:
        source = styles["Source Code"]
    source.font.name = "Consolas"
    source.font.size = Pt(7.5)
    source._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    source.paragraph_format.left_indent = Cm(0.25)
    source.paragraph_format.right_indent = Cm(0.25)
    source.paragraph_format.space_before = Pt(2)
    source.paragraph_format.space_after = Pt(3)
    source.paragraph_format.line_spacing = 1.0

    if "Caption" in styles:
        caption = styles["Caption"]
        caption.font.name = "Times New Roman"
        caption.font.size = Pt(9)
        caption.font.italic = False
        caption.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
        caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        caption.paragraph_format.space_before = Pt(3)
        caption.paragraph_format.space_after = Pt(8)


def style_document(path: Path) -> None:
    document = Document(path)
    document.core_properties.title = "Ablation test and architecture comparison"
    document.core_properties.subject = "FlowPilot NewGen 2.0 benchmark"
    document.core_properties.author = "FlowPilot"
    document.core_properties.keywords = "FlowPilot, ablation, architecture comparison, benchmark"

    for section in document.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.9)
        section.bottom_margin = Cm(1.9)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)
        if not section.footer.paragraphs:
            section.footer.add_paragraph()
        add_page_number(section.footer.paragraphs[0])

    configure_styles(document)

    for paragraph in document.paragraphs:
        has_drawing = bool(paragraph._p.xpath(".//w:drawing"))
        if has_drawing:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(3)

        if paragraph.style and paragraph.style.name == "Source Code":
            p_pr = paragraph._p.get_or_add_pPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F2F2F2")
            p_pr.append(shading)

    for table in document.tables:
        table.autofit = True
        if table.rows:
            set_repeat_table_header(table.rows[0])
            for cell in table.rows[0].cells:
                set_cell_shading(cell, "D9E2F3")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_font(run, "Arial", 8.5, bold=True)
        for row in table.rows[1:]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        set_run_font(run, "Times New Roman", 8.5)

    text_width = min(
        section.page_width - section.left_margin - section.right_margin
        for section in document.sections
    )
    for shape in document.inline_shapes:
        ratio = text_width / shape.width
        shape.width = int(shape.width * ratio)
        shape.height = int(shape.height * ratio)

    document.save(path)


def build(input_path: Path, output_path: Path) -> None:
    if not input_path.exists():
        raise FileNotFoundError(input_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    transformed = docx_markdown(input_path.read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory(prefix="flowpilot_esi_docx_") as tmp:
        temporary_markdown = Path(tmp) / "section.md"
        temporary_markdown.write_text(transformed, encoding="utf-8")
        subprocess.run(
            [
                "pandoc",
                str(temporary_markdown),
                "--from=gfm+tex_math_dollars",
                "--to=docx",
                f"--resource-path={input_path.parent}",
                "--standalone",
                "--output",
                str(output_path),
            ],
            check=True,
        )
    style_document(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    build(args.input.resolve(), args.output.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
