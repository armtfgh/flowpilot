"""Assemble FlowPilot ESI Figures S1-S10 and captions into a review document."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "deliverables" / "flowpilot_esi_figures_s1_s10_20260811"
CAPTIONS = PACKAGE / "documentation" / "captions.md"
OUTPUT = PACKAGE / "FlowPilot_ESI_Figures_S1-S10.docx"


def parse_captions() -> dict[int, str]:
    text = CAPTIONS.read_text(encoding="utf-8")
    matches = list(re.finditer(r"^## Figure S(\d+)\s*$", text, flags=re.MULTILINE))
    captions: dict[int, str] = {}
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        captions[int(match.group(1))] = text[match.end() : end].strip()
    return captions


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def image_dimensions(path: Path, max_width: float = 7.0, max_height: float = 8.35) -> tuple[float, float]:
    with Image.open(path) as image:
        width_px, height_px = image.size
    ratio = width_px / height_px
    width = min(max_width, max_height * ratio)
    height = width / ratio
    if height > max_height:
        height = max_height
        width = height * ratio
    return width, height


def build() -> Path:
    captions = parse_captions()
    if sorted(captions) != list(range(1, 11)):
        raise ValueError(f"Expected captions S1-S10, found {sorted(captions)}")

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(16)
    styles["Title"].font.bold = True

    header = section.header.paragraphs[0]
    header.text = "FlowPilot | Electronic Supporting Information | Figures S1-S10"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = None

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)

    for number in range(1, 11):
        if number > 1:
            document.add_section(WD_SECTION.NEW_PAGE)

        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
        heading.paragraph_format.space_after = Pt(4)
        run = heading.add_run(f"Figure S{number}")
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(12)

        image_path = PACKAGE / "figures" / f"Figure_S{number:02d}.png"
        width, height = image_dimensions(image_path)
        figure_paragraph = document.add_paragraph()
        figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figure_paragraph.paragraph_format.space_after = Pt(4)
        figure_paragraph.add_run().add_picture(str(image_path), width=Inches(width), height=Inches(height))

        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        caption.paragraph_format.space_before = Pt(2)
        caption.paragraph_format.space_after = Pt(0)
        caption.paragraph_format.line_spacing = 1.0
        caption_run = caption.add_run(captions[number])
        caption_run.font.name = "Arial"
        caption_run.font.size = Pt(8)

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
