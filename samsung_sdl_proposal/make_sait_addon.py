"""SAIT interview addon slides — matches the plain template style
(white bg, bold black title top-left, double rule). Paste into the main deck."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BLACK = RGBColor(0x11, 0x11, 0x11)
DARK = RGBColor(0x22, 0x2A, 0x35)
GRAY = RGBColor(0x55, 0x5E, 0x6A)
TEAL = RGBColor(0x1F, 0x7A, 0x6E)
RED = RGBColor(0xB4, 0x2E, 0x2E)
AMBER = RGBColor(0xB97, 0x0, 0x0) if False else RGBColor(0xB9, 0x77, 0x00)
LIGHT = RGBColor(0xF2, 0xF4, 0xF6)
LINE = RGBColor(0xC9, 0xCF, 0xD8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

SW, SH = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def txt(s, x, y, w, h, paras, size=12, color=GRAY, align=PP_ALIGN.LEFT,
        space_after=5, line_spacing=1.08):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        if isinstance(para, str):
            para = [(para, {})]
        for t, o in para:
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.italic = o.get("italic", False)
            r.font.color.rgb = o.get("color", color)
    return tb


def rect(s, x, y, w, h, fill=LIGHT, line=None, line_w=0.75, round_=True):
    sp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        x, y, w, h)
    if round_:
        sp.adjustments[0] = 0.06
    sp.fill.solid()
    sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    return sp


def header(s, title):
    """Template look: bold black title top-left + double rule."""
    txt(s, Inches(0.35), Inches(0.12), Inches(12.6), Inches(0.5),
        [[(title, {"size": 19, "bold": True, "color": BLACK})]])
    rect(s, Inches(0.35), Inches(0.62), Inches(12.63), Pt(2.2), fill=BLACK, round_=False)
    rect(s, Inches(0.35), Inches(0.68), Inches(12.63), Pt(0.9), fill=BLACK, round_=False)


def card(s, x, y, w, h, title, lines, accent=TEAL, tsize=12.5, bsize=11):
    rect(s, x, y, w, h, fill=WHITE, line=LINE)
    rect(s, x, y, Inches(0.08), h, fill=accent)
    paras = [[(title, {"size": tsize, "bold": True, "color": DARK})]]
    for ln in lines:
        if isinstance(ln, tuple):
            lead, rest = ln
            paras.append([("• ", {"size": bsize, "color": accent, "bold": True}),
                          (lead, {"size": bsize, "bold": True, "color": DARK}),
                          (rest, {"size": bsize, "color": GRAY})])
        else:
            paras.append([("• ", {"size": bsize, "color": accent, "bold": True}),
                          (ln, {"size": bsize, "color": GRAY})])
    txt(s, x + Inches(0.22), y + Inches(0.1), w - Inches(0.35), h - Inches(0.2),
        paras, space_after=4)


def foot(s, note):
    txt(s, Inches(0.35), Inches(7.05), Inches(12.6), Inches(0.35),
        [[(note, {"size": 9.5, "italic": True, "color": GRAY})]])


# ================= P1-A: requirements answered with numbers =================
s = slide()
header(s, "Problem 1 — Answering the requirements directly (indicative numbers)")
rows = [
    ("Requirement", "Naive fully-automated SDL", "This proposal (two-tier, human-in-the-loop)"),
    ("High-\nthroughput",
     "Bottlenecked anyway: probing/anneal\nrobotics limit end-to-end rate",
     "Fast tier: 10–15 structural fingerprints/day + in-situ data every cycle (runs 24/7 unattended). "
     "Expensive tier: ~2x fewer electrical tests per insight (shown in silico) — effective throughput "
     "is measured in insights/month, not wafers/day"),
    ("Footprint",
     "150–250 m2 vacuum-integrated cluster\n+ custom transfer robotics",
     "One lab bay, ~60–80 m2: ALD + in-situ ellipsometry, XRR/GIXRD, RTP, probe station — all standard "
     "commercial tools, no custom integration"),
    ("Cost",
     "$8–15M class, 3+ years of robotics\nintegration before first value\n(cf. Ada: ~$8M for one property)",
     "$1.5–2.5M capex + 2–3 FTE. No bespoke robotics in Phases 0–1; automation added in Phase 2 only "
     "where data proves it pays"),
    ("Period",
     "Value arrives only after full\nintegration succeeds",
     "Phase 0 (6 mo): decision engine validated on existing/retrospective data. "
     "Phase 1 (by month 12–18): first virtual-metrology model + mapped process window — "
     "each phase delivers standalone value"),
]
y = Inches(0.95)
widths = [Inches(1.55), Inches(4.1), Inches(7.0)]
heights = [Inches(0.42), Inches(1.28), Inches(1.05), Inches(1.28), Inches(1.28)]
for r, (c0, c1, c2) in enumerate(rows):
    x = Inches(0.35)
    hh = heights[r]
    for c, (cw, val) in enumerate(zip(widths, (c0, c1, c2))):
        if r == 0:
            rect(s, x, y, cw, hh, fill=DARK, round_=False)
            txt(s, x + Inches(0.08), y + Inches(0.06), cw - Inches(0.16), hh,
                [[(val, {"size": 11.5, "bold": True, "color": WHITE})]])
        else:
            fill = WHITE if c != 1 else LIGHT
            rect(s, x, y, cw, hh, fill=fill, line=LINE, round_=False)
            col = DARK if c == 0 else (GRAY if c == 1 else DARK)
            bold = (c == 0)
            txt(s, x + Inches(0.08), y + Inches(0.05), cw - Inches(0.16), hh,
                [[(val, {"size": 10 if c else 11, "bold": bold, "color": col})]],
                line_spacing=1.04, space_after=2)
        x += cw
    y += hh
foot(s, "Numbers are indicative planning estimates for discussion (tool list & vendor quotes in backup) — "
        "the design choice they reflect: spend capital on measurements and intelligence, not on robotizing steps "
        "a human does better during R&D.")

# ================= P1-B: fab connection / industrial payoff =================
s = slide()
header(s, "Problem 1 — From SDL to the fab: what Samsung gets beyond one material")
txt(s, Inches(0.35), Inches(0.82), Inches(12.6), Inches(0.35),
    [[("The SDL's product is not a hero film — it is a TRANSLATOR (fast optical/structural data -> electrical "
       "performance) plus a map of where that translator is valid. That is fab language: virtual metrology.",
       {"size": 12.5, "bold": True, "color": TEAL})]])
cw, ch, gap = Inches(3.05), Inches(2.5), Inches(0.12)
x0, y0 = Inches(0.35), Inches(1.35)
card(s, x0, y0, cw, ch, "1 · R&D acceleration (now)",
     [("2x fewer electrical-test cycles ", "to a mapped process window (in-silico result) — directly shortens "
       "dielectric screening campaigns"),
      ("Candidates arrive with confidence labels, ", "not just best guesses — fewer surprises downstream")],
     accent=TEAL)
card(s, x0 + (cw + gap), y0, cw, ch, "2 · Pilot-line transfer (next)",
     [("The trained translator deploys ", "against inline ellipsometry/XRR that production lines already run "
       "on every wafer"),
      ("Validity map tells engineers ", "WHEN inline data can replace destructive electrical/TEM sampling — "
       "and when it cannot (risk control, not just prediction)")],
     accent=AMBER)
card(s, x0 + 2 * (cw + gap), y0, cw, ch, "3 · Compounding data asset",
     [("Every sample carries full provenance: ", "recipe -> tool logs -> fingerprint -> anneal -> electrical"),
      ("Dataset + models transfer ", "to the next dielectric program (doped HfO2, FE-HZO, new precursors) — "
       "the asset outlives any single target spec")],
     accent=TEAL)
card(s, x0 + 3 * (cw + gap) + Inches(0.06), y0, Inches(3.15), ch, "Where it lands (examples)",
     [("DRAM capacitor dielectrics: ", "EOT scaling needs new doped ZrO2/HfO2 stacks every node"),
      ("GAA / advanced logic gate stacks: ", "interface + leakage co-optimization"),
      ("FE-HZO memory: ", "wake-up/endurance vs anneal — same two-tier problem"),
      ("Same engine, new objectives: ", "only the metrology list changes")],
     accent=DARK)
rect(s, Inches(0.35), Inches(4.15), Inches(12.63), Inches(0.02), fill=LINE, round_=False)
txt(s, Inches(0.35), Inches(4.3), Inches(12.6), Inches(2.3), [
    [("KPI view (how I would be measured): ", {"size": 12, "bold": True, "color": DARK}),
     ("electrical-test wafers per learned process window (target: -50%);  time from 'new candidate material' "
      "to 'mapped safe operating window' (target: months, not years);  fraction of routine electrical tests "
      "replaced by trusted virtual metrology (target grows every quarter);  dataset reuse across programs "
      "(second program should start at ~30% lower cost).",
      {"size": 12, "color": GRAY})],
], line_spacing=1.15)
foot(s, "This is the industrialization argument: the algorithm is the cheap part — the durable value is a "
        "validated translator + provenance-complete data, both of which plug into existing fab practice (inline "
        "metrology, APC/FDC), no new fab hardware required.")

# ================= P1-C: risks =================
s = slide()
header(s, "Problem 1 — Risks I have designed for (and what stays honest)")
risks = [
    ("Tool drift & metrology noise can mimic physics",
     "Scheduled reference depositions + 2 process replicates per batch; noise is an explicit, learned term in "
     "the model. In the benchmark, removing replicates collapsed hidden-zone detection to chance — this is "
     "load-bearing, not bookkeeping.", TEAL),
    ("Probe/anneal automation is where thin-film SDLs die",
     "Human-in-the-loop tier + pre-patterned test substrates / Hg probe from day one. Robotics only in Phase 2, "
     "only where Phase-1 data proves the ROI.", TEAL),
    ("Lab test structure is not the device stack",
     "Position outputs as candidate windows; Phase-2 transfer validation on device-representative stacks before "
     "any claim reaches a product team.", AMBER),
    ("Algorithm can lock in on a wrong 'untrustworthy' label",
     "Found in silico: an early wrong noise estimate can freeze exploration. Fixed with an exploration safeguard "
     "(one random pick per batch). Cost to find it in simulation: days. In the lab: months.", AMBER),
    ("Deposition temperature changes are slow (30-60 min stabilization)",
     "Cost-aware scheduling batches experiments by thermal budget — the optimizer pays a modeled price for "
     "temperature moves.", TEAL),
]
y = Inches(0.95)
for title, mit, acc in risks:
    rect(s, Inches(0.35), y, Inches(12.63), Inches(1.08), fill=WHITE, line=LINE)
    rect(s, Inches(0.35), y, Inches(0.08), Inches(1.08), fill=acc)
    txt(s, Inches(0.6), y + Inches(0.07), Inches(12.2), Inches(0.4),
        [[(title, {"size": 12, "bold": True, "color": DARK})]])
    txt(s, Inches(0.6), y + Inches(0.42), Inches(12.2), Inches(0.6),
        [[(mit, {"size": 10.5, "color": GRAY})]], line_spacing=1.05)
    y += Inches(1.18)
foot(s, "The honest boundary: the in-silico result validates the decision engine, not the material physics — "
        "the first Phase-1 milestone is measuring the real fingerprint-to-electrical signal strength.")

# ================= P2-A: 3 requirements -> 3 answers =================
s = slide()
header(s, "Problem 2 — The three requirements map to three architectural decisions")
cw, ch = Inches(4.05), Inches(4.6)
x0, y0, gap = Inches(0.35), Inches(1.0), Inches(0.22)
card(s, x0, y0, cw, ch,
     "REQ 1 · Individual AND integrated operation",
     [("Hierarchical control planes: ", "device layer -> per-SDL orchestrator -> fleet layer. Each SDL is fully "
       "functional standalone; the fleet layer only coordinates."),
      ("Shared modules (robot arms, furnaces, analyzers) ", "are RESOURCES with a reservation-based scheduler — "
       "no workflow owns a device, it books capabilities."),
      ("Graceful degradation: ", "if the fleet layer is down, every SDL keeps running its local queue. "
       "Integration is additive, never a dependency."),
      ("Fab analogy: ", "this is MES + equipment-level control, the pattern Samsung already trusts at scale — "
       "applied to research labs.")],
     accent=TEAL, tsize=12.5, bsize=10.5)
card(s, x0 + cw + gap, y0, cw, ch,
     "REQ 2 · Usable by non-SDL experts",
     [("Researchers compose CAPABILITIES, not devices: ", "a no-code canvas of verbs ('anneal at 500 C', "
       "'measure thickness') — the scheduler resolves which physical tool executes."),
      ("Recipe templates per domain ", "(thin film, battery, organic): start from a working experiment, edit "
       "parameters, not code."),
      ("Dry-run by default: ", "every workflow executes first against the digital twin (simulated devices) — "
       "errors surface before hardware moves."),
      ("Safety interlocks & approval gates ", "are declared in the device contract, enforced by the "
       "orchestrator — not by user discipline.")],
     accent=AMBER, tsize=12.5, bsize=10.5)
card(s, x0 + 2 * (cw + gap), y0, cw, ch,
     "REQ 3 · HW add/change/remove, minimal burden",
     [("Device-plugin contract: ", "a containerized driver + a machine-readable capability descriptor "
       "(verbs, parameter ranges, units, safety limits, calibration state)."),
      ("Register -> everything updates automatically: ", "scheduler sees a new resource, GUI auto-generates its "
       "panel, provenance logging attaches — zero orchestrator code changes."),
      ("Workflows bind to capabilities, not device IDs: ", "swapping a furnace vendor is a config change; "
       "no workflow is rewritten."),
      ("Versioned configs + twin regression test ", "before a changed device re-enters production use.")],
     accent=DARK, tsize=12.5, bsize=10.5)
rect(s, Inches(0.35), Inches(5.8), Inches(12.63), Inches(0.9), fill=LIGHT)
txt(s, Inches(0.6), Inches(5.92), Inches(12.2), Inches(0.7), [
    [("One sentence: ", {"size": 12, "bold": True, "color": DARK}),
     ("standardize the CONTRACT (capability descriptor + communication grammar), containerize the drivers, "
      "centralize only scheduling and data — then labs scale like fleets, not like one-off projects.",
      {"size": 12, "color": GRAY})]], line_spacing=1.15)
foot(s, "Standards I would build on rather than reinvent: SiLA2 / OPC-UA for lab devices (the 'SECS/GEM of the "
        "lab'), gRPC for transport, containerized microservices per module.")

# ================= P2-B: new tool onboarding lifecycle =================
s = slide()
header(s, "Problem 2 — 'Minimal burden' made concrete: a new instrument joins in ~1 day")
steps = [
    ("1 · Wrap", "Vendor driver goes into a standard SDK container — whatever the tool speaks "
     "(RS-232, Modbus, TCP, vendor DLL), the container's outward face is the universal grammar.", TEAL),
    ("2 · Declare", "Capability descriptor (YAML): verbs it offers, parameter ranges & units, timing, "
     "safety limits, consumables, calibration schedule. This file IS the integration.", TEAL),
    ("3 · Register", "Container announces itself to the orchestrator. Scheduler now sees a new bookable "
     "resource; data layer attaches provenance logging automatically.", AMBER),
    ("4 · Verify", "Digital-twin dry-run + supervised first runs (shadow mode). Device is 'trusted' only "
     "after passing — same discipline as fab equipment qualification.", AMBER),
    ("5 · Use", "GUI panel auto-generated from the descriptor; existing workflows that request matching "
     "capabilities can now be scheduled onto it. No workflow rewrites anywhere.", DARK),
]
x = Inches(0.35)
y = Inches(1.05)
w = Inches(2.42)
for i, (t, body, acc) in enumerate(steps):
    card(s, x + i * (w + Inches(0.13)), y, w, Inches(3.3), t, [body],
         accent=acc, tsize=13, bsize=10.5)
    if i < 4:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                x + (i + 1) * (w + Inches(0.13)) - Inches(0.16),
                                y + Inches(1.45), Inches(0.2), Inches(0.22))
        ar.fill.solid(); ar.fill.fore_color.rgb = GRAY; ar.line.fill.background()
        ar.shadow.inherit = False
rect(s, Inches(0.35), Inches(4.65), Inches(12.63), Inches(1.95), fill=WHITE, line=LINE)
txt(s, Inches(0.6), Inches(4.8), Inches(12.15), Inches(1.7), [
    [("Why this is the industrially credible answer:  ", {"size": 12, "bold": True, "color": DARK}),
     ("integration cost is the #1 reason lab-automation programs stall. Making the capability descriptor the "
      "single integration artifact converts 'integration projects' (weeks of bespoke code, one engineer per "
      "instrument) into 'configuration work' (one file + one container). It also future-proofs the fleet: ",
      {"size": 12, "color": GRAY})],
    [("removal ", {"size": 12, "bold": True, "color": DARK}),
     ("is just deregistration (scheduler routes around it);  ", {"size": 12, "color": GRAY}),
     ("replacement ", {"size": 12, "bold": True, "color": DARK}),
     ("is a descriptor diff;  ", {"size": 12, "color": GRAY}),
     ("upgrades ", {"size": 12, "bold": True, "color": DARK}),
     ("are versioned and twin-tested before re-entering service — the same lifecycle discipline fabs apply to "
      "production equipment, scaled down to research.", {"size": 12, "color": GRAY})],
], line_spacing=1.15)
foot(s, "This slide operationalizes the microservice diagram: the architecture is the same — this is the "
        "day-in-the-life proof that the burden is actually minimal.")

# ================= P2-C: scheduling + data backbone =================
s = slide()
header(s, "Problem 2 — The two hard problems under the hood: shared-resource scheduling & data")
card(s, Inches(0.35), Inches(1.0), Inches(6.15), Inches(4.5),
     "Scheduling shared modules (the real multi-SDL problem)",
     [("Reservation-based allocation: ", "workflows request capability + time window; the scheduler books "
       "physical resources — shared robot arms and furnaces stop being collision points."),
      ("Priorities & preemption policy: ", "campaign deadlines, maintenance windows, and calibration runs "
       "coexist by policy, not by hallway negotiation."),
      ("Deadlock avoidance: ", "all-or-nothing resource acquisition per workflow step (no workflow holds the "
       "robot while waiting for a busy furnace)."),
      ("Fairness across SDLs: ", "per-lab budgets/quotas so one hot project cannot starve the others."),
      ("Failure containment: ", "a faulted device is quarantined; queued steps re-route to equivalent "
       "capabilities or pause cleanly — the fleet never cascades.")],
     accent=TEAL, tsize=12.5, bsize=11)
card(s, Inches(6.75), Inches(1.0), Inches(6.2), Inches(4.5),
     "One data backbone for every SDL (thin film, battery, organic...)",
     [("Sample-centric provenance graph: ", "recipe -> device logs -> measurements -> model decisions — one "
       "queryable record per sample, across labs and modalities."),
      ("Shared schema, domain-specific vocabularies: ", "units, materials, and method ontologies per domain "
       "plug into one common core — integration without forcing uniformity."),
      ("Streaming telemetry with drift alarms: ", "FDC-style monitoring for research tools (reference-run "
       "tracking catches drift before it poisons a campaign)."),
      ("Decisions are data too: ", "every algorithm suggestion is logged with model version + inputs — "
       "full auditability, reproducible campaigns, and the training corpus for the next generation of models."),
      ("This is the asset that compounds: ", "hardware depreciates; the provenance-complete dataset and the "
       "models trained on it appreciate.")],
     accent=DARK, tsize=12.5, bsize=11)
rect(s, Inches(0.35), Inches(5.75), Inches(12.63), Inches(1.0), fill=LIGHT)
txt(s, Inches(0.6), Inches(5.88), Inches(12.15), Inches(0.8), [
    [("Positioning: ", {"size": 12, "bold": True, "color": DARK}),
     ("the GUI and microservices make the system usable; the scheduler and the data backbone make it an "
      "INSTITUTIONAL capability — multiple labs, one operational discipline, one growing asset. That is the "
      "difference between a lab demo and infrastructure a company runs for ten years.",
      {"size": 12, "color": GRAY})]], line_spacing=1.15)
foot(s, "Both components are deliberately boring technology (reservation schedulers, message queues, versioned "
        "schemas) — novelty lives in the decision algorithms, reliability lives here.")

out = "/home/amirreza/Documents/codes/Flow Agent/samsung_sdl_proposal/SAIT_addon_slides.pptx"
prs.save(out)
print("saved", out)
