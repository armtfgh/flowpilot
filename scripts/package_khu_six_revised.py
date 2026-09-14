"""Create a portable reviewed PPTX and tables without changing the KHU source deck."""
import argparse
import json
from pathlib import Path
import shutil
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image
from scripts.check_khu_six_revised import check


def num(x):
    return '--' if x is None else f'{x:.5g}' if isinstance(x,(int,float)) else str(x)


class Deck:
    def __init__(self):
        self.p=Presentation();self.p.slide_width=Inches(16);self.p.slide_height=Inches(9)
    def text(self,s,x,y,w,h,value,size=19,color='263238',bold=False):
        box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=box.text_frame;tf.word_wrap=True
        for i,line in enumerate(str(value).split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
            p.space_after=Pt(5)
        return box
    def slide(self,title,subtitle=''):
        s=self.p.slides.add_slide(self.p.slide_layouts[6])
        self.text(s,.55,.3,14.9,.65,title,27,bold=True)
        self.text(s,.55,.98,14.9,.5,subtitle,14,'51636C')
        self.text(s,.55,8.55,14.9,.28,'FlowPilot | KHU revised inventory and response sets | 14 September 2026 | Screening proposals, not measured yields',11,'51636C')
        return s
    def table(self,s,headers,rows,y=1.7,widths=None,font=17):
        n=len(rows)
        height=max(.75,(n+1)*.48)
        table=s.shapes.add_table(n+1,len(headers),Inches(.55),Inches(y),Inches(14.9),Inches(height)).table
        if widths:
            for c,w in zip(table.columns,widths):c.width=Inches(w)
        for i,row in enumerate([headers,*rows]):
            for j,value in enumerate(row):
                cell=table.cell(i,j);cell.text=str(value);cell.margin_left=Inches(.09);cell.margin_right=Inches(.06)
                cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('1D5965' if i==0 else 'EDF3F4' if i%2 else 'FFFFFF')
                for p in cell.text_frame.paragraphs:
                    p.font.name='Arial';p.font.size=Pt(font);p.font.bold=i==0
                    p.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '263238')
        return height
    def picture(self,s,path):
        with Image.open(path) as im:w,h=im.size
        scale=min(14.9/w,6.7/h)
        s.shapes.add_picture(str(path),Inches((16-w*scale)/2),Inches(1.65+(6.7-h*scale)/2),width=Inches(w*scale),height=Inches(h*scale))


def main():
    parser=argparse.ArgumentParser();parser.add_argument('root',type=Path);args=parser.parse_args()
    out=args.root/'presentation';out.mkdir(exist_ok=True)
    d=Deck();summary=[]
    title=d.slide('KHU revised flow designs','Two chemistries | Three response sets each | Original files retained')
    d.text(title,.7,2,14.5,5.8,
        'Inputs: exact batch text and revised response slides 5, 8, 11, 16, 19 and 22.\n'
        'Models: Claude Opus 4.6 upstream; Claude Sonnet 4.6 downstream. Twelve connected-process candidates reviewed per run.\n'
        'Inventory: revised workbook, pump-tubing guide and separately sourced no-inline-degasser constraint.\n'
        'Checks: stage arithmetic, component stoichiometry, pump/module resources, light compatibility, gas addition stage and check valve.\n'
        'Gas residence times are nominal V / (liquid flow + inlet/STP gas flow). They are not measured hydrodynamic residence times.\n'
        'No measured flow yields or conversions were supplied for these six revised requests.',24)
    overview=d.slide('Six revised sets | Condition summary','Latest independently checked attempt per set; identical outcomes are retained when selected by the pipeline')
    batch_shown=set()
    for case in sorted(args.root.glob('figure*_set*')):
        results=sorted(case.glob('attempt*/result.json'))
        attempts=[(p,check(p)) for p in results]
        good=[(p,r) for p,r in attempts if r['all_passed']]
        chosen=(good or attempts)[-1] if attempts else None
        folder=chosen[0].parent if chosen else sorted(case.glob('attempt*'))[-1]
        if not chosen and not (folder/'failure.txt').exists():
            continue
        source=json.loads((folder/'provided_input.json').read_text())
        title=f'Figure {source["figure"]} | Revised set {source["set"]}'
        if source['figure'] not in batch_shown:
            batch=d.slide(f'Figure {source["figure"]} | Initial batch protocol','Original batch protocol text from the supplied KHU PPT; not the previous flow design')
            d.text(batch,.7,1.7,14.5,6.5,source['protocol'],20)
            batch_shown.add(source['figure'])
        inputslide=d.slide(title+' | Input',f'Exact revised response text from KHU PPT slide {source["source_slide"]}; full batch text in supplied input files')
        y=1.7
        for q,answer in source['answers'].items():
            d.text(inputslide,.65,y,2,.45,q,17,'1D5965',True)
            h=max(.65,len(answer)/125*.29)
            d.text(inputslide,2.65,y,12.65,h+.3,answer,17)
            y+=h+.33
        if not chosen:
            fail=d.slide(title+' | Run failed','No numerical result was produced; no design is represented as successful.')
            d.text(fail,.7,2,14.5,5,(folder/'failure.txt').read_text()[-2500:],16)
            summary.append({'case':case.name,'passed':False,'attempt':folder.name});continue
        path,review=chosen;r=json.loads(path.read_text())
        target=out/case.name;target.mkdir(exist_ok=True)
        for name in ['provided_input.json','intake_package.json','request.json','inventory_profile.json','result.json','independent_checks.json','stage_summary.csv','component_feeds.csv','summary.json','display_revision_provenance.json','reused_generation_provenance.json','reused_council_provenance.json']:
            if (folder/name).exists():shutil.copy2(folder/name,target/name)
        image=folder/'gui_export/process.png'
        label='Checked screening proposal' if review['all_passed'] else 'DIAGNOSTIC ONLY: checks failed'
        top=d.slide(title+' | Process topology',f'{label} | {folder.name} | Generic pump symbols; exact equipment labels below')
        if image.exists():
            shutil.copy2(image,target/'topology.png');d.picture(top,image)
            svg=folder/'gui_export/process.svg'
            if svg.exists():shutil.copy2(svg,target/'topology.svg')
        else:d.text(top,.7,2,14.5,5,'No executable topology. Blocking reasons: '+str(review['failed']),22,'A93232')
        stage=d.slide(title+' | Stage conditions',label)
        rows=review['stage_rows']
        d.table(stage,['Stage','Reactor / ID','T (C)','Liquid\nmL/min','Gas inlet/STP\nmL/min','Nominal time\nmin','BPR\nbar'],[
            [x['stage'],f'{num(x["volume_mL"])} mL {x["material"]}\n{num(x["ID_mm"])} mm',num(x['temperature_C']),num(x['liquid_flow_mL_min']),num(x['gas_inlet_STP_mL_min']),num(x['nominal_inlet_residence_min']),num(x['BPR_bar'])] for x in rows],widths=[.8,3.1,1.1,2.2,2.7,3,2])
        d.text(stage,.7,4.05,14.5,2.4,'\n'.join(f'Stage {x["stage"]}: {x["reactor"]}; '+(f'{x["light"]}; {num(x["wavelength_nm"])} nm' if x['wavelength_nm'] else 'thermal operation, no light source') for x in rows),20)
        d.text(stage,.7,6.5,14.5,1.6,'Timing: V / cumulative liquid flow for liquid-only stages; V / (cumulative liquid + inlet/STP gas flow) for gas-fed stages. This reporting basis does not demonstrate conversion or physical gas-liquid contact time.\nPressure basis and equipment ratings must be confirmed before experimental operation.',17,'51636C')
        components=review['component_rows']
        for start in range(0,len(components),6):
            feed=d.slide(title+' | Feed composition',f'{label} | Component-level concentration and molar flow; gas flow referenced to 273.15 K and 1.01325 bar')
            subset=components[start:start+6]
            d.table(feed,['Stream / stage','Component','C (M)','Q (mL/min)','Molar flow\nmmol/min','Equiv.'],[
                [f'{x["stream"]} / {x["stage"]}',x['component'],num(x['concentration_M']),num(x['flow_mL_min']),num(x['molar_flow_mmol_min']),num(x['equivalents'])] for x in subset],widths=[1.4,6,1.4,1.8,2.5,1.8],font=15)
            pumps={x['equipment_id']:x['name'] for x in json.loads((folder/'inventory_profile.json').read_text())['lab_inventory']['pumps']}
            streams=r.get('proposal',{}).get('streams',[])
            d.text(feed,.7,6.3,14.5,1.8,'\n'.join(f'Stream {x["stream_label"]}: {pumps.get(x.get("pump_equipment_id"),x.get("pump_equipment_id"))}; solvent: {x.get("solvent") or "gas"}' for x in streams),16)
        audit=d.slide(title+' | Review and rationale',f'{label} | Full model conversations and every attempt remain in the campaign directory')
        rationale=r.get('scientific_assessment',{}).get('chief',{}).get('justification','No chief rationale available.')
        d.text(audit,.7,1.7,14.5,3,str(rationale)[:2200],18)
        d.text(audit,.7,5.1,14.5,2.8,
            f'Independent checks: {sum(review["checks"].values())}/{len(review["checks"])} passed.\n'
            f'Failed checks: {", ".join(review["failed"]) or "None"}.\n'+
            ('Pending laboratory confirmation: pressurized air delivery; check-valve maximum pressure; MFC reference/quantity; buffer formulation; full feed compatibility.\n' if source['figure']==5 else 'Pending laboratory confirmation: full feed compatibility with selected pump tubing; stock solution solubility; pressure reference and heating arrangement.\n')+
            'Preparation quantities, procedures, uncertainty and source provenance are retained in each result JSON. No experimental yield claim is made.',18)
        preparations=[x for x in r.get('final_design',{}).get('operating_procedure',[]) if x.get('section')=='preparation']
        for start in range(0,len(preparations),3):
            prep=d.slide(title+' | Feed preparation',label+' | Model/software-generated preparation for chemist review')
            d.text(prep,.7,1.7,14.5,6.4,'\n\n'.join(x['instruction'] for x in preparations[start:start+3]),19)
        calls=[json.loads(line) for line in (folder/'llm_calls.jsonl').read_text().splitlines()] if (folder/'llm_calls.jsonl').exists() else []
        tokens={k:sum(x.get('usage',{}).get(k,0) or 0 for x in calls) for k in ['input_tokens','output_tokens']}
        runstats=json.loads((folder/'summary.json').read_text())
        reuse=json.loads((folder/'reused_generation_provenance.json').read_text()) if (folder/'reused_generation_provenance.json').exists() else None
        display_only=(folder/'display_revision_provenance.json').exists()
        usage_basis='Fresh council-review calls only; upstream and initial proposal reused' if reuse else 'Original generation run; subsequent diagram refresh uses no model calls' if display_only else 'This complete generation run'
        if (folder/'reused_council_provenance.json').exists():
            usage_basis='Fresh Skeptic/Chief calls only; generation and four exact-matched specialist reviews reused'
        (target/'usage.json').write_text(json.dumps(dict(tokens=tokens,elapsed_seconds=runstats.get('elapsed_seconds'),model_calls=len(calls),usage_basis=usage_basis,reused_generation_provenance=reuse),indent=2))
        d.text(audit,.7,7.63,14.5,.32,usage_basis,13,'51636C')
        d.text(audit,.7,8.03,14.5,.4,f'Model calls: {len(calls)} | Input tokens: {tokens["input_tokens"]:,} | Output tokens: {tokens["output_tokens"]:,} | Runtime: {num(runstats.get("elapsed_seconds",0)/60)} min',14,'51636C')
        summary.append({'case':case.name,'passed':review['all_passed'],'attempt':folder.name,'stages':rows,'failed':review['failed']})
    (out/'selected_results.json').write_text(json.dumps(summary,indent=2))
    overview_rows=[]
    for row in summary:
        stages=row.get('stages',[])
        overview_rows.append([row['case'].replace('_',' '),'PASS' if row['passed'] else 'REVIEW',
            ' / '.join(num(s['volume_mL']) for s in stages),
            ' / '.join(num(s['nominal_inlet_residence_min']) for s in stages),
            ' / '.join(num(s['temperature_C']) for s in stages),
            num(stages[0]['BPR_bar']) if stages else '--'])
    if overview_rows:
        d.table(overview,['Case','Checks','V1 / V2 (mL)','Time1 / Time2 (min)','T1 / T2 (C)','BPR (bar)'],overview_rows,
                widths=[3,1.4,2.4,3.8,2.5,1.8],font=17)
    d.text(overview,.7,6.8,14.5,1.2,'PASS denotes the listed software checks, not experimental validation. No buffer molarity is inferred from pH. Unconfirmed laboratory ratings and mixture compatibility remain explicit pre-run requirements.',18,'51636C')
    shutil.copy2(ROOT/'inventory_khu/KHU_inventory_20260914_v5.json',out/'KHU_inventory_20260914_v5.json')
    d.p.save(out/'KHU_revised_six_designs.pptx')
    print(out/'KHU_revised_six_designs.pptx');print('Slides',len(d.p.slides))


if __name__=='__main__':main()
