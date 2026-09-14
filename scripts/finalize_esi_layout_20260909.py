"""Update cached contents-page numbers from a rendered revision and audit it."""

import hashlib
import json
import re

import fitz
from docx import Document
from docx.oxml.ns import nsmap, qn
from lxml import etree

from revise_current_esi_20260909 import OUTPUT, OUT, SOURCE


def xpath(node, expression):
    return etree.XPath(expression, namespaces=nsmap)(node)


def normalized(value):
    return re.sub(r"\s+", " ", value).strip()


def main():
    doc = Document(OUTPUT)
    pdf = fitz.open(OUT / "esi_revised_20260909.pdf")
    pages = [normalized(p.get_text()) for p in pdf]
    records = []
    for hyperlink in xpath(doc.element, ".//w:sdt//w:hyperlink[@w:anchor]"):
        anchor = hyperlink.get(qn("w:anchor"))
        bookmark = xpath(doc.element, f'.//w:bookmarkStart[@w:name="{anchor}"]')[0]
        heading = normalized("".join(xpath(bookmark, "ancestor::w:p[1]//w:t/text()")))
        if not heading:
            heading = re.sub(r"^\d+(?:\.\d+)*\s+", "", xpath(hyperlink, ".//w:t/text()")[0])
        assert heading
        hits = [i + 1 for i, page in enumerate(pages) if i > 1 and heading in page]
        assert hits, heading
        nodes = xpath(hyperlink, ".//w:t")
        assert nodes[-1].text.isdigit()
        records.append({"heading": heading, "old_page": nodes[-1].text, "page": hits[0]})
        nodes[-1].text = str(hits[0])
    assert len(records) == 23
    while not doc.paragraphs[-1].text and not doc.paragraphs[-1]._p.xpath(".//w:drawing | .//w:sectPr"):
        p = doc.paragraphs[-1]._p
        p.getparent().remove(p)
    doc.save(OUTPUT)
    checks = json.loads((OUT / "revision_checks.json").read_text())
    checks["output_sha256"] = hashlib.sha256(OUTPUT.read_bytes()).hexdigest()
    checks["source_unchanged"] = hashlib.sha256(SOURCE.read_bytes()).hexdigest() == checks["source_sha256"]
    assert checks["source_unchanged"]
    reopened = Document(OUTPUT)
    for name, expected in checks["raw_blocks_verified"].items():
        values, active, identity = [], False, None
        for p in reopened.paragraphs:
            markers = p._p.xpath(f'.//w:bookmarkStart[@w:name="{name}"]')
            if markers:
                active, identity = True, markers[0].get(qn("w:id"))
            if active:
                values.append(p.text)
                if p._p.xpath(f'.//w:bookmarkEnd[@w:id="{identity}"]'):
                    break
        assert hashlib.sha256("\n".join(values).encode()).hexdigest() == expected["sha256"]
    checks["contents_page_updates"] = records
    (OUT / "revision_checks.json").write_text(json.dumps(checks, indent=2) + "\n")
    print(json.dumps(records, indent=2))


if __name__ == "__main__":
    main()
