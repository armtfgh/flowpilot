"""Extract structured review material without changing source documents."""
from pathlib import Path
from hashlib import sha256
import json
import shutil
from zipfile import ZipFile
from lxml import etree as E
from docx import Document
import fitz

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'manuscript/revision_20260922'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}


def main():
    review = OUT / 'source_review'
    review.mkdir(parents=True, exist_ok=True)
    records = []
    for name in ('manuscript_text_revised_round2.docx', 'esi.docx', 'manuscript.docx'):
        source = ROOT / 'manuscript' / name
        doc = Document(source)
        paragraphs = []
        for i, p in enumerate(doc.paragraphs):
            paragraphs.append(dict(index=i, style=p.style.name, text=p.text,
                drawings=p._p.xpath('.//a:blip/@r:embed'),
                runs=[dict(text=r.text, font=r.font.name, size=r.font.size.pt if r.font.size else None,
                           superscript=r.font.superscript, highlight=str(r.font.highlight_color),
                           bold=r.font.bold, italic=r.font.italic) for r in p.runs]))
        tables = [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
        data = dict(source=str(source), sha256=sha256(source.read_bytes()).hexdigest(), paragraphs=paragraphs,
                    tables=tables, sections=[dict(width=s.page_width.inches, height=s.page_height.inches,
                        left=s.left_margin.inches, right=s.right_margin.inches) for s in doc.sections])
        (review / (source.stem + '.json')).write_text(json.dumps(data, indent=2, ensure_ascii=False))
        lines = [f'# {name}', '']
        for p in paragraphs:
            if p['text'] or p['drawings']:
                lines.append(f'[{p["index"]}] ({p["style"]}) {p["text"]}')
                if p['drawings']:
                    lines.append('DRAWING: ' + ', '.join(p['drawings']))
        for i, t in enumerate(tables):
            lines.extend(['', f'TABLE {i}', *[' | '.join(row) for row in t]])
        (review / (source.stem + '.txt')).write_text('\n'.join(lines), encoding='utf-8')
        with ZipFile(source) as z:
            rels = E.fromstring(z.read('word/_rels/document.xml.rels'))
            mapping = {r.get('Id'): r.get('Target') for r in rels}
            media = {k: sha256(z.read(k)).hexdigest() for k in z.namelist() if k.startswith('word/media/')}
            (review / (source.stem + '_media.json')).write_text(json.dumps(dict(relationships=mapping, media=media), indent=2))
        records.append(dict(file=name, sha256=data['sha256'], paragraphs=len(paragraphs), tables=len(tables), images=len(media)))
    pdf = fitz.open(ROOT / 'manuscript/preprint.pdf')
    (review / 'preprint.txt').write_text('\n'.join(f'PAGE {i+1}\n{p.get_text()}' for i,p in enumerate(pdf)), encoding='utf-8')
    pages = []
    for i, page in enumerate(pdf):
        if 'Figure 5.' in page.get_text() or 'Figure 6.' in page.get_text():
            page.get_pixmap(matrix=fitz.Matrix(1.7,1.7)).save(review / f'preprint_page_{i+1:02d}.png')
            pages.append(i+1)
    (review / 'source_manifest.json').write_text(json.dumps(dict(documents=records, preprint_figure_pages=pages), indent=2))
    print(json.dumps(dict(documents=records, preprint_figure_pages=pages), indent=2))


if __name__ == '__main__':
    main()
